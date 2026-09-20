"""FastAPI application."""

from __future__ import annotations

import re
import shutil
import threading
from collections import defaultdict
from datetime import date, datetime, timedelta
from urllib.parse import urlencode
import subprocess
import sys
from pathlib import Path

from fastapi import FastAPI, Form, HTTPException, Request, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from . import (bootstrap, auth, bondmap, config, itemissue, leaves,
               pi as pi_mod,
               promote as promote_mod, reports_api, reports_pdf,
               targets as targets_mod)
from .jobs import STORE, Job, JobContext, JobStatus, run_pipeline
from .pipelines import (
    MONTH_NAMES,
    STREAMS,
    UploadRejected,
    canonical_secondary_name,
    canonical_shop_sales_name,
    forbidden_in_pipelines,
    preflight,
    validate_warehouse_name,
)

app = FastAPI(title="KSD Report Transformer")
templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))


def _trim0(text: str) -> str:
    """Drop a decimal point that has nothing but zeros after it.

    "5.00" is 5 with noise on the end, and a column of them is noise all the
    way down. Right to left, so "1,000.00" loses its zeros and then its point
    and stops at the nought it needs - it does not become "1,".
    """
    text = str(text)
    return text.rstrip("0").rstrip(".") if "." in text else text


def _cases(value, places: int = 0):
    """Round half away from zero, the way the office rounds.

    Jinja's own |round filter is Python's round(), which rounds a tie to the
    even number: |round on 20.5 gave 20 where the browser and Excel both print
    21 for the same figure. Templates use this instead.
    """
    from decimal import Decimal, ROUND_HALF_UP
    step = Decimal(1).scaleb(-places)
    got = Decimal(str(float(value or 0))).quantize(step, rounding=ROUND_HALF_UP)
    return int(got) if places == 0 else float(got)


templates.env.filters["trim0"] = _trim0
templates.env.filters["cases"] = _cases

# The brand mark, and anywhere else the UI needs a file rather than markup.
# Mounted rather than inlined so the browser caches it once instead of
# re-downloading it inside every page.
_STATIC = Path(__file__).parent / "static"
_STATIC.mkdir(exist_ok=True)
app.mount("/static", StaticFiles(directory=str(_STATIC)), name="static")

MAX_UPLOAD_BYTES = 80 * 1024 * 1024

# The upload page is organised by DATA TYPE, which is how the office thinks
# about it, rather than by pipeline. Several cards can feed the same pipeline:
# a daily shop-sales file and a cumulative-period file are the same stream with
# a different date mode, and making that a card each is clearer than a free-text
# "cumulative period" box people have to know the convention for.
#
# mode: "day"   -> one date
#       "range" -> from/to, sent to the pipeline as a cumulative period
#       "files" -> no date, the export carries its own
UPLOAD_CARDS = [
    {"id": "wh_stock", "title": "Warehouse Physical Stock", "stream": "warehouse_stock",
     "mode": "files", "icon": "warehouse", "ready": True},
    {"id": "shop_daily", "title": "Shop Sales - Daily", "stream": "shop_sales_daily",
     "mode": "day", "icon": "shop", "ready": True},
    {"id": "shop_cum", "title": "Shop Sales - Cumulative", "stream": "shop_sales_cumulative",
     "mode": "range", "icon": "calendar", "ready": True},
    {"id": "secondary", "title": "Secondary Sales - Daily", "stream": "secondary_sales",
     "mode": "day", "icon": "truck", "ready": True},
    {"id": "item_issue", "title": "Secondary Sales - Analysis", "stream": "item_issue",
     "mode": "batch", "icon": "truck", "ready": True},
    {"id": "pi_variance", "title": "Purchase Instruction", "stream": "purchase_instruction",
     "mode": "batch", "icon": "clipboard", "ready": True},
]

LEAVE_CARDS = [
    {"title": "Warehouse Leaves", "icon": "calendar-off",
     "href": "/settings/leaves/warehouse",
     "blurb": "Days a warehouse did not issue"},
    {"title": "Shop Leaves", "icon": "calendar-off",
     "href": "/settings/leaves/shop",
     "blurb": "Days shops were shut - dry days, hartals"},
]

# When the process started. Shown in the header so it is obvious at a glance
# whether the server you are looking at is the one you just restarted.
STARTED_AT = datetime.now().strftime("%d %b %H:%M")


@app.exception_handler(404)
async def page_not_found(request: Request, exc):
    """A missing PAGE gets an explanation; a missing API call stays JSON.

    Templates reload on refresh but routes do not, so a sidebar link can point
    at a route the running process has never heard of. Raw {"detail":"Not Found"}
    gives the operator nothing to act on; this says which half is stale.
    """
    wants_json = (request.url.path.startswith(("/api/", "/jobs/"))
                  or "application/json" in request.headers.get("accept", ""))
    if wants_json or not current_user(request):
        # A 404 raised with something to say keeps saying it: "Nothing is filed
        # for January 2026" is an answer, "Not Found" is a shrug.
        said = getattr(exc, "detail", None)
        return JSONResponse({"detail": said or "Not Found"}, status_code=404)
    return templates.TemplateResponse(
        request, "notfound.html",
        {"user": current_user(request), "page": "", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
        status_code=404,
    )


@app.middleware("http")
async def no_cache_html(request: Request, call_next):
    """Never let a browser cache a page of this app.

    All the CSS is inline, so a cached HTML page shows an old interface with no
    hint that anything is stale — which is indistinguishable from the code not
    having been deployed.
    """
    response = await call_next(request)
    if "text/html" in response.headers.get("content-type", ""):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
    return response


def _drain_pending() -> None:
    """Clear builds still sitting at the approval step that no longer exists.

    Deliberately NOT promoted. A build that has been waiting since before this
    restart was made from an older snapshot of the data, and writing it over
    whatever is live now could quietly undo newer work. Discarding costs one
    re-upload; promoting a stale build costs a day of numbers. The scratch
    output stays on disk either way, so nothing is actually destroyed.
    """
    stuck = [j for j in STORE.recent(200) if j.status is JobStatus.AWAITING_APPROVAL]
    for job in stuck:
        job.status = JobStatus.DISCARDED
        job.error = "Left over from the old approval step — upload it again and it will save itself."
        STORE.persist(job)
    if stuck:
        print(f"[startup] discarded {len(stuck)} build(s) left at the old approval step")


@app.on_event("startup")
def _startup() -> None:
    # A server's disk comes up bare. Lay out the folders the build scripts
    # expect and seed the reference material before anything reads it.
    try:
        seeded = bootstrap.ensure_workspace()
        if seeded:
            print("[startup] seeded " + ", ".join(seeded))
    except OSError as exc:
        print(f"[startup] PROBLEM: could not prepare the workspace — {exc}")

    config.JOBS_ROOT.mkdir(parents=True, exist_ok=True)
    STORE.load_existing()
    _drain_pending()

    problems = config.config_problems() + forbidden_in_pipelines()
    app.state.problems = problems
    # Pipeline preflight only makes sense once the workspace exists.
    if not problems:
        app.state.problems = preflight()

    for p in app.state.problems:
        print(f"[startup] PROBLEM: {p}")

    # Every report is a read over uploaded workbooks, and openpyxl is the whole
    # cost. It is paid once per file and then cached, but somebody has to pay
    # it - and it should not be the first person to open a report after a
    # deploy. A thread does the reading while the server is coming up.
    threading.Thread(target=reports_api.warm_cache, name="warm", daemon=True).start()


# ---------------------------------------------------------------------------
# Auth plumbing
# ---------------------------------------------------------------------------


def current_user(request: Request) -> str | None:
    return auth.read_session(request.cookies.get(auth.COOKIE_NAME))


def require_user(request: Request) -> str:
    user = current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="Sign in first.")
    return user


@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, error: str = ""):
    return templates.TemplateResponse(
        request, "login.html", {"error": error, "first_run": not auth.load_users()})


# A server has no terminal to run bootstrap_user.py in, so the first password
# is set through the app itself. It is open only while there are NO accounts at
# all, and only to an address already on the allowlist - so it is not a sign-up
# page, it is the one-time handover of an install to its owner.
@app.get("/first-run", response_class=HTMLResponse)
def first_run_form(request: Request, error: str = ""):
    if auth.load_users():
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "first_run.html", {"error": error, "allowed": sorted(config.ALLOWED_EMAILS)})


@app.post("/first-run")
def first_run(request: Request, email: str = Form(...), password: str = Form(...),
              confirm: str = Form("")):
    if auth.load_users():
        return RedirectResponse("/login", status_code=303)

    def again(msg: str):
        return templates.TemplateResponse(
            request, "first_run.html",
            {"error": msg, "allowed": sorted(config.ALLOWED_EMAILS)}, status_code=400)

    if password != confirm:
        return again("Those two passwords are not the same.")
    try:
        auth.set_password(email, password)
    except ValueError as exc:
        return again(str(exc))
    # The install belongs to whoever sets it up. Everybody added later can read
    # every report; only this account decides who they are.
    auth.set_owner(email)
    return RedirectResponse("/login?error=Password+set.+Sign+in+with+it.", status_code=303)


@app.post("/login")
def login(request: Request, email: str = Form(...), password: str = Form(...)):
    who = auth.authenticate(email, password)
    if not who:
        # first_run belongs on THIS render too. Without it, one failed attempt
        # hid the only way to create the first account - which is exactly when
        # somebody signing in to a brand new install needs to see it.
        return templates.TemplateResponse(
            request,
            "login.html",
            {"error": ("No account has been set up on this install yet — "
                       "set your password first."
                       if not auth.load_users()
                       else "That email and password combination was not accepted."),
             "first_run": not auth.load_users()},
            status_code=401,
        )
    response = RedirectResponse("/", status_code=303)
    response.set_cookie(
        auth.COOKIE_NAME,
        auth.issue_session(who),
        httponly=True,
        samesite="lax",
        secure=config.COOKIE_SECURE,
        max_age=config.SESSION_MAX_AGE_SECONDS,
    )
    return response


# ---------------------------------------------------------------------------
# People
# ---------------------------------------------------------------------------
#
# Adding a colleague used to mean editing render.yaml and redeploying. It is a
# normal Tuesday thing to do, so it belongs in the app - but with one hand on
# it. The owner adds and removes accounts; everybody else can see who holds
# one and nothing more. Before this, any account could remove any other, the
# owner's included, and the only way back was wiping the auth file on Render.


@app.get("/settings/people", response_class=HTMLResponse)
def people_page(request: Request, note: str = "", error: str = ""):
    user = require_user(request)
    return templates.TemplateResponse(
        request, "settings_people.html",
        {"user": user, "page": "people", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", []),
         "people": sorted(auth.load_users()), "owner": auth.owner(),
         "can_manage": auth.is_owner(user), "note": note, "error": error},
    )


@app.post("/settings/people/add")
def people_add(request: Request, email: str = Form(...), password: str = Form(...),
               confirm: str = Form("")):
    me = require_user(request)
    email = (email or "").strip().lower()

    def back(**kw):
        return RedirectResponse("/settings/people?" + urlencode(kw), status_code=303)

    if not auth.is_owner(me):
        return back(error="Only the owner of this install can add an account.")
    if password != confirm:
        return back(error="Those two passwords are not the same.")
    if email in auth.load_users():
        return back(error=f"{email} already has an account.")
    try:
        auth.invite(email)          # allowed from now on
        auth.set_password(email, password)
    except ValueError as exc:
        return back(error=str(exc))
    return back(note=f"{email} can sign in now. Send them the password yourself — "
                     "it is not stored anywhere readable.")


@app.post("/settings/people/remove")
def people_remove(request: Request, email: str = Form(...)):
    me = require_user(request)
    email = (email or "").strip().lower()

    def back(**kw):
        return RedirectResponse("/settings/people?" + urlencode(kw), status_code=303)

    if not auth.is_owner(me):
        return back(error="Only the owner of this install can remove an account.")
    if email == me:
        return back(error="You cannot remove your own account.")
    # Belt and braces: the owner is the only account that cannot be removed,
    # so an install can never be left with nobody able to manage it.
    if auth.is_owner(email):
        return back(error="The owner's account cannot be removed.")
    users = auth.load_users()
    users.pop(email, None)
    auth.save_users(users)
    auth.save_invited([e for e in auth.load_invited() if e != email])
    return RedirectResponse(
        "/settings/people?" + urlencode({"note": f"{email} can no longer sign in."}),
        status_code=303)


# ---------------------------------------------------------------------------
# Manage leaves
#
# A per-day figure divides by the days the thing could actually sell. Until
# now that divisor was the calendar, so a month with three dry days read as a
# month of bad trading. These two calendars are where the closed days live.
# ---------------------------------------------------------------------------

_LEAVE_LOOK = {
    "shop": {
        "heading": "Shop Leaves",
        "noun_pl": "shops",
        "blurb": "A day every KSBC outlet was shut - a dry day, a hartal, a "
                 "local festival. Pick the days on the calendar and write down "
                 "why. The 1st of the month repeats with one tick.",
        "feeds": "Used by <b>Shop Sales Analysis</b>: the per-day rate divides "
                 "by the days shops could trade, not by every day in the window. "
                 "The report says what it divided by and which days came out.",
    },
    "warehouse": {
        "heading": "Warehouse Leaves",
        "noun_pl": "warehouses",
        "blurb": "A day the KSBC warehouses did not issue. Pick the days on the "
                 "calendar and write down why.",
        "feeds": "Recorded, and shown here. <b>No report divides by it yet</b> - "
                 "say the word and secondary-sales dispatch rates can use it.",
    },
}


def _leave_month(month: str) -> tuple[int, int]:
    today = date.today()
    try:
        year, mon = int(str(month)[:4]), int(str(month)[5:7])
        date(year, mon, 1)
        return year, mon
    except (ValueError, TypeError):
        return today.year, today.month


@app.get("/settings/leaves/{kind}", response_class=HTMLResponse)
def leaves_page(request: Request, kind: str, month: str = "",
                note: str = "", error: str = ""):
    user = require_user(request)
    look = _LEAVE_LOOK.get(kind)
    if look is None:
        raise HTTPException(404, "There is no such leave calendar.")

    year, mon = _leave_month(month)
    today = date.today()
    rows = leaves.in_month(kind, year, mon)

    by_day: dict[str, list] = defaultdict(list)
    for row in rows:
        by_day[row["date"]].append(row)

    days = []
    for day in leaves.month_days(year, mon):
        here = by_day.get(day.isoformat(), [])
        said = [r.get("reason") for r in here if r.get("reason")]
        title = ("Shut" + (" - " + "; ".join(said) if said else "")) if here else "Open"
        days.append({"iso": day.isoformat(), "dom": day.day,
                     "all": bool(here), "today": day == today, "title": title})

    recorded = []
    for row in sorted(rows, key=lambda r: r["date"]):
        day = date.fromisoformat(row["date"])
        recorded.append({
            "id": row.get("id", ""), "reason": row.get("reason", ""),
            "nice": f"{day.day} {MONTH_NAMES[day.month - 1][:3].title()}",
            "who": f"All {look['noun_pl']}",
        })

    first = date(year, mon, 1)
    prev = first - timedelta(days=1)
    nxt = date(year + (mon == 12), (mon % 12) + 1, 1)
    return templates.TemplateResponse(
        request, "settings_leaves.html",
        {"user": user, "page": f"leaves_{kind}", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", []),
         "kind": kind, "heading": look["heading"], "noun_pl": look["noun_pl"],
         "blurb": look["blurb"], "feeds": look["feeds"],
         "ym": f"{year:04d}-{mon:02d}",
         "month_label": f"{MONTH_NAMES[mon - 1].title()} {year}",
         "prev_ym": f"{prev.year:04d}-{prev.month:02d}",
         "next_ym": f"{nxt.year:04d}-{nxt.month:02d}",
         "lead": first.weekday(), "days": days,
         "recorded": recorded, "note": note, "error": error},
    )


@app.post("/settings/leaves/{kind}/add")
def leaves_add(request: Request, kind: str, days: str = Form(""),
               reason: str = Form(""), monthly: str = Form(""),
               month: str = Form("")):
    user = require_user(request)
    if kind not in leaves.KINDS:
        raise HTTPException(404, "There is no such leave calendar.")
    got = leaves.add(kind, [d for d in days.split(",") if d.strip()],
                     reason, user, monthly=bool(monthly))
    back = {"month": month}
    if "error" in got:
        back["error"] = got["error"]
    else:
        said = f"{got['added']} day{'' if got['added'] == 1 else 's'} marked closed"
        if got["already"]:
            said += f" ({got['already']} already on the calendar)"
        back["note"] = said + "."
    return RedirectResponse(f"/settings/leaves/{kind}?" + urlencode(back), status_code=303)


@app.post("/settings/leaves/{kind}/remove")
def leaves_remove(request: Request, kind: str, id: str = Form(...),
                  month: str = Form("")):
    require_user(request)
    if kind not in leaves.KINDS:
        raise HTTPException(404, "There is no such leave calendar.")
    got = leaves.remove(kind, id)
    back = {"month": month}
    back["error" if "error" in got else "note"] = got.get("error", "Leave removed.")
    return RedirectResponse(f"/settings/leaves/{kind}?" + urlencode(back), status_code=303)


@app.post("/logout")
def logout():
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie(auth.COOKIE_NAME)
    return response


# ---------------------------------------------------------------------------
# Pages
# ---------------------------------------------------------------------------


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request,
        "index.html",
        {
            "user": user,
            "page": "upload",
            "started_at": STARTED_AT,
            "streams": list(STREAMS.values()),
            "streams_by_key": STREAMS,
            "cards": UPLOAD_CARDS,
            "leave_cards": LEAVE_CARDS,
            "jobs": STORE.recent(),
            "problems": getattr(app.state, "problems", []),
            "today": date.today().isoformat(),
        },
    )


@app.get("/healthz")
def healthz():
    problems = getattr(app.state, "problems", [])
    return JSONResponse(
        {"ok": not problems, "problems": problems},
        status_code=200 if not problems else 503,
    )


# ---------------------------------------------------------------------------
# Uploads and builds
# ---------------------------------------------------------------------------


def _resolve_shop_sales_workbook(month: int, year: int) -> Path:
    """Find the month's analysis workbook the driver should extend.

    A full-month file wins over a mid-month one. If neither exists the month
    has not been bootstrapped — say so plainly rather than letting the driver
    fail deep inside with a confusing message.
    """
    folder = STREAMS["shop_sales"].input_path()
    month_u = MONTH_NAMES[month - 1]

    full = folder / f"{month_u} SHOP SALES ANALYSIS.xlsx"
    if full.is_file():
        return full

    mid = sorted(
        folder.glob(f"{month_u} 1st - *ANALYSIS.xlsx"),
        key=lambda p: p.stat().st_mtime,
    )
    if mid:
        return mid[-1]

    # No workbook for this month yet. The driver cannot create one, so clone the
    # structural template — a workbook with every data sheet stripped out. The
    # month therefore starts empty and fills up only from what is uploaded.
    return _bootstrap_month(month_u, folder)


def _bootstrap_month(month_u: str, folder: Path) -> Path:
    """Start an empty month from the stripped template, via the locked script."""
    template = config.WORKSPACE_ROOT / "_seed" / "SHOP SALES TEMPLATE.xlsx"
    if not template.is_file():
        raise UploadRejected(
            f"No {month_u} workbook exists yet, and there is no shop-sales structure "
            "template to start one from. Shop Sales is the one pipeline that extends a "
            "month rather than building it, so it needs that template once. Re-run "
            "run_local.sh to build it from your own workbook."
        )

    out = folder / f"{month_u} 1st - 1st ANALYSIS.xlsx"
    script = config.SCRIPTS_DIR / "ksbc_bootstrap_month.py"
    proc = subprocess.run(
        [sys.executable, str(script), "--month", month_u,
         "--from", str(template), "--out", str(out), "--folder", str(folder)],
        capture_output=True, text=True, timeout=300,
    )
    if not out.is_file():
        raise UploadRejected(
            f"Could not start {month_u}: {(proc.stderr or proc.stdout or '').strip()[:300]}"
        )
    return out


def _save_upload(upload: UploadFile, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    size = 0
    with destination.open("wb") as fh:
        while chunk := upload.file.read(1024 * 1024):
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                fh.close()
                destination.unlink(missing_ok=True)
                raise UploadRejected(
                    f"'{upload.filename}' is larger than {MAX_UPLOAD_BYTES // (1024*1024)}MB."
                )
            fh.write(chunk)


@app.post("/build/{stream_key}")
async def build(
    request: Request,
    stream_key: str,
    files: list[UploadFile] = File(...),
    covers_date: str = Form(""),
    covers_from: str = Form(""),
    cumulative: str = Form(""),
):
    user = require_user(request)
    stream = STREAMS.get(stream_key)
    if stream is None:
        raise HTTPException(404, "Unknown report stream.")
    if getattr(app.state, "problems", []):
        raise HTTPException(503, "Service is not correctly configured; see /healthz.")

    job = STORE.create(stream_key, user)
    scratch_dir = STORE.dir_for(job)
    ctx = JobContext(job_id=job.id, stream_key=stream_key, scratch_dir=scratch_dir)

    try:
        placed: list[Path] = []
        receipt: dict | None = None

        for upload in files:
            suffix = Path(upload.filename or "").suffix.lower()
            if suffix not in stream.extensions:
                raise UploadRejected(
                    f"'{upload.filename}' is a {suffix or 'unknown'} file; "
                    f"{stream.label} expects {' or '.join(stream.extensions)}."
                )

        if stream_key == "warehouse_stock":
            for upload in files:
                name = validate_warehouse_name(upload.filename or "")
                target = stream.input_path() / name
                _save_upload(upload, target)
                placed.append(target)
            # Bevco writes the report date into the filename and the build
            # aborts on a mixed batch, so the first name speaks for all of them.
            # A date typed in the dialog wins: the operator is looking at the
            # file, and a renamed export should not silently file itself wrong.
            typed = _parse_date(covers_date) if covers_date else None
            job.covers = (typed.isoformat() if typed else
                          _warehouse_covers(job.uploaded_names or
                                            [p.name for p in placed]) or "")

        elif stream_key == "shop_sales_daily":
            # Standalone day: name it canonically and store it. No month
            # workbook to extend, so no contiguity question to answer.
            covers = _parse_date(covers_date)
            if len(files) != 1:
                raise UploadRejected("Upload one day at a time.")
            name = canonical_shop_sales_name(
                files[0].filename or "",
                day=covers.day, month=covers.month, year=covers.year,
            )
            target = stream.input_path() / name
            _save_upload(files[0], target)
            placed.append(target)
            job.covers = covers.isoformat()

        elif stream_key == "shop_sales":
            covers = _parse_date(covers_date)
            workbook = _resolve_shop_sales_workbook(covers.month, covers.year)
            ctx.meta["analysis_workbook"] = workbook
            job.summary["analysis_workbook"] = str(workbook)

            cum_range = _parse_cumulative(cumulative)
            if cum_range and len(files) != 1:
                raise UploadRejected("Upload one file at a time for a cumulative period.")
            # Answer the contiguity question now, in a second, instead of after
            # a three-minute build that the driver would abort anyway.
            _check_shop_day_gap(workbook, MONTH_NAMES[covers.month - 1],
                                covers.day, cum_range)

            for upload in files:
                name = canonical_shop_sales_name(
                    upload.filename or "",
                    day=covers.day,
                    month=covers.month,
                    year=covers.year,
                    cumulative=cum_range,
                )
                target = stream.input_path() / name
                _save_upload(upload, target)
                placed.append(target)

            if cum_range:
                job.covers = covers.replace(day=cum_range[0]).isoformat()
                job.covers_to = covers.replace(day=cum_range[1]).isoformat()
            else:
                job.covers = covers.isoformat()

        elif stream_key == "shop_sales_cumulative":
            # A period, stored whole. The two cumulative PDFs are drawn from
            # this file on demand, so nothing is built here and nothing can go
            # stale between the raw and the report.
            covers = _parse_date(covers_date)
            cum_range = _parse_cumulative(cumulative)
            if len(files) != 1:
                raise UploadRejected("Upload one cumulative period at a time.")
            if not cum_range:
                raise UploadRejected(
                    "A cumulative upload needs the period it covers - pick a From and a To date.")
            # Year-stamped, unlike the script-facing daily names: these files
            # are kept forever in one folder, and 'september 1-16' alone would
            # have next September overwrite this one.
            name = (f"{covers.year}-{covers.month:02d} {MONTH_NAMES[covers.month - 1]} "
                    f"{cum_range[0]}-{cum_range[1]} CUMULATIVE.xlsx")
            target = stream.input_path() / name
            target.parent.mkdir(parents=True, exist_ok=True)
            _save_upload(files[0], target)
            placed.append(target)
            job.covers = covers.replace(day=cum_range[0]).isoformat()
            job.covers_to = covers.replace(day=cum_range[1]).isoformat()

        elif stream_key == "secondary_sales":
            covers = _parse_date(covers_date)
            if len(files) != 1:
                raise UploadRejected("Secondary sales raws are uploaded one day at a time.")
            name = canonical_secondary_name(day=covers.day, month=covers.month, year=covers.year)
            target = stream.input_path() / name
            _save_upload(files[0], target)
            placed.append(target)
            job.covers = covers.isoformat()
            # What the file actually holds. A KSBC secondary pull is cumulative,
            # so one upload can carry a fortnight of dispatch; saying so here is
            # the difference between a report that looks invented and one that
            # can be traced back to the raw it came from.
            job.spans_from = _first_dispatch_day(target)

        elif stream_key == "purchase_instruction":
            # The month is picked in the dialog, but the files still name their
            # own Report Month - and ~295 of them agreeing beats one typed date.
            # So the pick is a check: pi.store() files the batch under the month
            # the files declare and refuses it if that is not the month picked.
            expect = ""
            if covers_date:
                try:
                    expect = _parse_date(covers_date).strftime("%Y-%m")
                except UploadRejected:
                    expect = ""
            staged = []
            for upload in files:
                target = scratch_dir / (upload.filename or "pi.xls")
                _save_upload(upload, target)
                staged.append(target)
            got = pi_mod.store(staged, expect=expect)
            if "error" in got:
                raise UploadRejected(got["error"])
            placed = sorted(pi_mod.month_dir(got["month"]).glob("*.xls*"))
            job.covers = f"{got['month']}-01"
            job.note = (f"{got['label']}: {got['ok']} shops with an instruction"
                        + (f", {got['blank']} blank" if got["blank"] else "")
                        + (f", {got['failed']} unreadable" if got["failed"] else ""))
            # What landed, file by file. The operator drops 295 files in one go;
            # a count they can check is the difference between a report they
            # trust and one they hope about.
            receipt = got
            job.summary["receipt"] = got

        elif stream_key == "item_issue":
            # Every export states the range it covers, and twenty-eight of them
            # agreeing beats a typed date - so the period picked on the screen
            # is a check on that, and the answer when an export names no range.
            expect = None
            if covers_from and covers_date:
                a, b = _parse_date(covers_from), _parse_date(covers_date)
                if b < a:
                    a, b = b, a
                expect = (a, b)
            staged = []
            for upload in files:
                target = scratch_dir / (upload.filename or "issue.xls")
                _save_upload(upload, target)
                staged.append(target)
            got = itemissue.store(staged, expect=expect)
            if "error" in got:
                raise UploadRejected(got["error"])
            placed = sorted((itemissue.root() / got["period"]).glob("*.xls*"))
            job.covers = got["period"].split("_")[1]
            job.note = (f"{got['label']}: {got['warehouses']} warehouses"
                        + (f", {got['undated']} without a stated period"
                           if got.get("undated") else ""))

        ctx.uploaded = placed
        job.uploaded_names = [p.name for p in placed]
        # Keep the raws with the job. The live input folder is shared and the
        # warehouse build deletes what it consumes, so this is the only copy
        # that survives — and history is worth little if you cannot get back
        # the file a day was built from.
        _archive_raws(job, placed)
        STORE.persist(job)

    except UploadRejected as exc:
        job.status = JobStatus.FAILED
        job.error = str(exc)
        STORE.persist(job)
        return JSONResponse({"job_id": job.id, "error": str(exc)}, status_code=400)

    threading.Thread(target=run_pipeline, args=(job, ctx), daemon=True).start()
    reply = {"job_id": job.id}
    if receipt:
        reply["receipt"] = receipt
    return JSONResponse(reply, status_code=202)


_WH_DATE_RE = re.compile(r"^Report[ _](\d{1,2})-([A-Za-z]{3})-(\d{4})", re.IGNORECASE)
_MON3 = {m[:3].upper(): i for i, m in enumerate(MONTH_NAMES, 1)}


def _warehouse_covers(names: list[str]) -> str:
    """The report date written into a Bevco export's filename, as ISO."""
    for name in names:
        m = _WH_DATE_RE.match(name)
        if not m:
            continue
        month = _MON3.get(m.group(2).upper())
        if not month:
            continue
        try:
            return date(int(m.group(3)), month, int(m.group(1))).isoformat()
        except ValueError:
            continue
    return ""


_DAY_SHEET_RE = r"^{month}\s+(\d+)(?:-(\d+))?$"
_CUM_SHEET_RE = r"^{month}\s+(\d+)-(\d+)\s+CUMULATIVE$"


def _shop_days_covered(workbook: Path, month_u: str) -> int:
    """Last day of the month this workbook already holds, 0 if none.

    Mirrors what ksbc_daily_update.py works out for itself, so the app can give
    the same answer before spending the time.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(workbook, read_only=True)
    except OSError:
        return 0
    last = 0
    for name in wb.sheetnames:
        for pattern in (_CUM_SHEET_RE, _DAY_SHEET_RE):
            m = re.match(pattern.format(month=month_u), name.strip(), re.IGNORECASE)
            if m:
                end = int(m.group(2) or m.group(1))
                last = max(last, end)
                break
    wb.close()
    return last


def _check_shop_day_gap(workbook: Path, month_u: str, day: int,
                        cum_range: tuple[int, int] | None) -> None:
    covered = _shop_days_covered(workbook, month_u)
    start = cum_range[0] if cum_range else day
    if start <= covered + 1:
        return

    missing = f"{covered + 1}" if covered + 1 == start - 1 else f"{covered + 1}–{start - 1}"
    month_t = month_u.title()
    if covered == 0:
        raise UploadRejected(
            f"{month_t} has nothing in it yet, and this file covers day {start}. "
            f"Shop sales is built day by day, so it cannot start at {start} — the "
            f"month's totals would be missing {month_t} {missing}. Either upload "
            f"those days first, or upload the {month_t} 1–{start - 1} period export "
            f"on the Shop Sales - Cumulative card and then this one."
        )
    raise UploadRejected(
        f"{month_t} currently covers up to day {covered}, and this file covers day "
        f"{start}. {month_t} {missing} would be missing, which would understate the "
        f"month. Upload the missing day(s) first."
    )


def _first_dispatch_day(raw: Path) -> str:
    """The earliest dispatch date inside a secondary raw, as ISO. '' if unreadable."""
    try:
        dates = [l["date"] for l in reports_api.load_dispatch_lines(raw) if l.get("date")]
    except Exception:
        return ""
    return min(dates).isoformat() if dates else ""


def _archive_raws(job: Job, placed: list[Path]) -> None:
    keep = STORE.dir_for(job) / "raw"
    keep.mkdir(parents=True, exist_ok=True)
    saved = 0
    for src in placed:
        try:
            shutil.copy2(src, keep / src.name)
            saved += 1
        except OSError:
            # A missing archive copy must never fail the upload itself — but
            # the flag has to stay honest, or the screen offers a download
            # that 404s.
            pass
    job.raw_kept = saved == len(placed) and saved > 0


def _rounded(value: str) -> bool:
    """Whole cases unless the page says otherwise.

    Round off is what the office reads out, so it is the default: a request
    that says nothing gets whole figures, and only an explicit 0 - which every
    switch on every report sends when it is turned off - asks for the two
    decimals underneath.
    """
    return str(value).strip().lower() not in ("0", "false", "no", "off")


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError:
        raise UploadRejected("Tell the app which date this export covers.")


def _parse_cumulative(value: str) -> tuple[int, int] | None:
    if not value.strip():
        return None
    m = re.match(r"^\s*(\d{1,2})\s*-\s*(\d{1,2})\s*$", value)
    if not m:
        raise UploadRejected("A cumulative period looks like '1-16' or '17-30'.")
    return int(m.group(1)), int(m.group(2))


# ---------------------------------------------------------------------------
# Job status, approval, download
# ---------------------------------------------------------------------------


@app.get("/jobs/{job_id}")
def job_status(request: Request, job_id: str):
    require_user(request)
    job = STORE.get(job_id)
    if job is None:
        raise HTTPException(404, "No such build.")
    return JSONResponse(job.to_dict())


@app.get("/jobs/{job_id}/log", response_class=HTMLResponse)
def job_log(request: Request, job_id: str):
    require_user(request)
    job = STORE.get(job_id)
    if job is None:
        raise HTTPException(404, "No such build.")
    log = STORE.dir_for(job) / "build.log"
    text = log.read_text() if log.is_file() else "(no log)"
    return HTMLResponse(f"<pre style='white-space:pre-wrap;font:12px/1.5 ui-monospace,monospace;padding:16px'>{_escape(text)}</pre>")


def _escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


@app.post("/jobs/{job_id}/approve")
def approve(request: Request, job_id: str):
    require_user(request)
    job = STORE.get(job_id)
    if job is None:
        raise HTTPException(404, "No such build.")
    try:
        target = promote_mod.promote(job)
    except promote_mod.PromotionError as exc:
        return JSONResponse({"error": str(exc)}, status_code=409)
    return JSONResponse({"ok": True, "promoted_to": str(target)})


@app.post("/jobs/{job_id}/discard")
def discard(request: Request, job_id: str):
    require_user(request)
    job = STORE.get(job_id)
    if job is None:
        raise HTTPException(404, "No such build.")
    try:
        promote_mod.discard(job)
    except promote_mod.PromotionError as exc:
        return JSONResponse({"error": str(exc)}, status_code=409)
    return JSONResponse({"ok": True})


@app.post("/api/uploads/{stream_key}/delete")
def delete_upload_date(request: Request, stream_key: str, day: str = Form(...),
                       day_to: str = Form("")):
    """Remove an uploaded date: its runs, its raws, and the data it produced.

    Deleting an upload has to mean the reports stop showing it. The raws are
    the truth here; a month's analysis workbook is something the build made out
    of them, so once one of its days is deleted the workbook no longer
    describes anything that was uploaded and it goes too. Whatever raws remain
    answer on their own, which is how a month rebuilds itself down to nothing.

    Nothing is erased: the workbook is moved to _deleted/ in the workspace,
    where it can be fetched back by hand if this was a mistake.
    """
    require_user(request)
    if stream_key not in STREAMS:
        raise HTTPException(404, "No such data type.")
    stream = STREAMS[stream_key]

    # A cumulative upload is identified by its whole period. Matching on the
    # first day alone meant deleting 1-2 September took 1-3 and 1-4 with it -
    # every window that happened to start on the same day.
    def same_window(j) -> bool:
        if (j.covers or j.created_at[:10]) != day:
            return False
        return (j.covers_to or "") == day_to if day_to else True

    jobs = [j for j in STORE.recent(500)
            if j.stream_key == stream_key and same_window(j)]
    if not jobs:
        raise HTTPException(404, "Nothing uploaded for that period.")

    removed = 0
    for job in jobs:
        for name in job.uploaded_names:
            live = stream.input_path() / name
            try:
                if live.is_file():
                    live.unlink()
                    removed += 1
            except OSError:
                pass
        shutil.rmtree(STORE.dir_for(job), ignore_errors=True)
        STORE.forget(job.id)

    note = ""
    if stream_key == "warehouse_stock":
        rows, books = _purge_stock_date(day)
        note = (f"Removed {rows} history row(s) and {books} day workbook(s) — "
                "the stock report will stop showing that date.")
    else:
        moved = _retire_derived(stream_key, day)
        left = _uploads_left(stream_key)
        if moved:
            note = (f"Removed {removed} raw file(s) and retired {len(moved)} built "
                    f"workbook(s) ({', '.join(moved)}) to _deleted/. ")
            note += ("The reports now show only what is still uploaded."
                     if left else "Nothing is uploaded for this report any more, "
                                  "so it will read as empty.")
        else:
            note = ("Removed the raw file(s). The reports now show only what is "
                    "still uploaded.")

    return JSONResponse({"ok": True, "runs": len(jobs), "raws": removed, "note": note})


# Which built workbooks belong to which stream. They are outputs, not sources:
# every one of them was made from raws that were uploaded, so when those raws
# go, so does the workbook - otherwise a report keeps quoting a file nobody can
# trace back to an upload.
_DERIVED = {
    "secondary_sales": ("Secondary sales", "*SECONDARY SALES ANALYSIS.xlsx"),
    "shop_sales": ("KSBC shop sales", "*SHOP SALES ANALYSIS.xlsx"),
}


def _retire_derived(stream_key: str, day: str) -> list[str]:
    """Move that month's built workbooks out of the way. Returns their names."""
    where = _DERIVED.get(stream_key)
    if not where:
        return []
    folder = config.CLAUDE_ROOT / where[0]
    if not folder.is_dir():
        return []

    try:
        month_u = MONTH_NAMES[date.fromisoformat(day).month - 1]
    except (ValueError, IndexError):
        return []

    bin_dir = config.WORKSPACE_ROOT / "_deleted" / datetime.now().strftime("%Y%m%d_%H%M%S")
    moved: list[str] = []
    for path in sorted(folder.glob(where[1])):
        if month_u not in path.name.upper():
            continue
        try:
            bin_dir.mkdir(parents=True, exist_ok=True)
            shutil.move(str(path), str(bin_dir / path.name))
            moved.append(path.name)
        except OSError:
            pass
    return moved


def _uploads_left(stream_key: str) -> bool:
    """Is anything still uploaded for this stream?"""
    return any(j.stream_key == stream_key for j in STORE.recent(500))


def _purge_stock_date(day: str) -> tuple[int, int]:
    """Drop one report date from the warehouse stock history and day workbooks."""
    import csv

    rows_removed = 0
    hist = config.CLAUDE_ROOT / "Warehouse stock" / "_history"
    if hist.is_dir():
        for path in sorted(hist.glob("*.csv")):
            try:
                with path.open(newline="") as fh:
                    rows = list(csv.reader(fh))
            except OSError:
                continue
            if not rows:
                continue
            header = rows[0]
            if "date" not in header:
                continue
            ix = header.index("date")
            kept = [r for r in rows[1:] if not (len(r) > ix and r[ix] == day)]
            dropped = len(rows) - 1 - len(kept)
            if not dropped:
                continue
            try:
                with path.open("w", newline="") as fh:
                    writer = csv.writer(fh)
                    writer.writerow(header)
                    writer.writerows(kept)
                rows_removed += dropped
            except OSError:
                continue

    books = 0
    try:
        d = date.fromisoformat(day)
    except ValueError:
        return rows_removed, books
    target = (config.CLAUDE_ROOT / "Warehouse stock" /
              f"{MONTH_NAMES[d.month - 1]} {_day_ordinal(d.day)} WAREHOUSE STOCK.xlsx")
    try:
        if target.is_file():
            target.unlink()
            books += 1
    except OSError:
        pass
    return rows_removed, books


def _day_ordinal(n: int) -> str:
    if 10 <= n % 100 <= 20:
        return f"{n}th"
    return f"{n}{ {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th') }"


@app.get("/jobs/{job_id}/raw")
def download_raw(request: Request, job_id: str, name: str = ""):
    """The raw file(s) this upload was built from — one file, or all as a zip."""
    require_user(request)
    job = STORE.get(job_id)
    if job is None:
        raise HTTPException(404, "No such build.")

    folder = STORE.dir_for(job) / "raw"
    have = [n for n in job.uploaded_names if (folder / n).is_file()]
    if not have:
        raise HTTPException(404, "The raw files for this upload are no longer on disk.")

    if name:
        if name not in job.uploaded_names:
            raise HTTPException(404, "No such file in this upload.")
        one = folder / name
        if not one.is_file():
            raise HTTPException(404, "That file is no longer on disk.")
        return FileResponse(one, filename=name)

    if len(have) == 1:
        return FileResponse(folder / have[0], filename=have[0])

    import tempfile
    import zipfile

    label = STREAMS[job.stream_key].label
    day = job.covers or job.created_at[:10]
    out = Path(tempfile.mkdtemp()) / f"{label} — {day} (raw files).zip"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in have:
            z.write(folder / n, n)
    return FileResponse(out, filename=out.name)


@app.get("/jobs/{job_id}/download")
def download(request: Request, job_id: str, name: str = ""):
    require_user(request)
    job = STORE.get(job_id)
    if job is None:
        raise HTTPException(404, "No such build.")

    if name:
        # Only ever serve a file this job actually produced, by its own listed
        # name — never a path assembled from user input.
        if name not in job.artifacts:
            raise HTTPException(404, "No such file in this build.")
        artifact = STORE.dir_for(job) / name
        if not artifact.is_file():
            raise HTTPException(404, "That file is no longer on disk.")
        return FileResponse(artifact, filename=name)

    if not job.output_path:
        raise HTTPException(404, "Nothing to download.")
    path = Path(job.output_path)
    if not path.is_file():
        raise HTTPException(404, "The output file is no longer on disk.")

    stream = STREAMS[job.stream_key]
    nice = f"{stream.label} — {job.created_at[:10]}.xlsx"
    return FileResponse(path, filename=nice)


# ---------------------------------------------------------------------------
# Reports (read-only)
# ---------------------------------------------------------------------------


@app.get("/reports/brandwise", response_class=HTMLResponse)
def brandwise_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "report_brandwise.html",
        {"user": user, "page": "brandwise", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


def _report_args(date_from: str, date_to: str):
    def d(v):
        try:
            return date.fromisoformat(v) if v else None
        except ValueError:
            return None
    return d(date_from), d(date_to)


@app.get("/api/reports/brandwise")
def brandwise_data(
    request: Request,
    view: str = "bond",
    date_from: str = "",
    date_to: str = "",
    bond: str = "",
    warehouse: str = "",
    round_off: str = "",
):
    require_user(request)
    round_off = _rounded(round_off)
    if view not in ("bond", "warehouse", "shop"):
        raise HTTPException(400, "Unknown view.")
    f, t = _report_args(date_from, date_to)
    return JSONResponse(reports_api.brandwise(
        view=view, date_from=f, date_to=t, bond=bond,
        warehouse=warehouse, round_off=bool(round_off),
    ))


@app.get("/api/reports/brandwise/shops")
def brandwise_shops_data(
    request: Request,
    view: str = "bond",
    key: str = "",
    date_from: str = "",
    date_to: str = "",
    brands: str = "",
    round_off: str = "",
):
    """One group's shops, for the drill-down under a bond or warehouse row."""
    require_user(request)
    round_off = _rounded(round_off)
    if view not in ("bond", "warehouse"):
        raise HTTPException(400, "Unknown view.")
    f, t = _report_args(date_from, date_to)
    cols = [b for b in brands.split("\u001f") if b] if brands else None
    return JSONResponse(reports_api.brandwise_shops(
        view=view, key=key, date_from=f, date_to=t,
        round_off=bool(round_off), brands=cols,
    ))


def centre_all(wb) -> None:
    """Every cell in the book, centred - the office reads these side by side.

    Each export grew its own alignment rules: the stock sheet left-aligned
    item names, liquidation right-aligned its figures, the analysis sheet
    left-aligned the first column and centred the rest. Printed and laid next
    to each other they did not line up, and the rule people actually wanted
    turned out to be the simple one. Vertical placement and wrapping are left
    as each sheet set them - those are about row height, not about columns.
    """
    from openpyxl.styles import Alignment

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                a = c.alignment
                c.alignment = Alignment(horizontal="center",
                                        vertical=a.vertical or "center",
                                        wrap_text=a.wrap_text,
                                        indent=0, shrink_to_fit=a.shrink_to_fit,
                                        text_rotation=a.text_rotation)


@app.get("/reports/brandwise/export.xlsx")
def brandwise_xlsx(
    request: Request,
    view: str = "bond",
    date_from: str = "",
    date_to: str = "",
    bond: str = "",
    warehouse: str = "",
    round_off: str = "",
):
    """The cumulative sheet as a workbook, laid out as the PDF prints it.

    Same two bands, the same navy header with gold labels, the same zebra body
    with dimmed zeros, cluster subtotals on navy and one gold-on-navy total -
    so the workbook and the PDF of the same view cannot be told apart.
    """
    require_user(request)
    round_off = _rounded(round_off)
    f, t = _report_args(date_from, date_to)
    data = reports_api.brandwise(view=view, date_from=f, date_to=t, bond=bond,
                                 warehouse=warehouse, round_off=bool(round_off))
    if "error" in data:
        raise HTTPException(404, data["error"])

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    NAVY, GOLD = "FF0A294F", "FFFFBD30"
    INK, ZEBRA, ZERO = "FF1B2A4A", "FFF5F7FB", "FFC7CDD8"
    HAIR = Side(style="thin", color="FFDCE1EA")
    BOX = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)

    brands = data["brands"]
    span = len(brands) + 2                     # label + brands + TOTAL
    last_col = get_column_letter(span)

    wb = Workbook()
    ws = wb.active
    ws.title = f"CUMULATIVE {view.upper()}"[:31]
    ws.sheet_view.showGridLines = False

    def paint(row, col, fill=None, colour=INK, bold=False, size=10,
              align="center", fmt=None, value=None, wrap=False):
        c = ws.cell(row=row, column=col)
        if value is not None:
            c.value = value
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(bold=bold, size=size, color=colour)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        c.border = BOX
        if fmt:
            c.number_format = fmt
        return c

    def nice(value):
        # The window arrives as dates, the span as ISO text - the band prints
        # either.
        if not value:
            return ""
        d = value if isinstance(value, date) else date.fromisoformat(str(value)[:10])
        return f"{d.day} {reports_pdf._MONTH_ABBR[d.month - 1]} {d.year}"

    span_dates = data.get("span") or {}
    period = f"{nice(f or span_dates.get('min'))} to {nice(t or span_dates.get('max'))}"
    scope = f"{data['label'].upper()} VIEW"
    for extra in (bond, warehouse):
        if extra:
            scope += f"  \u00b7  {extra}"

    # ---- the two bands ----
    ws.merge_cells(f"A1:{last_col}1")
    for col in range(1, span + 1):
        paint(1, col, NAVY)
    paint(1, 1, NAVY, GOLD, True, 16, value="K.S DISTILLERY")
    ws.row_dimensions[1].height = 30

    half = max(2, span - 2)
    for row, height in ((2, 22), (3, 20)):
        for col in range(1, span + 1):
            paint(row, col, GOLD)
        ws.row_dimensions[row].height = height
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=half)
    ws.merge_cells(start_row=2, start_column=half + 1, end_row=2, end_column=span)
    paint(2, 1, GOLD, NAVY, True, 11, "left", value="SECONDARY SALES - CUMULATIVE")
    paint(2, half + 1, GOLD, NAVY, True, 10, "right", value=period)
    ws.merge_cells(f"A3:{last_col}3")
    paint(3, 1, GOLD, NAVY, True, 11, value=scope)

    # ---- the header ----
    head = 4
    paint(head, 1, NAVY, GOLD, True, 9, "left", value=data["label"].upper(), wrap=True)
    for i, name in enumerate(brands, start=2):
        paint(head, i, NAVY, GOLD, True, 9, value=name, wrap=True)
    paint(head, span, NAVY, GOLD, True, 9, value="TOTAL", wrap=True)
    ws.row_dimensions[head].height = 40

    # ---- the body ----
    # A zero is dimmed rather than dropped, exactly as the PDF prints it: the
    # eye should land on the figures that moved.
    FIG = "#,##0" if round_off else "#,##0.##"
    r = head
    stripe = 0
    for row in data["rows"]:
        r += 1
        cluster = row.get("kind") == "cluster"
        if cluster:
            fill, colour, bold = NAVY, GOLD, True
        else:
            fill = ZEBRA if stripe % 2 == 0 else "FFFFFFFF"
            colour, bold = INK, False
            stripe += 1
        paint(r, 1, fill, colour, True, 10, "left", value=str(row["name"]))
        for i, brand in enumerate(brands, start=2):
            v = row["cells"].get(brand, 0) or 0
            paint(r, i, fill, colour if (v or cluster) else ZERO, bold, 10,
                  fmt=FIG, value=v)
        paint(r, span, fill, colour, True, 10, fmt=FIG, value=row.get("total", 0))
        ws.row_dimensions[r].height = 17

    # ---- the total ----
    # The PDF rules the total band off in gold. Without it the last cluster and
    # the grand total are two navy rows running into each other.
    RULE = Side(style="medium", color=GOLD)
    r += 1
    paint(r, 1, NAVY, GOLD, True, 10, "left", value="GRAND TOTAL")
    for i, brand in enumerate(brands, start=2):
        paint(r, i, NAVY, GOLD, True, 10, fmt=FIG,
              value=data["grand"]["cells"].get(brand, 0) or 0)
    paint(r, span, NAVY, GOLD, True, 10, fmt=FIG, value=data["grand"]["total"])
    for col in range(1, span + 1):
        c = ws.cell(row=r, column=col)
        c.border = Border(left=HAIR, right=HAIR, top=RULE, bottom=RULE)
    ws.row_dimensions[r].height = 19

    # ---- the shape of the page ----
    ws.column_dimensions["A"].width = 26
    for col in range(2, span):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions[last_col].width = 12
    ws.freeze_panes = ws.cell(row=head + 1, column=2)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4,
                                  header=0.2, footer=0.2)
    ws.print_area = f"A1:{last_col}{r}"
    ws.print_title_rows = f"1:{head}"

    tmp = Path(tempfile.mkdtemp()) / "Secondary Sales - Cumulative.xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


@app.get("/reports/brandwise/export.pdf")
def brandwise_pdf(
    request: Request,
    scope: str = "cluster",
    cluster: int = 1,
    view: str = "bond",
    date_from: str = "",
    date_to: str = "",
    bond: str = "",
    warehouse: str = "",
    round_off: str = "",
):
    """Two shapes of PDF from one endpoint.

    scope="cluster"  one ASM cluster, in the fixed warehouse-per-page layout the
                     field team already receives. Built by the same script the
                     Secondary Sales pipeline runs, so a PDF downloaded here is
                     byte-for-byte the kind of artefact a build produces.
    scope="current"  exactly what the page is showing — the chosen view, the
                     live filters and date range, cluster subtotals and grand
                     total. Rendered from the same brandwise() result that fed
                     the table, so the two cannot drift apart.
    """
    require_user(request)
    round_off = _rounded(round_off)

    # Only the official cluster books need the month's workbook - they are built
    # by the same script the pipeline runs. The current view renders from the
    # figures already on screen, which the raws alone can answer for.
    workbook = reports_api.find_secondary_workbook()
    if workbook is None and scope != "current":
        raise HTTPException(404, "No secondary-sales workbook has been built yet. "
                                 "The current view still exports.")

    import subprocess, tempfile
    from .pipelines import APP_REPORTS

    outdir = Path(tempfile.mkdtemp())

    if scope == "current":
        if view not in ("bond", "warehouse", "shop"):
            raise HTTPException(400, "Unknown view.")
        f, t = _report_args(date_from, date_to)
        data = reports_api.brandwise(view=view, date_from=f, date_to=t, bond=bond,
                                     warehouse=warehouse, round_off=bool(round_off))
        if "error" in data:
            raise HTTPException(404, data["error"])
        if not data["rows"]:
            raise HTTPException(404, "Nothing matches those filters.")

        name = f"Secondary Sales - Cumulative ({data['label']} View).pdf"
        out = reports_pdf.build_view_pdf(
            data, outdir / name, round_off=bool(round_off),
            filters={"bond": bond, "warehouse": warehouse,
                     "date_from": date_from, "date_to": date_to},
        )
        return FileResponse(out, filename=out.name)

    # scope="clusters" is all three at once. One request, one folder - the
    # page used to fire three downloads a second apart and hope the browser
    # allowed them all.
    wanted = [1, 2, 3] if scope == "clusters" else [cluster]
    if any(c not in (1, 2, 3) for c in wanted):
        raise HTTPException(400, "Cluster must be 1, 2 or 3.")

    built = []
    for c in wanted:
        into = outdir / f"c{c}"
        into.mkdir(parents=True, exist_ok=True)
        argv = ["python3", str(APP_REPORTS / "build_secondary_brandwise_pdfs.py"),
                "--base", str(config.CLAUDE_ROOT), "--workbook", str(workbook),
                "--outdir", str(into), "--cluster", str(c)]
        if date_from:
            argv += ["--from", date_from]
        if date_to:
            argv += ["--to", date_to]

        proc = subprocess.run(argv, capture_output=True, text=True,
                              timeout=config.STEP_TIMEOUT_SECONDS)
        if proc.returncode != 0:
            raise HTTPException(500, f"PDF build failed: {(proc.stderr or proc.stdout)[-400:]}")
        built += list(into.glob("*.pdf"))

    if not built:
        where = "any cluster" if len(wanted) > 1 else f"cluster {wanted[0]}"
        raise HTTPException(404, f"No dispatches for {where} in the selected range.")
    span = _asked_period(date_from, date_to, "")
    label = f" ({span['short']})" if span else ""
    return _as_folder(built, f"Secondary Sales - Cumulative{label}")



def _as_folder(files: list, name: str):
    """Several PDFs as one zip, because fifteen downloads is not a delivery.

    A report that splits into a file per cluster - or per bond - used to fire
    one download per file, spaced out because browsers throttle them, and the
    person ended up hunting fifteen PDFs through their Downloads folder. One
    archive opens as one folder with everything in it, named and in order.
    """
    import tempfile, zipfile

    keep = [Path(f) for f in files if Path(f).is_file()]
    if not keep:
        raise HTTPException(404, "Nothing to export for that selection.")
    if len(keep) == 1:
        return FileResponse(keep[0], filename=keep[0].name)

    safe = _safe_name(name)
    bundle = Path(tempfile.mkdtemp()) / f"{safe}.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(keep, key=lambda x: x.name):
            z.write(f, f"{safe}/{f.name}")
    return FileResponse(bundle, filename=bundle.name)


@app.get("/api/jobs")
def jobs_list(request: Request, stream: str = ""):
    """Recent builds, optionally for one stream.

    The upload page reads this instead of scraping a table out of the DOM, so
    the History dialog and the in-progress strip stay correct even though the
    page no longer renders a build list of its own.
    """
    require_user(request)
    # History is a permanent record: when a stream is named, every run it has
    # ever had comes back, not a recent window. The upload page asks without a
    # stream and only needs enough to colour the cards.
    limit = 5000 if stream else 60
    jobs = [j for j in STORE.recent(limit) if not stream or j.stream_key == stream]
    return JSONResponse({
        "jobs": [j.to_dict() for j in jobs],
        "labels": {k: v.label for k, v in STREAMS.items()},
    })


# ---------------------------------------------------------------------------
# Warehouse Stock Report
# ---------------------------------------------------------------------------


@app.get("/status-calendar", response_class=HTMLResponse)
def status_calendar(request: Request):
    """Which days the reports can answer for, and which they cannot."""
    user = require_user(request)
    return templates.TemplateResponse(
        request, "status_calendar.html",
        {"user": user, "page": "calendar", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", []),
         "calendar": reports_api.upload_calendar()},
    )


@app.get("/reports/warehouse-stock", response_class=HTMLResponse)
def stock_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "report_stock.html",
        {"user": user, "page": "stock", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


@app.get("/api/reports/warehouse-stock")
def stock_data(request: Request, as_of: str = "", cluster: str = "", warehouse: str = ""):
    require_user(request)
    cl = int(cluster) if cluster in ("1", "2", "3") else None
    return JSONResponse(reports_api.warehouse_stock(as_of=as_of, cluster=cl,
                                                    warehouse=warehouse))


# ---------------------------------------------------------------------------
# Shop sales - cumulative and analysis
# ---------------------------------------------------------------------------
#
# Both are drawn from the stored cumulative raw on demand rather than baked at
# upload time. That is what makes a period from months ago still openable, and
# it means a change to Bond Mapping shows up in every past period at once
# instead of only in whatever is rebuilt next.


def _analysis_module():
    """The PDF builder, imported as a module so the page and the PDF agree."""
    import importlib.util
    from .pipelines import APP_REPORTS
    spec = importlib.util.spec_from_file_location(
        "build_shop_analysis_pdf", APP_REPORTS / "build_shop_analysis_pdf.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _period_or_404(period: str) -> dict:
    chosen = reports_api.cumulative_period(period)
    if chosen is None:
        raise HTTPException(404, "No cumulative period has been uploaded yet.")
    return chosen


def _asked_dates(date_from: str, date_to: str, period: str = ""):
    """The window the caller actually asked for, so a refusal can name it."""
    if not (date_from and date_to) and period and ".." in period:
        date_from, date_to = period.split("..", 1)
    if not (date_from and date_to):
        return None
    try:
        start = datetime.strptime(date_from, "%Y-%m-%d").date()
        end = datetime.strptime(date_to, "%Y-%m-%d").date()
    except ValueError:
        return None
    return (start, end) if start <= end else (end, start)


def _asked_period(date_from: str, date_to: str, period: str = ""):
    got = _asked_dates(date_from, date_to, period)
    return reports_api.period_for(*got) if got else None


def _suggested_window():
    """The widest answerable window, for a screen that has to refuse one."""
    widest = reports_api.widest_window()
    if not widest:
        return None
    return {"from": widest[0].isoformat(), "to": widest[1].isoformat(),
            "short": reports_api.period_for(*widest)["short"]}


def _window(date_from: str, date_to: str, period: str = ""):
    """Whatever the caller asked for, as a window and a source.

    A calendar range is the primary input now; `period` is still honoured so an
    older bookmark or a link from another screen keeps working.
    """
    if not (date_from and date_to) and period and ".." in period:
        date_from, date_to = period.split("..", 1)
    if date_from and date_to:
        try:
            start = datetime.strptime(date_from, "%Y-%m-%d").date()
            end = datetime.strptime(date_to, "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(400, "Those dates are not readable.")
        if end < start:
            start, end = end, start
        return reports_api.resolve_window(start, end)

    # Default to the widest window the uploads can answer for, rather than the
    # newest file: with 1-16 and 17-23 on disk, 1-23 is the report you want.
    widest = reports_api.widest_window()
    if widest is None:
        return {"error": "Nothing has been uploaded yet."}
    return reports_api.resolve_window(*widest)


def _window_data(win: dict, cluster: int, bond: str, warehouse: str = "",
                 group_by: str = "bond") -> dict:
    """The grouping for a resolved window, whatever its source."""
    if win["source"] == "cumulative":
        return reports_api.shop_cumulative(win["period"], cluster or None, bond,
                                           warehouse, group_by)
    slim = {k: win["period"][k] for k in ("key", "short", "long", "days")}
    return reports_api._group_shops(win["shops"], slim, cluster or None, bond,
                                    warehouse, group_by)


@app.get("/reports/shop-cumulative", response_class=HTMLResponse)
def shop_cumulative_page(request: Request, date_from: str = "", date_to: str = "",
                         period: str = ""):
    user = require_user(request)
    periods = reports_api.cumulative_periods()
    # Seed the page with the same window the API would default to - the widest
    # the uploads can answer for. Taking the newest file instead would open on
    # 17-23 while a 1-16 sits next to it, which is not the report anyone wants.
    # An explicit window in the link wins, so a shared URL opens on its period.
    win = _window(date_from, date_to, period)
    chosen = None if "error" in win else win["period"]
    if chosen is None:
        asked = _asked_period(date_from, date_to, period)
        if asked:
            chosen = asked
    return templates.TemplateResponse(
        request, "report_shop_cumulative.html",
        {"user": user, "page": "shop_cumulative", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", []),
         "periods": periods, "chosen": chosen},
    )


@app.get("/api/shop-cumulative")
def shop_cumulative_api(request: Request, date_from: str = "", date_to: str = "",
                        period: str = "", cluster: int = 0, bond: str = "",
                        warehouse: str = "", group_by: str = "bond"):
    """Bond (or warehouse) rows with their shops underneath, for any window."""
    require_user(request)
    win = _window(date_from, date_to, period)
    if "error" in win:
        # A window we cannot tile is a normal thing to ask for, not a failure:
        # say what is missing, and offer the widest window that IS answerable
        # so the screen has something to act on rather than a dead end.
        asked = _asked_period(date_from, date_to, period)
        return JSONResponse({"error": win["error"], "calendar": _calendar_days(),
                             "asked": {"short": asked["short"]} if asked else None,
                             "need": win.get("need", []),
                             "suggest": _suggested_window()})

    data = _window_data(win, cluster, bond, warehouse, group_by)
    if "error" not in data:
        data["source"] = win["source"]
        data["chain"] = win.get("chain", [])
        data["source_block"] = _shop_source(win)
        data["all_warehouses"] = data.pop("warehouses", [])
        data["cluster"] = cluster
        data["bond"] = (bond or "").upper()
        data["warehouse"] = (warehouse or "").upper()
        data["calendar"] = _calendar_days()
    return JSONResponse(data)



def _shop_source(win: dict, prev: dict | None = None,
                 prev_span: tuple | None = None) -> dict:
    """Which uploads answered a shop window - and the one it is set against.

    A window to the 18th is normally 1-16 stitched to 17-18, because KSBC caps
    a pull at sixteen days. That is worth being able to check on any report
    that reads these files, so both shop reports build it the same way.
    """
    now = reports_api.src_windows(win.get("chain", []))
    per = win.get("period") or {}
    hole = reports_api.src_gap(per.get("start"), per.get("end"),
                               reports_api.src_chain_days(win.get("chain", [])))
    legs = [reports_api.src_leg("Shop sales (KSBC)",
                                now + ([hole] if hole else []),
                                per.get("short", ""))]
    if prev and prev.get("chain"):
        legs.append(reports_api.src_leg(
            "Compared against", reports_api.src_windows(prev["chain"]),
            prev["period"]["short"] if prev.get("period") else "", "green"))
    elif prev_span:
        # The month back was asked for and nothing answers it. Every row on
        # this report prints a dash in its last-month column because of that,
        # and this is the only place that says why.
        legs.append(reports_api.src_leg(
            "Compared against", [reports_api.src_missing(*prev_span)],
            "", "green"))
    return reports_api.src_block(legs)


def _calendar_days() -> dict:
    """What the calendar can offer: whole uploaded periods, and single days."""
    return {
        "periods": [{"from": p["start"].isoformat(), "to": p["end"].isoformat(),
                     "short": p["short"]} for p in reports_api.cumulative_periods()],
        "days": sorted(d.isoformat() for d in reports_api.shop_day_files()),
    }


@app.get("/api/shop-cumulative/shop")
def shop_cumulative_detail(request: Request, code: str, date_from: str = "",
                           date_to: str = "", period: str = ""):
    """One shop's brand and pack lines, fetched when its row is opened."""
    require_user(request)
    win = _window(date_from, date_to, period)
    if "error" in win:
        return JSONResponse({"error": win["error"]})
    cur = win["period"]
    return JSONResponse(reports_api.shop_detail(cur["start"], cur["end"], code))


@app.get("/reports/shop-cumulative/export.xlsx")
def shop_cumulative_xlsx(request: Request, date_from: str = "", date_to: str = "",
                         period: str = "", cluster: int = 0, bond: str = "",
                         warehouse: str = "", group_by: str = "bond",
                         round_off: str = ""):
    """The screen, as a workbook: bonds grouped, shops collapsible under them.

    The earlier flat dump pivoted well and read badly - four hundred rows with
    the bond repeated on every one, and nothing to tell a shop line from a
    total. This keeps the grouping people actually work with and lets Excel
    fold each bond away, while the autofilter still gives the flat view to
    anyone who wants to pivot.
    """
    require_user(request)
    round_off = _rounded(round_off)
    win = _window(date_from, date_to, period)
    if "error" in win:
        raise HTTPException(404, win["error"])
    chosen = win["period"]
    data = _window_data(win, cluster, bond, warehouse, group_by)
    if "error" in data:
        raise HTTPException(404, data["error"])
    if not data["bonds"]:
        raise HTTPException(404, "Nothing matches that filter.")

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    NAVY, GOLD, PAPER = "FF0A294F", "FFFFBD30", "FFF5F7FC"
    INK, HAIR = "FF28324A", "FFD9DEE9"
    thin = Side(style="thin", color=HAIR)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    # Eight thousand rows is sixty-five thousand cells, and openpyxl charges
    # for every style OBJECT, not for every cell it is put on: building a new
    # Font per cell took twenty-six seconds to write this book. Built once
    # here and shared, it is a couple.
    F_SHOP = Font(bold=True, size=10, color=INK)
    F_BRAND = Font(bold=True, size=9, color=INK)
    F_PACK = Font(size=9, color="FF6B7280")
    F_BAND = Font(bold=True, color=GOLD, size=10)
    F_GRAND = Font(bold=True, color=NAVY, size=11)
    FILL_PAPER = PatternFill("solid", fgColor=PAPER)
    FILL_NAVY = PatternFill("solid", fgColor=NAVY)
    FILL_GOLD = PatternFill("solid", fgColor=GOLD)
    AL_RIGHT = Alignment(horizontal="right")
    AL_MID = Alignment(horizontal="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "SHOP SALES CUMULATIVE"
    grouping = "WAREHOUSE" if group_by == "warehouse" else "BOND"
    headings = (["CLUSTER"] if grouping == "BOND" else []) + [
        grouping, "SHOP CODE", "SHOP", "OPENING", "RECEIPT", "SALES", "CLOSING"]
    last_col = get_column_letter(len(headings))

    # A title anyone can read six months later, without opening the filename.
    ws.merge_cells(f"A1:{last_col}1")
    t = ws["A1"]
    t.value = f"SHOP SALES CUMULATIVE   ·   {chosen['long']}"
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.font = Font(bold=True, color=GOLD, size=13)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    scope = (bond.title() if bond else warehouse.title() if warehouse
             else f"Cluster {cluster}" if cluster else "All bonds")
    ws.merge_cells(f"A2:{last_col}2")
    sub = ws["A2"]
    sub.value = (f"{scope}   ·   {data['shop_count']} shops in {len(data['bonds'])} bonds"
                 f"   ·   cases")
    sub.fill = PatternFill("solid", fgColor=GOLD)
    sub.font = Font(bold=True, color=NAVY, size=10)
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.append(headings)
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=GOLD, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = box
    ws.row_dimensions[3].height = 22

    # ws[ws.max_row] looks harmless and is quadratic: max_row walks the sheet,
    # so dressing eight thousand rows one at a time took twenty-two seconds to
    # write this book. The row number is known - it is counted here - and the
    # cells are addressed directly, which is the same work done once.
    wide = len(headings)
    figures_at = wide - 3          # the four measures always close the row
    figures = "#,##0" if round_off else "#,##0.##"
    at = 3

    # Setting a cell's font, fill, border and format costs a recursive hash of
    # each of those objects - openpyxl looks them up in the workbook's style
    # tables - and sixty-five thousand cells is a hundred and eighty thousand
    # of those hashes, which was the whole two seconds. Each distinct look is
    # therefore built once, on a scratch sheet, and what the cells are handed
    # afterwards is the finished index, not the objects.
    from copy import copy as _copy
    scratch = wb.create_sheet("_styles")

    def look(font, fill=None, align=None, fmt="General"):
        c = scratch.cell(row=1, column=1)
        c.border = box
        c.font = font
        c.fill = fill if fill is not None else PatternFill()
        c.alignment = align if align is not None else Alignment()
        c.number_format = fmt
        return _copy(c._style)

    LOOKS = {}
    for name, font, fill in (("shop", F_SHOP, None), ("shopband", F_SHOP, FILL_PAPER),
                             ("brand", F_BRAND, None), ("pack", F_PACK, None),
                             ("band", F_BAND, FILL_NAVY), ("grand", F_GRAND, FILL_GOLD)):
        LOOKS[name] = (look(font, fill, AL_MID),            # the key columns
                       look(font, fill),                    # the shop name
                       look(font, fill, AL_RIGHT, figures)) # the four measures

    def dress(row: int, name: str, level: int = 0) -> None:
        keyed, plain, figure = LOOKS[name]
        for col in range(1, wide + 1):
            ws.cell(row=row, column=col)._style = _copy(
                figure if col >= figures_at else keyed if col < 4 else plain)
        if level:
            ws.row_dimensions[row].outlineLevel = level


    detail = reports_api.window_lines(chosen["start"], chosen["end"])
    lines = {} if "error" in detail else detail["lines"]
    by_shop: dict = {}
    for key, values in lines.items():
        shop_code, brand, pack = key.split("|", 2)
        by_shop.setdefault(shop_code, {}).setdefault(brand, {})[pack] = values

    for g in data["bonds"]:
        for i, shop in enumerate(g["shops"]):
            code = shop["code"]
            ws.append([g["cluster"] or "", g["bond"],
                       int(code) if code.isdigit() and not code.startswith("0") else code,
                       shop["name"], *shop["values"]])
            at += 1
            # Three collapsible levels, the way the screen folds: bond, then
            # shop, then the brand and pack lines under it.
            dress(at, "shopband" if i % 2 else "shop", 1)

            for brand in sorted(by_shop.get(code, {})):
                packs = by_shop[code][brand]
                sub = [sum(packs[pk][k] for pk in packs) for k in range(4)]
                ws.append(["", "", "", f"    {brand}", *[round(v, 2) for v in sub]])
                at += 1
                dress(at, "brand", 2)
                for pack in sorted(packs):
                    ws.append(["", "", "", f"        {pack}",
                               *[round(v, 2) for v in packs[pack]]])
                    at += 1
                    dress(at, "pack", 3)

        ws.append([g["cluster"] or "", g["bond"], "", f"{g['bond'].title()} total",
                   *g["totals"]])
        at += 1
        dress(at, "band")

    # The grand total rounds once from the unrounded figures, never from the
    # bond totals - adding fifteen rounded numbers drifts.
    ws.append(["", "", "", "GRAND TOTAL", *data["total"]])
    at += 1
    dress(at, "grand")

    widths = {"A": 9, "B": 18, "C": 12, "D": 36, "E": 13, "F": 13, "G": 13, "H": 13}
    for col, wide in widths.items():
        ws.column_dimensions[col].width = wide
    wb.remove(scratch)
    ws.freeze_panes = "E4"
    ws.auto_filter.ref = f"A3:{last_col}{at - 1}"
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_view.showGridLines = False

    tmp = Path(tempfile.mkdtemp()) / f"Shop Sales Cumulative - {scope} ({chosen['short']}).xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


@app.get("/reports/shop-cumulative/export.pdf")
def shop_cumulative_pdf(request: Request, date_from: str = "", date_to: str = "",
                        period: str = "", bond: str = "", cluster: int = 0,
                        warehouse: str = "", group_by: str = "bond",
                        scope: str = "", round_off: str = ""):
    """Two shapes from one endpoint.

    scope="current"  the screen, printed - the bonds you filtered to with their
                     shops under them, as one document.
    anything else    the office's own books: one PDF per bond, one page per
                     shop, zipped when more than one bond is asked for.
    """
    require_user(request)
    round_off = _rounded(round_off)
    win = _window(date_from, date_to, period)
    if "error" in win:
        raise HTTPException(404, win["error"])
    chosen = win["period"]

    import json, subprocess, tempfile, zipfile
    from .pipelines import APP_REPORTS

    outdir = Path(tempfile.mkdtemp())

    if scope == "current":
        label = (bond.title() if bond else warehouse.title() if warehouse
                 else f"Cluster {cluster}" if cluster else "All bonds")
        data = _window_data(win, cluster, bond, warehouse, group_by)
        if "error" in data or not data.get("bonds"):
            raise HTTPException(404, data.get("error", "Nothing matches that filter."))
        out = outdir / f"Shop Sales Cumulative - {label} ({chosen['short']}).pdf"
        reports_pdf.build_cumulative_view_pdf(data, out, scope=label,
                                              round_off=bool(round_off))
        return FileResponse(out, filename=out.name)

    argv = ["python3", str(APP_REPORTS / "build_shop_cumulative_pdfs.py"),
            "--outdir", str(outdir), "--period", chosen["long"]]
    # The office books ignored the switch entirely and always printed the exact
    # figure, so the same report read 103.5 in the PDF and 104 on the screen.
    if round_off:
        argv.append("--round-off")
    if win["source"] == "cumulative":
        argv += ["--raw", str(chosen["path"])]
    else:
        # A window stitched from several files has no file of its own, so the
        # shop lines go to the builder as data rather than as a path.
        feed = outdir / "_window.json"
        feed.write_text(json.dumps(reports_api.shop_range_lines(
            chosen["start"], chosen["end"])))
        argv += ["--lines", str(feed)]

    proc = subprocess.run(argv, capture_output=True, text=True,
                          timeout=config.STEP_TIMEOUT_SECONDS)
    if proc.returncode != 0:
        raise HTTPException(500, f"PDF build failed: {(proc.stderr or proc.stdout)[-400:]}")

    pdfs = sorted(outdir.glob("*.pdf"))
    if not pdfs:
        raise HTTPException(404, "Nothing in that period mapped to a bond.")

    if bond and bond.lower() != "all":
        want = f"Shop Sales Cumulative - {bond.title()}.pdf"
        for f in pdfs:
            if f.name.lower() == want.lower():
                return FileResponse(f, filename=f.name)
        raise HTTPException(404, f"No pages for {bond} in that period.")

    if cluster:
        keep = {b.title().lower() for b, c in reports_api.cluster_of_bond().items()
                if c == cluster}
        pdfs = [f for f in pdfs if f.stem.split(" - ", 1)[-1].lower() in keep]
        if not pdfs:
            raise HTTPException(404, f"No bonds in cluster {cluster} for that period.")

    stamp = f"{chosen['start'].isoformat()} to {chosen['end'].isoformat()}"
    label = f"Cluster {cluster} " if cluster else ""
    bundle = outdir / f"Shop Sales Cumulative {label}({stamp}).zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for f in pdfs:
            z.write(f, f.name)
    return FileResponse(bundle, filename=bundle.name)


def _bond_totals(win: dict) -> dict:
    """bond -> [opening, receipt, sales, closing], whichever way it was sourced."""
    data = _window_data(win, 0, "")
    return {g["bond"]: g["totals"] for g in data.get("bonds", [])}


def _previous_window(period: dict):
    """The same day-window one month back, from a cumulative or the dailies.

    Matched on the day numbers rather than on elapsed days: 'the first half of
    last month' is the comparison the trade actually makes.
    """
    span = _previous_span(period)
    if not span:
        return None
    prev = reports_api.resolve_window(*span)
    return None if "error" in prev else prev


def _previous_span(period: dict):
    """The window a month back as asked for - uploaded or not.

    _previous_window answers None for a month nobody uploaded, which is the
    right answer for the figures and the wrong one for the Source: a
    comparison against a month with no files behind it is exactly the thing
    worth saying out loud, and to say it the panel needs the dates that were
    asked for rather than the ones that were found.
    """
    start, end = period["start"], period["end"]
    month = start.month - 1 or 12
    year = start.year - (1 if start.month == 1 else 0)
    try:
        return date(year, month, start.day), date(year, month, end.day)
    except ValueError:
        return None


def _analysis_basis(chosen: dict, prev: dict | None) -> dict:
    """How many days each bond could actually trade in the window.

    The per-day rate used to divide by the calendar, so three dry days read as
    three days of bad trading. It now divides by the days shops could open -
    the window, less the days the leave calendar says they were shut. Shops
    shut together, so this is one number for the whole book.
    """
    start, end = chosen["start"], chosen["end"]
    shut = leaves.closed_days("shop", start, end)
    basis = {
        "span": leaves.span_days(start, end),
        "open": leaves.open_days(start, end),
        "closed": sorted(d.isoformat() for d in shut),
    }
    if prev:
        p_start, p_end = prev["period"]["start"], prev["period"]["end"]
        basis["prev_open"] = leaves.open_days(p_start, p_end)
    else:
        # No window to compare against - last month's cumulative was never
        # uploaded. This used to borrow THIS window's day count, which made
        # the per-day rate divide a last-month total of nothing by a real
        # number of days: every bond read "last month 0" and posted the whole
        # of this month as a gain. Nought days says there is no comparison,
        # and every row prints a dash instead of a figure nobody can stand
        # behind.
        basis["prev_open"] = 0
    return basis


def _analysis_rows(win: dict, cluster: int, bond: str, round_off: bool = False):
    """The rows the screen shows, the PDF prints and the workbook exports."""
    mod = _analysis_module()
    places = 0 if round_off else 2
    chosen = win["period"]
    prev = _previous_window(chosen)
    now = _bond_totals(win)
    before = _bond_totals(prev) if prev else {}
    basis = _analysis_basis(chosen, prev)
    raw = mod.build_rows(now, before, basis["open"], basis["prev_open"],
                         cluster, bond, places=places)
    rows = [{
        "kind": r["kind"], "label": r["label"], "cells": r["cells"],
        "net_pct": mod.pct(r["net_pct"]), "sell": mod.pct(r["sell"]),
        "amber": r["sell"] is not None and r["sell"] >= mod.SELL_AMBER_AT,
        "cm": r["cm"], "lm": r["lm"],
        "up": (r["trend"] or 0) >= 0,
        "trend": (None if r["trend"] is None
                  else mod.whole(abs(r["trend"]), places)),
    } for r in raw]
    return rows, (prev["period"] if prev else None), sorted(now), basis


@app.get("/reports/shop-analysis", response_class=HTMLResponse)
def shop_analysis_page(request: Request, date_from: str = "", date_to: str = "",
                       period: str = "", cluster: int = 0, bond: str = "",
                       round_off: str = ""):
    user = require_user(request)
    round_off = _rounded(round_off)
    periods = reports_api.cumulative_periods()
    win = _window(date_from, date_to, period)
    chosen = None if "error" in win else win["period"]

    rows, prev, bonds, basis = [], None, [], None
    if chosen:
        rows, prev, bonds, basis = _analysis_rows(win, cluster, bond, bool(round_off))
    return templates.TemplateResponse(
        request, "report_shop_analysis.html",
        {"user": user, "page": "shop_analysis", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", []),
         "periods": periods, "chosen": chosen, "prev": prev, "rows": rows,
         "basis": basis,
         "bonds": bonds, "cluster": cluster, "bond": bond.upper(),
         "round_off": bool(round_off),
         "calendar": _calendar_days(),
         "source": win.get("source", ""), "chain": win.get("chain", []),
         "source_block": (_shop_source(win, _previous_window(chosen),
                                       _previous_span(chosen))
                          if chosen else {"legs": []}),
         "no_data": win.get("error", ""),
         "need": win.get("need", []),
         "asked": _asked_period(date_from, date_to, period),
         "suggest": _suggested_window()},
    )


@app.get("/reports/shop-analysis/export.xlsx")
def shop_analysis_xlsx(request: Request, date_from: str = "", date_to: str = "",
                       period: str = "", cluster: int = 0, bond: str = "",
                       round_off: str = ""):
    require_user(request)
    round_off = _rounded(round_off)
    win = _window(date_from, date_to, period)
    if "error" in win:
        raise HTTPException(404, win["error"])
    chosen = win["period"]
    rows, _prev, _bonds, _basis = _analysis_rows(win, cluster, bond, bool(round_off))
    if not rows:
        raise HTTPException(404, "Nothing matches that filter.")

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    navy, gold = "FF0A294F", "FFFFBD30"
    wb = Workbook()
    ws = wb.active
    ws.title = "SHOPSALES COMPARATIVE"
    ws.append(["BOND", "OPENING", "RECEIPT", "SALES", "CLOSING", "STOCK NET",
               "STOCK NET %", "SELL-THROUGH %", "AVG/DAY CM", "AVG/DAY LM", "TREND"])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(bold=True, color=gold, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for r in rows:
        ws.append([r["label"], *r["cells"], r["net_pct"], r["sell"],
                   r["cm"] if r["cm"] is not None else "-",
                   r["lm"] if r["lm"] is not None else "-",
                   "-" if r["trend"] is None
                   else ("+" if r["up"] else "-") + str(r["trend"])])
        if r["kind"] != "bond":
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(bold=True, color=gold, size=10)
    figures = "#,##0" if round_off else "#,##0.##"
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
            if isinstance(cell.value, (int, float)):
                cell.number_format = figures

    ws.column_dimensions["A"].width = 22
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 13
    ws.freeze_panes = "B2"

    tmp = Path(tempfile.mkdtemp()) / f"Shop Sales Analysis ({chosen['short']}).xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


@app.get("/reports/shop-analysis/export.pdf")
def shop_analysis_pdf(request: Request, date_from: str = "", date_to: str = "",
                      period: str = "", cluster: int = 0, bond: str = "",
                      round_off: str = ""):
    require_user(request)
    round_off = _rounded(round_off)
    win = _window(date_from, date_to, period)
    if "error" in win:
        raise HTTPException(404, win["error"])
    chosen = win["period"]
    prev = _previous_window(chosen)

    import json, subprocess, tempfile
    from .pipelines import APP_REPORTS

    outdir = Path(tempfile.mkdtemp())
    # Totals rather than a path: the window may have been rebuilt from a dozen
    # daily files, and either way the builder should draw the same figures the
    # screen just showed instead of re-deriving them from a different read.
    totals = outdir / "_totals.json"
    totals.write_text(json.dumps(_bond_totals(win)))
    basis = _analysis_basis(chosen, prev)
    argv = ["python3", str(APP_REPORTS / "build_shop_analysis_pdf.py"),
            "--totals", str(totals), "--outdir", str(outdir),
            "--period", chosen["short"], "--days", str(basis["open"])]
    if prev:
        before = outdir / "_prev.json"
        before.write_text(json.dumps(_bond_totals(prev)))
        argv += ["--prev-totals", str(before), "--prev-days", str(basis["prev_open"])]
    if cluster:
        argv += ["--cluster", str(cluster)]
    if bond:
        argv += ["--bond", bond]
    if not round_off:
        argv += ["--places", "2"]

    proc = subprocess.run(argv, capture_output=True, text=True,
                          timeout=config.STEP_TIMEOUT_SECONDS)
    if proc.returncode != 0:
        raise HTTPException(500, f"PDF build failed: {(proc.stderr or proc.stdout)[-400:]}")

    out = outdir / "Shop Sales Analysis.pdf"
    if not out.is_file():
        raise HTTPException(404, "Nothing in that period mapped to a bond.")
    name = f"Shop Sales Analysis ({chosen['short']}).pdf"
    return FileResponse(out, filename=name)


# ---------------------------------------------------------------------------
# Liquidation Summary
# ---------------------------------------------------------------------------


def _liquidation(date_from: str, date_to: str, period: str = "",
                 prev_from: str = "", prev_to: str = ""):
    """The scorecard for a window, plus the labels every export needs.

    The comparison window is normally the same days a month back. When one is
    named outright the two sides can be any two windows - 1-16 August against
    1-20 October - and the labels stop pretending otherwise: a column that
    would read AUG against AUG says AUG 1-16 and AUG 20-31 instead.
    """
    win = _window(date_from, date_to, period)
    if "error" in win:
        return {"error": win["error"]}
    cur = win["period"]

    chosen = None
    if prev_from and prev_to:
        try:
            a = date.fromisoformat(prev_from)
            b = date.fromisoformat(prev_to)
        except ValueError:
            return {"error": "Those comparison dates are not readable."}
        chosen = (min(a, b), max(a, b))

    data = reports_api.liquidation(cur["start"], cur["end"], chosen)
    if "error" in data:
        return data
    prev = data["previous"]
    data["custom_prev"] = bool(chosen)

    def days_of(win_):
        return (win_["start"].day, win_["end"].day)

    # Two windows over the same month, or over different lengths, cannot both
    # be called by their month alone.
    plain = bool(prev) and days_of(cur) == days_of(prev) and \
        cur["start"].month != prev["start"].month

    def head(win_):
        mon = win_["start"].strftime("%b").upper()
        if plain:
            return mon
        if win_["start"].month == win_["end"].month:
            return f'{mon} {win_["start"].day}-{win_["end"].day}'
        return f'{win_["start"].day} {mon}-{win_["end"].day} {win_["end"].strftime("%b").upper()}'

    data["headings"] = [head(cur), head(prev) if prev else "PREV", "CS", "%"]
    if plain:
        data["subtitle"] = (
            f'{cur["start"].strftime("%B").upper()} vs '
            f'{prev["start"].strftime("%B").upper()} {cur["start"].year}'
            f'  \u00b7  DAYS {cur["start"].day}-{cur["end"].day}  \u00b7  cases')
    else:
        def span_label(win_):
            a, b = win_["start"], win_["end"]
            mon, mon_b = a.strftime("%b").upper(), b.strftime("%b").upper()
            if a == b:
                return f"{a.day} {mon} {a.year}"
            if (a.month, a.year) == (b.month, b.year):
                return f"{a.day}-{b.day} {mon} {a.year}"
            return f"{a.day} {mon}-{b.day} {mon_b} {b.year}"

        against = span_label(prev) if prev else "LAST MONTH"
        data["subtitle"] = f'{span_label(cur)} vs {against}  \u00b7  cases'
    return data


@app.get("/reports/liquidation", response_class=HTMLResponse)
def liquidation_page(request: Request, date_from: str = "", date_to: str = "",
                     period: str = "", round_off: str = "",
                     prev_from: str = "", prev_to: str = ""):
    user = require_user(request)
    round_off = _rounded(round_off)
    data = _liquidation(date_from, date_to, period, prev_from, prev_to)
    return templates.TemplateResponse(
        request, "report_liquidation.html",
        {"user": user, "page": "liquidation", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", []),
         "calendar": _calendar_days(),
         "groups": ["Shop liquidation (KSBC)", "Secondary sales",
                    "Fed / Bar invoice", "Total liquidation"],
         "data": data, "error": data.get("error", ""),
         "round_off": bool(round_off),
         "prev_from": prev_from, "prev_to": prev_to},
    )


@app.get("/reports/liquidation/export.pdf")
def liquidation_pdf(request: Request, date_from: str = "", date_to: str = "",
                    period: str = "", round_off: str = "",
                    prev_from: str = "", prev_to: str = ""):
    require_user(request)
    round_off = _rounded(round_off)
    data = _liquidation(date_from, date_to, period, prev_from, prev_to)
    if "error" in data:
        raise HTTPException(404, data["error"])

    import json, subprocess, tempfile
    from .pipelines import APP_REPORTS

    outdir = Path(tempfile.mkdtemp())
    feed = outdir / "scorecard.json"
    feed.write_text(json.dumps({"rows": data["rows"], "headings": data["headings"]}, default=str))
    out = outdir / f"Liquidation Summary ({data['period']['short']}).pdf"
    proc = subprocess.run(
        ["python3", str(APP_REPORTS / "build_liquidation_pdf.py"),
         "--data", str(feed), "--out", str(out), "--title", data["subtitle"]]
        + (["--round"] if round_off else []),
        capture_output=True, text=True, timeout=config.STEP_TIMEOUT_SECONDS)
    if proc.returncode != 0:
        raise HTTPException(500, f"PDF build failed: {(proc.stderr or proc.stdout)[-400:]}")
    return FileResponse(out, filename=out.name)


@app.get("/reports/liquidation/export.xlsx")
def liquidation_xlsx(request: Request, date_from: str = "", date_to: str = "",
                     period: str = "", round_off: str = "",
                     prev_from: str = "", prev_to: str = ""):
    """The scorecard as a workbook, cell for cell as the PDF prints it.

    Measured off the office's own sheet rather than styled by eye: its navy and
    gold, its Segoe UI, its hairlines inside a block and gold rules around one,
    its zebra, and the four colours a delta can take - two on a white row, two
    on a dark one, because red on navy cannot be read.

    A delta keeps its number. The arrow lives in the number format and the
    colour in the font, so the cell still sums, sorts and filters.
    """
    require_user(request)
    round_off = _rounded(round_off)
    data = _liquidation(date_from, date_to, period, prev_from, prev_to)
    if "error" in data:
        raise HTTPException(404, data["error"])

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    # ---- the office's palette, read out of its own workbook ----
    NAVY, GOLD = "FF0B2C52", "FFFAAF19"
    SLATE, SLATE_2 = "FF2C3540", "FF3E4957"
    ZEBRA, WHITE = "FFF5F7FC", "FFFFFFFF"
    NAME_INK, FIG_INK = "FF282828", "FF0F192D"
    UP_LIGHT, DOWN_LIGHT = "FF3F8600", "FFCF1322"
    UP_DARK, DOWN_DARK = "FF52C41A", "FFFF7875"
    FACE = "Segoe UI"

    HAIR = Side(style="thin", color="FFC7C7C7")
    RULE = Side(style="medium", color=GOLD)

    now_label, was_label = data["headings"][0], data["headings"][1]
    groups = ["SHOP LIQUIDATION (KSBC)", "SECONDARY SALES",
              "FED / BAR INVOICE", "TOTAL LIQUIDATION"]

    # A gutter column between blocks: navy, gold-ruled both sides, barely wide
    # enough to be a channel rather than a column.
    first = [2 + i * 5 for i in range(len(groups))]      # B, G, L, Q
    gutters = [c - 1 for c in first[1:]]
    span = first[-1] + 3
    last_col = get_column_letter(span)

    wb = Workbook()
    ws = wb.active
    ws.title = "Liquidation Scorecard"
    ws.sheet_view.showGridLines = False

    def edges(col):
        """Gold down a block's outer edges, hairline between its columns."""
        left = RULE if (col == 1 or col in first or col in gutters) else HAIR
        right = RULE if (col == 1 or col - 3 in first or col in gutters) else HAIR
        return left, right

    def paint(row, col, fill=None, colour=FIG_INK, bold=False, size=9.5,
              align="center", fmt=None, value=None, wrap=False, strong=False,
              plain=False, channel=False):
        c = ws.cell(row=row, column=col)
        if value is not None:
            c.value = value
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(name=FACE, bold=bold, size=size, color=colour)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if not plain:
            left, right = edges(col)
            # The channel is one strip from the header to the last row. A top
            # or bottom edge on it would saw it into twenty-odd pieces.
            band = None if channel else (RULE if strong else HAIR)
            c.border = Border(left=left, right=right, top=band, bottom=band)
        if fmt:
            c.number_format = fmt
        return c

    # ---- the two bands ----
    ws.merge_cells(f"A1:{last_col}1")
    for col in range(1, span + 1):
        paint(1, col, NAVY, plain=True)
    paint(1, 1, NAVY, GOLD, True, 18, value="K.S DISTILLERY", plain=True)
    ws.row_dimensions[1].height = 36

    half = span // 2
    for col in range(1, span + 1):
        paint(2, col, GOLD, plain=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=half)
    ws.merge_cells(start_row=2, start_column=half + 1, end_row=2, end_column=span)
    paint(2, 1, GOLD, NAVY, True, 12, "left", value="LIQUIDATION SUMMARY",
          plain=True)
    paint(2, half + 1, GOLD, NAVY, True, 12, "right", value=data["subtitle"],
          plain=True)
    ws.row_dimensions[2].height = 24

    # ---- the grouped header, two tiers as the sheet prints it ----
    ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
    paint(3, 1, NAVY, GOLD, True, 9.5, value="Bond")
    paint(4, 1, NAVY)
    for g, name in zip(first, groups):
        ws.merge_cells(start_row=3, start_column=g, end_row=3, end_column=g + 3)
        for col in range(g, g + 4):
            paint(3, col, NAVY)
        paint(3, g, NAVY, GOLD, True, 9.5, value=name, wrap=True)
        for col, label in zip(range(g, g + 4),
                              (now_label, was_label, "\u25b2 CS", "\u25b2 %")):
            paint(4, col, NAVY, GOLD, True, 9.5, value=label, wrap=True)
    for col in gutters:
        paint(3, col, NAVY, channel=True)
        paint(4, col, NAVY, channel=True)
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 18

    # ---- the rows ----
    # The arrow is the number format's doing and the sign is the font's, so a
    # delta reads as the sheet prints it and still behaves as a number.
    places = "0" if round_off else "0.00"
    FIG = f'{places};-{places};"-"'
    CS_FMT = f'"\u25b2 "{places};"\u25bc -"{places};"\u25b2 "{places}'
    pc = "0%" if round_off else "0.0%"
    PC_FMT = f'"\u25b2 "{pc};"\u25bc -"{pc};"\u25b2 "{pc}'

    r = 4
    stripe = 0
    for row in data["rows"]:
        r += 1
        kind = row.get("kind", "bond")
        strong = kind in ("total", "average")
        if kind == "cluster":
            fill, name_ink, fig_ink = NAVY, GOLD, GOLD
            up, down, bold, size = UP_DARK, DOWN_DARK, True, 10
        elif kind == "total":
            fill, name_ink, fig_ink = SLATE, WHITE, WHITE
            up, down, bold, size = UP_DARK, DOWN_DARK, True, 10
        elif kind == "average":
            fill, name_ink, fig_ink = SLATE_2, WHITE, WHITE
            up, down, bold, size = UP_DARK, DOWN_DARK, True, 10
        else:
            fill = ZEBRA if stripe % 2 == 0 else WHITE
            name_ink, fig_ink = NAME_INK, FIG_INK
            up, down, bold, size = UP_LIGHT, DOWN_LIGHT, False, 9.5
            stripe += 1

        paint(r, 1, fill, name_ink, bold, size, "left",
              value=str(row["label"]), strong=strong)
        for g, (now, was) in enumerate(row["blocks"]):
            col = first[g]
            paint(r, col, fill, fig_ink, bold, size, fmt=FIG,
                  value=float(now), strong=strong)
            paint(r, col + 1, fill, fig_ink, bold, size, fmt=FIG,
                  value=float(was), strong=strong)

            # Two empty months have no story, so they read as a dash rather
            # than as a confident nought.
            if not now and not was:
                paint(r, col + 2, fill, fig_ink, bold, size, value="-", strong=strong)
            else:
                d = float(now) - float(was)
                paint(r, col + 2, fill, up if d >= 0 else down, True, size,
                      fmt=CS_FMT, value=round(d, 2), strong=strong)

            # A per cent needs something to divide by.
            if not was:
                paint(r, col + 3, fill, fig_ink, bold, size, value="-", strong=strong)
            else:
                pct = (float(now) - float(was)) / float(was)
                paint(r, col + 3, fill, up if pct >= 0 else down, True, size,
                      fmt=PC_FMT, value=round(pct, 4), strong=strong)
        for col in gutters:
            paint(r, col, NAVY, channel=True)
        ws.row_dimensions[r].height = 20

    # One merged cell per channel - across the header, then down the body - so
    # each reads as a single strip of navy rather than a stack of cells.
    for col in gutters:
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        ws.merge_cells(start_row=5, start_column=col, end_row=r, end_column=col)
        paint(3, col, NAVY, channel=True)
        paint(5, col, NAVY, channel=True)

    # ---- the shape of the page ----
    ws.column_dimensions["A"].width = 24
    for g in first:
        for col, width in zip(range(g, g + 4), (11.5, 11.5, 13, 10.5)):
            ws.column_dimensions[get_column_letter(col)].width = width
    for col in gutters:
        ws.column_dimensions[get_column_letter(col)].width = 1.0
    ws.freeze_panes = ws.cell(row=5, column=2)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4,
                                  header=0.2, footer=0.2)
    ws.print_area = f"A1:{last_col}{r}"
    ws.print_title_rows = "1:4"

    tmp = Path(tempfile.mkdtemp()) / f"Liquidation Summary ({data['period']['short']}).xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


@app.get("/reports/warehouse-stock/export.pdf")
def stock_pdf(request: Request, scope: str = "cluster", cluster: str = "1",
              as_of: str = "", warehouse: str = ""):
    """scope='cluster' -> one cluster; scope='current' -> whatever is filtered."""
    require_user(request)
    import tempfile

    outdir = Path(tempfile.mkdtemp())

    # scope="clusters" builds all three in one go and hands back a folder,
    # rather than firing three downloads a second apart.
    if scope == "clusters":
        built = []
        for c in (1, 2, 3):
            got = reports_api.warehouse_stock(as_of=as_of, cluster=c)
            if "error" in got or not got["warehouses"]:
                continue
            out = outdir / f"Warehouse Stock Report - Cluster {c}.pdf"
            reports_pdf.build_stock_pdf(got, out)
            built.append(out)
        if not built:
            raise HTTPException(404, "No stock rows for any cluster on that date.")
        return _as_folder(built, "Warehouse Stock Report - All Clusters")

    if scope == "current":
        cl = int(cluster) if cluster in ("1", "2", "3") else None
        data = reports_api.warehouse_stock(as_of=as_of, cluster=cl, warehouse=warehouse)
        label = warehouse or (f"Cluster {cl}" if cl else "All Warehouses")
    else:
        if cluster not in ("1", "2", "3"):
            raise HTTPException(400, "Cluster must be 1, 2 or 3.")
        data = reports_api.warehouse_stock(as_of=as_of, cluster=int(cluster))
        label = f"Cluster {cluster}"

    if "error" in data:
        raise HTTPException(404, data["error"])
    if not data["warehouses"]:
        raise HTTPException(404, f"No stock rows for {label} on that date.")

    out = outdir / f"Warehouse Stock Report - {label}.pdf"
    reports_pdf.build_stock_pdf(data, out)
    return FileResponse(out, filename=out.name)


@app.get("/reports/warehouse-stock/export.xlsx")
def stock_xlsx(request: Request, as_of: str = "", cluster: str = "", warehouse: str = ""):
    """The stock position as a workbook, in the same dress as its PDF.

    The earlier version was a flat dump with a navy strip on top: no title
    anyone could read six months later, no borders, nothing to tell a stock
    line from a warehouse total except its colour, and no grand total at all.
    This keeps the grouping people actually work with - every warehouse folds
    away - while the autofilter still gives the flat view to anyone who wants
    to pivot.
    """
    require_user(request)
    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    cl = int(cluster) if cluster in ("1", "2", "3") else None
    data = reports_api.warehouse_stock(as_of=as_of, cluster=cl, warehouse=warehouse)
    if "error" in data:
        raise HTTPException(404, data["error"])
    if not data["warehouses"]:
        raise HTTPException(404, "No stock rows for that filter.")

    NAVY, GOLD, PAPER = "FF0A294F", "FFFFBD30", "FFF5F7FC"
    INK, HAIR, ZERO = "FF28324A", "FFD9DEE9", "FF9AA3B4"
    thin = Side(style="thin", color=HAIR)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    headings = ["WAREHOUSE", "ITEM NAME", "PACK", "PHYSICAL", "ALLOTABLE", "PENDING"]
    last_col = get_column_letter(len(headings))

    wb = Workbook()
    ws = wb.active
    ws.title = "WAREHOUSE STOCK"

    # A stock position is a photograph of one day, so the day belongs in the
    # title - a sheet filed without it is worth nothing a week later.
    day = data["date"]
    try:
        day = date.fromisoformat(data["date"]).strftime("%d %b %Y")
    except (TypeError, ValueError):
        pass
    ws.merge_cells(f"A1:{last_col}1")
    t = ws["A1"]
    t.value = f"WAREHOUSE STOCK REPORT   ·   {day}"
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.font = Font(bold=True, color=GOLD, size=13)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # "All warehouses · 28 warehouses" said it twice. The scope only earns a
    # word of its own when it is actually narrowing something.
    houses = len(data["warehouses"])
    items = sum(len(w["rows"]) for w in data["warehouses"])
    scope = warehouse.title() if warehouse else f"Cluster {cl}" if cl else ""
    count = (f"{houses} warehouse{'' if houses == 1 else 's'}   ·   " if not warehouse else "")
    ws.merge_cells(f"A2:{last_col}2")
    sub = ws["A2"]
    sub.value = ((f"{scope}   ·   " if scope else "")
                 + f"{count}{items} items   ·   cases")
    sub.fill = PatternFill("solid", fgColor=GOLD)
    sub.font = Font(bold=True, color=NAVY, size=10)
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.append(headings)
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=GOLD, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = box
    ws.row_dimensions[3].height = 22

    for wh in data["warehouses"]:
        for i, r in enumerate(wh["rows"]):
            ws.append([wh["name"], r["brand"], r["pack"],
                       r["physical"], r["allotable"], r["pending"]])
            row = ws.max_row
            # One outline level, so a warehouse folds to its total the way the
            # screen folds it.
            ws.row_dimensions[row].outlineLevel = 1
            band = PatternFill("solid", fgColor=PAPER) if i % 2 else None
            for cell in ws[row]:
                cell.border = box
                cell.font = Font(size=10, color=INK)
                if band:
                    cell.fill = band
            # A nil reads as nil, not as a figure: dimmed here, the way the
            # PDF greys it, so a column of zeros does not compete with stock.
            for c in range(4, 7):
                if not (ws.cell(row=row, column=c).value or 0):
                    ws.cell(row=row, column=c).font = Font(size=10, color=ZERO)

        ws.append([wh["name"], "TOTAL", "", wh["totals"]["physical"],
                   wh["totals"]["allotable"], wh["totals"]["pending"]])
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color=GOLD, size=10)
            cell.border = box

    grand = data["grand"]
    ws.append(["", "GRAND TOTAL", "", grand["physical"],
               grand["allotable"], grand["pending"]])
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.font = Font(bold=True, color=NAVY, size=11)
        cell.border = box

    for row in ws.iter_rows(min_row=4, min_col=4):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0"
    for row in ws.iter_rows(min_row=4, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = Alignment(horizontal="left" if cell.column == 2 else "center")

    for col, wide in {"A": 20, "B": 38, "C": 11, "D": 13, "E": 13, "F": 13}.items():
        ws.column_dimensions[col].width = wide
    ws.freeze_panes = "D4"
    ws.auto_filter.ref = f"A3:{last_col}{ws.max_row - 1}"
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_view.showGridLines = False

    # Printed, it should come out as one readable column of pages rather than
    # a sheet cut down the middle.
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)

    label = warehouse or (f"Cluster {cl}" if cl else "All Warehouses")
    tmp = Path(tempfile.mkdtemp()) / f"Warehouse Stock Report - {label}.xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


# ---------------------------------------------------------------------------
# Daily grids: Shop Sales - Daily and Secondary Sales - Daily
#
# One screen and one PDF shape serve both — the reports differ only in where
# the cases come from, so the route takes the kind and everything else is
# shared. Keeping them apart would mean maintaining the same grid twice.
# ---------------------------------------------------------------------------

_DAILY_PAGES = {"shop": "shop_daily", "secondary": "secondary_daily"}


def _daily_kind(path_kind: str) -> str:
    if path_kind not in reports_api.DAILY_KINDS:
        raise HTTPException(404, "No such report.")
    return path_kind


@app.get("/reports/daily/{kind}", response_class=HTMLResponse)
def daily_page(request: Request, kind: str):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    _daily_kind(kind)
    return templates.TemplateResponse(
        request, "report_daily.html",
        {"user": user, "page": _DAILY_PAGES[kind], "kind": kind,
         "title": reports_api.DAILY_KINDS[kind]["title"],
         "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


@app.get("/api/reports/daily/{kind}")
def daily_data(request: Request, kind: str, date_from: str = "", date_to: str = "",
               cluster: str = "", view: str = "bond", round_off: str = ""):
    require_user(request)
    round_off = _rounded(round_off)
    _daily_kind(kind)
    if view not in ("bond", "warehouse"):
        raise HTTPException(400, "Unknown view.")
    cl = int(cluster) if cluster in ("1", "2", "3") else None
    return JSONResponse(reports_api.daily_grid(kind, date_from=date_from,
                                               date_to=date_to, cluster=cl, view=view,
                                               round_off=bool(round_off)))


@app.get("/reports/daily/{kind}/export.pdf")
def daily_pdf(request: Request, kind: str, date_from: str = "", date_to: str = "",
              cluster: str = "", view: str = "bond", round_off: str = ""):
    require_user(request)
    round_off = _rounded(round_off)
    _daily_kind(kind)
    import tempfile

    cl = int(cluster) if cluster in ("1", "2", "3") else None
    data = reports_api.daily_grid(kind, date_from=date_from, date_to=date_to,
                                  cluster=cl, view=view, round_off=bool(round_off))
    if "error" in data:
        raise HTTPException(404, data["error"])

    suffix = f" (Cluster {cl})" if cl else ""
    out = Path(tempfile.mkdtemp()) / f"{data['title']}{suffix}.pdf"
    reports_pdf.build_daily_pdf(data, out, round_off=bool(round_off))
    return FileResponse(out, filename=out.name)


@app.get("/reports/daily/{kind}/export.xlsx")
def daily_xlsx(request: Request, kind: str, date_from: str = "", date_to: str = "",
               cluster: str = "", view: str = "bond", round_off: str = ""):
    """The daily grid as a workbook, in the same dress as its own PDF.

    It used to open on the bare column headings - no mark, no report name, no
    period - so a file saved to a folder said nothing about itself a month
    later, and a day nobody had uploaded read as a hard zero exactly like a
    day that traded nothing. This carries the PDF's title bands, splits the
    weekday from the date the way the screen does, and greys an uncovered day
    so it reads as absent rather than as nil.
    """
    require_user(request)
    round_off = _rounded(round_off)
    _daily_kind(kind)
    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    cl = int(cluster) if cluster in ("1", "2", "3") else None
    data = reports_api.daily_grid(kind, date_from=date_from, date_to=date_to,
                                  cluster=cl, view=view, round_off=bool(round_off))
    if "error" in data:
        raise HTTPException(404, data["error"])

    NAVY, GOLD, WHITE = "FF0A294F", "FFFFBD30", "FFFFFFFF"
    PAPER, INK, HAIR = "FFF5F7FC", "FF28324A", "FFD9DEE9"
    ZERO, MUTED = "FF9AA3B4", "FF6B7692"
    NC_BG, NC_INK = "FFF1F3F8", "FFC2C9D6"          # a day with no raw behind it
    NC_HEAD = "FF8FA0C4"

    thin = Side(style="thin", color=HAIR)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    days = data["days"]
    covered = set(data.get("covered") or [])
    missing = [d for d in days if d["iso"] not in covered]
    n_cols = len(days) + 2                           # label + days + total
    last = get_column_letter(n_cols)

    wb = Workbook()
    ws = wb.active
    ws.title = data["title"][:31]

    # --- the two bands the PDF wears -------------------------------------
    # How much room the bands actually have. A one-day sheet is three columns
    # wide, and a band set to one line there just clips its own title.
    span_chars = 22 + (len(days) + 1) * 7.5

    def band_height(text: str, per_line: float, base: float) -> float:
        lines = max(1, -(-len(text) // max(10, int(span_chars / per_line))))
        return base if lines == 1 else base * 0.8 + lines * 13

    ws.merge_cells(f"A1:{last}1")
    t = ws["A1"]
    t.value = "K.S DISTILLERY"
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.font = Font(bold=True, color=GOLD, size=14)
    t.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = band_height("K.S DISTILLERY", 1.15, 30)

    # A window of one day is one date, not the same date printed twice.
    per = data["period"]
    label = per["label"]
    if per.get("from") and per["from"] == per.get("to"):
        label = label.split(" - ")[0]

    ws.merge_cells(f"A2:{last}2")
    sub = ws["A2"]
    sub.value = f"{data['title'].upper()}   \u00b7   {label}"
    sub.fill = PatternFill("solid", fgColor=GOLD)
    sub.font = Font(bold=True, color=NAVY, size=11)
    sub.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = band_height(sub.value, 0.92, 22)

    # --- what you are looking at, in one line ----------------------------
    bits = [f"Grouped by {data.get('group_label', 'Bond').lower()}",
            f"Cluster {cl}" if cl else "All clusters",
            "cases, rounded to whole numbers" if round_off else "cases, exact"]
    if missing:
        # Sixteen day numbers in a row is a wall; "1-16, 18" is the same fact.
        nums, runs = [int(d["dom"]) for d in missing], []
        for n in nums:
            if runs and n == runs[-1][1] + 1:
                runs[-1][1] = n
            else:
                runs.append([n, n])
        shown = ", ".join(str(a) if a == b else f"{a}\u2013{b}" for a, b in runs)
        bits.append(f"no raw uploaded for {shown} \u2014 shown as zero")
    if data.get("unmapped"):
        bits.append(f"{len(data['unmapped'])} outlet(s) with no bond in master data")
    ws.merge_cells(f"A3:{last}3")
    note = ws["A3"]
    note.value = "   \u00b7   ".join(bits)
    note.font = Font(size=9, color=MUTED, italic=True)
    note.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = band_height(note.value, 0.62, 16)

    # --- header, weekday over date, the way the screen prints it ---------
    ws.merge_cells(f"A4:A5")
    ws.cell(4, 1, data.get("group_label", "Bond").upper())
    ws.merge_cells(f"{last}4:{last}5")
    ws.cell(4, n_cols, "TOTAL")
    for i, d in enumerate(days):
        ws.cell(4, i + 2, d["dow"])
        ws.cell(5, i + 2, int(d["dom"]) if d["dom"].isdigit() else d["dom"])
    for r in (4, 5):
        for c in range(1, n_cols + 1):
            cell = ws.cell(r, c)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = box
            iso = days[c - 2]["iso"] if 2 <= c <= len(days) + 1 else None
            dim = iso is not None and iso not in covered
            cell.font = Font(bold=True, size=9 if r == 4 else 10,
                             color=NC_HEAD if dim else GOLD)
    ws.cell(4, 1).font = Font(bold=True, size=10, color=GOLD)
    ws.cell(4, n_cols).font = Font(bold=True, size=10, color=GOLD)
    ws.row_dimensions[4].height = 17
    ws.row_dimensions[5].height = 17

    # --- the grid --------------------------------------------------------
    figures = "#,##0" if round_off else "#,##0.##"
    stripe = 0
    at = 5
    for r in data["rows"]:
        at += 1
        ws.append([r["label"]] + list(r["cells"]) + [r["total"]])
        kind_ = r["kind"]
        if kind_ == "bond":
            stripe += 1
            band = PAPER if stripe % 2 == 0 else None
        else:
            stripe = 0
            band = None
        for c in range(1, n_cols + 1):
            cell = ws.cell(at, c)
            cell.border = box
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
            if c > 1:
                cell.number_format = figures
            iso = days[c - 2]["iso"] if 2 <= c <= len(days) + 1 else None
            dim = iso is not None and iso not in covered
            nil = c > 1 and not (cell.value or 0)
            if kind_ == "cluster":
                cell.fill = PatternFill("solid", fgColor="FF16305F" if dim else NAVY)
                cell.font = Font(bold=True, size=10,
                                 color="FF7C89A6" if (dim or nil) else GOLD)
            elif kind_ == "grand":
                cell.fill = PatternFill("solid", fgColor="FFF0D08A" if dim else GOLD)
                cell.font = Font(bold=True, size=10,
                                 color="FF9A8340" if (dim or nil) else NAVY)
            else:
                if dim:
                    cell.fill = PatternFill("solid", fgColor=NC_BG)
                elif band:
                    cell.fill = PatternFill("solid", fgColor=band)
                cell.font = Font(size=10, bold=c == n_cols,
                                 color=NC_INK if dim else ZERO if nil else INK)

    # --- shape -----------------------------------------------------------
    ws.freeze_panes = "B6"
    # Wide enough for the longest name on the sheet - a warehouse view carries
    # longer labels than a bond one - but not so wide it pushes the days off.
    widest = max((len(str(r["label"])) for r in data["rows"]), default=18)
    ws.column_dimensions["A"].width = min(34, max(22, widest + 3))
    for i in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 7.5
    ws.column_dimensions[last].width = 9.5

    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
    ws.sheet_view.showGridLines = False

    suffix = f" (Cluster {cl})" if cl else ""
    out = Path(tempfile.mkdtemp()) / f"{data['title']}{suffix}.xlsx"
    centre_all(wb)
    wb.save(out)
    return FileResponse(out, filename=out.name)


# ---------------------------------------------------------------------------
# Settings: bond mapping
# ---------------------------------------------------------------------------


@app.get("/settings/bond-mapping", response_class=HTMLResponse)
def bond_mapping_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "settings_bonds.html",
        {"user": user, "page": "bonds", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


@app.get("/api/bond-mapping")
def bond_mapping_data(request: Request):
    require_user(request)
    return JSONResponse(bondmap.load_shops())


@app.post("/api/bond-mapping")
async def bond_mapping_save(request: Request):
    """Save shop -> bond edits, the cluster split, or both."""
    require_user(request)
    body = await request.json()

    result: dict = {}
    clusters = body.get("clusters")
    if isinstance(clusters, dict):
        bondmap.save_clusters(clusters)
        result["clusters_saved"] = True

    changes = body.get("bonds") or {}
    if changes:
        written = bondmap.save_bonds(changes)
        if "error" in written:
            raise HTTPException(400, written["error"])
        result.update(written)
        # The reports cache master data on its mtime, so the next read picks
        # the new mapping up on its own.

    return JSONResponse({"ok": True, **result})


# ---------------------------------------------------------------------------
# Target vs achievement
# ---------------------------------------------------------------------------
#
# The targets are typed here rather than uploaded, because they are not KSBC
# data: nobody exports them, they are a decision taken in a meeting and revised
# mid-month. The achievement beside them is total liquidation - shop sales plus
# the FED and BAR invoices - which reports_api assembles from the raws already
# in the workspace.


def _tva_window(date_from: str, date_to: str):
    """The window to report on: what was asked, or the current month so far."""
    asked = _asked_dates(date_from, date_to)
    if asked:
        return asked
    widest = reports_api.widest_window()
    if not widest:
        return None
    start, end = widest
    return (max(start, end.replace(day=1)), end)


@app.get("/reports/target-achievement", response_class=HTMLResponse)
def target_achievement_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "report_target.html",
        {"user": user, "page": "target", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


@app.get("/api/target-achievement")
def target_achievement_data(request: Request, date_from: str = "", date_to: str = "",
                            cluster: int = 0, month: str = "", round_off: str = ""):
    require_user(request)
    round_off = _rounded(round_off)
    window = _tva_window(date_from, date_to)
    if window is None:
        return JSONResponse({"error": "No shop sales have been uploaded yet, so there is "
                                      "nothing to measure a target against."})
    data = reports_api.target_vs_achievement(
        window[0], window[1], cluster=cluster or None, month=month,
        round_off=bool(round_off))
    data.setdefault("bonds", sorted({b for members in reports_api.bond_clusters().values()
                                     for b in members}))
    data["calendar"] = _calendar_days()
    data["suggest"] = _suggested_window()
    return JSONResponse(data)


@app.get("/api/targets")
def targets_read(request: Request, month: str = ""):
    require_user(request)
    month = month or date.today().strftime("%Y-%m")
    try:
        cells = targets_mod.load(month)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    return JSONResponse({
        "month": month,
        "bonds": cells,
        "families": [{"key": k, "label": l} for k, l in targets_mod.FAMILIES],
        "all_bonds": sorted({b for members in reports_api.bond_clusters().values()
                             for b in members}),
        "clusters": {str(c): sorted(m) for c, m in reports_api.bond_clusters().items()},
        "months": targets_mod.months(),
        "meta": targets_mod.meta(month),
    })


@app.post("/api/targets")
async def targets_write(request: Request):
    user = require_user(request)
    body = await request.json()
    month = str(body.get("month") or "")
    try:
        saved = targets_mod.save(month, body.get("bonds") or {}, by=user)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    return JSONResponse({"ok": True, **saved})


@app.get("/reports/target-achievement/export.xlsx")
def target_achievement_xlsx(request: Request, date_from: str = "", date_to: str = "",
                            cluster: int = 0, month: str = "", round_off: str = ""):
    """The sheet as it prints, in the shape the office already circulates."""
    require_user(request)
    round_off = _rounded(round_off)
    window = _tva_window(date_from, date_to)
    if window is None:
        raise HTTPException(404, "No shop sales have been uploaded yet.")
    data = reports_api.target_vs_achievement(
        window[0], window[1], cluster=cluster or None, month=month,
        round_off=bool(round_off))
    if "error" in data:
        raise HTTPException(404, data["error"])

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    # The same sheet as the PDF, cell for cell: the office should not be able
    # to tell which of the two it is looking at.
    NAVY, GOLD, GOLD_D = "FF0A294F", "FFFFBD30", "FFF2B22E"
    SHADE, INK, RED = "FFF2F2F2", "FF1B2A4A", "FFCF1322"
    LINE = Side(style="thin", color="FFC7C7C7")
    BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

    cols = data["columns"]
    labels = ["BOND", "CAT"] + [reports_pdf.TA_LABELS.get(c["key"], str(c["label"]).upper())
                                for c in cols] + ["GRAND TOTAL", "ACH %"]
    span = len(labels)
    last_col = get_column_letter(span)

    wb = Workbook()
    ws = wb.active
    ws.title = f"TGT vs ACH {data['month'][:4]}"[:31]
    ws.sheet_view.showGridLines = False

    def band(row, text, fill, colour, size, align="center", height=None):
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(f"A{row}:{last_col}{row}")
        for col in range(1, span + 1):
            c = ws.cell(row=row, column=col)
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(bold=True, size=size, color=colour)
            c.alignment = Alignment(horizontal=align, vertical="center")
        if height:
            ws.row_dimensions[row].height = height

    end_day = data["period"]["to"]
    as_on = (f"AS ON {int(end_day[8:10])} "
             f"{reports_pdf._MONTH_ABBR[int(end_day[5:7]) - 1]} {end_day[:4]}")
    title = "TARGET VS ACHIEVEMENT"
    if cluster in (1, 2, 3):
        title += f" - CLUSTER {cluster}"

    band(1, "K.S DISTILLERY", NAVY, GOLD, 18, height=32)
    band(2, title, GOLD, NAVY, 12, align="left", height=22)
    # the as-on date sits at the far end of the same gold band
    ws.unmerge_cells(f"A2:{last_col}2")
    ws.merge_cells(f"A2:{get_column_letter(span - 4)}2")
    ws.merge_cells(f"{get_column_letter(span - 3)}2:{last_col}2")
    # Merging re-anchors the range, and the fresh anchor comes back unpainted -
    # which left the right half of the gold band white.
    for col in range(1, span + 1):
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor=GOLD)
    right = ws.cell(row=2, column=span - 3, value=as_on)
    right.font = Font(bold=True, size=12, color=NAVY)
    right.alignment = Alignment(horizontal="right", vertical="center")

    # The header sits straight under the gold band: an empty row there reads as
    # a white gap the sheet never closes.
    head_row = 3
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=head_row, column=i, value=label)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, size=9.5, color="FFFFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[head_row].height = 30

    rows = data["rows"]
    has_bonds = any(r.get("kind") == "bond" for r in rows)
    r = head_row

    for idx, row in enumerate(rows):
        total = idx == len(rows) - 1
        # The same three tiers the PDF prints: a bond on the body's colours, a
        # cluster subtotal on navy with gold figures, and the one row the sheet
        # adds up to in gold.
        clus = not total and has_bonds and row.get("kind") == "cluster"
        strong = total or clus
        top, bottom = r + 1, r + 2
        for which, tag, rr, fill in (
                ("tgt", "TGT", top, GOLD_D if total else NAVY if clus else SHADE),
                ("ach", "ACH", bottom, GOLD if total else NAVY if clus else "FFFFFFFF")):
            line = [row["label"] if rr == top else None, tag]
            line += [float(row[which].get(c["key"], 0) or 0) for c in cols]
            line.append(float(row.get(which + "_total", 0) or 0))
            line.append((row["pct"] / 100) if (rr == top and row["pct"] is not None) else None)
            for i, value in enumerate(line, start=1):
                c = ws.cell(row=rr, column=i, value=value)
                c.border = BOX
                c.alignment = Alignment(horizontal="center", vertical="center")
                if i == 1:
                    c.fill = PatternFill("solid", fgColor=NAVY)
                    c.font = Font(bold=True, size=10,
                                  color=GOLD if total or clus else "FFFFFFFF")
                elif i == span:
                    c.fill = PatternFill("solid", fgColor=GOLD if total
                                         else NAVY if clus else "FFFFFFFF")
                    c.font = Font(bold=True, size=9.5,
                                  color=INK if total else GOLD if clus else RED)
                    c.number_format = "0.00%"
                else:
                    c.fill = PatternFill("solid", fgColor=fill)
                    c.font = Font(bold=strong or i in (2, span - 1), size=9.5,
                                  color=GOLD if clus else INK)
                    if i > 2:
                        c.number_format = "#,##0" if round_off else "#,##0.##"

        ws.merge_cells(start_row=top, start_column=1, end_row=bottom, end_column=1)
        ws.merge_cells(start_row=top, start_column=span, end_row=bottom, end_column=span)
        ws.row_dimensions[top].height = 18
        ws.row_dimensions[bottom].height = 18
        r = bottom

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 7
    for i in range(3, span - 1):
        ws.column_dimensions[get_column_letter(i)].width = 11.5
    ws.column_dimensions[get_column_letter(span - 1)].width = 13
    ws.column_dimensions[get_column_letter(span)].width = 10
    ws.freeze_panes = ws.cell(row=head_row + 1, column=3)

    # Twelve columns do not fit a portrait page, and nothing here said so, so
    # the sheet printed as two - six columns, then the rest, with the house
    # band sliced down the middle. It is one landscape page wide now, the
    # header repeats down a long one, and the print area stops where the
    # figures do.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4,
                                  header=0.2, footer=0.2)
    ws.print_area = f"A1:{last_col}{r}"
    ws.print_title_rows = f"1:{head_row}"

    label = data["period"]["short"].replace(" to ", " - ")
    tmp = Path(tempfile.mkdtemp()) / f"TARGET vs ACHIEVEMENT ({label}).xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


@app.get("/reports/target-achievement/export.pdf")
def target_achievement_pdf(request: Request, date_from: str = "", date_to: str = "",
                           cluster: int = 0, month: str = "", scope: str = "current",
                           round_off: str = ""):
    """Two ways out: what is on screen, or the set the office circulates.

    'all' is four sheets - one per cluster and the summary that sits on top of
    them - zipped, because that is how they go out together. 'current' is the
    page as it stands, cluster filter and all.
    """
    require_user(request)
    round_off = _rounded(round_off)
    if scope not in ("current", "all"):
        raise HTTPException(400, "Unknown scope.")
    window = _tva_window(date_from, date_to)
    if window is None:
        raise HTTPException(404, "No shop sales have been uploaded yet.")

    import tempfile, zipfile

    def grid(which: int | None):
        got = reports_api.target_vs_achievement(
            window[0], window[1], cluster=which, month=month,
            round_off=bool(round_off))
        if "error" in got:
            raise HTTPException(404, got["error"])
        return got

    out = Path(tempfile.mkdtemp())
    span = None

    if scope == "current":
        data = grid(cluster or None)
        span = data["period"]["short"].replace(" to ", " - ")
        name = f"CLUSTER {cluster}" if cluster in (1, 2, 3) else ""
        stem = f"TARGET vs ACHIEVEMENT{' - ' + name if name else ''} ({span})"
        pdf = reports_pdf.build_target_pdf(data, out / f"{stem}.pdf", scope=name,
                                           round_off=bool(round_off))
        return FileResponse(pdf, filename=pdf.name, media_type="application/pdf")

    whole = grid(None)
    span = whole["period"]["short"].replace(" to ", " - ")
    made: list[Path] = []
    for cid in (1, 2, 3):
        one = grid(cid)
        if not one["rows"]:
            continue
        made.append(reports_pdf.build_target_pdf(
            one, out / f"TARGET vs ACHIEVEMENT - CLUSTER {cid} ({span}).pdf",
            scope=f"CLUSTER {cid}", round_off=bool(round_off)))

    summary = [r for r in whole["rows"] if r.get("kind") in ("cluster", "grand")]
    made.append(reports_pdf.build_target_pdf(
        whole, out / f"TARGET vs ACHIEVEMENT - CLUSTER SUMMARY ({span}).pdf",
        rows=summary, scope="CLUSTER SUMMARY", round_off=bool(round_off)))

    bundle = out / f"TARGET vs ACHIEVEMENT ({span}).zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for pdf in made:
            z.write(pdf, pdf.name)
    return FileResponse(bundle, filename=bundle.name, media_type="application/zip")


# ---------------------------------------------------------------------------
# Secondary sales analysis - item issue
# ---------------------------------------------------------------------------


@app.get("/reports/secondary-analysis", response_class=HTMLResponse)
def item_issue_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "report_item_issue.html",
        {"user": user, "page": "item_issue", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


@app.get("/api/secondary-analysis")
def item_issue_data(request: Request, period: str = "", prior: str = "",
                    cluster: int = 0):
    require_user(request)
    return JSONResponse(reports_api.item_issue(period=period, prior=prior,
                                               cluster=cluster or None))


@app.post("/api/secondary-analysis/industry")
async def item_issue_industry(request: Request):
    """The industry figure is typed, because no export carries it."""
    user = require_user(request)
    body = await request.json()
    period = (body.get("period") or "").strip()
    if not period or not itemissue.period_of(period):
        raise HTTPException(400, "Which period is this industry figure for?")
    try:
        cases = float(body.get("cases") or 0)
        prior = float(body.get("prior") or 0)
    except (TypeError, ValueError):
        raise HTTPException(400, "Industry figures have to be numbers.")
    if cases < 0 or prior < 0:
        raise HTTPException(400, "Industry figures cannot be negative.")
    saved = itemissue.industry_set(period, cases, prior,
                                   by=getattr(user, "email", "") or str(user))
    return JSONResponse(saved)


@app.get("/reports/secondary-analysis/export.xlsx")
def item_issue_xlsx(request: Request, period: str = "", prior: str = "",
                    cluster: int = 0, round_off: str = ""):
    """The sheet as it prints, in the house colours."""
    require_user(request)
    round_off = _rounded(round_off)
    data = reports_api.item_issue(period=period, prior=prior, cluster=cluster or None)
    if "error" in data:
        raise HTTPException(404, data["error"])

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY, GOLD, INK, CREAM = "FF0A294F", "FFFFBD30", "FF1B2A4A", "FFFFFCF0"
    RED, GREEN = "FFB42318", "FF1B7F3B"
    LINE = Side(style="thin", color="FFC7C7C7")
    BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

    def day(iso):
        if not iso:
            return ""
        d = date.fromisoformat(iso)
        return f"{d.day} {reports_pdf._MONTH_ABBR[d.month - 1]} {d.year}"

    heads = (["WAREHOUSE"]
             + ["STN", "GTN", "TOTAL", "C FED", "BAR", "TO DATE"]
             + ["STN", "GTN", "TOTAL", "C FED", "BAR", "TO DATE"]
             + ["CASES", "%", "LAST MONTH"])
    span = len(heads)
    last_col = get_column_letter(span)

    wb = Workbook()
    ws = wb.active
    ws.title = "Secondary Analysis"
    ws.sheet_view.showGridLines = False

    def band(row, text, fill, colour, size, align="center", height=None):
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(f"A{row}:{last_col}{row}")
        for col in range(1, span + 1):
            c = ws.cell(row=row, column=col)
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(bold=True, size=size, color=colour)
            c.alignment = Alignment(horizontal=align, vertical="center")
        if height:
            ws.row_dimensions[row].height = height

    band(1, "K.S DISTILLERY", NAVY, GOLD, 18, height=32)
    band(2, f"SECONDARY SALES \u00b7 {data['period_label'].upper()}", GOLD, NAVY, 12,
         align="left", height=22)
    ws.unmerge_cells(f"A2:{last_col}2")
    ws.merge_cells(f"A2:{get_column_letter(span - 3)}2")
    ws.merge_cells(f"{get_column_letter(span - 2)}2:{last_col}2")
    right = ws.cell(row=2, column=span - 2, value=f"AS ON {day(data['as_on'])}")
    right.font = Font(bold=True, size=12, color=NAVY)
    right.alignment = Alignment(horizontal="right", vertical="center")

    # the two period banners over their blocks
    ws.cell(row=3, column=2, value=f"{data['period_label'].upper()}")
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=7)
    ws.cell(row=3, column=8,
            value=(data["prior_label"].upper() if data["prior"] else "NO PRIOR PULL"))
    ws.merge_cells(start_row=3, start_column=8, end_row=3, end_column=13)
    ws.cell(row=3, column=14, value="DIFFERENCE")
    ws.merge_cells(start_row=3, start_column=14, end_row=3, end_column=15)
    for col in range(1, span + 1):
        c = ws.cell(row=3, column=col)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, size=9.5, color=GOLD)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX

    for i, label in enumerate(heads, start=1):
        c = ws.cell(row=4, column=i, value=label)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, size=9.5, color="FFFFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[4].height = 26

    r = 4
    for row in data["rows"]:
        r += 1
        cur, pri = row["cur"], row["prior"]
        line = ([row["label"]]
                + [cur["stn"], cur["gtn"], cur["total"], cur["cfed"], cur["bar"], cur["all"]]
                + [pri["stn"], pri["gtn"], pri["total"], pri["cfed"], pri["bar"], pri["all"]]
                + [row["diff"],
                   (row["pct"] / 100) if row["pct"] is not None else None,
                   row["last_month"]])
        total = row["kind"] == "grand"
        band_row = row["kind"] == "cluster"
        for i, value in enumerate(line, start=1):
            c = ws.cell(row=r, column=i, value=value)
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if i == 1 else "center",
                                    vertical="center")
            if total:
                c.fill = PatternFill("solid", fgColor=NAVY)
                c.font = Font(bold=True, size=10, color=GOLD)
            elif band_row:
                c.fill = PatternFill("solid", fgColor=GOLD)
                c.font = Font(bold=True, size=10, color=INK)
            else:
                if 8 <= i <= 13 or i == span:
                    c.fill = PatternFill("solid", fgColor=CREAM)
                colour = INK
                if i == 14 and isinstance(value, (int, float)):
                    colour = GREEN if value > 0 else RED if value < 0 else INK
                c.font = Font(bold=(i in (7, 13)), size=10, color=colour)
            if i == 15:
                if value is not None:
                    c.number_format = "0.0%" if round_off else "0.00%"
            elif i > 1:
                c.number_format = "#,##0" if round_off else "#,##0.##"
        ws.row_dimensions[r].height = 17

    ws.column_dimensions["A"].width = 22
    for i in range(2, span + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.freeze_panes = ws.cell(row=5, column=2)

    tmp = (Path(tempfile.mkdtemp())
           / f"SECONDARY SALES ANALYSIS ({data['period_label']}).xlsx")
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)


# ---------------------------------------------------------------------------
# Purchase instruction
# ---------------------------------------------------------------------------


@app.get("/reports/pi-variance", response_class=HTMLResponse)
def pi_page(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(
        request, "report_pi.html",
        {"user": user, "page": "pi", "started_at": STARTED_AT,
         "problems": getattr(app.state, "problems", [])},
    )


def _pi_months(month: str = "", prior: str = "") -> tuple[str, str]:
    """The month to show and the one to compare it with.

    Default to the newest stored month against the one before it, because the
    only question anyone opens this report with is 'what changed'.
    """
    have = [m["key"] for m in pi_mod.months()]
    if not have:
        return "", ""
    month = month if month in have else have[0]
    if prior == "-":
        return month, ""
    if prior in have and prior != month:
        return month, prior
    later = [m for m in have if m < month]
    return month, (later[0] if later else "")


@app.get("/api/pi-receipt")
def pi_receipt(request: Request, month: str):
    """What landed for a month already filed.

    The upload writes a receipt and the history keeps it, but a run from
    before the receipt existed has none - and the question ("how many came
    back blank in July?") is worth answering either way. This reads the
    month's filed sheets back and says so.
    """
    require_user(request)
    got = pi_mod.receipt_for(month)
    if "error" in got:
        raise HTTPException(404, got["error"])
    return JSONResponse(got)


@app.get("/api/pi-variance")
def pi_data(request: Request, month: str = "", prior: str = "",
            view: str = "bond", cluster: int = 0):
    require_user(request)
    if view not in ("bond", "warehouse"):
        raise HTTPException(400, "Unknown view.")
    cur, prev = _pi_months(month, prior)
    if not cur:
        return JSONResponse({"error": "No purchase instruction has been uploaded yet. "
                                      "Upload a month's files on the Raw Data Upload page.",
                             "months": []})
    return JSONResponse(pi_mod.variance(cur, view=view,
                                        cluster=cluster or None, prior=prev))


def _safe_name(text: str) -> str:
    """A group name as a file name. KSBC writes '(NO 2)' and '/' into both."""
    clean = re.sub(r"[\\/:*?\"<>|]+", " ", str(text or "")).strip()
    return re.sub(r"\s+", " ", clean) or "report"


@app.get("/reports/pi-variance/export.pdf")
def pi_pdf(request: Request, month: str = "", view: str = "bond",
           cluster: int = 0, group: str = ""):
    """One sheet per bond - or per warehouse, whichever the page is grouped by.

    The office's own bond sheet: the bond's total at the top, then every shop
    under it, one block each, in the layout they already circulate. Ask for
    more than one and they come back zipped, because a bond sheet is a thing
    somebody forwards to one ASM - it is not a chapter of a book.
    """
    require_user(request)
    if view not in ("bond", "warehouse"):
        raise HTTPException(400, "Unknown view.")
    cur, _prev = _pi_months(month, "-")
    if not cur:
        raise HTTPException(404, "No purchase instruction has been uploaded yet.")
    data = pi_mod.variance(cur, view=view, cluster=cluster or None, prior="")
    if "error" in data:
        raise HTTPException(404, data["error"])

    wanted = [r for r in data["rows"] if r["kind"] == "group"]
    if group:
        wanted = [r for r in wanted if r["label"].lower() == group.strip().lower()]
    if not wanted:
        raise HTTPException(404, "Nothing matches that filter.")

    import tempfile, zipfile
    outdir = Path(tempfile.mkdtemp())
    label = data["month_label"]
    made: list[Path] = []

    for row in wanted:
        key = row["key"] or row["label"]
        shops = (data.get("shops") or {}).get(key) or []
        brands = [{"key": code, "label": label}
                  for code, label in sorted(reports_pdf.PB_BRAND.items(),
                                            key=lambda kv: kv[1])]
        out = outdir / f"Purchase Instruction - {_safe_name(row['label'])} ({label}).pdf"
        reports_pdf.build_pi_group_pdf(
            row["label"], label, brands, row,
            [{"label": s["name"], "cells": s["cells"], "total": s["total"]}
             for s in shops],
            out)
        made.append(out)

    if len(made) == 1:
        return FileResponse(made[0], filename=made[0].name)

    scope = (f"Cluster {cluster}" if cluster else
             "All warehouses" if view == "warehouse" else "All bonds")
    bundle = outdir / f"Purchase Instruction - {scope} ({label}).zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for f in made:
            z.write(f, f.name)
    return FileResponse(bundle, filename=bundle.name)


@app.get("/reports/pi-variance/export.xlsx")
def pi_xlsx(request: Request, month: str = "", prior: str = "",
            view: str = "bond", cluster: int = 0):
    """The workbook the office already circulates, cell for cell.

    Bond, then its shops, then the next bond; eight brands four columns wide;
    a zero prints as a dash. No cluster bands and no totals column - this is
    the shape people already read, and a better one they have to re-learn is
    not better.
    """
    require_user(request)
    cur, prev = _pi_months(month, prior)
    if not cur:
        raise HTTPException(404, "No purchase instruction has been uploaded yet.")
    data = pi_mod.variance(cur, view=view, cluster=cluster or None, prior=prev)
    if "error" in data:
        raise HTTPException(404, data["error"])

    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY, GOLD, INK, ROW, TOTAL, DASH = ("FF0B2C52", "FFFAAF19", "FF0F192D",
                                         "FFF5F7FC", "FF2C3540", "FF8C8C8C")
    thin, medium = Side(style="thin", color="FFE3E8F0"), Side(style="medium", color="FFC9D2E2")
    # Brands in the workbook's own order: alphabetical by the full KSBC name.
    brands = sorted(data["brands"], key=lambda b: b["full"])
    span = 1 + len(brands) * 4
    last = get_column_letter(span)

    wb = Workbook()
    ws = wb.active
    ws.title = "PI Variance"
    ws.sheet_view.showGridLines = False

    def band(row, text, fill, colour, size, align="center"):
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(f"A{row}:{last}{row}")
        for col in range(1, span + 1):
            c = ws.cell(row=row, column=col)
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(bold=True, size=size, color=colour)
            c.alignment = Alignment(horizontal=align, vertical="center")

    band(1, "K.S DISTILLERY", NAVY, GOLD, 18)
    band(2, f"PURCHASE INSTRUCTION \u00b7 1 {data['month_label'].upper()}", GOLD, NAVY, 12)
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 24

    ws.cell(row=3, column=1, value="Row Labels")
    ws.merge_cells("A3:A4")
    for i, b in enumerate(brands):
        col = 2 + i * 4
        ws.cell(row=3, column=col, value=b["full"])
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 3)
        for j, m in enumerate(("L3MS", "RL", "RQ", "MQ")):
            ws.cell(row=4, column=col + j, value=m)
    for r in (3, 4):
        ws.row_dimensions[r].height = 20
        for col in range(1, span + 1):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(bold=True, size=9.5, color=GOLD)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def line(label, cellset, fill, colour, bold, zero_colour=None):
        r = ws.max_row + 1
        ws.row_dimensions[r].height = 20
        ws.cell(row=r, column=1, value=label)
        for i, b in enumerate(brands):
            cell = cellset.get(b["code"], {})
            for j, m in enumerate(("l3ms", "rl", "rq", "mq")):
                v = cell.get(m, 0)
                ws.cell(row=r, column=2 + i * 4 + j, value=v if v else "-")
        for col in range(1, span + 1):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=fill)
            blank = col > 1 and c.value == "-"
            c.font = Font(bold=bold, size=9.5,
                          color=(zero_colour if (blank and zero_colour) else colour))
            c.alignment = Alignment(horizontal="left" if col == 1 else "center",
                                    vertical="center")
            c.number_format = "#,##0"
            c.border = Border(bottom=thin, left=medium if (col - 2) % 4 == 0 or col == 1 else thin)
        return r

    for row in data["rows"]:
        if row["kind"] != "group":
            continue
        line(row["label"], row["cells"], NAVY, GOLD, True)
        for shop in (data["shops"].get(row["key"]) or []):
            line(shop["name"], shop["cells"], ROW, INK, False, zero_colour=DASH)

    grand = next((r for r in data["rows"] if r["kind"] == "grand"), None)
    if grand:
        r = line("Grand Total", grand["cells"], TOTAL, "FFFFFFFF", True)
        ws.cell(row=r, column=1).font = Font(bold=True, size=9.5, color=GOLD)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12
    for i in range(3, span + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = "B5"

    # The blank-PI list rides along on its own sheet: it changes nothing about
    # the sheet above, and it is the one thing the figures cannot tell you.
    if data["blanks"]:
        b = wb.create_sheet("No instruction")
        b.append(["Bond", "Shop", f"{data['prior_label'] or 'Last month'} MQ"])
        for cell in b[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color=GOLD, size=10)
        for x in data["blanks"]:
            b.append([x["bond"], x["name"], x["prior_mq"] or ""])
        b.column_dimensions["A"].width = 20
        b.column_dimensions["B"].width = 38
        b.column_dimensions["C"].width = 16

    tmp = Path(tempfile.mkdtemp()) / f"pi_variance_report_{data['month']}.xlsx"
    centre_all(wb)
    wb.save(tmp)
    return FileResponse(tmp, filename=tmp.name)
