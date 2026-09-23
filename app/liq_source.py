"""The liquidation builder's sources, rebuilt from what this app holds.

The artifact this dashboard came from read one workbook per month: KSBC's
`<MONTH> ... ANALYSIS.xlsx`, a container whose sheets are the day exports
themselves (`SEPTEMBER 17`), the portal's block exports (`SEPTEMBER 1-16
CUMULATIVE`) and a roll-up of both (`SEPTEMBER 1-20 COMBINED`). That workbook
is assembled on Abhay's own machine and the server has never had one, because
the reports here read each day's export directly - which is why they work while
the dashboard had nothing to stand on.

The server does hold every grid that workbook is made of:

    KSBC shop sales/september 17th.xlsx                  one day
    KSBC shop sales/_cumulative/2026-09 ... CUMULATIVE.xlsx   one block
    Secondary sales/RAW DATA -SEPTEMBER 17TH ... .xlsx   one day of invoices

column for column the same exports. So rather than reimplement the
arithmetic - which would be a second definition of liquidation, free to drift
from the one Abhay checks his month against - this module dresses those files
as the sheets the builder already knows how to read. Every figure is then
computed by the original code, unchanged.

Two deliberate differences from a real month workbook:

  * A day or block sheet here carries only its lines that SOLD something.
    Both readers of those sheets sum Shop Out and skip the rest, and a day's
    export is mostly zeros - four thousand of its five and a half thousand
    lines on a normal day. Dropping them is what makes holding a month in
    memory reasonable.
  * The COMBINED sheet is computed here, block-true exactly as the builder's
    own window pass computes it: where a CUMULATIVE block covers days, it
    supersedes those days' own files, because KSBC's per-day case conversion
    does not sum to its block total.

What this cannot invent is a day nobody uploaded. A month shows the days the
app has, and says so through the same asOf / period labels as before.
"""
from __future__ import annotations

import logging
import re
from collections import defaultdict
from pathlib import Path

log = logging.getLogger("ksd.liquidation")

MONTHS_UP = ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE',
             'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']

# KSBC exports, in the shape the upload page files them under.
_DAY_RE = re.compile(r"^([a-z]+)\s+(\d{1,2})(?:st|nd|rd|th)\.xlsx$", re.I)
_CUM_RE = re.compile(r"([a-z]+)\s+(\d{1,2})\s*-\s*(\d{1,2})\s+CUMULATIVE\.xlsx$", re.I)
_SEC_RE = re.compile(r"^RAW DATA\s*-\s*([a-z]+)\s+(\d{1,2})(?:st|nd|rd|th)\s+SECONDARY",
                     re.I)

# Column positions in a KSBC export (SupplierWiseShopSaleReport). The builder
# reads these same positions out of a month workbook's day sheets.
C_SHOP, C_BRAND, C_PACK, C_BPC = 1, 4, 5, 6
C_OUT_CS, C_OUT_BTL, C_CLOSE_CS, C_CLOSE_BTL = 11, 12, 13, 14

# ... and in a secondary export (SupplierSalesAnalysisReport), which is the
# COMBINED DISPATCHES layout already.
S_LICENSEE, S_DATE, S_CASES = 6, 11, 12


# ---------------------------------------------------------------------------
# A workbook that is really a folder
# ---------------------------------------------------------------------------
class _Cell:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _Sheet:
    """Enough of an openpyxl worksheet for the builder's readers: iter_rows,
    max_row, and ws[1] for a header."""

    def __init__(self, rows: list):
        self.rows = rows

    @property
    def max_row(self) -> int:
        return len(self.rows)

    def iter_rows(self, min_row: int = 1, values_only: bool = True, **_kw):
        for row in self.rows[max(0, min_row - 1):]:
            yield row

    def __getitem__(self, n: int):
        return [_Cell(v) for v in self.rows[n - 1]]


class _Book:
    """Sheets named as a month workbook's are, materialised on first read."""

    def __init__(self, sheets: dict):
        # {name: callable -> rows}
        self._make = sheets
        self._got: dict = {}

    @property
    def sheetnames(self) -> list:
        return list(self._make)

    def __getitem__(self, name: str) -> _Sheet:
        if name not in self._got:
            self._got[name] = _Sheet(self._make[name]())
        return self._got[name]

    def close(self) -> None:
        """A real workbook holds a file open; this holds rows the next reader
        of the same month wants. Closing it would throw away the only reason
        it is cached."""

    @property
    def rowcount(self) -> int:
        return sum(len(s.rows) for s in self._got.values())


# ---------------------------------------------------------------------------
# Finding a month's exports
# ---------------------------------------------------------------------------
def _ksbc_files(root: Path, month: str):
    """({day: path}, [(from, to, path)]) for one month's KSBC exports."""
    folder = root / "KSBC shop sales"
    days, blocks = {}, []
    if not folder.is_dir():
        return days, blocks
    up = month.upper()
    for path in folder.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        m = _DAY_RE.match(path.name)
        if m and m.group(1).upper() == up:
            days[int(m.group(2))] = path
    cum = folder / "_cumulative"
    if cum.is_dir():
        for path in cum.glob("*.xlsx"):
            if path.name.startswith("~$"):
                continue
            m = _CUM_RE.search(path.name)
            if m and m.group(1).upper() == up:
                blocks.append((int(m.group(2)), int(m.group(3)), path))
    return days, _pick_blocks(blocks)


def pick_periods(blocks: list) -> list:
    """Public name for the overlap filter: liq_live needs it for the period
    SHEETS inside a real month workbook, which can overlap the same way."""
    return _pick_blocks(blocks)


def _pick_blocks(blocks: list) -> list:
    """Blocks that can be added together - no day counted twice.

    A month workbook's blocks tile the month (1-16, then 17-EOM) because that
    is how KSBC's portal exports them, and the builder adds them up. What gets
    uploaded here is cumulative-to-date: 1-8 one week, 1-16 the next, 1-22
    today, each one containing the last. Adding those counts most of the month
    five or six times over - 33,170 cs of KSBC tertiary against a daily trend
    totalling 5,790.

    So take the set of blocks that overlap nowhere and cover the most days:
    1-22 alone beats 1-8 + 1-16, and 1-16 + 17-23 beats 1-20. Days no block
    covers still come from their own day file.
    """
    if not blocks:
        return []
    spans = [(a, b, p, set(range(a, b + 1))) for a, b, p in blocks]

    def pick(rest, covered, used):
        """Best (days covered, fewest blocks) from here on."""
        if not rest:
            return len(covered), -len(used), used
        head, tail = rest[0], rest[1:]
        best = pick(tail, covered, used)
        if not (head[3] & covered):
            take = pick(tail, covered | head[3], used + [head])
            if take[:2] > best[:2]:
                best = take
        return best

    if len(spans) <= 12:
        spans.sort(key=lambda t: (t[0], -t[1]))
        chosen = pick(spans, set(), [])[2]
    else:
        # More period exports than anyone has reason to keep; widest-first is
        # close enough and does not blow up.
        chosen, covered = [], set()
        for sp in sorted(spans, key=lambda t: (t[0] - t[1], t[0])):
            if sp[3] & covered:
                continue
            chosen.append(sp)
            covered |= sp[3]

    keep = {id(c) for c in chosen}
    dropped = [getattr(sp[2], "name", str(sp[2])) for sp in spans if id(sp) not in keep]
    if dropped:
        log.info("liquidation: period export(s) inside another, skipped: %s",
                 ", ".join(sorted(dropped)))
    return sorted((a, b, p) for a, b, p, _ in chosen)


def _sec_files(root: Path, month: str) -> dict:
    """{day: path} for one month's secondary raw exports."""
    folder = root / "Secondary sales"
    out = {}
    if not folder.is_dir():
        return out
    up = month.upper()
    for path in folder.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        m = _SEC_RE.match(path.name)
        if m and m.group(1).upper() == up:
            out[int(m.group(2))] = path
    return out


def sources(root: Path, month: str) -> list:
    """Every file this month is built from - for the cache's fingerprint."""
    days, blocks = _ksbc_files(root, month)
    return (sorted(days.values())
            + sorted(p for _, _, p in blocks)
            + sorted(_sec_files(root, month).values()))


def ksbc_covered_to(root: Path, month: str) -> int:
    """The last day of the month the KSBC exports account for.

    A period export covers days that have no day file, and the builder's
    day-sheet scan cannot see inside one - so asking it alone how far the month
    has got answers 5 when a 1-22 block is sitting right there, and every
    per-day figure then divides 22 days of sales by 5. Invoice coverage is not
    counted here: that leg is dated by the invoices themselves, which the
    builder already reads.
    """
    days, blocks = _ksbc_files(root, month)
    return max([0] + list(days) + [b for _, b, _ in blocks])


def has_ksbc(root: Path, month: str) -> bool:
    days, blocks = _ksbc_files(root, month)
    return bool(days or blocks)


def has_secondary(root: Path, month: str) -> bool:
    return bool(_sec_files(root, month))


def months_with_data(root: Path) -> list:
    """Months with both legs somewhere on disk, in calendar order."""
    return [m for m in MONTHS_UP if has_ksbc(root, m) and has_secondary(root, m)]


def newest_month(root: Path):
    """The month whose newest KSBC export is newest - the working month."""
    best = None
    for m in MONTHS_UP:
        days, blocks = _ksbc_files(root, m)
        paths = list(days.values()) + [p for _, _, p in blocks]
        if not paths:
            continue
        stamp = max(p.stat().st_mtime for p in paths)
        if best is None or stamp > best[0]:
            best = (stamp, m)
    return best[1] if best else None


# ---------------------------------------------------------------------------
# Reading one export
# ---------------------------------------------------------------------------
def _sold_rows(path: Path) -> list:
    """A KSBC export's lines that sold something, as the export's own rows."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:                                    # noqa: BLE001
        # A half-uploaded or corrupt export is one day missing, not a dead
        # dashboard. zipfile raises things that are neither OSError nor
        # ValueError, hence the broad catch.
        log.info("liquidation: skipping %s (%s)", path.name, type(exc).__name__)
        return []
    try:
        ws = wb[wb.sheetnames[0]]
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= C_OUT_CS or row[C_SHOP] is None:
                continue
            try:
                cases = float(row[C_OUT_CS] or 0)
                btl = float(row[C_OUT_BTL] or 0) if len(row) > C_OUT_BTL else 0.0
            except (TypeError, ValueError):
                continue
            if cases or btl:
                out.append(row)
        return out
    finally:
        wb.close()


def _all_rows(path: Path) -> list:
    """Every line of a KSBC export, zeros included - for closing stock."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:                                    # noqa: BLE001
        log.info("liquidation: skipping %s (%s)", path.name, type(exc).__name__)
        return []
    try:
        ws = wb[wb.sheetnames[0]]
        return [r for r in ws.iter_rows(min_row=2, values_only=True)
                if r and len(r) > C_OUT_CS and r[C_SHOP] is not None]
    finally:
        wb.close()


def _sec_rows(path: Path) -> list:
    """A secondary export's dispatch lines - already the COMBINED layout."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:                                    # noqa: BLE001
        log.info("liquidation: skipping %s (%s)", path.name, type(exc).__name__)
        return []
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        return rows
    finally:
        wb.close()


def _cases(row) -> float:
    try:
        bpc = float(row[C_BPC] or 0)
        cs = float(row[C_OUT_CS] or 0)
        btl = float(row[C_OUT_BTL] or 0) if len(row) > C_OUT_BTL else 0.0
    except (TypeError, ValueError):
        return 0.0
    return cs + (btl / bpc if bpc else 0.0)


def _closing(row) -> float:
    try:
        bpc = float(row[C_BPC] or 0)
        cs = float(row[C_CLOSE_CS] or 0) if len(row) > C_CLOSE_CS else 0.0
        btl = float(row[C_CLOSE_BTL] or 0) if len(row) > C_CLOSE_BTL else 0.0
    except (TypeError, ValueError):
        return 0.0
    return cs + (btl / bpc if bpc else 0.0)


# ---------------------------------------------------------------------------
# The books
# ---------------------------------------------------------------------------
_COMBINED_HEAD = ("Shop Code", "Brand Name", "Packing",
                  "Sales (Cases)", "Closing (Cases)")
# A day or block sheet keeps the export's own layout, because that is what the
# builder reads out of a month workbook's day sheets - by position, not name.
_RAW_HEAD = ("Warehouse Name", "Shop Code", "Shop Name", "Product code",
             "Brand Name", "Packing", "Bottle Per Case",
             "Shop Opening Cases", "Shop Opening Bottles",
             "Shop In Cases", "Shop In Bottles",
             "Shop Out Cases", "Shop Out Bottles",
             "Shop Closing Cases", "Shop Closing Bottles")


def _combined(days: dict, blocks: list) -> list:
    """The month's roll-up: sales per shop x brand x pack, block-true, and the
    closing stock the latest export shows.

    Block-true means a CUMULATIVE block supersedes the day files it covers,
    which is what the builder's own window pass does with the same sheets:
    KSBC's per-day case conversion does not sum to its own block total.
    """
    covered = set()
    for a, b, _ in blocks:
        covered.update(range(a, b + 1))

    sales: dict = defaultdict(float)
    for a, b, path in blocks:
        for row in _sold_rows(path):
            sales[_key(row)] += _cases(row)
    for day, path in sorted(days.items()):
        if day in covered:
            continue
        for row in _sold_rows(path):
            sales[_key(row)] += _cases(row)

    # Closing is a position, not a total: it comes from the latest export that
    # covers the end of the period, and it needs the lines that never moved -
    # stock sitting still is the whole point of that panel.
    last = None
    last_day = max(days) if days else 0
    last_block = max((b for _, b, _ in blocks), default=0)
    if days and last_day >= last_block:
        last = days[last_day]
    elif blocks:
        last = max(blocks, key=lambda t: t[1])[2]
    closing: dict = defaultdict(float)
    if last is not None:
        for row in _all_rows(last):
            c = _closing(row)
            if c:
                closing[_key(row)] += c

    out = [list(_COMBINED_HEAD)]
    for key in set(sales) | set(closing):
        code, brand, pack = key
        out.append([code, brand, pack,
                    round(sales.get(key, 0.0), 2),
                    round(closing.get(key, 0.0), 2)])
    return out


def _key(row):
    code = row[C_SHOP]
    try:
        code = int(code)
    except (TypeError, ValueError):
        code = str(code).strip()
    return (code, row[C_BRAND], row[C_PACK])


def ksbc_book(root: Path, month: str):
    """The month's KSBC exports, named as a month workbook's sheets."""
    days, blocks = _ksbc_files(root, month)
    if not days and not blocks:
        return None
    up = month.upper()
    sheets = {}
    for day, path in sorted(days.items()):
        sheets[f"{up} {day}"] = (lambda p=path: [list(_RAW_HEAD)] + _sold_rows(p))
    for a, b, path in blocks:
        sheets[f"{up} {a}-{b} CUMULATIVE"] = (
            lambda p=path: [list(_RAW_HEAD)] + _sold_rows(p))
    last = max(list(days) + [b for _, b, _ in blocks])
    sheets[f"{up} 1-{last} COMBINED"] = lambda: _combined(days, blocks)
    return _Book(sheets)


def secondary_book(root: Path, month: str):
    """The month's invoice exports as one COMBINED DISPATCHES sheet."""
    files = _sec_files(root, month)
    if not files:
        return None

    def dispatches():
        head, body = None, []
        for _day, path in sorted(files.items()):
            rows = _sec_rows(path)
            if not rows:
                continue
            if head is None:
                head = list(rows[0])
            for row in rows[1:]:
                if not row or len(row) <= S_CASES or row[S_LICENSEE] is None:
                    continue
                try:
                    if not float(row[S_CASES] or 0):
                        continue
                except (TypeError, ValueError):
                    continue
                body.append(row)
        return [head or ["Licensee No."]] + body

    return _Book({f"{month.upper()} COMBINED DISPATCHES": dispatches})


# A payload reads its month four times over - the roll-up, the daily trend,
# the window breakdown, the stock panel - and a month is twenty-odd files.
# Keeping the last few books alive means each file is read once per build
# rather than four times; they hold only the lines that sold.
_CACHE: dict = {}
_KEEP = 2   # the current month and the prior one, each in two legs


def cached_book(root: Path, month: str, kind: str):
    """A book per (month, kind), reused while it is one of the last few."""
    files = sources(root, month)
    stamp = "|".join(f"{p.name}:{p.stat().st_size}:{int(p.stat().st_mtime)}"
                     for p in files)
    key = (str(root), month.upper(), kind)
    hit = _CACHE.get(key)
    if hit and hit[0] == stamp:
        return hit[1]
    book = (ksbc_book if kind == "ksbc" else secondary_book)(root, month)
    if book is None:
        return None
    _CACHE[key] = (stamp, book)
    while len(_CACHE) > _KEEP * 2:
        _CACHE.pop(next(iter(_CACHE)))
    return book


def forget() -> None:
    """Drop the held exports once a build is done.

    They exist to stop one payload reading the same twenty files four times
    over. Between builds they are tens of megabytes of rows the payload cache
    already has the answer to.
    """
    _CACHE.clear()
