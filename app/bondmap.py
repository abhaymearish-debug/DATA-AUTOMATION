"""Bond mapping: which shop belongs to which bond, and which bond to which cluster.

Two different kinds of fact live here, so they are stored differently.

  * Shop -> bond is KSBC's own reference data. It already lives in the 'Bond'
    column of MASTER DATA CONFIRMED.xlsx, which the 81 build scripts read for
    themselves. Editing anywhere else would give the app one answer and the
    workbooks another, so this writes back to that column — and takes a
    timestamped backup of the file first, every time.

  * Bond -> cluster is not in any workbook; it was a constant in the code. It
    becomes a small JSON the app owns, so the ASM split can be changed without
    a deploy. The built-in split is the fallback.

  * A new outlet - a KSBC shop, a Consumerfed outlet or a bar - is a new row
    in that same master sheet, written in the layout the build scripts read by
    position (Sl.No. first, then W Code, Warehouse, Code, Name, CAT, Field
    Staff, Bond, Status). Until it is there, its dispatches count in the
    network total and in no bond.
"""

from __future__ import annotations

import json
import re
import shutil
import threading
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl

from . import config

SHEET = "16-4-25"
BOND_COL = "Bond"

# Both writers rewrite the whole workbook, so two at once would lose one.
_WRITE_LOCK = threading.Lock()

# What master data's CAT column holds, and what the office calls each.
CATS = {"KSBC": "Shop", "FED": "Consumerfed", "BAR": "Bar"}
# KSBC's licensee numbers say what the outlet is in their first digit:
# 1xxxxx a KSBC shop, 2xxxxx a Consumerfed outlet, 4xxxxx a bar.
PREFIX_CAT = {"1": "KSBC", "2": "FED", "4": "BAR"}

# The split as it has always been. An override file replaces it wholesale.
DEFAULT_CLUSTERS: dict[str, list[str]] = {
    "1": ["ALAPPUZHA", "ATTINGAL", "KOLLAM", "KOTTARAKARA", "NEDUMANGAD", "PATHANAMTHITTA"],
    "2": ["ALUVA", "KOTTAYAM", "THODUPUZHA", "THRISSUR", "TRIPUNITHURA"],
    "3": ["KANNUR", "KOZHIKODE", "PALAKKAD", "PERINTHALMANNA"],
}


def clusters_file() -> Path:
    return config.WORKSPACE_ROOT / "_seed" / "bond_clusters.json"


def load_clusters() -> dict[str, list[str]]:
    path = clusters_file()
    if path.is_file():
        try:
            data = json.loads(path.read_text())
            out = {str(k): [str(b).strip().upper() for b in v if str(b).strip()]
                   for k, v in data.items() if str(k) in ("1", "2", "3")}
            if out:
                return {k: out.get(k, []) for k in ("1", "2", "3")}
        except (OSError, ValueError):
            pass
    return {k: list(v) for k, v in DEFAULT_CLUSTERS.items()}


def save_clusters(mapping: dict[str, list[str]]) -> None:
    """Replace the cluster split, having checked it is one.

    This is written from a request body and decides how every report groups
    its rows, so a malformed or empty body used to be able to blank the split
    with nothing to restore it from. Three checks and a backup: the bonds have
    to be bonds this install knows, no bond may sit in two clusters, and the
    result may not be empty. The previous file is kept beside the new one.
    """
    if not isinstance(mapping, dict):
        raise ValueError("The cluster mapping must be an object of 1/2/3 to bond lists.")

    clean = {k: sorted({str(b).strip().upper() for b in (mapping.get(k) or [])
                        if str(b).strip()})
             for k in ("1", "2", "3")}

    if not any(clean.values()):
        raise ValueError("That would leave every cluster empty.")

    seen: dict[str, str] = {}
    for k, bonds in clean.items():
        for b in bonds:
            if b in seen:
                raise ValueError(f"{b} is in cluster {seen[b]} and cluster {k}.")
            seen[b] = k

    known = {b.strip().upper() for bonds in DEFAULT_CLUSTERS.values() for b in bonds}
    known |= {b.strip().upper() for bonds in load_clusters().values() for b in bonds}
    strangers = sorted(set(seen) - known)
    if strangers:
        raise ValueError("Not a bond on this install: " + ", ".join(strangers))

    path = clusters_file()
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file():
        try:
            path.with_suffix(path.suffix + ".bak").write_text(path.read_text())
        except OSError:
            pass
    path.write_text(json.dumps(clean, indent=2))


def _sheet(wb):
    return wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[-1]]


def _header_row(ws) -> tuple[int, dict[str, int]]:
    """Find the header and its column numbers (1-based), by name not position."""
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
        names = [str(v).strip() if v is not None else "" for v in row]
        if "Shop Code" in names:
            return r, {n: i + 1 for i, n in enumerate(names) if n}
    raise RuntimeError("MASTER DATA has no 'Shop Code' header row.")


def load_shops() -> dict:
    """Every shop in master data, with the bond it is currently mapped to."""
    path = config.MASTER_DATA
    if not path.is_file():
        return {"error": f"MASTER DATA CONFIRMED.xlsx is not in the workspace."}

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = _sheet(wb)
    head_row, ix = _header_row(ws)

    def cell(row, name):
        i = ix.get(name)
        return "" if not i or i > len(row) or row[i - 1] is None else str(row[i - 1]).strip()

    shops = []
    for row in ws.iter_rows(min_row=head_row + 1, values_only=True):
        code = cell(row, "Shop Code")
        if not code:
            continue
        shops.append({
            "code": code,
            "name": cell(row, "Shop Name"),
            "warehouse": cell(row, "Warehouse Name"),
            "cat": cell(row, "CAT").upper(),
            "staff": cell(row, "Field Staff"),
            "bond": cell(row, "Bond").upper(),
            "status": cell(row, "Status"),
        })
    wb.close()

    clusters = load_clusters()
    bonds = sorted({s["bond"] for s in shops if s["bond"]}
                   | {b for v in clusters.values() for b in v})
    return {"shops": shops, "bonds": bonds, "clusters": clusters, "source": path.name,
            "warehouses": _warehouse_choices(shops), "staff": _staff_of_bond(shops),
            "cats": CATS, "prefix": PREFIX_CAT}


def save_bonds(changes: dict[str, str]) -> dict:
    """Write shop -> bond back into master data. Backs the file up first."""
    with _WRITE_LOCK:
        return _save_bonds(changes)


def _save_bonds(changes: dict[str, str]) -> dict:
    path = config.MASTER_DATA
    if not path.is_file():
        return {"error": "MASTER DATA CONFIRMED.xlsx is not in the workspace."}
    if not changes:
        return {"changed": 0, "backup": ""}

    lock = path.parent / f"~${path.name}"
    if lock.exists():
        return {"error": f"'{path.name}' is open in Excel. Close it and save again."}

    backup = path.parent / "_master_backups" / _backup(path)

    # Not read_only: this one writes. openpyxl is pinned at the version the
    # build scripts use, so a round-trip here is the same round-trip they do.
    wb = openpyxl.load_workbook(path)
    ws = _sheet(wb)
    head_row, ix = _header_row(ws)
    code_col, bond_col = ix.get("Shop Code"), ix.get(BOND_COL)
    if not code_col or not bond_col:
        wb.close()
        return {"error": "MASTER DATA is missing a 'Shop Code' or 'Bond' column."}

    changed, unknown = 0, []
    wanted = {str(k).strip(): str(v).strip().upper() for k, v in changes.items()}
    seen = set()
    for r in range(head_row + 1, ws.max_row + 1):
        code = ws.cell(row=r, column=code_col).value
        if code is None:
            continue
        code = str(code).strip()
        if code not in wanted:
            continue
        seen.add(code)
        cell = ws.cell(row=r, column=bond_col)
        if str(cell.value or "").strip().upper() != wanted[code]:
            cell.value = wanted[code]
            changed += 1
    unknown = sorted(set(wanted) - seen)

    wb.save(path)
    wb.close()
    return {"changed": changed, "unknown": unknown, "backup": backup.name}


# ---------------------------------------------------------------------------
# Adding an outlet
# ---------------------------------------------------------------------------


def _wh_label(full: str) -> str:
    """'WH-KOZHIKODE RFL9/2024/00001' -> 'KOZHIKODE', the way every report says it."""
    from . import reports_api
    return reports_api.canonical_warehouse(full) or str(full or "").strip()


def _warehouse_choices(shops: list[dict]) -> list[dict]:
    """One choice per warehouse, carrying what a new row there should say.

    A warehouse is written several ways in master (a new licence number, an
    old one), so the choice is the canonical name, and the row gets the
    spelling and W Code most of that warehouse's outlets already carry. The
    bond offered is the one most of them are mapped to - a suggestion, since
    AYARKKUNNAM, THIRUVALLA and a few others feed two bonds.
    """
    full, bond = defaultdict(Counter), defaultdict(Counter)
    for s in shops:
        if not s.get("warehouse"):
            continue
        label = _wh_label(s["warehouse"])
        full[label][s["warehouse"]] += 1
        if s.get("bond"):
            bond[label][s["bond"]] += 1
    return [{"label": label,
             "full": names.most_common(1)[0][0],
             "bond": bond[label].most_common(1)[0][0] if bond[label] else ""}
            for label, names in sorted(full.items())]


def _staff_of_bond(shops: list[dict]) -> dict[str, str]:
    """Field staff is assigned per bond, so a new outlet takes its bond's."""
    seen = defaultdict(Counter)
    for s in shops:
        if s.get("bond"):
            seen[s["bond"]][s.get("staff") or "VACANT"] += 1
    return {b: c.most_common(1)[0][0] for b, c in seen.items()}


def _extend_to(ws, last: int, new: int) -> None:
    """Stretch the sheet's filter and its CAT colouring over the new row."""
    tail = re.compile(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$")

    def grow(ref: str) -> str:
        m = tail.match(ref or "")
        if not m or int(m.group(4)) < last - 2:
            return ref
        return f"{m.group(1)}{m.group(2)}:{m.group(3)}{max(int(m.group(4)), new)}"

    if ws.auto_filter and ws.auto_filter.ref:
        ws.auto_filter.ref = grow(ws.auto_filter.ref)

    from openpyxl.formatting.formatting import ConditionalFormattingList
    old, fresh = ws.conditional_formatting, ConditionalFormattingList()
    for cf in old:
        ref = " ".join(grow(str(r)) for r in str(cf.sqref).split())
        for rule in cf.rules:
            fresh.add(ref, rule)
    ws.conditional_formatting = fresh


def add_outlet(entry: dict) -> dict:
    """Append one outlet to master data. Raises ValueError on a bad entry."""
    if not isinstance(entry, dict):
        raise ValueError("Expected the outlet's details.")
    code = re.sub(r"\s+", "", str(entry.get("code") or ""))
    if not re.fullmatch(r"\d{4,8}", code):
        raise ValueError("The code is the outlet's Licensee No. in the KSBC raw - "
                         "digits only, like 410013.")
    name = " ".join(str(entry.get("name") or "").split()).upper()
    if not name:
        raise ValueError("Give the outlet a name.")
    if len(name) > 80:
        raise ValueError("That name is longer than 80 characters - shorten it.")
    cat = str(entry.get("cat") or "").strip().upper()
    if cat not in CATS:
        raise ValueError("Pick what it is: Shop, Consumerfed or Bar.")
    bond = str(entry.get("bond") or "").strip().upper()
    wh_pick = str(entry.get("warehouse") or "").strip().upper()
    status = "Closed" if str(entry.get("status") or "").strip().lower() == "closed" else "Active"

    with _WRITE_LOCK:
        path = config.MASTER_DATA
        if not path.is_file():
            raise ValueError("MASTER DATA CONFIRMED.xlsx is not in the workspace.")
        if (path.parent / f"~${path.name}").exists():
            raise ValueError(f"'{path.name}' is open in Excel. Close it and try again.")

        current = load_shops()
        if current.get("error"):
            raise ValueError(current["error"])
        clash = next((s for s in current["shops"] if s["code"] == code), None)
        if clash:
            raise ValueError(f"{code} is already in master data - {clash['name']}"
                             f"{', ' + clash['bond'] + ' bond' if clash['bond'] else ''}.")
        if bond not in current["bonds"]:
            raise ValueError("Pick the bond it belongs to.")
        wh = next((w for w in current["warehouses"] if w["label"] == wh_pick), None)
        if wh is None:
            raise ValueError("Pick the KSBC warehouse that supplies it.")
        staff = current["staff"].get(bond, "VACANT")

        wb = openpyxl.load_workbook(path)
        ws = _sheet(wb)
        head_row, ix = _header_row(ws)
        need = ["Sl.No.", "W Code", "Warehouse Name", "Shop Code", "Shop Name",
                "CAT", "Field Staff", BOND_COL, "Status"]
        missing = [n for n in need if n not in ix]
        if missing:
            wb.close()
            raise ValueError("MASTER DATA is missing column(s): " + ", ".join(missing))

        last, top_sl = head_row, 0
        wcodes = Counter()
        for r in range(head_row + 1, ws.max_row + 1):
            if ws.cell(row=r, column=ix["Shop Code"]).value in (None, ""):
                continue
            last = r
            sl = ws.cell(row=r, column=ix["Sl.No."]).value
            try:
                top_sl = max(top_sl, int(sl))
            except (TypeError, ValueError):
                pass
            if ws.cell(row=r, column=ix["Warehouse Name"]).value == wh["full"]:
                wcodes[ws.cell(row=r, column=ix["W Code"]).value] += 1
        new = last + 1

        # Dress the row like the one two above it, so the banding carries on.
        model = new - 2 if new - 2 > head_row else last
        width = max(ix.values())
        for col in range(1, width + 1):
            src = ws.cell(row=model, column=col)
            if src.has_style:
                ws.cell(row=new, column=col)._style = copy(src._style)
        h = ws.row_dimensions[last].height or ws.row_dimensions[model].height
        if h:
            ws.row_dimensions[new].height = h

        values = {
            "Sl.No.": top_sl + 1,
            "W Code": wcodes.most_common(1)[0][0] if wcodes else "",
            "Warehouse Name": wh["full"],
            # Master holds codes as numbers; the readers all str() them.
            "Shop Code": int(code) if not code.startswith("0") else code,
            "Shop Name": name,
            "CAT": cat,
            "Field Staff": staff,
            BOND_COL: bond,
            "Status": status,
        }
        for k, v in values.items():
            ws.cell(row=new, column=ix[k]).value = v
        _extend_to(ws, last, new)

        backup = _backup(path)
        wb.save(path)
        wb.close()

    return {"shop": {"code": code, "name": name, "warehouse": wh["full"], "cat": cat,
                     "staff": staff, "bond": bond, "status": status},
            "backup": backup}


def _backup(path: Path) -> str:
    """Copy master aside before a write. Never over an earlier copy: two
    outlets added inside the same second used to share a name, and the second
    copy - already holding the first outlet - replaced the only clean one."""
    backups = path.parent / "_master_backups"
    backups.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = backups / f"{path.stem}.bak_{stamp}{path.suffix}"
    n = 2
    while dest.exists():
        dest = backups / f"{path.stem}.bak_{stamp}_{n}{path.suffix}"
        n += 1
    shutil.copy2(path, dest)
    return dest.name


def unmatched_outlets() -> dict:
    """Outlets KSBC dispatched to that master data has never heard of.

    The secondary build lists them on its UNMATCHED sheet; this is the same
    list, from the dispatch lines the app already reads, so the page can offer
    to add each one with what the raw already says about it.
    """
    from . import reports_api
    try:
        lines, _sources = reports_api.secondary_lines()
        known = set(reports_api.load_master())
    except Exception as exc:  # the page still works without the hint
        return {"outlets": [], "error": f"Could not read the dispatch data: {exc}"}

    found: dict[str, dict] = {}
    for l in lines:
        code = str(l.get("shop_code") or "").strip()
        if not code or code in known or not code.isdigit():
            continue
        o = found.setdefault(code, {"code": code, "name": "", "warehouse": "",
                                    "cases": 0.0, "first": None, "last": None})
        o["name"] = o["name"] or str(l.get("shop") or "").strip()
        o["warehouse"] = o["warehouse"] or str(l.get("warehouse") or "").strip()
        o["cases"] += float(l.get("cases") or 0)
        d = l.get("date")
        if d:
            o["first"] = min(o["first"] or d, d)
            o["last"] = max(o["last"] or d, d)

    out = []
    for o in sorted(found.values(), key=lambda o: -o["cases"]):
        out.append({**o, "cases": round(o["cases"], 2),
                    "cat": PREFIX_CAT.get(o["code"][:1], ""),
                    "first": o["first"].isoformat() if o["first"] else "",
                    "last": o["last"].isoformat() if o["last"] else ""})
    return {"outlets": out}
