#!/usr/bin/env python3
"""Shop Sales - Analysis: the one-page comparative sheet.

Fifteen bonds, grouped into their three clusters, each cluster subtotalled and
the whole page totalled - built from the same cumulative raw the per-bond PDFs
come from, so the two always agree.

Geometry was read out of the office's own PDF with a content-stream parse:

    page          595.276 x 841.89   (A4)
    title band    46  navy, 'K.S DISTILLERY' 20pt gold
    period band   34  gold, 13pt navy
    header        62  navy in two tiers, 8.5pt, gold rules 2.2 / 1.6
    body row      35.68, baseline row_top - (row_height / 2 + font * 0.35)
    sell-through  amber cell at 40% and above, pink below

Every figure is rounded half up from its own unrounded aggregate, never from
another rounded figure - which is why the bond column can sum to one less than
the total and still be right.

Usage:
    build_shop_analysis_pdf.py --raw <cumulative.xlsx> --outdir <dir>
                               [--prev <last month's cumulative.xlsx>]
                               [--period "1 Sep 2026 to 16 Sep 2026"] [--days 15]
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas as pdfcanvas

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app import reports_api  # noqa: E402

# ---------------------------------------------------------------- geometry --
PAGE_W, PAGE_H = 595.276, 841.89
TITLE_H, PERIOD_H, HEAD_TOP_H, HEAD_BOT_H = 46.0, 34.0, 30.0, 32.0
HEAD_H = HEAD_TOP_H + HEAD_BOT_H
ROW_H, BOTTOM = 35.68, 21.97
INSET = 12.0
BASE_TITLE, BASE_PERIOD = 29.0, 21.0
LEAD = 9.775                       # second line of a two-line heading
RULE_H, RULE_V = 2.2, 1.6

F_TITLE, F_PERIOD, F_HEAD, F_ROW = 20.0, 13.0, 8.5, 11.0
BOLD, BOOK = "Helvetica-Bold", "Helvetica"

EDGES = [0.0, 127.4144069, 180.5837040, 230.9703617, 271.7617110, 324.4512427,
         366.9697475, 409.4882523, 478.2019490, 511.5334163, 544.8648837, PAGE_W]
SELL_COL, AVG_COL = 7, 8           # index of SELL-THROUGH, and of the CM/LM/TREND group

NAVY    = colors.Color(0.04, 0.16, 0.31)
NAVY_T  = colors.Color(0.043, 0.161, 0.31)
GOLD    = colors.Color(1.0, 0.74, 0.19)
GOLD_T  = colors.Color(1.0, 0.741, 0.192)
PAPER   = colors.Color(0.96, 0.97, 0.99)
INK     = colors.Color(0.059, 0.098, 0.176)
AMBER_F = colors.Color(1.0, 0.88, 0.7)
PINK_F  = colors.Color(1.0, 0.8, 0.82)
AMBER_T = colors.Color(0.902, 0.318, 0.0)
RED_T   = colors.Color(0.776, 0.157, 0.157)
AMBER_C = colors.Color(1.0, 0.718, 0.302)      # on a navy cluster row
PINK_C  = colors.Color(0.898, 0.451, 0.451)
UP      = colors.Color(0.247, 0.525, 0.0)
DOWN    = colors.Color(0.812, 0.075, 0.133)
UP_C    = colors.Color(0.565, 0.933, 0.565)
DOWN_C  = colors.Color(1.0, 0.714, 0.757)

# The grid. One line per row in whatever reads as a seam on that row's own
# colour: a light rule over white is a scratch over navy, and over gold it
# disappears altogether.
GRID    = colors.Color(0.855, 0.875, 0.906)    # a bond row
GRID_C  = colors.Color(0.133, 0.251, 0.498)    # a navy cluster row
GRID_T  = colors.Color(0.847, 0.573, 0.094)    # the gold total row

HEADINGS = ["BOND", "OPENING", "RECEIPT", "SALES", "CLOSING",
            ["STOCK", "NET"], ["STOCK", "NET %"], ["SELL-", "THROUGH %"]]
SUBHEADS = ["CM", "LM", "TREND"]
SELL_AMBER_AT = 40.0               # per cent, on the unrounded figure

MEASURES = [("Shop Opening Cases", "Shop Opening Bottles"),
            ("Shop In Cases", "Shop In Bottles"),
            ("Shop Out Cases", "Shop Out Bottles"),
            ("Shop Closing Cases", "Shop Closing Bottles")]


def w(text, font, size):
    return pdfmetrics.stringWidth(str(text), font, size)


def whole(value: float, places: int = 0):
    """Half up, the way the office rounds - never banker's.

    `places` is what the round-off switch moves: nought for whole cases, two
    for the figure as it actually stands.
    """
    step = Decimal(1).scaleb(-places)
    got = Decimal(str(value)).quantize(step, rounding=ROUND_HALF_UP)
    return int(got) if places == 0 else float(got)


# -------------------------------------------------------------------- data --
def bond_totals(path: Path) -> dict[str, list[float]]:
    """bond -> [opening, receipt, sales, closing] in cases, unrounded."""
    master = reports_api.load_master()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(rows)]
    ix = {h: i for i, h in enumerate(header)}
    for col in ["Shop Code", "Bottle Per Case"] + [c for p in MEASURES for c in p]:
        if col not in ix:
            raise SystemExit(f"ERROR: the raw file has no '{col}' column.")

    out: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    for r in rows:
        if r is None or r[ix["Shop Code"]] is None:
            continue
        bond = str(master.get(str(r[ix["Shop Code"]]).strip(), {}).get("bond", "")).strip().upper()
        if not bond:
            continue
        bpc = r[ix["Bottle Per Case"]] or 0
        for k, (c_col, b_col) in enumerate(MEASURES):
            out[bond][k] += float(r[ix[c_col]] or 0) + (
                float(r[ix[b_col]] or 0) / bpc if bpc else 0.0)
    wb.close()
    return out


def build_rows(now: dict, prev: dict, days: int, prev_days: int,
               only_cluster: int = 0, only_bond: str = "",
               days_by: dict | None = None, prev_days_by: dict | None = None,
               places: int = 0):
    """Bond rows interleaved with their cluster subtotals, then the grand total.

    A cluster or bond filter narrows what is shown AND what is totalled, so the
    TOTAL line always ties out to the rows above it rather than to the book.

    `days_by` lets a bond divide by its own trading days. The trade shuts
    together in practice, so the app passes nothing and every row divides by
    `days`; it is here for the day that stops being true.
    """
    days_by = days_by or {}
    prev_days_by = prev_days_by or {}

    def span(label):
        return days_by.get(label, days), prev_days_by.get(label, prev_days)
    clusters = reports_api.bond_clusters()
    want = (only_bond or "").strip().upper()
    rows, grand, grand_prev = [], [0.0] * 4, 0.0

    for cid in sorted(clusters):
        if only_cluster and cid != only_cluster:
            continue
        names = [b for b in clusters[cid] if b in now and (not want or b == want)]
        if not names:
            continue
        sub, sub_prev = [0.0] * 4, 0.0
        for bond in sorted(names):
            v = now[bond]
            p = prev.get(bond, [0.0] * 4)[2]
            rows.append(make_row(bond, v, p, *span(bond), "bond", places))
            for k in range(4):
                sub[k] += v[k]
            sub_prev += p
        rows.append(make_row(f"CLUSTER - {cid}", sub, sub_prev, days, prev_days,
                             "cluster", places))
        for k in range(4):
            grand[k] += sub[k]
        grand_prev += sub_prev

    # A bond with no cluster still has to appear, or the total stops tying out.
    placed = {r["label"] for r in rows}
    for bond in sorted(b for b in now if b not in placed and not b.startswith("CLUSTER")):
        if bond in {b for bs in clusters.values() for b in bs}:
            continue
        if only_cluster or (want and bond != want):
            continue
        v = now[bond]
        rows.append(make_row(bond, v, prev.get(bond, [0.0] * 4)[2], *span(bond),
                             "bond", places))
        for k in range(4):
            grand[k] += v[k]
        grand_prev += prev.get(bond, [0.0] * 4)[2]

    rows.append(make_row("TOTAL", grand, grand_prev, days, prev_days, "total", places))
    return rows


def make_row(label, v, prev_sales, days, prev_days, kind, places: int = 0):
    opening, receipt, sales, closing = v
    net = closing - opening
    avail = opening + receipt
    cm = sales / days if days else 0.0
    # prev_days is nought only when there is no window to compare against at
    # all. A bond that genuinely sold nothing last month still has days, and
    # still reads zero - which is a fact. None is the other thing: nobody
    # knows, because the file was never uploaded.
    lm = prev_sales / prev_days if prev_days else None
    return {
        "kind": kind, "label": label,
        "cells": [whole(opening, places), whole(receipt, places), whole(sales, places),
                  whole(closing, places), whole(net, places)],
        "net_pct": (net / opening * 100) if opening else None,
        "sell": (sales / avail * 100) if avail else None,
        "cm": whole(cm, places),
        "lm": None if lm is None else whole(lm, places),
        "trend": None if lm is None else cm - lm,
    }


# ------------------------------------------------------------------ render --
def shown(v) -> str:
    """A cell as it prints: whole when it was rounded, up to two places when not.

    A figure that lands on a whole case prints as one - "5.00" is 5 with
    noise on the end, and a column of them is noise all the way down.
    """
    if not isinstance(v, float):
        return str(v)
    out = f"{v:.2f}"
    return out.rstrip("0").rstrip(".") if "." in out else out


def pct(value, places: int = 0) -> str:
    # whole() hands back a float at two places, so 62 arrives as 62.0 - and
    # "62.0%" is the same noise the cells had.
    return "-" if value is None else f"{shown(whole(value, places))}%"


def triangle(c, cx, base, up, colour):
    c.setFillColor(colour)
    c.setStrokeColor(colour)
    p = c.beginPath()
    if up:
        p.moveTo(cx, base + 5.72); p.lineTo(cx - 2.86, base); p.lineTo(cx + 2.86, base)
    else:
        p.moveTo(cx, base); p.lineTo(cx - 2.86, base + 5.72); p.lineTo(cx + 2.86, base + 5.72)
    p.close()
    c.drawPath(p, stroke=1, fill=1)


def draw_head(c) -> float:
    y = PAGE_H
    c.setFillColor(NAVY)
    c.rect(0, y - TITLE_H, PAGE_W, TITLE_H, stroke=0, fill=1)
    c.setFillColor(GOLD_T); c.setFont(BOLD, F_TITLE)
    c.drawCentredString(PAGE_W / 2, y - BASE_TITLE, "K.S DISTILLERY")
    y -= TITLE_H
    return y


def draw_period(c, y, period) -> float:
    c.setFillColor(GOLD)
    c.rect(0, y - PERIOD_H, PAGE_W, PERIOD_H, stroke=0, fill=1)
    c.setFillColor(NAVY_T); c.setFont(BOLD, F_PERIOD)
    c.drawString(INSET, y - BASE_PERIOD, "SHOP SALES - ANALYSIS")
    c.drawRightString(PAGE_W - INSET, y - BASE_PERIOD, period)
    return y - PERIOD_H


def draw_columns(c, y) -> float:
    """The two-tier header: eight columns, the last spanning CM / LM / TREND."""
    bottom = y - HEAD_H
    split = y - HEAD_TOP_H
    # One band, not two. Two navy rectangles meeting at the split left a
    # hairline seam along their shared edge - the reader saw a thin line ruled
    # straight through BOND, OPENING, RECEIPT and the rest, because the
    # single-line headings are centred on the whole header and sit across it.
    c.setFillColor(NAVY)
    c.rect(0, bottom, PAGE_W, HEAD_H, stroke=0, fill=1)

    c.setFont(BOLD, F_HEAD)
    base = y - (HEAD_H / 2 + F_HEAD * 0.35)
    for i, head in enumerate(HEADINGS):
        left, right = EDGES[i], EDGES[i + 1]
        c.setFillColor(colors.white)
        if i == 0:
            c.drawString(INSET, base, head)
        elif isinstance(head, list):
            for n, line in enumerate(head):
                c.drawCentredString((left + right) / 2, base - n * LEAD, line)
        else:
            c.drawCentredString((left + right) / 2, base, head)

    group_left = EDGES[AVG_COL]
    c.setFillColor(GOLD_T)
    c.drawCentredString((group_left + PAGE_W) / 2,
                        split + (HEAD_TOP_H / 2 - F_HEAD * 0.35), "AVERAGE SALES / DAY")
    c.setFillColor(colors.white)
    for n, sub in enumerate(SUBHEADS):
        left, right = EDGES[AVG_COL + n], EDGES[AVG_COL + n + 1]
        c.drawCentredString((left + right) / 2,
                            bottom + (HEAD_BOT_H / 2 - F_HEAD * 0.35), sub)

    c.setStrokeColor(GOLD)
    c.setLineWidth(RULE_H)
    c.line(0, y - 1.1, PAGE_W, y - 1.1)
    c.line(0, bottom + 1.1, PAGE_W, bottom + 1.1)
    c.setLineWidth(RULE_V)
    for x in EDGES[1:AVG_COL + 1]:
        c.line(x, y, x, bottom)
    for x in EDGES[AVG_COL + 1:-1]:
        c.line(x, split, x, bottom)
    c.line(group_left, split, PAGE_W, split)
    return bottom


def draw_row(c, y, h, row, stripe):
    kind = row["kind"]
    cluster, total = kind == "cluster", kind == "total"
    base = y - (h / 2 + F_ROW * 0.35)
    font = BOLD if (cluster or total) else BOOK

    if cluster:
        bg, ink = NAVY, GOLD_T
    elif total:
        bg, ink = GOLD, NAVY_T
    else:
        bg, ink = (PAPER if stripe % 2 == 0 else colors.white), INK

    amber = row["sell"] is not None and row["sell"] >= SELL_AMBER_AT
    for i in range(len(EDGES) - 1):
        left, right = EDGES[i], EDGES[i + 1]
        fill = bg
        if i == SELL_COL and not cluster and not total:
            fill = AMBER_F if amber else PINK_F
        c.setFillColor(fill)
        c.rect(left, y - h, right - left, h, stroke=0, fill=1)

    # Ruled both ways, under the figures rather than over them.
    c.setStrokeColor(GRID_C if cluster else (GRID_T if total else GRID))
    c.setLineWidth(0.5)
    for x in EDGES[1:-1]:
        c.line(x, y - h, x, y)
    c.line(0, y - h, PAGE_W, y - h)

    c.setFillColor(ink)
    size = F_ROW
    while w(row["label"], font, size) > EDGES[1] - INSET - 6 and size > 7:
        size -= 0.5
    c.setFont(font, size)
    c.drawString(INSET, base, row["label"])

    c.setFont(font, F_ROW)
    # No thousands separators: the office prints 2658, not 2,658.
    values = [shown(v) for v in row["cells"]] + [pct(row["net_pct"])]
    for i, text in enumerate(values, start=1):
        c.setFillColor(ink)
        c.drawCentredString((EDGES[i] + EDGES[i + 1]) / 2, base, text)

    if cluster:
        sell_ink = AMBER_C if amber else PINK_C
    else:
        sell_ink = AMBER_T if amber else RED_T
    c.setFillColor(sell_ink)
    c.setFont(BOLD, F_ROW)
    c.drawCentredString((EDGES[SELL_COL] + EDGES[SELL_COL + 1]) / 2, base, pct(row["sell"]))

    c.setFillColor(ink)
    c.setFont(font, F_ROW)
    for n, key in enumerate(("cm", "lm")):
        v = row[key]
        c.drawCentredString((EDGES[AVG_COL + n] + EDGES[AVG_COL + n + 1]) / 2,
                            base, "-" if v is None else str(v))

    trend = row["trend"]
    if trend is None:
        # No comparison, so no arrow: an arrow is a claim about a direction.
        c.setFillColor(ink)
        c.setFont(font, F_ROW)
        c.drawCentredString((EDGES[-2] + EDGES[-1]) / 2, base, "-")
        return
    rising = trend >= 0
    if cluster:
        tint = UP_C if rising else DOWN_C
    else:
        tint = UP if rising else DOWN
    # The arrow is the sign. Printing it again as + or - beside the arrow
    # said the same thing twice.
    text = f"{whole(abs(trend))}"
    c.setFont(BOLD, F_ROW)
    span = 5.72 + 3.3 + w(text, BOLD, F_ROW)
    left = (EDGES[-2] + EDGES[-1]) / 2 - span / 2
    triangle(c, left + 2.86, base + 0.99, rising, tint)
    c.setFillColor(tint)
    c.drawString(left + 5.72 + 3.3, base, text)


def build(rows, period, out: Path) -> None:
    c = pdfcanvas.Canvas(str(out), pagesize=(PAGE_W, PAGE_H))
    c.setTitle("Shop Sales - Analysis")
    y = draw_columns(c, draw_period(c, draw_head(c), period))
    h = min(ROW_H, (y - BOTTOM) / max(1, len(rows)))

    for i, row in enumerate(rows):
        draw_row(c, y, h, row, i)
        if row["kind"] != "bond":
            c.setStrokeColor(NAVY if row["kind"] == "total" else GOLD)
            c.setLineWidth(RULE_V)
            c.line(0, y, PAGE_W, y)
            c.line(0, y - h, PAGE_W, y - h)
        y -= h
    c.showPage()
    c.save()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", default="")
    ap.add_argument("--totals", default="",
                    help="bond -> [opening, receipt, sales, closing] as JSON, when the "
                         "window was rebuilt from the daily exports rather than read "
                         "from one cumulative file")
    ap.add_argument("--prev-totals", default="")
    ap.add_argument("--prev", default="", help="last month's cumulative, same window")
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--period", default="")
    ap.add_argument("--days", type=int, default=0)
    ap.add_argument("--cluster", type=int, default=0)
    ap.add_argument("--bond", default="")
    ap.add_argument("--prev-days", type=int, default=0)
    # Trading days per bond, when the leave calendar says they differ.
    ap.add_argument("--days-by", default="")
    ap.add_argument("--prev-days-by", default="")
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--places", type=int, default=0,
                    help="decimal places on the figures; 0 (the default) is "
                         "whole cases, as the office's own book prints them")
    a = ap.parse_args()

    import json
    days = a.days or 1

    if a.totals:
        now = {k: list(v) for k, v in json.loads(Path(a.totals).read_text()).items()}
    else:
        raw = Path(a.raw)
        if not raw.is_file():
            print(f"ERROR: cumulative raw not found: {raw}")
            return 2
        now = bond_totals(raw)
    if not now:
        print("ERROR: nothing in this window maps to a bond - check Bond Mapping in Settings.")
        return 3

    if a.prev_totals:
        prev = {k: list(v) for k, v in json.loads(Path(a.prev_totals).read_text()).items()}
    elif a.prev and Path(a.prev).is_file():
        prev = bond_totals(Path(a.prev))
    else:
        prev = {}

    # Borrow this window's day count only when there IS a last month to divide.
    # Without one, prev_days fell back to days and every bond divided a total
    # of nothing by six: LM printed 0 and TREND claimed the whole month as a
    # gain, on a comparison that was never uploaded. Nought days is how
    # make_row hears "nobody knows".
    prev_days = a.prev_days or (days if prev else 0)

    def _span_map(path):
        if not path or not Path(path).is_file():
            return {}
        return {str(k).upper(): int(v)
                for k, v in json.loads(Path(path).read_text()).items() if int(v) > 0}

    rows = build_rows(now, prev, days, prev_days, a.cluster, a.bond,
                      _span_map(a.days_by), _span_map(a.prev_days_by),
                      places=max(0, min(2, a.places)))
    if not rows or len(rows) == 1:
        print("ERROR: nothing matches that filter.")
        return 4
    outdir = Path(a.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    out = outdir / "Shop Sales Analysis.pdf"
    build(rows, a.period or "Shop Sales Analysis", out)

    grand = rows[-1]
    print(f"{len(now)} bond(s) over {days} day(s)  ->  {out.name}")
    print(f"  opening {grand['cells'][0]:,}  receipt {grand['cells'][1]:,}  "
          f"sales {grand['cells'][2]:,}  closing {grand['cells'][3]:,}")
    if a.verify:
        for r in rows:
            print(f"  {r['label']:<16} {r['cells'][0]:>6} {r['cells'][1]:>6} "
                  f"{r['cells'][2]:>6} {r['cells'][3]:>6} {r['cells'][4]:>6} "
                  f"{pct(r['net_pct']):>6} {pct(r['sell']):>5} "
                  f"{r['cm']:>4} {('-' if r['lm'] is None else r['lm']):>4} "
                  + ("     -" if r['trend'] is None else f"{r['trend']:+6.2f}"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
