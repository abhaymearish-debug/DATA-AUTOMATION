#!/usr/bin/env python3
"""FULL row-by-row audit: every shop data line in every PDF vs the workbook cached values."""
import openpyxl, glob, os, re, subprocess
BASE="/sessions/great-serene-rubin/mnt/Claude"
OUTDIR=os.path.join(BASE,"PURCHASE INSTRUCTION","BOND PI DRILL-DOWNS")
WB=max([c for c in glob.glob(BASE+"/PURCHASE INSTRUCTION/PI INSIGHTS - * & * *.xlsx") if not os.path.basename(c).startswith('~$')], key=os.path.getmtime)
m=re.search(r'PI INSIGHTS\s*-\s*([A-Za-z]+)\s*&\s*([A-Za-z]+)\s*(\d{4})', os.path.basename(WB))
PRIOR,CUR,YEAR=m.group(1),m.group(2),m.group(3)
wb=openpyxl.load_workbook(WB, data_only=True)
def txt(ws,r,c): v=ws.cell(r,c).value; return "" if v is None else str(v)
def pdftext(p): return subprocess.run(['pdftotext','-layout',p,'-'],capture_output=True,text=True).stdout
def toks(s): return re.findall(r'[+\-]?\d[\d,]*(?:\.\d+)?%?|·', s.replace('−','-'))
def fint(v):
    if v in (None,""): return None
    iv=int(round(float(v))); return None if iv==0 else f'{iv:,}'
def frate(v):
    if v in (None,""): return None
    return f'{float(v):.1f}'
def fdelta(v):
    if v in (None,""): return None
    iv=int(round(float(v))); return None if iv==0 else (f'+{iv:,}' if iv>0 else f'-{abs(iv):,}')
def fcov(v):
    if v in (None,""): return None
    return f'{float(v):.2f}'

bonds=[n for n in wb.sheetnames if txt(wb[n],1,1).endswith("PURCHASE INSTRUCTION DRILL-DOWN")]
def markers(ws):
    co=None
    for r in range(1,ws.max_row+1):
        if txt(ws,r,1).startswith("CEILING & OFFTAKE"): co=r; break
    last=0
    for r in range(1,ws.max_row+1):
        if any(txt(ws,r,c) for c in range(1,15)): last=r
    return co,last

TOTC=0; TOTB=0; BADS=[]
for name in bonds:
    ws=wb[name]; co,last=markers(ws)
    pdf=os.path.join(OUTDIR, f"{name} — PI DRILL-DOWN ({CUR.title()} vs {PRIOR.title()} {YEAR}).pdf")
    lines=pdftext(pdf).splitlines()
    # index shop-title lines in the PDF
    tidx=[k for k,l in enumerate(lines) if re.match(r'^\s*\d{3,5}\s*·',l)]
    cells=0; bad=0
    # walk workbook shop blocks
    r=co
    while r<=last:
        if re.match(r'^\s*\d{3,5}\s*·', txt(ws,r,1)):
            title=txt(ws,r,1); code=title.split('·')[0].strip(); place=title.split('·')[1].strip().split('  ')[0][:6]
            # data rows until TOTAL
            data=[]
            rr=r+2  # skip title + column-header row
            while rr<=last and txt(ws,rr,1).strip()!='TOTAL' and not re.match(r'^\s*\d{3,5}\s*·',txt(ws,rr,1)):
                nm=txt(ws,rr,1).strip()
                if nm: data.append((nm,rr))
                rr+=1
            # locate this block's PDF section
            start=next((k for k in tidx if code in lines[k] and place in lines[k]), None)
            if start is None: bad+=1; BADS.append(f"{name}:{title} block missing"); r+=1; continue
            nxt=next((k for k in tidx if k>start), len(lines))
            sect=lines[start:nxt]
            for nm,rr in data:
                exp=[frate(ws.cell(rr,2).value),fint(ws.cell(rr,3).value),fint(ws.cell(rr,4).value),
                     fint(ws.cell(rr,5).value),fint(ws.cell(rr,6).value),fdelta(ws.cell(rr,7).value),fcov(ws.cell(rr,8).value)]
                exp=[e for e in exp if e is not None]
                ln=next((l for l in sect if nm in l), None)
                if ln is None: bad+=len(exp); BADS.append(f"{name}:{title}:{nm} line missing"); continue
                got=toks(ln.replace(nm,''))
                for e in exp:
                    cells+=1
                    if e not in got: bad+=1; BADS.append(f"{name}:{title}:{nm} {e}∉{got}")
        r+=1
    TOTC+=cells; TOTB+=bad
    print(f"[{'OK ' if bad==0 else '!!!'}] {name:15} shop data cells {cells-bad}/{cells}")

print("\n==== FULL SHOP-ROW AUDIT ====")
print(f"shop DATA cells checked: {TOTC} | mismatches: {TOTB}")
print("RESULT:", "ALL SHOP ROWS MATCH ✅" if TOTB==0 else f"REVIEW ({TOTB})")
for x in BADS[:12]: print("   -",x)
