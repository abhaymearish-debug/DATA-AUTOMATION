#!/usr/bin/env python3
"""
Apply captured commitments to the K.S. Distillery Commitment Tracker workbook.

Used AFTER the photo-OCR step: vision extracts the handwritten Target Cases
from each signed PDF photo and writes them to a CSV. This script reads that
CSV and writes the numbers into the right cells of every bond's tab.

CSV schema  (with header row):
    bond,shop_code,brand,pack,target_cases
    KOTTAYAM,105028,OLD PEARL NO.1 MATURED XXX RUM,500 ML,8.0
    KOTTAYAM,105028,K.S 99 LIFE TIME MATURED XXX RUM,500 ML,6.0
    ...

A single CSV can hold commitments for all 15 bonds — group by bond column.

Also takes a meeting date that gets written into cell C3 of every bond tab.

Usage:
    python3 apply_commitments.py \
        --csv captured_commitments.csv \
        --meeting-date 2026-05-16 \
        --next-meeting-date 2026-06-15

If --next-meeting-date is omitted, defaults to meeting_date + 30 days.

After running this, run refresh_commitment_tracker.py to populate the
Sold + Days + Status columns and update the live artifact.
"""

from __future__ import annotations

import argparse
import os
import csv
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

_THIS = os.path.abspath(__file__)
CLAUDE_DIR = Path(os.path.abspath(os.path.join(os.path.dirname(_THIS), "..", "..")))
AGING_DIR  = CLAUDE_DIR / "Commitment tracking"

DEFAULT_CYCLE_DAYS = 30   # fallback for next-meeting-date if not supplied


def find_tracker():
    candidates = sorted(AGING_DIR.glob("COMMITMENT TRACKER - *.xlsx"),
                        key=lambda p: p.stat().st_mtime, reverse=True)
    candidates = [p for p in candidates if not p.name.startswith("~$")]
    if not candidates:
        raise SystemExit(f"No COMMITMENT TRACKER workbook in {AGING_DIR}")
    return candidates[0]


def index_bond_rows(ws):
    """Build a lookup {(shop_code, brand, pack) : row_number} for one bond tab."""
    idx = {}
    for r in range(7, ws.max_row + 1):
        shop = ws.cell(row=r, column=2).value
        sev  = ws.cell(row=r, column=4).value
        brand = ws.cell(row=r, column=5).value
        pack  = ws.cell(row=r, column=6).value
        if shop is None or sev is None or brand is None or pack is None:
            continue
        key = (str(shop).strip(), str(brand).strip(), str(pack).strip())
        idx[key] = r
    return idx


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--csv", required=True, help="Captured commitments CSV")
    p.add_argument("--meeting-date", required=True,
                   help="Meeting date — YYYY-MM-DD or DD-Mmm-YYYY")
    p.add_argument("--next-meeting-date",
                   help="Next meeting date — YYYY-MM-DD or DD-Mmm-YYYY. "
                        f"Defaults to meeting-date + {DEFAULT_CYCLE_DAYS} days.")
    p.add_argument("--tracker", help="Override path to tracker workbook")
    args = p.parse_args()

    def _parse(s, label):
        for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        raise SystemExit(f"Cannot parse {label}: {s!r}. Use YYYY-MM-DD or DD-Mmm-YYYY.")

    md = _parse(args.meeting_date, "meeting date")
    if args.next_meeting_date:
        nmd = _parse(args.next_meeting_date, "next-meeting date")
    else:
        nmd = md + timedelta(days=DEFAULT_CYCLE_DAYS)

    tracker_path = Path(args.tracker) if args.tracker else find_tracker()
    csv_path = Path(args.csv)
    if not csv_path.exists():
        raise SystemExit(f"CSV not found: {csv_path}")

    print(f"  Tracker            : {tracker_path.name}")
    print(f"  CSV                : {csv_path.name}")
    print(f"  Meeting date       : {md.isoformat()}")
    print(f"  Next meeting date  : {nmd.isoformat()}"
          + ("" if args.next_meeting_date else f"  (default = meeting + {DEFAULT_CYCLE_DAYS}d)"))
    print()

    # Group CSV rows by bond
    rows_by_bond = defaultdict(list)
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        required = {"bond", "shop_code", "brand", "pack", "target_cases"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise SystemExit(f"CSV missing columns: {sorted(missing)}")
        for row in reader:
            bond = (row["bond"] or "").strip().upper()
            if not bond:
                continue
            rows_by_bond[bond].append(row)

    wb = openpyxl.load_workbook(tracker_path)
    total_applied = 0
    total_unmatched = []

    for bond, items in rows_by_bond.items():
        if bond not in wb.sheetnames:
            print(f"  WARN: bond '{bond}' not in workbook — skipping {len(items)} rows")
            continue
        ws = wb[bond]
        # Set the meeting date in C3 (top-left of merged C3:D3)
        ws["C3"].value = md
        ws["C3"].number_format = 'dd-mmm-yyyy'
        # Set the next-meeting date in G3 (top-left of merged G3:H3)
        ws["G3"].value = nmd
        ws["G3"].number_format = 'dd-mmm-yyyy'

        idx = index_bond_rows(ws)
        applied = 0
        unmatched = []
        for row in items:
            key = (row["shop_code"].strip(),
                   row["brand"].strip(),
                   row["pack"].strip())
            r = idx.get(key)
            if r is None:
                unmatched.append(key)
                continue
            try:
                target = float(row["target_cases"]) if row["target_cases"] else None
            except ValueError:
                target = None
            ws.cell(row=r, column=11).value = target  # col K = Target Cases
            applied += 1
        print(f"  {bond:<14} → applied {applied}/{len(items)} commits"
              f"{f'  ·  unmatched: {len(unmatched)}' if unmatched else ''}")
        total_applied += applied
        total_unmatched.extend((bond, k) for k in unmatched)

    wb.save(tracker_path)
    print()
    print(f"Wrote {total_applied} commitments to {tracker_path.name}")

    if total_unmatched:
        print()
        print(f"⚠ {len(total_unmatched)} rows could not be matched (printed below).")
        print("  Check Shop Code / Brand / Pack spelling against the bond tab.")
        for bond, key in total_unmatched[:20]:
            print(f"    {bond}: shop={key[0]} brand={key[1]} pack={key[2]}")
        if len(total_unmatched) > 20:
            print(f"    ...and {len(total_unmatched) - 20} more")
        sys.exit(1)

    print()
    print("Next step:  python3 refresh_commitment_tracker.py")


if __name__ == "__main__":
    main()
