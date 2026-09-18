"""
KSBC shop-code alias guard.

Some shops appear in the legacy KSBC region sheets under a 6-digit label that
no longer matches the 5-digit code used in the raw SupplierWiseShopSaleReport
or MASTER DATA. The pipeline must resolve those legacy labels to their
canonical code before aggregating, otherwise the region sheet silently zeroes
out the shop's numbers even though the raw has them. That's the MUKKAM bug
(111014 on KOZHIKODE sheet → 11014 in raw) that surfaced on the April 1-20
build.

This module is the single source of truth for that mapping.

Usage in the KSBC pipeline (pseudo-step):

    from .claude.scripts.ksbc_alias_guard import (
        canonical_code, apply_aliases_to_region,
        validate_no_zero_drops, validate_no_double_counts,
    )

    # 1. BEFORE aggregation — normalise any legacy region-sheet shop codes
    #    so the aggregation key matches the raw data.
    apply_aliases_to_region(region_df)     # rewrites region_df['Shop Code']

    # 2. AFTER aggregation — bail out if any aliased row was zeroed while the
    #    raw has non-zero sales for its canonical code.
    validate_no_zero_drops(region_df, raw_df)

    # 3. AFTER BOND PERFORMANCE is built — bail out if BP disagrees with the
    #    region sheet totals (symptom of double-counting an aliased shop via
    #    both legacy and canonical codes on the BP aggregation path).
    validate_no_double_counts(bp_aggregates, region_totals)

    # 4. AFTER DETAIL sheets are built — bail out if any pack rows don't
    #    sum to their brand row (symptom of build_rollups / update_detail_sheet
    #    contract drift — the 11 May 2026 regression on shop 105020).
    validate_no_pack_brand_mismatch(wb)

If any check fails, the pipeline must stop and alert — do NOT write the
workbook. The zero-drop check catches the original MUKKAM bug (111014 silently
zeroed). The double-count check catches its inverse (MUKKAM counted twice on
BP while appearing once on the region sheet — the April 1-20 regression).

Extending: add to ALIAS_MAP below. Keys = legacy region-sheet label, values =
canonical master/raw code.
"""

from __future__ import annotations
import re
from typing import Iterable, Mapping

_PACK_RE = re.compile(r"^\s*\d+\s*ML\s*$", re.IGNORECASE)
_EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


# --- Alias registry ---------------------------------------------------------
# Legacy KOZHIKODE region-sheet label 111014 → canonical KSBC/master code 11014.
# Add future mismatches here; each entry is documented with the date found.
ALIAS_MAP: dict[int, int] = {
    111014: 11014,   # MUKKAM — found 21 Apr 2026 on April 1-20 build.
}


def canonical_code(code: int) -> int:
    """Return the canonical shop code for `code`, resolving through ALIAS_MAP."""
    return ALIAS_MAP.get(code, code)


def apply_aliases_to_region(rows: Iterable[Mapping]) -> int:
    """
    Rewrite the 'Shop Code' field in-place on every row that matches an alias.
    Returns the number of rows rewritten. Accepts a list of dicts (region sheet
    representation) or anything that exposes __setitem__ on 'Shop Code'.
    """
    rewrites = 0
    for row in rows:
        code = row.get("Shop Code")
        if code in ALIAS_MAP:
            row["Shop Code"] = ALIAS_MAP[code]
            rewrites += 1
    return rewrites


def validate_no_zero_drops(region_rows: Iterable[Mapping],
                           raw_rows: Iterable[Mapping],
                           tolerance: float = 0.01) -> None:
    """
    Fail fast if an aliased shop's data is lost on the region sheet.

    Hardened 1 Jun 2026. The old version did ONLY check (1) below, and because
    the alias is applied UPSTREAM (no legacy-coded row survives to the region
    sheet), the loop's `if canon == code: continue` skipped every row and the
    guard passed vacuously — it never confirmed the canonical row actually
    carried the data. Two checks now:

      (1) a surviving legacy-coded row is ~0 while raw has data for its
          canonical code (alias was NOT applied);
      (2) for every alias whose canonical code has raw sales, the canonical row
          must EXIST on a region sheet AND carry at least the raw's sales
          (within tolerance) — catches the canonical row being dropped or
          silently zeroed (the original MUKKAM symptom).

    Uses an epsilon compare (not exact ==0) so case-conversion rounding can't
    slip a near-zero drop past the check.
    """
    # 16 Jun 2026 (BUG 8): materialise both iterables — region_rows is walked
    # twice below, so a generator argument would make check (1) vacuous.
    region_rows = list(region_rows)
    raw_rows = list(raw_rows)
    # Raw sales by canonical code (raw_rows may already be canonicalised;
    # re-applying canonical_code is idempotent).
    raw_sales: dict[int, float] = {}
    for r in raw_rows:
        code = r.get("Shop Code")
        if code is None:
            continue
        try:
            canon = canonical_code(int(code))
        except (TypeError, ValueError):
            continue
        sales = r.get("Sales (Cases)") or r.get("Sales") or 0
        raw_sales[canon] = raw_sales.get(canon, 0) + (sales or 0)

    # Region sales collapsed by canonical code, plus a presence set.
    region_sales_by_canon: dict[int, float] = {}
    region_has_canon: set[int] = set()
    for row in region_rows:
        code = row.get("Shop Code")
        if code in (None, ''):
            continue
        try:
            canon = canonical_code(int(code))
        except (TypeError, ValueError):
            continue
        region_has_canon.add(canon)
        region_sales_by_canon[canon] = region_sales_by_canon.get(canon, 0) + (row.get("Sales") or 0)

    bad: list[str] = []

    # (1) Legacy-coded row survived with ~0 sales while raw has data.
    for row in region_rows:
        code = row.get("Shop Code")
        if code in (None, ''):
            continue
        try:
            ci = int(code)
        except (TypeError, ValueError):
            continue
        canon = canonical_code(ci)
        if canon == ci:
            continue
        if abs(row.get("Sales") or 0) < tolerance and raw_sales.get(canon, 0) > tolerance:
            bad.append(f"  legacy {ci} → canonical {canon}: raw has "
                       f"{raw_sales[canon]:.2f} cases but its region row is 0 "
                       f"(alias not applied?)")

    # (2) For every alias with raw data, the canonical row must exist + carry it.
    for legacy, canon in ALIAS_MAP.items():
        rs = raw_sales.get(canon, 0)
        if rs <= tolerance:
            continue   # nothing in raw for this shop this period — nothing to assert
        if canon not in region_has_canon:
            bad.append(f"  canonical {canon} (alias of {legacy}): raw has "
                       f"{rs:.2f} cases but NO region row exists for it "
                       f"(shop silently dropped)")
        elif region_sales_by_canon.get(canon, 0) + tolerance < rs:
            bad.append(f"  canonical {canon} (alias of {legacy}): raw has "
                       f"{rs:.2f} cases but region shows only "
                       f"{region_sales_by_canon.get(canon, 0):.2f} (data lost)")

    if bad:
        raise AliasGuardError(
            "Alias guard tripped — aliased-shop data lost on the region sheet:\n"
            + "\n".join(bad)
            + "\nFix the alias application before writing the workbook."
        )


def validate_no_double_counts(bond_aggregates: Mapping[str, Mapping[str, float]],
                              region_totals: Mapping[str, Mapping[str, float]],
                              tolerance: float = 0.01) -> None:
    """
    Fail fast if any bond's BOND PERFORMANCE aggregate disagrees with its
    region-sheet total — symptom of an aliased shop being counted twice (once
    via legacy code, once via canonical) on the BP path while the region sheet
    has it once.

    `bond_aggregates`: bond → {Opening, Receipts, Sales, Closing} as written to
        the BOND PERFORMANCE sheet.
    `region_totals`: bond → {Opening, Receipts, Sales, Closing} computed by
        summing the region sheet's shop rows.

    This is the inverse of `validate_no_zero_drops` — that one catches DROPS
    (alias not applied → region sums to 0 while raw has data). This one catches
    DUPLICATES (alias applied to region but BP/dashboard pulled from a parallel
    path that didn't dedupe → BP > region by exactly the aliased shop's row).

    Raises AliasGuardError with the offending bonds so the caller can abort
    before overwriting the live workbook.
    """
    bad: list[str] = []
    for bond, bp in bond_aggregates.items():
        rt = region_totals.get(bond)
        if rt is None:
            continue
        for field in ("Opening", "Receipts", "Sales", "Closing"):
            bp_v = float(bp.get(field, 0) or 0)
            rt_v = float(rt.get(field, 0) or 0)
            if abs(bp_v - rt_v) > tolerance:
                bad.append(
                    f"  {bond}.{field}: BOND PERFORMANCE = {bp_v:.2f}, "
                    f"region sheet = {rt_v:.2f} (diff {bp_v - rt_v:+.2f})"
                )
    if bad:
        raise AliasGuardError(
            "Alias guard tripped — BOND PERFORMANCE disagrees with region "
            "sheet totals (likely double-count via legacy + canonical code):\n"
            + "\n".join(bad)
            + "\nDedupe the BP/dashboard aggregation path before writing."
        )


def build_region_dupe_report(region_code_lists: Mapping[str, Iterable[int]]
                             ) -> list[tuple[str, int, int]]:
    """Detect duplicate canonical shop codes WITHIN a single region sheet.

    `region_code_lists`: bond → iterable of the raw shop codes physically
        present on that region sheet (in row order, legacy codes included).

    Each code is resolved through ALIAS_MAP, then counted per bond. Any
    canonical code appearing more than once on the same sheet means the shop
    is double-counted in that region's TOTAL (the MUKKAM 111014 + 11014 shape).

    Returns a list of (bond, canonical_code, count) for every duplicate. This
    is independent of any aggregate value, so it catches a duplicate row even
    when the values happen to net out or one copy is zero.
    """
    report: list[tuple[str, int, int]] = []
    for bond, codes in region_code_lists.items():
        counts: dict[int, int] = {}
        for c in codes:
            if c in (None, ''):
                continue
            try:
                canon = canonical_code(int(c))
            except (TypeError, ValueError):
                continue
            counts[canon] = counts.get(canon, 0) + 1
        for canon, n in counts.items():
            if n > 1:
                report.append((bond, canon, n))
    return report


def validate_against_independent_truth(
        region_totals: Mapping[str, Mapping[str, float]],
        bond_aggregates: Mapping[str, Mapping[str, float]],
        bond_truth: Mapping[str, Mapping[str, float]],
        region_dupe_report: Iterable[tuple[str, int, int]] = (),
        tolerance: float = 0.01) -> None:
    """Fail fast if the region-sheet totals OR the BOND PERFORMANCE aggregates
    disagree with an INDEPENDENT ground truth computed straight from the raw
    COMBINED rows summed by canonical code.

    This is the C2 fix (1 Jun 2026). The older `validate_no_double_counts`
    compares BP against the region totals — but BP is WRITTEN FROM the region
    totals, so any inflation present in the region totals themselves (e.g. an
    aliased MUKKAM row counted twice on the region sheet) appears identically
    on both sides and passes vacuously. By comparing BOTH paths against a
    third, alias-collapsed source (`bond_truth`), a double-count introduced
    anywhere upstream of the region sheet is caught.

    `bond_truth`: bond → {Opening,Receipts,Sales,Closing} aggregated from the
        canonical-coded COMBINED roll-up (build_rollups → shop_totals) mapped
        to bonds via master. One contribution per canonical code, so it can
        never double-count an aliased shop.
    `region_dupe_report`: output of build_region_dupe_report — any duplicate
        canonical code on a region sheet is a hard failure on its own.
    """
    bad: list[str] = []

    for bond, dupes in {(b, c): n for (b, c, n) in region_dupe_report}.items():
        bond_name, canon = bond
        bad.append(
            f"  {bond_name}: shop {canon} appears {dupes}× on the region sheet "
            f"(duplicate row → double-counted in TOTAL)"
        )

    def _cmp(label, agg):
        for bond, tr in bond_truth.items():
            av = agg.get(bond)
            if av is None:
                bad.append(f"  {label}: bond '{bond}' missing (truth has it)")
                continue
            for field in ("Opening", "Receipts", "Sales", "Closing"):
                a = float(av.get(field, 0) or 0)
                t = float(tr.get(field, 0) or 0)
                if abs(a - t) > tolerance:
                    bad.append(
                        f"  {label} {bond}.{field} = {a:.2f}, independent "
                        f"truth = {t:.2f} (diff {a - t:+.2f})"
                    )

    _cmp("region sheet", region_totals)
    _cmp("BOND PERFORMANCE", bond_aggregates)

    if bad:
        raise AliasGuardError(
            "Alias guard tripped — region/BP totals disagree with the "
            "independent canonical-coded COMBINED truth (double-count or "
            "drop somewhere upstream of the region sheet):\n"
            + "\n".join(bad)
            + "\nThe region sheet and BOND PERFORMANCE must both reconcile to "
            "the COMBINED roll-up summed by canonical code before writing."
        )


def validate_no_pack_brand_mismatch(wb, tolerance: float = 0.01) -> None:
    """Fail fast if any DETAIL sheet's pack-row B-E don't sum to the brand-row
    B-E above them. Catches the class of bug where pack values silently fail
    to populate while brand totals are correct — the regression Abhay spotted
    11 May 2026 on shop 105020 (BLENDER'S CHOICE brand=0.46/4/1.46/3 but all
    pack rows stuck at 0 because of a dict-key shape mismatch in build_rollups).

    For each shop block on each DETAIL sheet, walks the rows in order:
      - Banner row (header) → skipped.
      - Col-header row → skipped.
      - Brand row (col A is a brand name) → starts a new accumulator.
      - Pack row (col A matches `\\d+ ML`) → added to the current brand's
        pack-sum accumulator.
      - TOTAL row → end of block; final brand's sums get checked.

    For each brand, asserts:
        brand_row.{Opening,Receipts,Sales,Closing} ≈ sum(pack rows).{...}
    within `tolerance`. Any mismatch raises AliasGuardError with the exact
    sheet + row + brand + delta so the caller can abort before overwriting
    the live workbook.

    `wb` is an openpyxl Workbook (loaded with data_only=True to resolve any
    SUM/IFERROR formulas to their cached numeric values).
    """
    bad: list[str] = []

    def _is_header(v):
        return isinstance(v, str) and ' — ' in v and 'Staff:' in v

    def _is_pack(v):
        return isinstance(v, str) and bool(_PACK_RE.match(v))

    def _n(v):
        return float(v) if isinstance(v, (int, float)) else 0.0

    def _err_token(v):
        """Return the Excel error token if v is one, else None. Hardened
        1 Jun 2026: previously _n() silently coerced #DIV/0!/#REF!/etc to 0.0,
        which could mask a real value or fabricate a false mismatch. Now an
        error cell is surfaced as its own failure."""
        if isinstance(v, str):
            for tok in _EXCEL_ERRORS:
                if tok in v:
                    return tok
        return None

    for sn in wb.sheetnames:
        if not sn.endswith(' DETAIL'):
            continue
        ws = wb[sn]

        # Find all shop block (header_row, total_row) pairs
        headers = [r for r in range(1, ws.max_row + 1)
                   if _is_header(ws.cell(r, 1).value)]
        blocks = []
        for i, h in enumerate(headers):
            end = headers[i + 1] - 1 if i + 1 < len(headers) else ws.max_row
            for r in range(h + 2, end + 1):
                if ws.cell(r, 1).value == 'TOTAL':
                    blocks.append((h, r))
                    break

        # Vacuous-pass guard (hardened 1 Jun 2026): a real DETAIL sheet always
        # has shop blocks. If the banner format ever drifts so `_is_header`
        # matches nothing (or no TOTAL is found), the whole sheet would silently
        # pass the reconciliation. A DETAIL sheet that has data rows but yields
        # zero blocks is itself a failure — surface it.
        has_data = any(isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.strip()
                       for r in range(1, min(ws.max_row, 60) + 1))
        if has_data and not blocks:
            bad.append(f"  {sn}: no shop blocks detected (banner/TOTAL format "
                       f"drift?) — reconciliation could not run on this sheet")
            continue

        for header_r, total_r in blocks:
            shop_label = ws.cell(header_r, 1).value
            current_brand = None
            current_brand_row = None
            current_brand_vals = None   # (O, R, S, C) on the brand row
            pack_sum = [0.0, 0.0, 0.0, 0.0]

            def _flush():
                if current_brand is None or current_brand_vals is None:
                    return
                for j, field in enumerate(('Opening', 'Receipts', 'Sales', 'Closing')):
                    delta = current_brand_vals[j] - pack_sum[j]
                    if abs(delta) > tolerance:
                        bad.append(
                            f"  {sn} r{current_brand_row} (shop {str(shop_label)[:30]}): "
                            f"brand '{current_brand}' {field} = {current_brand_vals[j]:.2f} "
                            f"but Σ(pack rows) = {pack_sum[j]:.2f} "
                            f"(delta {delta:+.2f})"
                        )

            for r in range(header_r + 2, total_r):
                a = ws.cell(r, 1).value
                if a is None or (isinstance(a, str) and a.strip() == ''):
                    continue
                # Surface any Excel error cell in B-E rather than coercing to 0.
                for c in range(2, 6):
                    tok = _err_token(ws.cell(r, c).value)
                    if tok:
                        bad.append(f"  {sn} r{r} (shop {str(shop_label)[:30]}): "
                                   f"col {chr(64 + c)} holds {tok} — recalc/formula "
                                   f"error, value can't be reconciled")
                if _is_pack(a):
                    if current_brand is None:
                        continue
                    for j, c in enumerate(range(2, 6)):
                        pack_sum[j] += _n(ws.cell(r, c).value)
                else:
                    # Reached a new brand row — flush the previous brand
                    _flush()
                    current_brand = a
                    current_brand_row = r
                    current_brand_vals = tuple(_n(ws.cell(r, c).value) for c in range(2, 6))
                    pack_sum = [0.0, 0.0, 0.0, 0.0]
            # End-of-block — flush the last brand
            _flush()

    if bad:
        raise AliasGuardError(
            "Alias guard tripped — DETAIL sheet pack rows don't reconcile to brand "
            "totals (likely build_rollups/update_detail_sheet contract drift):\n"
            + "\n".join(bad)
            + "\nInvestigate before overwriting the live workbook."
        )


class AliasGuardError(RuntimeError):
    pass


# --- Standalone diagnostic --------------------------------------------------
if __name__ == "__main__":
    # Ad-hoc check against an analysis workbook. Usage:
    #   python3 ksbc_alias_guard.py "<path to analysis xlsx>"
    import sys
    from openpyxl import load_workbook

    if len(sys.argv) != 2:
        print("usage: ksbc_alias_guard.py <analysis.xlsx>")
        sys.exit(2)

    wb = load_workbook(sys.argv[1], data_only=True)
    # Raw sheet naming convention — try the common ones
    raw = None
    for cand in ("APRIL 1-20 COMBINED", "COMBINED"):
        if cand in wb.sheetnames:
            raw = wb[cand]; break
    if raw is None:
        # fall back to the first sheet that has 'COMBINED' in its name
        for name in wb.sheetnames:
            if "COMBINED" in name.upper():
                raw = wb[name]; break
    if raw is None:
        print("No combined raw sheet found; cannot validate.")
        sys.exit(3)

    raw_rows: list[dict] = []
    headers = None
    for row in raw.iter_rows(values_only=True):
        if headers is None:
            headers = list(row); continue
        raw_rows.append(dict(zip(headers, row)))

    # Walk every region sheet (those with a DETAIL partner), gathering shop
    # rows into ONE network-wide list.
    region_names = [n for n in wb.sheetnames
                    if f"{n} DETAIL" in wb.sheetnames and n != "MASTER DATA"]
    any_fail = False
    region_totals: dict[str, dict[str, float]] = {}
    all_region_rows: list[dict] = []          # GLOBAL — see note below
    region_code_lists: dict[str, list] = {}   # per-sheet raw code lists (dupe check)
    for rn in region_names:
        ws = wb[rn]
        op = rc = sl = cl = 0.0
        header_seen = False
        codes = []
        for row in ws.iter_rows(values_only=True):
            if not header_seen:
                if row and row[0] == "Shop Code":
                    header_seen = True
                continue
            if not row or row[0] in (None, "TOTAL"):
                continue
            all_region_rows.append({"Shop Code": row[0], "Sales": row[5] or 0})
            codes.append(row[0])
            op += float(row[3] or 0); rc += float(row[4] or 0)
            sl += float(row[5] or 0); cl += float(row[6] or 0)
        region_totals[rn] = {"Opening": op, "Receipts": rc, "Sales": sl, "Closing": cl}
        region_code_lists[rn] = codes

    # Zero-drop check — GLOBAL scope (fixed 3 Jun 2026).
    # Previously this ran PER region sheet against the network-wide `raw_rows`,
    # so every sheet that legitimately does NOT contain an aliased shop
    # false-tripped: MUKKAM (canonical 11014) lives only on KOZHIKODE, yet the
    # raw is network-wide, so all 14 other region sheets reported a phantom
    # "shop silently dropped". The aliased shop's region row exists on exactly
    # ONE sheet, so the presence test must span ALL sheets at once — exactly how
    # the driver invokes it (ksbc_daily_update.py builds `all_region_rows`
    # across every region before calling validate_no_zero_drops once).
    try:
        validate_no_zero_drops(all_region_rows, raw_rows)
    except AliasGuardError as e:
        any_fail = True
        print(str(e))

    # Duplicate canonical code WITHIN a single region sheet (MUKKAM 111014+11014
    # double-row shape). Workbook-only, independent of any aggregate value.
    dupes = build_region_dupe_report(region_code_lists)
    if dupes:
        any_fail = True
        for bond, canon, n in dupes:
            print(f"[{bond}] canonical shop {canon} appears {n}× on the region "
                  f"sheet (duplicate row → double-counted in TOTAL)")

    # Cross-check BOND PERFORMANCE against region totals (double-count guard).
    # Skip cluster banner rows (col A starts with "Cluster") and the TOTAL row —
    # both carry =SUM formulas, not bond aggregates.
    if "BOND PERFORMANCE" in wb.sheetnames:
        bp = wb["BOND PERFORMANCE"]
        bp_aggs: dict[str, dict[str, float]] = {}
        for r in range(4, bp.max_row + 1):
            name = bp.cell(r, 1).value
            if not name or name == "TOTAL":
                continue
            if isinstance(name, str) and name.strip().lower().startswith("cluster"):
                continue
            bp_aggs[str(name).strip()] = {
                "Opening":  float(bp.cell(r, 2).value or 0),
                "Receipts": float(bp.cell(r, 3).value or 0),
                "Sales":    float(bp.cell(r, 4).value or 0),
                "Closing":  float(bp.cell(r, 5).value or 0),
            }
        try:
            validate_no_double_counts(bp_aggs, region_totals)
        except AliasGuardError as e:
            any_fail = True
            print(str(e))

    # Pack-vs-brand reconciliation across every DETAIL sheet
    try:
        validate_no_pack_brand_mismatch(wb)
    except AliasGuardError as e:
        any_fail = True
        print(str(e))

    # 16 Jun 2026 (BUG 9): validate_against_independent_truth needs the driver's
    # canonical COMBINED roll-up, which this standalone CLI does not build, so it
    # is NOT run here. A clean result below does NOT rule out an upstream
    # double-count present identically in region totals AND BOND PERFORMANCE.
    print("NOTE: validate_against_independent_truth is driver-only (needs COMBINED); "
          "not checked in this standalone run.")
    if not any_fail:
        print("Alias guard: OK — no zeroed aliases, no double-counts, no pack/brand mismatches.")
    else:
        sys.exit(1)
