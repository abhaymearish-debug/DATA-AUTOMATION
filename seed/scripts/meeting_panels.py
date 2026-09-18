#!/usr/bin/env python3
import openpyxl, json, time, sys, re
from copy import copy as _cpy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import range_boundaries, get_column_letter
from openpyxl.worksheet.views import Selection
from openpyxl.formatting.formatting import ConditionalFormattingList
t0=time.time()
SRC=sys.argv[1]; OUT=sys.argv[2]
import os
DATA=os.environ.get('BOND_DATA') or (os.path.join(os.environ['MEETING_OUT'],'bond_data.json') if os.environ.get('MEETING_OUT') else '')
if not DATA or not os.path.exists(DATA): sys.exit('BOND_DATA not set / bond_data.json missing')
data=json.load(open(DATA))
META=data.get('__meta__',{'cur':'MAY 2026','prev':'APRIL 2026'})
BONDS=[b for b in ['KOLLAM','KOZHIKODE','ATTINGAL','PALAKKAD','KOTTARAKARA','KANNUR','ALAPPUZHA','NEDUMANGAD','ALUVA','PERINTHALMANNA','THODUPUZHA','PATHANAMTHITTA','KOTTAYAM','THRISSUR','TRIPUNITHURA'] if b in data]
NAVY_DEEP='FF0D1B4A'; NAVY_MID='FF1A237E'; NAVY_SOFT='FF263F80'
GOLD='FFFFD700'; WHITE='FFFFFFFF'; PALEBLUE='FFE5EAF0'; LIGHTGREY='FFF3F4F6'
GREEN='FF2E7D32'; RED='FFC62828'; AMBER='FFE65100'; GREY='FF6B7280'
TINT1='FFD6E4F7'  # blue  - shop liquidation
TINT2='FFD7EDDC'  # green - secondary sales
TINTL='FFFCEDBE'  # cream - total liquidation
TINT3='FFE5D9F6'  # lavender - brand & pack
TINT4='FFFAD9C6'  # peach - field coverage
T_HI=('FF1565C0','FFBBDEFB'); T_BAL=(GREEN,'FFDCEDC8'); T_INV=(AMBER,'FFFFE0B2'); T_CRI=(RED,'FFFFCDD2')
def f(h): return PatternFill('solid',fgColor=h)
thin=Side(style='thin',color='FFD0D0D0'); gold_b=Side(style='thin',color=GOLD); gold_m=Side(style='medium',color=GOLD)
C=Alignment(horizontal='center',vertical='center',wrap_text=True)
Lf=Alignment(horizontal='left',vertical='center',wrap_text=True,indent=1)
def fillrange(ws,r,c1,c2,fill=None,border=False,botgold=False,topgold=False):
    for c in range(c1,c2+1):
        cell=ws.cell(row=r,column=c); bd={}
        if fill: cell.fill=fill
        if border: bd=dict(left=thin,right=thin,top=thin,bottom=thin)
        if botgold: bd['bottom']=gold_b
        if topgold: bd['top']=gold_b
        if bd: cell.border=Border(**bd)
def banner(ws,r,text,c1=1,c2=12,fill=NAVY_DEEP,font=GOLD,sz=13,h=30,bold=True):
    ws.merge_cells(start_row=r,start_column=c1,end_row=r,end_column=c2)
    fillrange(ws,r,c1,c2,fill=f(fill),botgold=True)
    cell=ws.cell(row=r,column=c1,value=text); cell.font=Font(bold=bold,size=sz,color=font); cell.alignment=C
    ws.row_dimensions[r].height=h
def hdrrow(ws,r,layout,labels,fill=NAVY_SOFT,h=22):
    ws.row_dimensions[r].height=h; c=1
    for span,lab in zip(layout,labels):
        c2=c+span-1
        if span>1: ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c2)
        cell=ws.cell(row=r,column=c,value=lab); cell.font=Font(bold=True,size=10,color=WHITE); cell.alignment=C
        fillrange(ws,r,c,c2,fill=f(fill),botgold=True); c=c2+1
def trow(ws,r,layout,values,opts,h=19,herotot=False,body=None,heroend=12):
    ws.row_dimensions[r].height=h; c=1
    for span,val,o in zip(layout,values,opts):
        c2=c+span-1
        if span>1: ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c2)
        cell=ws.cell(row=r,column=c,value=val)
        cell.font=Font(bold=o.get('bold',False),size=o.get('sz',10),color=o.get('color','FF1A1A1A'),italic=o.get('italic',False))
        cell.alignment=Lf if o.get('left') else C
        if o.get('fmt'): cell.number_format=o['fmt']
        fl=f(o['fill']) if o.get('fill') else (f(body) if body else None)
        if herotot:
            for cc in range(c,c2+1):
                hc=ws.cell(row=r,column=cc); hc.fill=f(NAVY_DEEP)
                hc.border=Border(top=gold_m,bottom=gold_m,left=(gold_m if cc==1 else None),right=(gold_m if cc==heroend else None))
        else:
            fillrange(ws,r,c,c2,fill=fl,border=True)
        c=c2+1
def tier_of(st):
    if st>=0.8: return ('🚀 High Performance',T_HI)
    if st>=0.6: return ('✅ Balanced',T_BAL)
    if st>=0.4: return ('⚠️ Inventory Heavy',T_INV)
    return ('🚫 Critical Overstock',T_CRI)
def arrow(d): return '▲' if d>0.5 else ('▼' if d<-0.5 else '→')
def pdelta(n,o): return (n-o)/o if o else None   # D7: blank %delta for new-from-zero, not a misleading +0%
def dcol(v): return GREEN if (v or 0)>=0 else RED
L6=[2,1,1,1,1,3]   # 9-col region layout (Shop Code / Field Staff / Closing-vs-Sales removed, Abhay 2 Jul 2026)
GAP=16
def gap(ws,r,h=GAP,e=12):
    ws.row_dimensions[r].height=h
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=e)
    return r+1

def slim_region(ws,staff='VACANT'):
    """Slim a bond region sheet 12->9 cols (Abhay, 2 Jul 2026): drop Shop Code (A), Field Staff (C),
    Closing vs Sales % (I). Regenerates the shifted formulas (delete_cols does not rewrite refs),
    translates the rating CF range (J->G; rules are literal cellIs-equal, no col refs inside),
    remaps column widths, and re-labels the TOTAL row (old label sat in deleted col A)."""
    if str(ws.cell(row=4,column=1).value).strip()!='Shop Code': return   # already slim
    tot=None
    for rr in range(5,ws.max_row+1):
        if str(ws.cell(row=rr,column=1).value).strip().upper()=='TOTAL': tot=rr; break
    if tot is None: return
    oldw={L:(ws.column_dimensions[L].width if L in ws.column_dimensions else None) for L in 'ABCDEFGHIJKL'}
    def newcol(c):
        if c in (1,3,9): return None
        return c-(1 if c>1 else 0)-(1 if c>3 else 0)-(1 if c>9 else 0)
    saved=[(str(rng.sqref),list(rng.rules)) for rng in ws.conditional_formatting]
    ws.conditional_formatting=ConditionalFormattingList()
    t1old=str(ws.cell(row=1,column=1).value or '')   # title text lives in col A -- delete_cols eats it (caught 2 Jul 2026)
    for m in list(ws.merged_cells.ranges):
        if m.min_row<=3:
            try:
                ws.unmerge_cells(str(m))
            except KeyError:
                # overlapping/nested merge: interior cells already removed by an earlier
                # unmerge (openpyxl aborts _clean_merge_range mid-loop). Range is already
                # dropped from merged_cells; clear any survivors so no MergedCell lingers.
                for _row,_col in list(m.cells)[1:]:
                    ws._cells.pop((_row,_col),None)
    ws.delete_cols(9); ws.delete_cols(3); ws.delete_cols(1)
    m=re.search(r'\(KSBC,?\s*([^)]*)\)', t1old)
    per=m.group(1).strip() if m else META['cur']
    # masthead v2 (Abhay 2 Jul 2026): tall gold title on deep navy + ONE meta line on lighter navy + slim gold divider bar
    for rr,h,fl in ((1,40,NAVY_DEEP),(2,22,'FF13244A'),(3,6,GOLD)):
        ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=9)
        ws.row_dimensions[rr].height=h
        fillrange(ws,rr,1,9,fill=f(fl))
    c=ws.cell(row=1,column=1,value=f'{ws.title}  \u2014  SHOP-WISE PERFORMANCE')
    c.font=Font(bold=True,size=18,color=GOLD); c.alignment=C
    c=ws.cell(row=2,column=1,value=f'="KSBC   \u00b7   {per}   \u00b7   Sales Executive: {staff}   \u00b7   All figures in cases   \u00b7   Shops: "&COUNTA(A5:A{tot-1})')
    c.font=Font(size=10,color='FFA9B8D6'); c.alignment=C
    ws.cell(row=3,column=1).value=None
    ws.cell(row=tot,column=1,value='TOTAL')
    RT=('=IF(AND(B{r}=0,C{r}=0,D{r}=0,E{r}=0),"\u2014 No activity",IF(F{r}>=0.8,"\U0001F680 High Performance",'
        'IF(F{r}>=0.6,"\u2705 Balanced",IF(F{r}>=0.4,"\u26A0\uFE0F Inventory Heavy","\U0001F6AB Critical Overstock"))))')
    for rr in range(5,tot):
        if ws.cell(row=rr,column=1).value in (None,''): continue
        ws.cell(row=rr,column=6,value=f'=IFERROR(D{rr}/(B{rr}+C{rr}),0)')
        ws.cell(row=rr,column=7,value=RT.format(r=rr))
    for cidx,L in ((2,'B'),(3,'C'),(4,'D'),(5,'E')):
        ws.cell(row=tot,column=cidx,value=f'=SUM({L}5:{L}{tot-1})')
    ws.cell(row=tot,column=6,value=f'=IFERROR(D{tot}/(B{tot}+C{tot}),0)')
    ws.cell(row=tot,column=7,value=RT.format(r=tot))
    ws.cell(row=tot,column=8,value=f'=SUM(H5:H{tot-1})')
    for sq,rules in saved:
        parts=[]
        for piece in sq.split():
            c1,r1,c2,r2=range_boundaries(piece)
            n1=newcol(c1); n2=newcol(c2)
            if n1 is None and n2 is None: continue
            if n1 is None: n1=newcol(c1+1) or 1
            if n2 is None: n2=newcol(c2-1) or n1
            parts.append(f'{get_column_letter(n1)}{r1}:{get_column_letter(n2)}{r2}')
        if not parts: continue
        for rule in rules:
            ws.conditional_formatting.add(' '.join(parts),rule)
    for newL,oldL in (('A','B'),('B','D'),('C','E'),('D','F'),('E','G'),('F','H'),('G','J'),('H','K'),('I','L')):
        if oldw.get(oldL): ws.column_dimensions[newL].width=oldw[oldL]
        # KSBC region sheets HIDE Shop Code (A) + Field Staff (C) (16 Jul 2026). delete_cols moves
        # cells but NOT column_dimensions, so those hidden flags stay on physical A and C -- which
        # here hold Shop Name and Receipts. Every surviving column must be visible.
        ws.column_dimensions[newL].hidden=False
    for L in ('J','K','L'):
        if L in ws.column_dimensions: del ws.column_dimensions[L]
    ws.freeze_panes=None   # no freeze panes on bond sheets (Abhay 2 Jul 2026)
    ws.sheet_view.selection=[Selection(activeCell='A1', sqref='A1')]   # drop stale pane-selections (Excel 'Repaired' fix)
    return per

def fed_block(wsd,b,d):
    """CONSUMER-FED OUTLETS roster on each '<BOND> DETAIL' sheet (Abhay, 2 Jul 2026).
    Placed TOP-RIGHT in cols I:O, starting on the first shop block's banner row (Abhay: 'next to the 1st
    shop block', not at the bottom). Row heights are OWNED by the left-side shop blocks -- never set them here.
    Idempotent: removes a legacy bottom block (col-A banner) and clears/unmerges any previous I:O block."""
    rows=d.get('fed_outlets') or []
    # legacy bottom-placed block (col A banner) -- delete those rows
    mb=None
    for rr in range(1,wsd.max_row+1):
        v=wsd.cell(row=rr,column=1).value
        if v and str(v).startswith('\U0001F3EA'): mb=rr; break
    if mb:
        for rng in list(wsd.merged_cells.ranges):
            c1,r1,c2,r2=range_boundaries(str(rng))
            if r2>=mb-1: wsd.unmerge_cells(str(rng))
        wsd.delete_rows(mb-1, wsd.max_row-(mb-1)+1)
    # clear any prior top-right block
    for rng in list(wsd.merged_cells.ranges):
        c1,r1,c2,r2=range_boundaries(str(rng))
        if c2>=9: wsd.unmerge_cells(str(rng))
    _nofill=PatternFill(fill_type=None)
    for rr in range(1,61):
        for cc in range(9,16):
            cell=wsd.cell(row=rr,column=cc)
            cell.value=None; cell.fill=_nofill; cell.border=Border(); cell.font=Font()
    # anchor on the first shop banner row
    start=4
    for rr in range(2,min(30,wsd.max_row)+1):
        v=wsd.cell(row=rr,column=1).value
        if v and 'Staff:' in str(v): start=rr; break
    C1,C2=9,15   # I..O
    r=start
    n_act=sum(1 for x in rows if str(x.get('status','')).lower()=='active')
    wsd.merge_cells(start_row=r,start_column=C1,end_row=r,end_column=C2)
    fillrange(wsd,r,C1,C2,fill=f(NAVY_MID),botgold=True)
    c=wsd.cell(row=r,column=C1,value=f'\U0001F3EA  CONSUMER-FED OUTLETS \u2014 {b}   \u00b7   {n_act} active   \u00b7   {META["curm"]} vs {META["prevm"]}')
    c.font=Font(bold=True,size=12,color=WHITE); c.alignment=C; r+=1
    def hrow(rr,labels):
        for j,lab in enumerate(labels):
            cell=wsd.cell(row=rr,column=C1+j,value=lab)
            cell.font=Font(bold=True,size=10,color=WHITE); cell.alignment=C
        fillrange(wsd,rr,C1,C2,fill=f(NAVY_SOFT),botgold=True)
    def drow(rr,vals,opts):
        for j,(val,o) in enumerate(zip(vals,opts)):
            cell=wsd.cell(row=rr,column=C1+j,value=val)
            cell.font=Font(bold=o.get('bold',False),size=o.get('sz',10),color=o.get('color','FF1A1A1A'),italic=o.get('italic',False))
            cell.alignment=Lf if o.get('left') else C
            if o.get('fmt'): cell.number_format=o['fmt']
        fillrange(wsd,rr,C1,C2,fill=(f(opts[0]['fill']) if opts[0].get('fill') else None),border=True)
    if not rows:
        wsd.merge_cells(start_row=r,start_column=C1,end_row=r,end_column=C2)
        fillrange(wsd,r,C1,C2,fill=f(LIGHTGREY),border=True)
        c=wsd.cell(row=r,column=C1,value='No Consumer-fed outlets under this bond.')
        c.font=Font(size=10,italic=True,color=GREY); c.alignment=C
        return
    hrow(r,['Outlet',META['prevm'],META['curm'],'\u0394','% \u0394','% of Bond FED','Signal']); r+=1
    jt=sum(x['jun'] for x in rows); mt=sum(x['may'] for x in rows)
    for x in rows:
        mj=x['jun']; mm=x['may']; dv=mj-mm; pd=pdelta(mj,mm)
        zero=(mj==0 and mm==0)
        nm=f"{x['code']} \u2014 {x['name']}"
        closed=str(x.get('status','')).lower()!='active'
        if closed: nm+='   \u00b7 '+str(x.get('status',''))
        if mm==0 and mj>0: sig,sc=('NEW','FFB8860B')
        elif mm>0 and mj==0: sig,sc=('STOPPED',RED)
        elif zero: sig,sc=('\u00b7',GREY)
        else: sig,sc=(arrow(dv),dcol(dv))
        base=LIGHTGREY if zero else None
        fcol=GREY if zero else ('FFC62828' if closed else 'FF1A1A1A')
        drow(r,[nm,mm,mj,(dv if not zero else ''),(pd if (pd is not None and not zero) else ''),((mj/jt) if (jt and mj>0) else ''),sig],
          [dict(left=True,sz=9,italic=zero,color=fcol,fill=base),dict(fmt='#,##0',italic=zero,color=fcol),
           dict(fmt='#,##0',bold=not zero,italic=zero,color=fcol),dict(fmt='+#,##0;-#,##0',color=(GREY if zero else dcol(dv))),
           dict(fmt='+0%;-0%',color=(GREY if zero else dcol(pd))),dict(fmt='0%'),dict(bold=True,sz=10,color=sc)]); r+=1
    dvT=jt-mt; pdT=pdelta(jt,mt); DK='FF374151'; BR=('FF81C784' if dvT>=0 else 'FFE57373')
    drow(r,[f'TOTAL \u00b7 {len(rows)} outlets',mt,jt,dvT,(pdT if pdT is not None else ''),(1.0 if jt else 0),arrow(dvT)],
      [dict(left=True,bold=True,sz=11,color=WHITE,fill=DK),dict(fmt='#,##0',bold=True,color=WHITE),
       dict(fmt='#,##0',bold=True,color=WHITE),dict(fmt='+#,##0;-#,##0',bold=True,color=BR),
       dict(fmt='+0%;-0%',bold=True,color=BR),dict(fmt='0%',bold=True,color=WHITE),dict(bold=True,sz=11,color=BR)])
    for L,w in (('I',36),('J',9),('K',9),('L',9),('M',9),('N',12),('O',11)):
        wsd.column_dimensions[L].width=w
    wsd.sheet_view.showGridLines=False

def detail_masthead(wsd,b,staff,per):
    """DETAIL-sheet masthead (Abhay 2 Jul 2026) -- same design as the region masthead v2, full width A:O
    so it spans both the shop blocks and the top-right Consumer-fed roster."""
    for m in list(wsd.merged_cells.ranges):
        c1,r1,c2,r2=range_boundaries(str(m))
        if r1<=3: wsd.unmerge_cells(str(m))
    for rr,h,fl in ((1,40,NAVY_DEEP),(2,22,'FF13244A'),(3,6,GOLD)):
        wsd.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=15)
        wsd.row_dimensions[rr].height=h
        fillrange(wsd,rr,1,15,fill=f(fl))
    c=wsd.cell(row=1,column=1,value=f'{b}  \u2014  SHOP DETAILS')
    c.font=Font(bold=True,size=18,color=GOLD); c.alignment=C
    c=wsd.cell(row=2,column=1,value=f'KSBC   \u00b7   {per}   \u00b7   Sales Executive: {staff}   \u00b7   Brand & Pack breakdown by shop   \u00b7   All figures in cases')
    c.font=Font(size=10,color='FFA9B8D6'); c.alignment=C
    wsd.cell(row=3,column=1).value=None
    wsd.freeze_panes=None   # no freeze panes on DETAIL sheets (Abhay 2 Jul 2026)
    wsd.sheet_view.selection=[Selection(activeCell='A1', sqref='A1')]   # drop stale pane-selections (Excel 'Repaired' fix)

# ---- DAILY TREND relocated to the right of the shop table (Abhay, 1 Aug 2026) ----
# The region sheets inherit a 'DAILY TREND + CUMULATIVE RACE' block from the KSBC workbook,
# stacked BELOW the shop TOTAL. Abhay wants it BESIDE the shop table instead. This also
# repairs the chart: slim_region deletes 3 columns, which moved the cumulative helper series
# from M/N to J/K, but openpyxl does not rewrite chart refs -- so the race chart was pointing
# at empty cells and plotted NOTHING. Series are repointed here.
TR_GUT, TR_DAY, TR_CUR, TR_PRV, TR_DLT, TR_H1, TR_H2, TR_CHART = 10,11,12,13,14,15,16,17
TR_TOP = 4          # trend banner sits level with the shop table's column-header row
TR_MAST_END = 24    # masthead is widened to span the shop table + trend + chart

def move_trend_right(ws):
    tot=None
    for rr in range(5,ws.max_row+1):
        if str(ws.cell(row=rr,column=1).value).strip().upper()=='TOTAL': tot=rr; break
    if tot is None: return
    # The banner TEXT is BLANK on meeting sheets: it lived in col A of the KSBC sheet and
    # slim_region's delete_cols(1) ate it (same class as the 2 Jul title-row regression), so
    # JUNE/JULY shipped this table under an unlabelled navy band. Anchor on the header row
    # instead, and rebuild the banner text at the new position below.
    hdr=None
    for rr in range(tot+1,ws.max_row+1):
        if (str(ws.cell(row=rr,column=1).value).strip()=='Day'
                and str(ws.cell(row=rr,column=4).value).strip().endswith('cs')):
            hdr=rr; break
    if hdr is None: return         # already moved / nothing to do (idempotent)
    tb=hdr-1
    tr=None
    for rr in range(hdr+1,ws.max_row+1):
        if str(ws.cell(row=rr,column=1).value).strip().upper()=='TOTAL': tr=rr; break
    if tr is None: return
    avg=tr+1; n=avg-tb+1
    # ---- capture values + styles + heights before anything moves ----
    SRC2DST=((1,TR_DAY),(2,TR_CUR),(3,TR_PRV),(4,TR_DLT),(10,TR_H1),(11,TR_H2))
    cap={}; hts={}
    for off in range(n):
        rr=tb+off
        for sc,dc in SRC2DST:
            c=ws.cell(row=rr,column=sc)
            cap[(off,dc)]=(c.value,_cpy(c._style),c.number_format)
        d=ws.row_dimensions.get(rr); hts[off]=None if d is None else d.height
    bars={}
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            if rule.type=='dataBar':
                col=str(rng.sqref).strip()[:1]
                if col=='B': bars[TR_CUR]=rule
                elif col=='C': bars[TR_PRV]=rule
    chart=ws._charts[0] if ws._charts else None
    # ---- remove the old block (banner..avg plus the spacer row above it) ----
    for m in list(ws.merged_cells.ranges):
        c1,r1,c2,r2=range_boundaries(str(m))
        if r2>=tb-1 and r1<=avg: ws.unmerge_cells(str(m))
    ws.delete_rows(tb-1, avg-(tb-1)+1)
    # delete_rows shifts the KSBC SALES MIX chrome up into the deep-dive's landing zone, so
    # banner() would hit a MergedCell at TOTAL+2. Clear cols A-I below the shop TOTAL; the
    # deep-dive owns that area and the relocated trend lives in cols J+ (untouched here).
    for m in list(ws.merged_cells.ranges):
        c1,r1,c2,r2=range_boundaries(str(m))
        if r1>tot and c1<=9:
            try: ws.unmerge_cells(str(m))
            except KeyError:
                for _row,_col in list(m.cells)[1:]: ws._cells.pop((_row,_col),None)
    # delete_rows MOVES cell objects without re-typing them, so MergedCell instances from the
    # removed merges land in the deep-dive's zone as orphans (no merge range covers them, but
    # .value is still read-only). Pop them so ws.cell() rebuilds a real Cell.
    for rr in range(tot+1, ws.max_row+1):
        for cc in range(1,10):
            c=ws._cells.get((rr,cc))
            if c is None: continue
            if c.__class__.__name__=='MergedCell': ws._cells.pop((rr,cc),None)
            elif c.value is not None: c.value=None
    # drop every dataBar rule: the trend pair is re-added below, and the KSBC SALES MIX
    # leftovers (B/E) would otherwise paint stray bars over the deep-dive tables.
    keep=[(str(r.sqref),[x for x in r.rules if x.type!='dataBar']) for r in ws.conditional_formatting]
    ws.conditional_formatting=ConditionalFormattingList()
    for sq,rules in keep:
        for x in rules: ws.conditional_formatting.add(sq,x)
    # ---- write the block in its new home ----
    for off in range(n):
        rr=TR_TOP+off
        for _,dc in SRC2DST:
            v,st,nf=cap[(off,dc)]
            c=ws.cell(row=rr,column=dc,value=v); c._style=st; c.number_format=nf
        if rr>tot and hts[off]: ws.row_dimensions[rr].height=hts[off]   # never restyle the shop table's rows
    ws.merge_cells(start_row=TR_TOP,start_column=TR_DAY,end_row=TR_TOP,end_column=TR_DLT)
    fillrange(ws,TR_TOP,TR_DAY,TR_DLT,fill=f('FF263F80'))          # merge first, then style
    bc=ws.cell(row=TR_TOP,column=TR_DAY,
               value='\U0001F4C5 DAILY TREND \u2014 %s vs %s (same day)'
                     % (META['cur'].split()[0], META['prev'].split()[0]))
    bc.font=Font(bold=True,size=12,color='FFFFD54F'); bc.alignment=C
    if (ws.row_dimensions[TR_TOP].height or 0) < 26: ws.row_dimensions[TR_TOP].height=26
    nd0,nd1=TR_TOP+2,TR_TOP+n-3        # first/last day row in the new position
    for dc,w in ((TR_GUT,2.5),(TR_DAY,13),(TR_CUR,9.5),(TR_PRV,9.5),(TR_DLT,9.5),(TR_H1,9),(TR_H2,9)):
        ws.column_dimensions[get_column_letter(dc)].width=w
    for dc in (TR_CUR,TR_PRV):
        if dc in bars:
            L=get_column_letter(dc); ws.conditional_formatting.add(f'{L}{nd0}:{L}{nd1}',bars[dc])
    # ---- re-anchor the chart and repoint its series (the M/N -> J/K slim bug) ----
    if chart is not None:
        q=f"'{ws.title}'"
        cats=f'{q}!${get_column_letter(TR_DAY)}${nd0}:${get_column_letter(TR_DAY)}${nd1}'
        for s,hc in zip(chart.series,(TR_H1,TR_H2)):
            L=get_column_letter(hc)
            if s.val is not None and s.val.numRef is not None: s.val.numRef.f=f'{q}!${L}${nd0}:${L}${nd1}'
            if s.cat is not None:
                if s.cat.numRef is not None: s.cat.numRef.f=cats
                if getattr(s.cat,'strRef',None) is not None: s.cat.strRef.f=cats
        chart.anchor=f'{get_column_letter(TR_CHART)}{TR_TOP}'
        chart.width, chart.height = 12.2, 7.4
    # ---- widen the masthead so the band spans the whole sheet ----
    for m in list(ws.merged_cells.ranges):
        c1,r1,c2,r2=range_boundaries(str(m))
        if r1<=3 and r2<=3: ws.unmerge_cells(str(m))
    for rr,fl in ((1,NAVY_DEEP),(2,'FF13244A'),(3,GOLD)):
        ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=TR_MAST_END)
        fillrange(ws,rr,1,TR_MAST_END,fill=f(fl))


wb=openpyxl.load_workbook(SRC)
print('loaded',round(time.time()-t0,1),'s')
for b in BONDS:
    ws=wb[b]; ws.sheet_view.showGridLines=False; d=data[b]; tm=d['tert_may']; ta=d['tert_apr']; sm=d['sec_may']; sa=d['sec_apr']
    _staff=(d.get('vis') or {}).get('staff') or 'VACANT'
    _per=slim_region(ws, staff=_staff) or META['cur']
    move_trend_right(ws)      # DAILY TREND + race chart go beside the shop table (1 Aug 2026)
    # delete existing deep-dive (master banner -> end)
    mb=None
    for rr in range(1,ws.max_row+1):
        v=ws.cell(row=rr,column=1).value
        if v and str(v).startswith('📊  BOND DEEP-DIVE'): mb=rr; break
    if mb:
        for rng in list(ws.merged_cells.ranges):
            c1,r1,c2,r2=range_boundaries(str(rng))
            if r2>=mb-1: ws.unmerge_cells(str(rng))
        ws.delete_rows(mb-1, ws.max_row-(mb-1)+1)  # also remove the spacer above
    # find shop TOTAL row
    totrow=ws.max_row
    for rr in range(1,ws.max_row+1):
        if str(ws.cell(row=rr,column=1).value).strip().upper()=='TOTAL': totrow=rr
    r=totrow+2
    banner(ws,r,f'📊  BOND DEEP-DIVE — {b}     ·     '+META['cur']+'  vs  '+META['prev'],c2=9,fill=NAVY_DEEP,sz=14,h=34); r+=1
    r=gap(ws,r,10,9)
    # ① SHOP LIQUIDATION
    banner(ws,r,'SHOP LIQUIDATION (KSBC)',c2=9,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    hdrrow(ws,r,L6,['Metric',META['prevm'],META['curm'],'Δ','% Δ','Signal / Rating']); r+=1
    ml,(mfg,mbg)=tier_of(tm['st'])
    rowsT=[('Closing',ta.get('C',0),tm['C'],'#,##0',False,None),
      ('Active shops',ta.get('nshops',tm['nshops']),tm['nshops'],'0',False,None),   # prev-month count, master-rebucketed (2 Jul 2026)
      ('Opening',ta.get('O',0),tm['O'],'#,##0',False,None),
      ('Receipts',ta.get('R',0),tm['R'],'#,##0',False,None),
      ('Sales (cases)',ta.get('S',0),tm['S'],'#,##0',True,PALEBLUE),
      ('Sell-Through %',ta.get('st',0),tm['st'],'0%',True,PALEBLUE)]
    for nm,av,mv,fmt,em,fillc in rowsT:
        if fmt=='0%':
            dv=mv-av
            trow(ws,r,L6,[nm,av,mv,dv,'',ml],[dict(left=True,bold=em,fill=fillc),dict(fmt='0%',fill=fillc),
              dict(fmt='0%',bold=True,fill=fillc),dict(fmt='+0%;-0%',color=dcol(dv),fill=fillc),dict(fill=fillc),
              dict(bold=True,sz=9,color=mfg,fill=mbg)],body=TINT1)
        elif nm=='Active shops':
            trow(ws,r,L6,[nm,av,mv,'—','—','—'],[dict(left=True),dict(fmt='0'),dict(fmt='0',bold=True),dict(),dict(),dict(color=GREY)],body=TINT1)
        else:
            dv=mv-av; pd=pdelta(mv,av)
            trow(ws,r,L6,[nm,av,mv,dv,pd,arrow(dv)],[dict(left=True,bold=em,fill=fillc),dict(fmt='#,##0',fill=fillc),
              dict(fmt='#,##0',bold=em,fill=fillc),dict(fmt='+#,##0;-#,##0',color=dcol(dv),fill=fillc),
              dict(fmt='+0%;-0%',color=dcol(pd),fill=fillc),dict(bold=True,sz=11,color=dcol(dv),fill=fillc)],body=TINT1)
        r+=1
    r=gap(ws,r,e=9)
    # ② SECONDARY SALES (no liquidation row here)
    banner(ws,r,'SECONDARY SALES',c2=9,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    hdrrow(ws,r,L6,['Channel',META['prevm'],META['curm'],'Δ','% Δ','Share of Secondary']); r+=1
    sect_m=sm.get('total',0)
    rowsS=[('KSBC dispatch',sa.get('ksbc',0),sm.get('ksbc',0),sm.get('ksbc',0)/sect_m if sect_m else 0,None),
      ('Consumer-fed (invoice)',sa.get('fed',0),sm.get('fed',0),sm.get('fed',0)/sect_m if sect_m else 0,None),
      ('BAR (invoice)',sa.get('bar',0),sm.get('bar',0),sm.get('bar',0)/sect_m if sect_m else 0,None),
      ('Secondary TOTAL',sa.get('total',0),sm.get('total',0),1.0,PALEBLUE)]
    for nm,av,mv,share,fillc in rowsS:
        dv=mv-av; pd=pdelta(mv,av); em=fillc is not None
        trow(ws,r,L6,[nm,av,mv,dv,pd,share],[dict(left=True,bold=em,fill=fillc),dict(fmt='#,##0',fill=fillc),
          dict(fmt='#,##0',bold=em,fill=fillc),dict(fmt='+#,##0;-#,##0',color=dcol(dv),fill=fillc),
          dict(fmt='+0%;-0%',color=dcol(pd),fill=fillc),dict(fmt='0%',fill=fillc)],body=TINT2); r+=1
    r+=1
    # ✦ TOTAL LIQUIDATION rollup block (its own section)
    banner(ws,r,'✦  TOTAL LIQUIDATION — SHOP LIQUIDATION + INVOICE SALES',c2=9,fill=NAVY_DEEP,font=GOLD,sz=11,h=26); r+=1
    hdrrow(ws,r,L6,['Component',META['prevm'],META['curm'],'Δ','% Δ','% of Liquidation'],fill=NAVY_SOFT); r+=1
    shop_a=ta.get('S',0); shop_m=tm['S']; fed_a=sa.get('fed',0); fed_m=sm.get('fed',0); bar_a=sa.get('bar',0); bar_m=sm.get('bar',0)
    liq_a=shop_a+fed_a+bar_a; liq_m=shop_m+fed_m+bar_m
    comp=[('Shop liquidation (KSBC, tertiary)',shop_a,shop_m),
          ('Consumer-fed (invoice)',fed_a,fed_m),
          ('BAR (invoice)',bar_a,bar_m)]
    for nm,av,mv in comp:
        dv=mv-av; pd=pdelta(mv,av); share=mv/liq_m if liq_m else 0
        trow(ws,r,L6,[nm,av,mv,dv,pd,share],[dict(left=True),dict(fmt='#,##0'),dict(fmt='#,##0'),
          dict(fmt='+#,##0;-#,##0',color=dcol(dv)),dict(fmt='+0%;-0%',color=dcol(pd)),dict(fmt='0%')],body=TINTL); r+=1
    # hero total
    dv=liq_m-liq_a; pd=pdelta(liq_m,liq_a)
    trow(ws,r,L6,['TOTAL LIQUIDATION ✦',liq_a,liq_m,dv,pd,1.0],
        [dict(left=True,bold=True,sz=12,color=GOLD),dict(fmt='#,##0',bold=True,sz=12,color=GOLD),
         dict(fmt='#,##0',bold=True,sz=12,color=GOLD),dict(fmt='+#,##0;-#,##0',bold=True,sz=12,color=GOLD),
         dict(fmt='+0%;-0%',bold=True,sz=12,color=GOLD),dict(fmt='0%',bold=True,sz=12,color=GOLD)],h=27,herotot=True,heroend=9); r+=1
    r=gap(ws,r,e=9)
    # ③ BRAND & PACK
    banner(ws,r,'BRAND & PACK MIX — TOTAL LIQUIDATION, '+META['curm'],c2=9,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    LB=[2,1,1,2,1,2]
    hdrrow(ws,r,LB,['Brand','Liquidation','% Bond','Top Pack','Closing','Status']); r+=1
    for i,br in enumerate(d['brands']):
        sv=br['sales']; nm=br['brand']
        if sv<=0: status,scol=('🚫 Non-mover',RED)
        elif i==0: status,scol=('⭐ Lead brand','FFB8860B')
        elif br['pct']>=0.15: status,scol=('Core',GREEN)
        else: status,scol=('',GREY)
        trow(ws,r,LB,[nm,sv,br['pct'],br['toppack'],br['close'],status],[dict(left=True,bold=(i==0),sz=9),
          dict(fmt='#,##0.0',bold=(i==0)),dict(fmt='0%'),dict(sz=9),dict(fmt='#,##0.0'),dict(bold=bool(status),sz=9,color=scol)],body=TINT3); r+=1
    LP=[3,3,3]; hdrrow(ws,r,LP,['Pack Size','Liquidation (cs)','% of Bond'],fill=NAVY_SOFT); r+=1
    ptot=sum(p[1] for p in d['packs']) or 1
    for pn,pv in d['packs']:
        if pv<=0: continue
        trow(ws,r,LP,[pn,pv,pv/ptot],[dict(left=True,sz=9),dict(fmt='#,##0.0'),dict(fmt='0%')],body=TINT3); r+=1
    r=gap(ws,r,e=9)
    # ④ FIELD COVERAGE
    banner(ws,r,'FIELD COVERAGE — SHOP VISITS, '+META['curm'],c2=9,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    nshops=tm['nshops']; visited=tm['visited']; cov=visited/nshops if nshops else 0; v=d.get('vis')
    L2=[2,1,1,1,2,2]
    hdrrow(ws,r,L2,['Total Visits','Shops Visited','Coverage %','Unvisited','Avg Visits / Shop','Assessment'],h=30); r+=1
    if tm['visits']==0:
        trow(ws,r,L2,['—','—','—','—','—','No visit log captured for this bond in '+META['curm']],
          [dict(bold=True),dict(),dict(),dict(),dict(),dict(color=RED,sz=9,bold=True)],h=20,body=TINT4); r+=1
    else:
        lab='✅ Strong' if cov>=0.9 else ('⚠️ Partial' if cov>=0.6 else '🚫 Low')
        cc=T_BAL if cov>=0.9 else (T_INV if cov>=0.6 else T_CRI)
        avg_vps=tm['visits']/visited if visited else 0
        trow(ws,r,L2,[tm['visits'],visited,cov,nshops-visited,avg_vps,lab],
          [dict(fmt='#,##0',bold=True),dict(fmt='0'),dict(fmt='0%',bold=True),dict(fmt='0'),dict(fmt='0.0'),
           dict(bold=True,color=cc[0],fill=cc[1])],h=20,body=TINT4); r+=1
        if v:
            ad=v['active_days']; avd=(round(float(v['avg_visits_day']),1) if v['avg_visits_day'] else '—'); ad=ad if ad else '—'
            hdrrow(ws,r,L2,['Sales Executive','Active Field Days','Avg Visits / Day','Avg Time / Visit','Field Time (h)','Warehouse Stops'],fill=NAVY_SOFT,h=30); r+=1
            trow(ws,r,L2,[v['staff'],ad,avd,v['avg_time'],v['field_h'],v['wh_stops']],
              [dict(bold=True,sz=9),dict(fmt=('0' if ad!='—' else None)),dict(fmt=('0.0' if avd!='—' else None),bold=True),dict(sz=9),dict(fmt='#,##0.0'),dict(fmt='0')],h=20,body=TINT4); r+=1
    ws.row_dimensions[r].height=10
    dn=b+' DETAIL'
    if dn in wb.sheetnames:
        fed_block(wb[dn], b, d)
        detail_masthead(wb[dn], b, _staff, _per)

# ---- NETWORK DEEP-DIVE on BOND PERFORMANCE (Abhay-requested, 2 Jul 2026) ----
# Cluster grouping is the locked BOND PERFORMANCE clusterwise layout (CLAUDE.md 25 Apr 2026).
CLUSTERS=[('Cluster 1',['KOLLAM','ALAPPUZHA','ATTINGAL','KOTTARAKARA','NEDUMANGAD','PATHANAMTHITTA']),
          ('Cluster 2',['THRISSUR','KOTTAYAM','TRIPUNITHURA','THODUPUZHA','ALUVA']),
          ('Cluster 3',['PERINTHALMANNA','KANNUR','KOZHIKODE','PALAKKAD'])]
if 'BOND PERFORMANCE' in wb.sheetnames:
    ws=wb['BOND PERFORMANCE']; ws.sheet_view.showGridLines=False
    # hero title band (Abhay 2 Jul 2026) -- restyle row 1 in place, NO row inserts (cluster SUM formulas must not shift)
    for m in list(ws.merged_cells.ranges):
        if m.min_row==1 and m.max_row==1: ws.unmerge_cells(str(m))
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=8)
    ws.row_dimensions[1].height=40
    fillrange(ws,1,1,8,fill=f(NAVY_DEEP),botgold=True)
    _t=ws.cell(row=1,column=1)
    _t.font=Font(bold=True,size=16,color=GOLD); _t.alignment=C
    mb=None
    for rr in range(1,ws.max_row+1):
        v=ws.cell(row=rr,column=1).value
        if v and str(v).startswith('\U0001F4CA  NETWORK DEEP-DIVE'): mb=rr; break
    if mb:
        for rng in list(ws.merged_cells.ranges):
            c1,r1,c2,r2=range_boundaries(str(rng))
            if r2>=mb-1: ws.unmerge_cells(str(rng))
        ws.delete_rows(mb-1, ws.max_row-(mb-1)+1)
    totrow=ws.max_row
    for rr in range(1,ws.max_row+1):
        if str(ws.cell(row=rr,column=1).value).strip().upper()=='TOTAL': totrow=rr
    def gap8(r,h=GAP):
        ws.row_dimensions[r].height=h
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
        return r+1
    L8=[2,1,1,1,1,2]
    def tsum(bs,key,cur=True):
        return sum((data[b]['tert_may'] if cur else data[b]['tert_apr']).get(key,0) for b in bs)
    def ssum(bs,key,cur=True):
        return sum((data[b]['sec_may'] if cur else data[b]['sec_apr']).get(key,0) for b in bs)
    r=totrow+2
    banner(ws,r,'\U0001F4CA  NETWORK DEEP-DIVE \u2014 ALL BONDS     \u00b7     '+META['cur']+'  vs  '+META['prev'],c2=8,fill=NAVY_DEEP,sz=14,h=34); r+=1
    r=gap8(r,10)
    # 1: NETWORK SHOP LIQUIDATION
    banner(ws,r,'SHOP LIQUIDATION (KSBC) \u2014 NETWORK',c2=8,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    hdrrow(ws,r,L8,['Metric',META['prevm'],META['curm'],'\u0394','% \u0394','Signal / Rating']); r+=1
    nmv={k:tsum(BONDS,k) for k in ('O','R','S','C')}; nav={k:tsum(BONDS,k,False) for k in ('O','R','S','C')}
    st_m=nmv['S']/(nmv['O']+nmv['R']) if (nmv['O']+nmv['R']) else 0
    st_a=nav['S']/(nav['O']+nav['R']) if (nav['O']+nav['R']) else 0
    nsh_m=sum(data[b]['tert_may']['nshops'] for b in BONDS)
    nsh_a=sum(data[b]['tert_apr'].get('nshops',data[b]['tert_may']['nshops']) for b in BONDS)
    ml,(mfg,mbg)=tier_of(st_m)
    for nm,av,mv,fmt,em,fillc in [('Closing',nav['C'],nmv['C'],'#,##0',False,None),
      ('Active shops',nsh_a,nsh_m,'0',False,None),
      ('Opening',nav['O'],nmv['O'],'#,##0',False,None),
      ('Receipts',nav['R'],nmv['R'],'#,##0',False,None),
      ('Sales (cases)',nav['S'],nmv['S'],'#,##0',True,PALEBLUE),
      ('Sell-Through %',st_a,st_m,'0%',True,PALEBLUE)]:
        if fmt=='0%':
            dv=mv-av
            trow(ws,r,L8,[nm,av,mv,dv,'',ml],[dict(left=True,bold=em,fill=fillc),dict(fmt='0%',fill=fillc),
              dict(fmt='0%',bold=True,fill=fillc),dict(fmt='+0%;-0%',color=dcol(dv),fill=fillc),dict(fill=fillc),
              dict(bold=True,sz=9,color=mfg,fill=mbg)],body=TINT1)
        elif nm=='Active shops':
            trow(ws,r,L8,[nm,av,mv,('\u2014' if av==mv else mv-av),'\u2014','\u2014'],
              [dict(left=True),dict(fmt='0'),dict(fmt='0',bold=True),dict(fmt=('+0;-0' if av!=mv else None)),dict(),dict(color=GREY)],body=TINT1)
        else:
            dv=mv-av; pd=pdelta(mv,av)
            trow(ws,r,L8,[nm,av,mv,dv,pd,arrow(dv)],[dict(left=True,bold=em,fill=fillc),dict(fmt='#,##0',fill=fillc),
              dict(fmt='#,##0',bold=em,fill=fillc),dict(fmt='+#,##0;-#,##0',color=dcol(dv),fill=fillc),
              dict(fmt='+0%;-0%',color=dcol(pd),fill=fillc),dict(bold=True,sz=11,color=dcol(dv),fill=fillc)],body=TINT1)
        r+=1
    r=gap8(r)
    # 2: CLUSTER PERFORMANCE (KSBC shop sales)
    banner(ws,r,'CLUSTER PERFORMANCE \u2014 SHOP SALES (KSBC)',c2=8,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    hdrrow(ws,r,L8,['Cluster',META['prevm'],META['curm'],'\u0394','% \u0394','Sell-Thr% \u00b7 Rating']); r+=1
    for cname,mem in CLUSTERS:
        aS=tsum(mem,'S',False); mS=tsum(mem,'S')
        stc=tsum(mem,'S')/(tsum(mem,'O')+tsum(mem,'R')) if (tsum(mem,'O')+tsum(mem,'R')) else 0
        lab,(fg,bg)=tier_of(stc); dv=mS-aS; pd=pdelta(mS,aS)
        trow(ws,r,L8,[cname,aS,mS,dv,pd,f'{stc:.0%}  \u00b7  '+lab],[dict(left=True,bold=True),dict(fmt='#,##0'),
          dict(fmt='#,##0',bold=True),dict(fmt='+#,##0;-#,##0',color=dcol(dv)),dict(fmt='+0%;-0%',color=dcol(pd)),
          dict(bold=True,sz=9,color=fg,fill=bg)],body=TINT2); r+=1
    dvN=nmv['S']-nav['S']; pdN=pdelta(nmv['S'],nav['S'])
    trow(ws,r,L8,['NETWORK',nav['S'],nmv['S'],dvN,pdN,f'{st_m:.0%}  \u00b7  '+ml],
      [dict(left=True,bold=True,fill=PALEBLUE),dict(fmt='#,##0',bold=True,fill=PALEBLUE),dict(fmt='#,##0',bold=True,fill=PALEBLUE),
       dict(fmt='+#,##0;-#,##0',bold=True,color=dcol(dvN),fill=PALEBLUE),dict(fmt='+0%;-0%',bold=True,color=dcol(pdN),fill=PALEBLUE),
       dict(bold=True,sz=9,color=mfg,fill=mbg)],body=TINT2); r+=1
    r=gap8(r)
    # 3: TOTAL LIQUIDATION BY CLUSTER
    banner(ws,r,'\u2726  TOTAL LIQUIDATION BY CLUSTER \u2014 SHOP + CONSUMER-FED + BAR',c2=8,fill=NAVY_DEEP,font=GOLD,sz=11,h=26); r+=1
    hdrrow(ws,r,L8,['Cluster',META['prevm'],META['curm'],'\u0394','% \u0394','Share of Network']); r+=1
    liqN_m=nmv['S']+ssum(BONDS,'fed')+ssum(BONDS,'bar'); liqN_a=nav['S']+ssum(BONDS,'fed',False)+ssum(BONDS,'bar',False)
    for cname,mem in CLUSTERS:
        a=tsum(mem,'S',False)+ssum(mem,'fed',False)+ssum(mem,'bar',False)
        m=tsum(mem,'S')+ssum(mem,'fed')+ssum(mem,'bar')
        dv=m-a; pd=pdelta(m,a)
        trow(ws,r,L8,[cname,a,m,dv,pd,(m/liqN_m if liqN_m else 0)],[dict(left=True,bold=True),dict(fmt='#,##0'),
          dict(fmt='#,##0',bold=True),dict(fmt='+#,##0;-#,##0',color=dcol(dv)),dict(fmt='+0%;-0%',color=dcol(pd)),
          dict(fmt='0%')],body=TINTL); r+=1
    dvL=liqN_m-liqN_a; pdL=pdelta(liqN_m,liqN_a)
    trow(ws,r,L8,['NETWORK TOTAL LIQUIDATION \u2726',liqN_a,liqN_m,dvL,pdL,1.0],
        [dict(left=True,bold=True,sz=12,color=GOLD),dict(fmt='#,##0',bold=True,sz=12,color=GOLD),
         dict(fmt='#,##0',bold=True,sz=12,color=GOLD),dict(fmt='+#,##0;-#,##0',bold=True,sz=12,color=GOLD),
         dict(fmt='+0%;-0%',bold=True,sz=12,color=GOLD),dict(fmt='0%',bold=True,sz=12,color=GOLD)],h=27,herotot=True,heroend=8); r+=1
    r=gap8(r)
    # 4: BOND MOVERS
    banner(ws,r,'BOND MOVERS \u2014 SHOP SALES vs '+META['prevm'].upper(),c2=8,fill=NAVY_MID,font=WHITE,sz=11,h=24); r+=1
    hdrrow(ws,r,L8,['Bond',META['prevm'],META['curm'],'\u0394','% \u0394','Signal']); r+=1
    dlt=sorted(((b,data[b]['tert_apr'].get('S',0),data[b]['tert_may']['S']) for b in BONDS),key=lambda x:-(x[2]-x[1]))
    movers=[(x,'\u25b2 Rising',GREEN) for x in dlt[:3]]+[(x,'\u25bc Falling',RED) for x in sorted(dlt[-3:],key=lambda x:(x[2]-x[1]))]
    for (b,aS,mS),siglab,sigc in movers:
        dv=mS-aS; pd=pdelta(mS,aS)
        trow(ws,r,L8,[b,aS,mS,dv,pd,siglab],[dict(left=True,bold=True),dict(fmt='#,##0'),dict(fmt='#,##0',bold=True),
          dict(fmt='+#,##0;-#,##0',color=dcol(dv)),dict(fmt='+0%;-0%',color=dcol(pd)),dict(bold=True,color=sigc)],body=TINT3); r+=1
    ws.row_dimensions[r].height=10; r+=1
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
    fc=ws.cell(row=r,column=1,value='Cluster grouping as per the table above  \u00b7  figures in cases  \u00b7  Total Liquidation = Shop + Consumer-fed + BAR (KSBC dispatch excluded \u2014 no double-count)')
    fc.font=Font(size=8,italic=True,color=GREY); fc.alignment=C
    print('BOND PERFORMANCE network deep-dive appended')

print('built, saving',round(time.time()-t0,1),'s')
wb.save(OUT); print('saved',round(time.time()-t0,1),'s')
