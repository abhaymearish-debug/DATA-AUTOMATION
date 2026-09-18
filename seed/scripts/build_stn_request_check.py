#!/usr/bin/env python3
"""
STN REQUEST CHECK — latest closing stock + PI RL / RQ / MQ for every line
on one or more STN REQUEST template workbooks.

    python3 build_stn_request_check.py --out OUT.xlsx [--base "<Claude folder>"]
                                       [--label "3 Sep 2026"] FILE.xlsx [FILE.xlsx ...]

One sheet per request file (named by the warehouse(s) on it), in the exact
layout of `STN tracking/STN REQUEST - ALAPPUZHA 25 Aug 2026 - RL RQ MQ + CLOSING.xlsx`
(the hand-built precedent):

  Shop · Outlet · Brand · Pack · STN Req · Closing · RL · RQ · MQ ·
  Brand Closing · Brand STN · After STN · vs MQ · Prev 3-Mo Sale (btl) · Note

Sources (auto-discovered under --base):
  * Closing  = newest KSBC daily raw sheet in the newest `KSBC shop sales/`
               workbook (cases + bottles/BPC folded), per shop × product.
  * RL/RQ/MQ = newest `PURCHASE INSTRUCTION/PI INSIGHTS - <PRIOR> & <CUR> <YEAR>.xlsx`,
               sheet RAW PI DATA, rows for the CURRENT month (brand-level).
  * Outlet names / bond / staff from MASTER DATA CONFIRMED.xlsx (sheet 16-4-25).
"""
import argparse, glob, os, re, sys
from collections import OrderedDict, defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_stn_tracker import canon_brand, canon_pack, CANON  # noqa: E402

# ---------------------------------------------------------------- constants
NAVY, NAVY_MID, GOLD_DIM, STEEL = "FF0D1B4A", "FF263F80", "FFFFD54F", "FFB9C4E8"
GREY_TOTAL, ZEBRA, AMBER, GRID, GRID_DK = "FF374151", "FFEDF1F8", "FFFFE0B2", "FFD5DBE8", "FF8FA8D8"
RED, GREEN, BLUE, INK, INK_SOFT = "FFC62828", "FF2E7D32", "FF1565C0", "FF0D1B4A", "FF44506E"
FONT = "Aptos Narrow"
NUM2 = '#,##0.00;-#,##0.00;"·"'
NUM0 = "#,##0"
MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
          "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]

# request-template brand banner -> canonical (build_stn_tracker vocabulary)
TEMPLATE_BRANDS = {
    "BCB CLASSIC": "BCB CLASSIC", "BLENDER'S CHOICE": "BLENDER'S CHOICE",
    "MORNING WALKER": "MORNING WALKER", "CCB": "CHAIRMAN'S CHOICE",
    "MAGIC BLEND": "MAGIC BLEND RUM", "KS 99": "KS 99",
    "OLD PEARL": "OLD PEARL", "ROYAL OLD FORT": "ROYAL OLD FORT",
}
# canonical -> PI RAW 'Brand' short code
PI_CODE = {
    "BCB CLASSIC": "BCB", "BLENDER'S CHOICE": "BLENDERS", "MORNING WALKER": "MWB",
    "CHAIRMAN'S CHOICE": "CC", "MAGIC BLEND RUM": "MBR", "KS 99": "KS 99",
    "OLD PEARL": "OLD PEARL", "ROYAL OLD FORT": "ROF",
}
BPC = {"180 ML": 48, "375 ML": 24, "500 ML": 18, "750 ML": 12, "1000 ML": 9}
PACK_ORDER = {"180 ML": 0, "375 ML": 1, "500 ML": 2, "750 ML": 3, "1000 ML": 4}
BRAND_ORDER = list(TEMPLATE_BRANDS.values())


def thin(color=GRID):
    return Side(style="thin", color=color)


def med(color=GRID):
    return Side(style="medium", color=color)


# ---------------------------------------------------------------- discovery
def newest_ksbc_workbook(base):
    folder = os.path.join(base, "KSBC shop sales")
    cands = []
    for f in glob.glob(os.path.join(folder, "*ANALYSIS.xlsx")):
        name = os.path.basename(f).upper()
        m = re.match(r"([A-Z]+)\b", name)
        if not m or m.group(1) not in MONTHS or name.startswith("~$"):
            continue
        cands.append((os.path.getmtime(f), f))
    if not cands:
        sys.exit("No KSBC analysis workbook found")
    # newest by month index first, then mtime
    def key(t):
        name = os.path.basename(t[1]).upper()
        return (MONTHS.index(re.match(r"([A-Z]+)", name).group(1)), t[0])
    return max(cands, key=key)[1]


def newest_daily_sheet(wb):
    """Newest '<MONTH> <N>' or '<MONTH> <A>-<B>' daily raw sheet (not CUMULATIVE/COMBINED)."""
    best = None
    for sn in wb.sheetnames:
        m = re.fullmatch(r"([A-Z]+) (\d{1,2})(?:-(\d{1,2}))?", sn.strip())
        if not m or m.group(1) not in MONTHS:
            continue
        last = int(m.group(3) or m.group(2))
        k = (MONTHS.index(m.group(1)), last)
        if best is None or k > best[0]:
            best = (k, sn)
    if best is None:
        sys.exit("No daily raw sheet found in KSBC workbook")
    (mi, d), sn = best
    return sn, mi, d


def newest_pi_workbook(base):
    cands = glob.glob(os.path.join(base, "PURCHASE INSTRUCTION", "PI INSIGHTS - * & * *.xlsx"))
    cands = [c for c in cands if not os.path.basename(c).startswith("~$")]
    if not cands:
        sys.exit("No PI INSIGHTS workbook found")
    def key(f):
        m = re.search(r"PI INSIGHTS - (\w+) & (\w+) (\d{4})", os.path.basename(f))
        return (int(m.group(3)), MONTHS.index(m.group(2).upper()))
    return max(cands, key=key)


# ---------------------------------------------------------------- loaders
def load_master(base):
    wb = openpyxl.load_workbook(os.path.join(base, "MASTER DATA CONFIRMED.xlsx"), read_only=True)
    ws = wb["16-4-25"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and "Shop Code" in [str(v).strip() for v in r if v])
    hdr = [str(h).strip().lower() if h else "" for h in rows[hi]]
    rows = rows[hi:]
    def col(*names):
        for n in names:
            for i, h in enumerate(hdr):
                if h == n:
                    return i
        return None
    ci = {"code": col("shop code", "code"), "name": col("shop name", "name"),
          "staff": col("field staff", "staff"), "bond": col("bond"), "status": col("status"),
          "type": col("cat", "type", "category")}
    out = {}
    for r in rows[1:]:
        if ci["code"] is None or r[ci["code"]] is None:
            continue
        try:
            code = int(str(r[ci["code"]]).strip())
        except ValueError:
            continue
        out[code] = {k: (str(r[i]).strip() if i is not None and r[i] is not None else "")
                     for k, i in ci.items()}
    return out


def load_closing(base):
    path = newest_ksbc_workbook(base)
    wb = openpyxl.load_workbook(path, read_only=True)
    sn, mi, day = newest_daily_sheet(wb)
    ws = wb[sn]
    hdr = None
    closing = defaultdict(float)       # (code, canon_brand, pack) -> cs
    listed = set()                     # (code, canon_brand, pack) present on the report
    shops_seen = set()
    for r in ws.iter_rows(values_only=True):
        if hdr is None:
            if r and r[0] == "Warehouse Name":
                hdr = [str(h).strip() for h in r]
                ix = {h: i for i, h in enumerate(hdr)}
            continue
        if not r or r[1] is None:
            continue
        try:
            code = int(str(r[ix["Shop Code"]]).strip())
        except ValueError:
            continue
        shops_seen.add(code)
        brand = canon_brand(r[ix["Brand Name"]])
        pack = canon_pack(r[ix["Packing"]])
        if brand is None or pack is None:
            continue
        bpc = r[ix["Bottle Per Case"]] or BPC[pack]
        cs = float(r[ix["Shop Closing Cases"]] or 0) + float(r[ix["Shop Closing Bottles"]] or 0) / bpc
        closing[(code, brand, pack)] += cs
        listed.add((code, brand, pack))
    asof = date(date.today().year, mi + 1, day)
    return closing, listed, shops_seen, asof, os.path.basename(path), sn


def load_pi(base):
    path = newest_pi_workbook(base)
    m = re.search(r"PI INSIGHTS - (\w+) & (\w+) (\d{4})", os.path.basename(path))
    cur_month, year = m.group(2), int(m.group(3))
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["RAW PI DATA"]
    hdr = None
    pi = {}       # (short_code, PI brand code) -> dict
    pi_shops = set()
    for r in ws.iter_rows(values_only=True):
        if hdr is None:
            if r and r[0] == "Month":
                hdr = [str(h).strip() for h in r]
                ix = {h: i for i, h in enumerate(hdr)}
            continue
        if not r or r[0] is None:
            continue
        if str(r[ix["Month"]]).strip().upper() != cur_month.upper():
            continue
        short = str(r[ix["Shop Code"]]).strip()
        pi_shops.add(short)
        pi[(short, str(r[ix["Brand"]]).strip())] = {
            "name": r[ix["Shop Name"]], "prev": r[ix["Prev 3-Mo Sale (btl)"]],
            "rl": r[ix["RL (cs)"]], "rq": r[ix["RQ (cs)"]], "mq": r[ix["MQ (cs)"]],
        }
    return pi, pi_shops, cur_month, year, os.path.basename(path)


# ---------------------------------------------------------------- request parser
def parse_request(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    # row 3 = brand banners (merged), row 4 = packs
    banners = {}
    cur = None
    for c in range(3, ws.max_column + 1):
        v = ws.cell(3, c).value
        if v is not None and str(v).strip() and str(v).strip().upper() != "TOTAL":
            cur = TEMPLATE_BRANDS.get(str(v).strip().upper()) or canon_brand(v)
            if cur is None:
                sys.exit(f"{os.path.basename(path)}: unknown brand banner {v!r}")
        pk = ws.cell(4, c).value
        if cur and pk is not None:
            pack = canon_pack(pk)
            if pack is None:
                sys.exit(f"{os.path.basename(path)}: unknown pack {pk!r}")
            banners[c] = (cur, pack)
    total_col = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(3, c).value or "").strip().upper() == "TOTAL":
            total_col = c
    meta = {"warehouse": ws["C2"].value, "date": ws["J2"].value, "exec": ws["Q2"].value}
    shops = OrderedDict()     # code -> {"place":..., "lines": [(brand, pack, cs)]}
    for r in range(5, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        if isinstance(a, str) and a.strip().upper().startswith("TOTAL"):
            break
        try:
            short = int(float(a))
        except (TypeError, ValueError):
            continue
        place = (ws.cell(r, 2).value or "").strip()
        lines = OrderedDict()
        for c, (brand, pack) in banners.items():
            v = ws.cell(r, c).value
            if v in (None, "", 0):
                continue
            lines[(brand, pack)] = lines.get((brand, pack), 0) + float(v)
        rowtot = ws.cell(r, total_col).value if total_col else None
        s = sum(lines.values())
        if rowtot is not None and abs(float(rowtot) - s) > 0.01:
            sys.exit(f"{os.path.basename(path)}: shop {short} row total {rowtot} ≠ parsed {s}")
        if lines:
            shops[short] = {"place": place, "lines": lines}
    return meta, shops


def resolve_code(short, master):
    """Field short code -> master code via '10' prefix (5002 -> 105002); abort if ambiguous."""
    hits = [c for c in master if str(c).endswith(str(short)) and master[c]["type"].upper() == "KSBC"]
    pref = [c for c in hits if c == int("10" + str(short))]
    if pref:
        return pref[0]
    if len(hits) == 1:
        return hits[0]
    sys.exit(f"shop {short}: {'ambiguous' if hits else 'not found'} in master {hits}")


# ---------------------------------------------------------------- writer
def write_sheet(wb, title, meta, shops, ctx, label):
    master, closing, listed, shops_seen, asof = ctx["master"], ctx["closing"], ctx["listed"], ctx["shops_seen"], ctx["asof"]
    pi, pi_shops, pi_month = ctx["pi"], ctx["pi_shops"], ctx["pi_month"]
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.showGridLines = False
    widths = [9, 25, 19, 8, 10, 10, 7, 7, 7, 13, 11, 11, 9, 11, 34]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    NC = len(widths)

    def band(row, text, size, bold, italic, color, align, h):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NC)
        for c in range(1, NC + 1):
            cell = ws.cell(row, c)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
            cell.alignment = Alignment(horizontal=align, vertical="center", indent=1 if align != "center" else 0)
        ws.cell(row, 1).value = text
        ws.row_dimensions[row].height = h

    # header block
    n_lines = sum(len(s["lines"]) for s in shops.values())
    n_cs = sum(sum(s["lines"].values()) for s in shops.values())
    codes = [resolve_code(s, master) for s in shops]
    bonds = sorted({master[c]["bond"] for c in codes})
    staff = sorted({master[c]["staff"] for c in codes if master[c]["staff"]})
    places = sorted({s["place"].strip().title() for s in shops.values() if s["place"]})
    wh = meta["warehouse"] or " / ".join(places)
    when = meta["date"]
    when_txt = when.strftime("%d/%m/%Y") if hasattr(when, "strftime") else (str(when) if when else f"received {label}")
    band(1, "STN REQUEST — RL / RQ / MQ & LATEST CLOSING STOCK", 18, True, False, "FFFFFFFF", "center", 34)
    band(2, f"{str(wh).upper()}   ·   bond {' / '.join(bonds)}   ·   STN {when_txt}   ·   {' / '.join(staff) or 'VACANT'}"
            f"   ·   {len(shops)} outlets · {n_lines} SKU lines · {n_cs:,.0f} cs", 10, False, True, GOLD_DIM, "left", 17)
    band(3, f"RL / RQ / MQ = {pi_month.title()} {ctx['pi_year']} Purchase Instruction (brand-level, covers all packs)   ·   "
            f"Closing stock as on {asof.strftime('%d %b %Y')}   ·   all figures in cases", 10, False, False, STEEL, "right", 18)
    for c in range(1, NC + 1):
        ws.cell(3, c).border = Border(bottom=Side(style="medium", color="FFFFB300"))

    hdrs = ["Shop", "Outlet", "Brand", "Pack\n(ml)", "STN Req\n(cs)", "Closing\n(cs)", "RL\n(cs)", "RQ\n(cs)", "MQ\n(cs)",
            "Brand Closing\n(cs)", "Brand STN\n(cs)", "After STN\n(cs)", "vs MQ\n(cs)", "Prev 3-Mo\nSale (btl)", "Note"]
    HR = 5
    for i, h in enumerate(hdrs, 1):
        c = ws.cell(HR, i, h)
        c.fill = PatternFill("solid", fgColor=NAVY_MID)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin())
    ws.row_dimensions[HR].height = 34

    r = HR + 1
    tot_req = tot_close = tot_bclose = 0.0
    warnings = []
    for si, (short, s) in enumerate(shops.items()):
        code = resolve_code(short, master)
        m = master[code]
        pi_name = next((v["name"] for (sc, b), v in pi.items() if sc == str(short)), None)
        outlet = pi_name or re.sub(r"^\d+-\s*", "", m["name"]).title()
        if code not in shops_seen:
            warnings.append(f"{short} {outlet}: not on the {asof:%d %b} KSBC report — closing shown as 0")
        zebra = ZEBRA if si % 2 else None
        # group lines by brand in template order, packs in pack order
        by_brand = OrderedDict()
        for (brand, pack), cs in sorted(s["lines"].items(), key=lambda kv: (BRAND_ORDER.index(kv[0][0]), PACK_ORDER[kv[0][1]])):
            by_brand.setdefault(brand, []).append((pack, cs))
        block_start = r
        for brand, packs in by_brand.items():
            p = pi.get((str(short), PI_CODE[brand]))
            b_close = sum(closing.get((code, brand, pk), 0.0) for pk in BPC)
            b_stn = sum(cs for _, cs in packs)
            after = b_close + b_stn
            off_indent = p is None or (p["mq"] or 0) == 0
            note_brand = None
            if str(short) not in pi_shops:
                note_brand = f"NO {pi_month[:3].upper()} PI — shop has no {pi_month.title()} PI file"
            elif p is None:
                note_brand = f"OFF-INDENT — no line on {pi_month[:3].title()} PI"
            elif (p["mq"] or 0) == 0:
                note_brand = f"OFF-INDENT — MQ 0 on {pi_month[:3].title()} PI"
            first = r
            for pack, cs in packs:
                cl = closing.get((code, brand, pack), 0.0)
                note = note_brand if r == first else None
                if (code, brand, pack) not in listed and code in shops_seen:
                    note = (note + "  ·  " if note else "") + "pack not listed at this shop"
                vals = [code if r == block_start else None, outlet if r == block_start else None,
                        brand if brand != "MAGIC BLEND RUM" else "MAGIC BLEND", int(pack.split()[0]), cs, round(cl, 2)]
                if r == first:
                    vals += [p["rl"] if p else 0, p["rq"] if p else 0, p["mq"] if p else 0,
                             round(b_close, 2), b_stn, round(after, 2), round(after - (p["mq"] if p else 0), 2),
                             p["prev"] if p else None]
                else:
                    vals += [None] * 8
                vals.append(note)
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(r, ci, v)
                    c.font = Font(name=FONT, size=10)
                    c.alignment = Alignment(horizontal="left" if ci in (2, 3, 15) else "center", vertical="center")
                    if zebra:
                        c.fill = PatternFill("solid", fgColor=zebra)
                    top = med() if r == block_start else thin()
                    c.border = Border(top=top, bottom=thin(), left=thin(GRID_DK if ci == 3 else GRID), right=thin())
                    if ci in (5, 6, 10, 11, 12, 13):
                        c.number_format = NUM2
                    elif ci in (7, 8, 9, 14):
                        c.number_format = NUM0
                ws.cell(r, 5).font = Font(name=FONT, size=10, bold=True, color=BLUE)
                if r == first:
                    d = after - (p["mq"] if p else 0)
                    ws.cell(r, 13).font = Font(name=FONT, size=10, bold=True, color=RED if d > 0 else GREEN)
                    if off_indent:
                        for ci in (7, 8, 9, 10, 11, 12, 13):
                            ws.cell(r, ci).fill = PatternFill("solid", fgColor=AMBER)
                if note:
                    ws.cell(r, 15).font = Font(name=FONT, size=9, bold=True, color=RED if "INDENT" in note or "NO " in note else INK_SOFT)
                ws.row_dimensions[r].height = 17
                tot_req += cs
                tot_close += cl
                r += 1
            if len(packs) > 1:
                for ci in range(7, 15):
                    ws.merge_cells(start_row=first, start_column=ci, end_row=r - 1, end_column=ci)
            tot_bclose += b_close
        # shop id cells: bold navy, merged down the block
        for ci in (1, 2):
            c = ws.cell(block_start, ci)
            c.font = Font(name=FONT, size=11, bold=True, color=INK)
            if r - 1 > block_start:
                ws.merge_cells(start_row=block_start, start_column=ci, end_row=r - 1, end_column=ci)
    last = r - 1
    # total row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    for ci in range(1, NC + 1):
        c = ws.cell(r, ci)
        c.fill = PatternFill("solid", fgColor=GREY_TOTAL)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(top=med())
    ws.cell(r, 1, f"TOTAL — {len(shops)} outlets  ·  {n_lines} SKU lines").alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.cell(r, 5, round(tot_req, 2)).number_format = NUM0
    ws.cell(r, 6, round(tot_close, 2)).number_format = NUM2
    ws.cell(r, 10, round(tot_bclose, 2)).number_format = NUM2
    ws.row_dimensions[r].height = 26
    ws.auto_filter.ref = f"A{HR}:O{last}"
    ws.freeze_panes = f"C{HR + 1}"
    ws.print_title_rows = f"{HR}:{HR}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # how to read
    r += 2
    ws.cell(r, 1, "HOW TO READ THIS").font = Font(name=FONT, size=10, italic=True, color=INK)
    ws.row_dimensions[r].height = 16
    notes = [
        f"RL / RQ / MQ come from KSBC's {pi_month.title()} {ctx['pi_year']} Purchase Instruction and are BRAND-level — one set per shop × brand, covering all pack sizes together. They are repeated across that brand's pack rows.",
        "MQ (= RL + RQ) is KSBC's replenish-to target for the brand at that shop. 'After STN' = brand closing stock + total STN cases for that brand; 'vs MQ' is how far that lands above (red) or below (green) the MQ.",
        f"Closing (cs) is the shop's own closing stock for that exact pack as on {asof.strftime('%d %b %Y')}, from the KSBC daily shop-sale report. Loose bottles are folded into cases.",
        "Most STN lines land above MQ by design — an STN is the supplier's lever to place stock the shop's own re-order quantity will not pull. The figure is shown so the size of the overshoot is visible, not as a warning.",
        f"OFF-INDENT (amber) means KSBC gives that brand no ceiling at this shop in {pi_month.title()} — there will be no auto-replenishment there until offtake earns the line back.",
        f"Prev 3-Mo Sale is the bottle offtake KSBC used to set {pi_month.title()}'s levels.",
    ]
    for t in notes:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
        c = ws.cell(r, 1, t)
        c.font = Font(name=FONT, size=10, color=INK_SOFT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 26
    return {"outlets": len(shops), "lines": n_lines, "cs": n_cs, "closing": tot_close,
            "brand_closing": tot_bclose, "warnings": warnings, "bonds": bonds}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--base", default=os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
    ap.add_argument("--label", default=date.today().strftime("%d %b %Y"))
    a = ap.parse_args()

    master = load_master(a.base)
    closing, listed, shops_seen, asof, ksbc_file, ksbc_sheet = load_closing(a.base)
    pi, pi_shops, pi_month, pi_year, pi_file = load_pi(a.base)
    ctx = dict(master=master, closing=closing, listed=listed, shops_seen=shops_seen, asof=asof,
               pi=pi, pi_shops=pi_shops, pi_month=pi_month, pi_year=pi_year)
    print(f"closing  : {ksbc_file} → sheet {ksbc_sheet} (as on {asof})")
    print(f"PI       : {pi_file} → {pi_month} rows ({len(pi)} lines / {len(pi_shops)} shops)")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    for f in a.files:
        meta, shops = parse_request(f)
        places = []
        for s in shops.values():
            p = s["place"].strip().upper()
            if p and p not in places:
                places.append(p)
        title = " + ".join(places) or os.path.splitext(os.path.basename(f))[0]
        base_title, k = title, 2
        while title[:31] in used:
            title = f"{base_title} ({k})"; k += 1
        used.add(title[:31])
        info = write_sheet(wb, title, meta, shops, ctx, a.label)
        print(f"sheet {title[:31]!r}: {info['outlets']} outlets · {info['lines']} lines · {info['cs']:.0f} cs STN · "
              f"closing {info['closing']:.2f} cs (pack) / {info['brand_closing']:.2f} cs (brand) · bond {'/'.join(info['bonds'])}")
        for w in info["warnings"]:
            print("   ⚠", w)
    wb.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
