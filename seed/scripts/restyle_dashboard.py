#!/usr/bin/env python3
"""Rebuild the DASHBOARD sheet using the KSBC shop sales dark-navy style."""
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# Auto-detect session root: <session>/mnt/Claude/.claude/scripts/restyle_dashboard.py
_THIS = os.path.abspath(__file__)
_CLAUDE = os.path.abspath(os.path.join(os.path.dirname(_THIS), "..", ".."))
_SESSION_ROOT = os.path.abspath(os.path.join(_CLAUDE, "..", ".."))
SCRATCH = f"{_SESSION_ROOT}/sec_scratch.xlsx"

# KSBC palette
DARK_BG     = "FF0B1929"   # page bg
PANEL       = "FF1E3A5F"   # KPI card bg
ROW_A       = "FF122240"   # zebra A
ROW_B       = "FF1E3A5F"   # zebra B
CYAN        = "FF00D4FF"
CYAN2       = "FF00D5FF"
GREEN       = "FF00F5A0"
AMBER       = "FFFFD166"
AMBER2      = "FFFCA33D"
PURPLE      = "FFA3A1FB"
PINK        = "FFFF7AB6"
LIGHTCYAN   = "FF66D9EF"
LIGHTSUB    = "FF7FDBFF"
CORAL       = "FFFF6B6B"
MUTED       = "FFA8C5E0"
WHITE       = "FFFFFFFF"

FONT = "Segoe UI"

wb = openpyxl.load_workbook(SCRATCH)

# Derive MONTH and latest_day from COMBINED DISPATCHES sheet (month-agnostic)
_comb_sheet = [s for s in wb.sheetnames if s.endswith("COMBINED DISPATCHES")][0]
MONTH = _comb_sheet.replace(" COMBINED DISPATCHES", "")
_comb_probe = wb[_comb_sheet]
_days = set()
for _r in range(2, _comb_probe.max_row+1):
    _d = _comb_probe.cell(row=_r, column=12).value
    if _d:
        try:
            _days.add(int(str(_d).split("-")[0]))
        except Exception:
            pass
LATEST_DAY = max(_days) if _days else 0
# Derive YEAR from the COMBINED dispatch dates (audit fix #1 — was hardcoded 2026).
_years = set()
for _r in range(2, _comb_probe.max_row + 1):
    _dv = _comb_probe.cell(row=_r, column=12).value
    if _dv:
        try:
            _years.add(int(str(_dv).split("-")[2]))
        except Exception:
            pass
YEAR = max(_years) if _years else datetime.now().year

# Delete existing DASHBOARD and recreate as first
if "DASHBOARD" in wb.sheetnames:
    del wb["DASHBOARD"]
ws = wb.create_sheet("DASHBOARD", 0)
ws.sheet_view.showGridLines = False

# Blanket dark background across visible area
bg_fill = PatternFill("solid", fgColor=DARK_BG)
panel_fill = PatternFill("solid", fgColor=PANEL)

MAX_ROW = 70
MAX_COL = 30  # A..AD
for r in range(1, MAX_ROW+1):
    for c in range(1, MAX_COL+1):
        cell = ws.cell(row=r, column=c)
        cell.fill = bg_fill

# Uniform column widths so the 5 KPI panels render with equal width.
col_widths = {"A": 3.34}
for cl in "BCDEFGHIJKLMNOPQRSTUV":
    col_widths[cl] = 13.0
# Widen bond-name columns slightly for rows 19-34 and league table
col_widths["C"] = 16.0
col_widths["L"] = 16.0
col_widths["W"] = 13.0
for cl, w in col_widths.items():
    ws.column_dimensions[cl].width = w

# Row heights
row_heights = {1:12, 2:37.5, 3:21.75, 4:13.5, 5:18, 6:13.5, 7:18, 8:31.5, 9:15.75,
               10:13.5, 11:18, 12:15.75, 13:31.5, 14:15.75, 15:13.5, 16:21.75, 17:25.5,
               18:21.75, 19:21.75, 20:21.75, 21:21.75, 22:21.75, 23:21.75, 24:21.75,
               25:21.75, 26:21.75, 27:21.75, 28:21.75, 29:21.75, 30:21.75, 31:21.75,
               32:21.75, 33:21.75, 34:21.75, 35:12, 36:12, 37:12,
               38:25.5, 39:21.75, 40:21.75, 41:21.75, 42:21.75, 43:21.75, 44:21.75,
               45:12, 46:12, 47:25.5, 48:24, 49:22, 50:22, 51:22, 52:22, 53:22,
               54:22, 55:22, 56:22, 57:22, 58:22, 59:22, 60:22, 61:22, 62:22, 63:22,
               64:12, 65:12, 66:18}
for r, h in row_heights.items():
    ws.row_dimensions[r].height = h

# -------- Load aggregates from other sheets to rebuild dashboard --------
# BOND PERFORMANCE sheet holds per-bond totals
bp = wb["BOND PERFORMANCE"]
bond_data = []  # list of dicts
for r in range(4, 25):
    name = bp.cell(row=r, column=1).value
    if name in (None, "TOTAL", "UNMATCHED / Pending master"):
        continue
    k = bp.cell(row=r, column=2).value or 0
    f = bp.cell(row=r, column=3).value or 0
    b = bp.cell(row=r, column=4).value or 0
    tot = bp.cell(row=r, column=5).value or 0
    bond_data.append({"bond": name, "k": k, "f": f, "b": b, "tot": tot, "cash": f+b})

# Totals
total_ksbc = sum(x["k"] for x in bond_data)
total_fed = sum(x["f"] for x in bond_data)
total_bar = sum(x["b"] for x in bond_data)
total_all = total_ksbc + total_fed + total_bar
# Read real totals from BP total row to catch unmatched
for r in range(4, 25):
    if bp.cell(row=r, column=1).value == "TOTAL":
        grand_k = bp.cell(row=r, column=2).value
        grand_f = bp.cell(row=r, column=3).value
        grand_b = bp.cell(row=r, column=4).value
        grand_total = bp.cell(row=r, column=5).value
        break

cashflow = grand_f + grand_b
pct_fed = grand_f/grand_total if grand_total else 0
pct_bar = grand_b/grand_total if grand_total else 0

# Outlets count from combined dispatches
comb = wb[_comb_sheet]
licensees = set()
for r in range(2, comb.max_row+1):
    v = comb.cell(row=r, column=7).value  # Licensee No. col G
    if v: licensees.add(str(v).strip())
n_outlets = len(licensees)

# Top brand
brd = wb["BRAND PERFORMANCE"]
top_brand = brd.cell(row=4, column=1).value
top_brand_cases = brd.cell(row=4, column=5).value

# Peak day
dt = wb["DAILY TREND"]
peak_day_name = None
peak_day_cases = 0
for r in range(4, dt.max_row+1):
    d = dt.cell(row=r, column=1).value
    v = dt.cell(row=r, column=3).value or 0
    if d == "TOTAL" or d is None: continue
    if v > peak_day_cases:
        peak_day_cases = v
        peak_day_name = d

# Unmatched info
un_count = 0
un_cases = 0
if "UNMATCHED" in wb.sheetnames:
    um = wb["UNMATCHED"]
    for r in range(5, um.max_row+1):
        v = um.cell(row=r, column=3).value
        if v:
            un_count += 1
            un_cases += v

# -------- Title & subtitle --------
ws.merge_cells("B2:W2")
c = ws["B2"]
c.value = "K.S. DISTILLERY — SECONDARY SALES DASHBOARD"
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=26, bold=True, color=CYAN)
c.alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("B3:W3")
c = ws["B3"]
c.value = f"{MONTH} 1–{LATEST_DAY}, {YEAR}  |  Warehouse → outlets (KSBC dispatch + invoice cashflow: Consumer fed + BAR)  |  All figures in cases"
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=11, color=LIGHTSUB)
c.alignment = Alignment(horizontal="center", vertical="center")

# -------- KPI row 7/8/9 --------
kpis = [
    ("B", "TOTAL SECONDARY SALES", grand_total, "all outlets, all cases", CYAN),
    ("F", "INVOICE SALE (FED AND BAR)", cashflow, "Consumer fed + BAR", GREEN),
    ("J", "KSBC DISPATCH", grand_k, "pipeline to KSBC shops", AMBER),
    ("N", "CONSUMER FED", grand_f, f"{pct_fed*100:.1f}% of total", PURPLE),
    ("R", "BAR", grand_b, f"{pct_bar*100:.1f}% of total", PINK),
]
for start_col, label, value, sub, accent in kpis:
    sc = openpyxl.utils.column_index_from_string(start_col)
    ec = sc + 3  # 4-col wide panels → 5 panels across B:U, evenly spaced
    end_col = get_column_letter(ec)
    ws.merge_cells(f"{start_col}7:{end_col}7")
    c = ws[f"{start_col}7"]; c.value = label
    c.fill = panel_fill
    c.font = Font(name=FONT, size=10, bold=True, color=accent)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"{start_col}8:{end_col}8")
    c = ws[f"{start_col}8"]; c.value = value
    c.fill = panel_fill
    c.font = Font(name=FONT, size=24, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.number_format = "#,##0"
    ws.merge_cells(f"{start_col}9:{end_col}9")
    c = ws[f"{start_col}9"]; c.value = sub
    c.fill = panel_fill
    c.font = Font(name=FONT, size=9, color=MUTED)
    c.alignment = Alignment(horizontal="center", vertical="center")

# -------- Callout row 12-14: 3 callouts --------
# Determine top bonds
bond_by_total = sorted(bond_data, key=lambda x: -x["tot"])
bond_by_cash = sorted(bond_data, key=lambda x: -x["cash"])

top_sales = bond_by_total[0]
top_cash = bond_by_cash[0]

callouts = [
    ("B", "🏆 TOP BOND — SECONDARY SALES", top_sales["bond"], f"{top_sales['tot']:,.0f} cases dispatched", GREEN),
    ("H", "💰 TOP BOND — CASHFLOW (FED+BAR)", top_cash["bond"], f"{top_cash['cash']:,.0f} cashflow cases", CYAN),
    ("N", "⭐ TOP BRAND", top_brand, f"{top_brand_cases:,.0f} cases", AMBER),
]
for start_col, label, big, sub, accent in callouts:
    sc = openpyxl.utils.column_index_from_string(start_col)
    ec = sc + 5  # span 6 cols
    end_col = get_column_letter(ec)
    # Row 12: label with colored fill
    ws.merge_cells(f"{start_col}12:{end_col}12")
    c = ws[f"{start_col}12"]; c.value = label
    c.fill = PatternFill("solid", fgColor=accent)
    c.font = Font(name=FONT, size=10, bold=True, color=DARK_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")
    # Row 13: big name
    ws.merge_cells(f"{start_col}13:{end_col}13")
    c = ws[f"{start_col}13"]; c.value = big
    c.fill = PatternFill("solid", fgColor=ROW_A)
    c.font = Font(name=FONT, size=20, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    # Row 14: accent sub
    ws.merge_cells(f"{start_col}14:{end_col}14")
    c = ws[f"{start_col}14"]; c.value = sub
    c.fill = PatternFill("solid", fgColor=ROW_A)
    c.font = Font(name=FONT, size=11, bold=True, color=accent)
    c.alignment = Alignment(horizontal="center", vertical="center")

# -------- Section header row 17 — spans both breakdown tables (B:R) --------
ws.merge_cells("B17:R17")
c = ws["B17"]; c.value = "📊 BOND-WISE BREAKDOWN"
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=14, bold=True, color=CYAN)
c.alignment = Alignment(horizontal="center", vertical="center")

# Row 18: left = SALES BY BOND (Cases), right = CASHFLOW BY BOND (FED+BAR)
ws.merge_cells("B18:I18")
c = ws["B18"]; c.value = "SECONDARY SALES BY BOND (Cases)"
c.fill = PatternFill("solid", fgColor=ROW_A)
c.font = Font(name=FONT, size=13, bold=True, color=GREEN)
c.alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("K18:R18")
c = ws["K18"]; c.value = "CASHFLOW (FED+BAR) BY BOND"
c.fill = PatternFill("solid", fgColor=ROW_A)
c.font = Font(name=FONT, size=13, bold=True, color=CYAN)
c.alignment = Alignment(horizontal="center", vertical="center")

# Row 19 headers
left_headers = [("B", "#"), ("C", "Bond"), ("I", "Cases")]
right_headers = [("K", "#"), ("L", "Bond"), ("R", "Cases")]

for col, h in left_headers + right_headers:
    c = ws[f"{col}19"]
    c.value = h
    c.fill = PatternFill("solid", fgColor=PANEL)
    c.font = Font(name=FONT, size=10, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")

# Rows 20-34: bond data (left = by total, right = by cashflow)
for i in range(15):
    r = 20 + i
    zebra = ROW_A if i % 2 == 0 else ROW_B
    # left: sales rank
    bd = bond_by_total[i]
    ws[f"B{r}"].value = i+1
    ws[f"C{r}"].value = bd["bond"]
    # merge D:H
    ws.merge_cells(f"D{r}:H{r}")
    ws[f"I{r}"].value = bd["tot"]
    # right: cashflow rank
    bd2 = bond_by_cash[i]
    ws[f"K{r}"].value = i+1
    ws[f"L{r}"].value = bd2["bond"]
    ws.merge_cells(f"M{r}:Q{r}")
    ws[f"R{r}"].value = bd2["cash"]

    for col in ["B","C","I","K","L","R"]:
        c = ws[f"{col}{r}"]
        c.fill = PatternFill("solid", fgColor=zebra)
        c.font = Font(name=FONT, size=10, color=WHITE)
        if col in ("C","L"):
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        elif col in ("B","K"):
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
            c.number_format = "#,##0.00"
    # Also fill the merged D:H and M:Q with zebra
    for col_letter in "DEFGH":
        ws[f"{col_letter}{r}"].fill = PatternFill("solid", fgColor=zebra)
    for col_letter in "MNOPQ":
        ws[f"{col_letter}{r}"].fill = PatternFill("solid", fgColor=zebra)

# -------- Top outlets (rows 38-44) --------
# Read region sheets for per-outlet cases, plus cat info from master
from collections import defaultdict
# Reload outlet agg via region sheets
outlets = []  # (code, name, bond, cat, cases)
BONDS = ["KOLLAM","KOZHIKODE","ATTINGAL","PALAKKAD","KOTTARAKARA","KANNUR","ALAPPUZHA",
         "NEDUMANGAD","ALUVA","PERINTHALMANNA","THODUPUZHA","PATHANAMTHITTA",
         "KOTTAYAM","THRISSUR","TRIPUNITHURA"]
for bond in BONDS:
    s = wb[bond]
    for r in range(5, s.max_row+1):
        code = s.cell(row=r, column=1).value
        if code is None or code == "": continue
        name = s.cell(row=r, column=2).value
        if name == "TOTAL": continue
        staff = s.cell(row=r, column=3).value
        cat = s.cell(row=r, column=4).value
        cases = s.cell(row=r, column=5).value or 0
        outlets.append({"code": code, "name": name, "bond": bond, "cat": cat, "cases": cases})

outlets_sorted = sorted(outlets, key=lambda x: -x["cases"])
top5 = outlets_sorted[:5]

# ==== Top 5 tables — layout with merged Outlet column and gap before Bond ====
# Left table cols: B(Rank) | C:D(Outlet merged) | E(gap) | F(Bond) | G(CAT) | H:I(Cases merged)
# Right table cols: K(Rank) | L:M(Outlet merged) | N(gap) | O(Bond) | P(CAT) | Q:R(Cases merged)

ws.merge_cells("B38:I38")
c = ws["B38"]; c.value = "🚀 TOP 5 OUTLETS BY SECONDARY SALES"
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=13, bold=True, color=GREEN)
c.alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("K38:R38")
c = ws["K38"]; c.value = "💎 TOP 5 INVOICE OUTLETS (FED+BAR)"
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=13, bold=True, color=CYAN)
c.alignment = Alignment(horizontal="center", vertical="center")

# Row 39 headers — each "Outlet"/"Cases" cell is merged across two cols; E and N are gap
ws.merge_cells("C39:D39"); ws.merge_cells("H39:I39")
ws.merge_cells("L39:M39"); ws.merge_cells("Q39:R39")
l_hdrs = [("B","Rank"),("C","Outlet"),("F","Bond"),("G","CAT"),("H","Cases")]
r_hdrs = [("K","Rank"),("L","Outlet"),("O","Bond"),("P","CAT"),("Q","Cases")]
for col, h in l_hdrs:
    c = ws[f"{col}39"]; c.value = h
    c.fill = PatternFill("solid", fgColor=DARK_BG)
    c.font = Font(name=FONT, size=10, bold=True, color=GREEN)
    c.alignment = Alignment(horizontal="center", vertical="center")
for col, h in r_hdrs:
    c = ws[f"{col}39"]; c.value = h
    c.fill = PatternFill("solid", fgColor=DARK_BG)
    c.font = Font(name=FONT, size=10, bold=True, color=CYAN)
    c.alignment = Alignment(horizontal="center", vertical="center")

# Top 5 by total (left)
for i, o in enumerate(top5):
    r = 40 + i
    zebra = ROW_A if i % 2 == 0 else ROW_B
    ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"H{r}:I{r}")
    ws[f"B{r}"].value = i+1
    ws[f"C{r}"].value = o["name"]
    ws[f"F{r}"].value = o["bond"]
    ws[f"G{r}"].value = o["cat"]
    ws[f"H{r}"].value = o["cases"]
    # Fill EVERY cell in this row B:I (so gap E shows the zebra, not dark bg)
    for col in ["B","C","D","E","F","G","H","I"]:
        c = ws[f"{col}{r}"]
        c.fill = PatternFill("solid", fgColor=zebra)
        c.font = Font(name=FONT, size=11, color=WHITE)
    ws[f"B{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"F{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"G{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"H{r}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws[f"H{r}"].number_format = "#,##0.00"

# Top 5 invoice (FED+BAR only, right)
top5_inv = [o for o in outlets_sorted if o["cat"] in ("FED","BAR")][:5]
for i, o in enumerate(top5_inv):
    r = 40 + i
    zebra = ROW_A if i % 2 == 0 else ROW_B
    ws.merge_cells(f"L{r}:M{r}"); ws.merge_cells(f"Q{r}:R{r}")
    ws[f"K{r}"].value = i+1
    ws[f"L{r}"].value = o["name"]
    ws[f"O{r}"].value = o["bond"]
    ws[f"P{r}"].value = o["cat"]
    ws[f"Q{r}"].value = o["cases"]
    for col in ["K","L","M","N","O","P","Q","R"]:
        c = ws[f"{col}{r}"]
        c.fill = PatternFill("solid", fgColor=zebra)
        c.font = Font(name=FONT, size=11, color=WHITE)
    ws[f"K{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"L{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"O{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"P{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"Q{r}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws[f"Q{r}"].number_format = "#,##0.00"

# -------- Bond league table (rows 47-63) --------
# Bond League title — merge over exactly the league-table width (B:I, 8 cols)
ws.merge_cells("B47:I47")
c = ws["B47"]; c.value = "🏅 BOND LEAGUE TABLE"
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=14, bold=True, color=CYAN)
c.alignment = Alignment(horizontal="center", vertical="center")

# Header row 48
league_hdrs = [("B","#"),("C","Bond"),("D","Outlets"),("E","KSBC"),("F","FED"),
               ("G","BAR"),("H","Total"),("I","% Total")]
for col, h in league_hdrs:
    c = ws[f"{col}48"]; c.value = h
    c.fill = PatternFill("solid", fgColor=DARK_BG)
    c.font = Font(name=FONT, size=11, bold=True, color=CYAN2)
    c.alignment = Alignment(horizontal="center", vertical="center")

# Count outlets per bond from master (active only)
# Read region sheets (already has active outlets)
outlets_per_bond = defaultdict(int)
for o in outlets:
    outlets_per_bond[o["bond"]] += 1

for i, bd in enumerate(bond_by_total):
    r = 49 + i
    zebra = ROW_A if i % 2 == 0 else ROW_B
    pct_of_total = bd["tot"]/grand_total if grand_total else 0
    cells = {
        "B": i+1,
        "C": bd["bond"],
        "D": outlets_per_bond.get(bd["bond"], 0),
        "E": bd["k"],
        "F": bd["f"],
        "G": bd["b"],
        "H": bd["tot"],
        "I": pct_of_total,
    }
    for col, v in cells.items():
        c = ws[f"{col}{r}"]
        c.value = v
        c.fill = PatternFill("solid", fgColor=zebra)
        c.font = Font(name=FONT, size=10, color=WHITE)
        if col == "C":
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        elif col == "B":
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        if col in ("E","F","G","H"):
            c.number_format = "#,##0.00"
        if col == "I":
            c.number_format = "0.0%"

# Footnote row 66 — centered over the league table width (B:I)
ws.merge_cells("B66:I66")
c = ws["B66"]
note = f"Dashboard generated from {MONTH} 1–{LATEST_DAY} COMBINED secondary sales data  •  Drill into region sheets for outlet-level detail  •  {un_count} unmatched licensee(s) flagged in UNMATCHED sheet" if un_count else f"Dashboard generated from {MONTH} 1–{LATEST_DAY} COMBINED secondary sales data  •  Drill into region sheets for outlet-level detail"
c.value = note
c.fill = PatternFill("solid", fgColor=DARK_BG)
c.font = Font(name=FONT, size=10, italic=True, color=MUTED)
c.alignment = Alignment(horizontal="center", vertical="center")

# Save
wb.save(SCRATCH)
print(f"Saved: {SCRATCH}")
print(f"Totals: grand={grand_total}, cashflow={cashflow}, ksbc={grand_k}, fed={grand_f}, bar={grand_b}")
print(f"Top sales bond: {top_sales['bond']} ({top_sales['tot']:,.0f})")
print(f"Top cashflow bond: {top_cash['bond']} ({top_cash['cash']:,.0f})")
print(f"Top brand: {top_brand} ({top_brand_cases})")
print(f"Peak day: {peak_day_name} ({peak_day_cases})")
print(f"Top 5 outlets: {[(o['name'], o['cases']) for o in top5]}")
print(f"Top 5 invoice outlets: {[(o['name'], o['cases']) for o in top5_inv]}")
