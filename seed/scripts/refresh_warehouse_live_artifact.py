#!/usr/bin/env python3
"""Rebuild the ksd-warehouse-live artifact HTML by embedding the latest
stock_history.csv + inbound_history.csv as a JS PAYLOAD constant.

The artifact iframe can't call workspace.bash (sandboxed), so the data is
baked into the HTML at refresh time. The warehouse-stock scheduled task
runs this script after every successful BUILT, then calls
mcp__cowork__update_artifact with the path printed on the last stdout line.

Locked: 12 May 2026. Updated 17 May 2026 to additionally embed a per-warehouse
3-month average secondary-sales aggregate so the artifact can render a
"Stock vs Sales Demand" section. The 3 months auto-roll: they are the most
recent COMPLETE months before the latest stock date (was hardcoded Feb–Apr
2026 until 17 Jun 2026). Single source of truth for the artifact HTML is
.claude/scripts/warehouse_live_template.html — edit there if the layout changes.

Usage:
    python3 refresh_warehouse_live_artifact.py [--out PATH]

Prints:
    Multi-line summary, then a final line:
        OUT: <absolute path to refreshed HTML>
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from pathlib import Path
from datetime import datetime, date


def claude_root() -> Path:
    """Locate the Claude workspace folder (parent of .claude/)."""
    here = Path(__file__).resolve()
    # .../<Claude>/.claude/scripts/<this file>
    return here.parent.parent.parent


def history_dir(root: Path) -> Path:
    return root / "Warehouse stock" / "_history"


def template_path() -> Path:
    return Path(__file__).resolve().parent / "warehouse_live_template.html"


def default_out_path() -> Path:
    """Prefer session outputs dir for scheduled runs; fall back to a script-
    sibling path for ad-hoc invocations."""
    home = os.environ.get("HOME", "")
    if home:
        outputs = Path(home) / "mnt" / "outputs"
        if outputs.is_dir():
            return outputs / "warehouse_live.html"
    return Path(__file__).resolve().parent / "_warehouse_live_refreshed.html"


def read_stock(path: Path):
    out = []
    with path.open() as f:
        for row in csv.DictReader(f):
            out.append({
                "d": row["date"],
                "w": row["warehouse"],
                "p": int(float(row["physical"] or 0)),
                "a": int(float(row["allotable"] or 0)),
                "n": int(float(row["pending"] or 0)),
            })
    return out


def read_inbound(path: Path):
    out = []
    if not path.exists():
        return out
    with path.open() as f:
        for row in csv.DictReader(f):
            out.append({
                "d": row["date"],
                "w": row["warehouse"],
                "in": int(float(row["inbound_cases"] or 0)),
                "out": int(float(row["dispatched_cases"] or 0)),
            })
    return out


def read_brand_trend(path: Path):
    """Network-level per-day brand aggregates from brand_pack_history.csv.
    Compact rows [date, brand, phys, allot, pend] so the artifact can render
    the Brand Mix panel + 7-day brand movers (added 9 Jul 2026).
    Returns [] if the CSV doesn't exist yet."""
    if not path.exists():
        return []
    agg = {}
    with path.open() as f:
        for row in csv.DictReader(f):
            try:
                phys = float(row["physical"] or 0)
                allot = float(row["allotable"] or 0)
                pend = float(row["pending"] or 0)
            except (TypeError, ValueError):
                continue
            key = (row["date"], row["brand"])
            cur = agg.setdefault(key, [0.0, 0.0, 0.0])
            cur[0] += phys
            cur[1] += allot
            cur[2] += pend
    return [[d, b, round(v[0]), round(v[1]), round(v[2])]
            for (d, b), v in sorted(agg.items())]


def read_pack_trend(path: Path):
    """Network-level per-day PACK aggregates from brand_pack_history.csv
    (added 9 Jul 2026 PM for the pack-level Brand & Pack Mix panel).
    Compact rows [date, pack, phys, allot, pend]; [] if the CSV is absent."""
    if not path.exists():
        return []
    agg = {}
    with path.open() as f:
        for row in csv.DictReader(f):
            try:
                phys = float(row["physical"] or 0)
                allot = float(row["allotable"] or 0)
                pend = float(row["pending"] or 0)
            except (TypeError, ValueError):
                continue
            key = (row["date"], row["pack"])
            cur = agg.setdefault(key, [0.0, 0.0, 0.0])
            cur[0] += phys
            cur[1] += allot
            cur[2] += pend
    return [[d, k, round(v[0]), round(v[1]), round(v[2])]
            for (d, k), v in sorted(agg.items())]


def read_mix_by_wh(path: Path):
    """Per-WAREHOUSE per-day brand and pack aggregates (added 21 Jul 2026 so the
    Brand & Pack Mix panel can be filtered by warehouse/cluster, not just read
    at network level).

    Index-compressed because the flat form is ~9,900 rows: names are hoisted
    into `wh` / `brand` / `pack` / `date` lists and every row carries integer
    positions into them. All-zero rows are dropped. Roughly 125 KB for a month
    of history against ~800 KB flat.

        {"date": [...], "wh": [...], "brand": [...], "pack": [...],
         "b": [[dateI, whI, brandI, phys, allot, pend], ...],
         "p": [[dateI, whI, packI,  phys, allot, pend], ...]}

    Returns None if the CSV doesn't exist yet -- the artifact then falls back to
    the network-level brandTrend/packTrend and hides the scope control."""
    if not path.exists():
        return None
    rows = []
    with path.open() as f:
        for row in csv.DictReader(f):
            try:
                rows.append((row["date"], row["warehouse"], row["brand"], row["pack"],
                             float(row["physical"] or 0), float(row["allotable"] or 0),
                             float(row["pending"] or 0)))
            except (TypeError, ValueError):
                continue
    if not rows:
        return None

    dates = sorted({r[0] for r in rows})
    whs = sorted({r[1] for r in rows})
    brands = sorted({r[2] for r in rows})
    packs = sorted({r[3] for r in rows}, key=_pack_ml)
    di = {v: i for i, v in enumerate(dates)}
    wi = {v: i for i, v in enumerate(whs)}
    bi = {v: i for i, v in enumerate(brands)}
    pi = {v: i for i, v in enumerate(packs)}

    def fold(key_idx, lookup):
        agg = {}
        for r in rows:
            k = (di[r[0]], wi[r[1]], lookup[r[key_idx]])
            cur = agg.setdefault(k, [0.0, 0.0, 0.0])
            cur[0] += r[4]
            cur[1] += r[5]
            cur[2] += r[6]
        out = []
        for k in sorted(agg):
            v = [round(x) for x in agg[k]]
            if any(v):                      # a zeroed line carries no signal
                out.append(list(k) + v)
        return out

    return {"date": dates, "wh": whs, "brand": brands, "pack": packs,
            "b": fold(2, bi), "p": fold(3, pi)}


def _pack_ml(pack: str) -> int:
    """180 ML -> 180, so packs sort by size rather than alphabetically."""
    digits = "".join(ch for ch in str(pack) if ch.isdigit())
    return int(digits) if digits else 9999


def read_sku_detail(path: Path):
    """Latest-date brand × pack × warehouse detail from brand_pack_history.csv
    (written by build_warehouse_stock.py since 10 Jun 2026). Returns
    {"date": iso, "byWarehouse": {WH: [[brand, pack, phys, allot, pend], ...]}}
    or None if the CSV doesn't exist yet (first capture = next stock upload)."""
    if not path.exists():
        return None
    rows = []
    with path.open() as f:
        for row in csv.DictReader(f):
            rows.append(row)
    if not rows:
        return None
    latest = max(r["date"] for r in rows)
    by_wh = {}
    for r in rows:
        if r["date"] != latest:
            continue
        try:
            phys = round(float(r["physical"] or 0))
            allot = round(float(r["allotable"] or 0))
            pend = round(float(r["pending"] or 0))
        except (TypeError, ValueError):
            continue
        by_wh.setdefault(r["warehouse"], []).append(
            [r["brand"], r["pack"], phys, allot, pend])
    for wh in by_wh:
        by_wh[wh].sort(key=lambda x: -x[2])  # physical desc
    return {"date": latest, "byWarehouse": by_wh}


# --- Secondary-sales 3-month per-warehouse aggregation -----------------------
#
# Pulls warehouse-level dispatch totals (Issue Cases) from the per-month
# Secondary Sales analysis workbooks. Sheet name varies slightly between
# months; we accept any sheet ending in "COMBINED DISPATCHES" or any sheet
# named "<MONTH> RAW" — both have the same column layout (col 13 = Issue
# Cases, col 2 = Warehouse Name).
#
# Warehouse Name is shaped `WH-<CITY> <license-id-fragment>`. We strip to the
# CITY token so it matches the stock_history.csv `warehouse` field exactly.
# Verified 17 May 2026: all 28 warehouses on the stock side align to the
# 28 warehouse-name buckets on the dispatch side with no manual aliasing.

MONTH_ABBR = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
MONTH_FULL_UP = ['JANUARY','FEBRUARY','MARCH','APRIL','MAY','JUNE',
                 'JULY','AUGUST','SEPTEMBER','OCTOBER','NOVEMBER','DECEMBER']


def recent_complete_months(as_of_iso, n=3):
    """The n most recent COMPLETE calendar months before as_of's month,
    chronological. e.g. as_of 2026-06-16 -> [(2026,3),(2026,4),(2026,5)].
    Cross-year safe (wraps the year)."""
    dt = datetime.strptime(as_of_iso, "%Y-%m-%d")
    y, mo = dt.year, dt.month
    out = []
    for _ in range(n):
        mo -= 1
        if mo == 0:
            mo = 12; y -= 1
        out.append((y, mo))
    out.reverse()
    return out


def sales_period_label(months_ym):
    """'Mar–May 2026 avg' from the (year,month) list actually loaded."""
    if not months_ym:
        return "3-mo avg"
    (y0, m0), (y1, m1) = months_ym[0], months_ym[-1]
    if y0 == y1:
        return f"{MONTH_ABBR[m0-1]}–{MONTH_ABBR[m1-1]} {y1} avg"
    return f"{MONTH_ABBR[m0-1]} {y0}–{MONTH_ABBR[m1-1]} {y1} avg"


def _warehouse_city(name):
    if not isinstance(name, str):
        return None
    m = re.match(r"\s*WH-([A-Z]+)\b", name.strip().upper())
    return m.group(1) if m else None


def _candidate_dispatch_sheet(sheet_names, month: str):
    upper = month.upper()
    for s in sheet_names:
        if s.upper() == f"{upper} COMBINED DISPATCHES":
            return s
    for s in sheet_names:
        if s.upper() == f"{upper} RAW":
            return s
    return None


def _find_secondary_month_file(secondary_dir: Path, month_upper: str):
    """Return (path, kind) for a month's Secondary analysis workbook: the clean
    full-month file if present, else the newest mid-month
    '<MONTH> 1st - Nth SECONDARY SALES ANALYSIS.xlsx' file, else (None, None).

    Why the mid-month fallback (added 10 Jul 2026): the 3-mo cover average only
    ever asks for COMPLETE past months, but at the start of each month the just-
    finished month's workbook may not yet be renamed from its mid-month form to
    the clean '<MONTH> SECONDARY SALES ANALYSIS.xlsx'. A completed past month's
    mid-file still carries the WHOLE month's dispatches (the rename is cosmetic),
    so falling back to it keeps the denominator at 3 months instead of silently
    dropping to 2 — which was skewing cover/tiers high for the first days of
    every month."""
    full = secondary_dir / f"{month_upper} SECONDARY SALES ANALYSIS.xlsx"
    if full.exists():
        return full, 'full'
    mids = sorted(
        secondary_dir.glob(f"{month_upper} 1st - *SECONDARY SALES ANALYSIS.xlsx"),
        key=lambda p: p.stat().st_mtime, reverse=True)
    if mids:
        return mids[0], 'mid'
    return None, None


def build_sales_avg(root: Path, as_of_iso: str):
    """Return per-warehouse 3-mo avg monthly sales (cases) plus per-month
    breakdown for tooltips. The 3 months are the most recent COMPLETE months
    before as_of and auto-roll every month (was hardcoded Feb–Apr until 17 Jun
    2026). Falls back to empty dict if openpyxl or files are missing — the
    artifact still renders, just without the demand section."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("WARN: openpyxl not available — skipping 3-mo sales aggregation",
              file=sys.stderr)
        return {"byWarehouse": {}, "monthsUsed": [], "periodLabel": "3-mo avg"}

    secondary_dir = root / "Secondary sales"
    candidates = recent_complete_months(as_of_iso)   # [(year, month_idx), ...]
    by_wh = {}
    months_used = []        # UPPER month names actually loaded
    months_used_ym = []     # (year, month_idx) actually loaded
    for (yy, mm) in candidates:
        m = MONTH_FULL_UP[mm - 1]
        path, kind = _find_secondary_month_file(secondary_dir, m)
        if path is None:
            print(f"WARN: no Secondary workbook (full or mid-month) for {m} "
                  f"— skipping; 3-mo cover will use fewer months",
                  file=sys.stderr)
            continue
        if kind == 'mid':
            print(f"NOTE: {m} cover sourced from mid-month file '{path.name}' "
                  f"(full-month not yet renamed) — a completed past month's "
                  f"mid-file still holds the whole month's dispatches",
                  file=sys.stderr)
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet_name = _candidate_dispatch_sheet(wb.sheetnames, m)
        if not sheet_name:
            print(f"WARN: no dispatch sheet in {path.name} — skipping {m}",
                  file=sys.stderr)
            wb.close()
            continue
        ws = wb[sheet_name]
        rows_seen = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 13:
                continue
            wh_city = _warehouse_city(row[1])
            if not wh_city:
                continue
            try:
                cases = int(float(row[12] or 0))
            except (TypeError, ValueError):
                cases = 0
            if cases <= 0:
                continue
            bucket = by_wh.setdefault(wh_city, {})
            bucket[m] = bucket.get(m, 0) + cases
            rows_seen += 1
        wb.close()
        months_used.append(m)
        months_used_ym.append((yy, mm))
        print(f"  sales[{m}]: {rows_seen:,} dispatch rows from '{sheet_name}'",
              file=sys.stderr)

    if not months_used:
        return {"byWarehouse": {}, "monthsUsed": [], "periodLabel": "3-mo avg"}

    label = sales_period_label(months_used_ym)
    divisor = len(months_used)
    out_by_wh = {}
    for wh, months in by_wh.items():
        total = sum(months.values())
        out_by_wh[wh] = {
            "avg": int(round(total / divisor)),
            "total": total,
            "months": {m: months.get(m, 0) for m in months_used},
        }
    return {
        "byWarehouse": out_by_wh,
        "monthsUsed": months_used,
        "periodLabel": label,
    }


def _bpc_from_pack(pack) -> int:
    m = re.search(r"(\d+)", str(pack or ""))
    ml = int(m.group(1)) if m else 0
    return {180: 48, 375: 24, 500: 18, 750: 12, 1000: 9}.get(ml, 12)


def _parse_gtn_date(v):
    if isinstance(v, datetime):
        return v.date()
    s = str(v or "").strip()
    m = re.match(r"(\d{1,2})-(\d{1,2})-(\d{4})", s)
    if m:
        try:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date()
        except ValueError:
            return None
    return None


def _shop_cat_map(root: Path):
    """shop code -> 'k' (KSBC) or 'i' (CFD/BAR invoice) from MASTER DATA
    CONFIRMED.xlsx sheet 16-4-25 (title row 1, header row 2: Shop Code idx 3,
    CAT idx 5). Non-fatal — empty map means the first-digit fallback rules
    (1xx=KSBC, 2xx=FED, 4xx=BAR) carry the split alone."""
    out = {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(root / "MASTER DATA CONFIRMED.xlsx",
                           read_only=True, data_only=True)
        ws = wb["16-4-25"] if "16-4-25" in wb.sheetnames else wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row or len(row) < 6 or row[3] is None:
                continue
            code = str(row[3]).strip()
            if code.endswith(".0"):
                code = code[:-2]
            cat = str(row[5] or "").strip().upper()
            if code and cat:
                out[code] = 'k' if cat == 'KSBC' else 'i'
        wb.close()
    except Exception as e:
        print(f"WARN: shop CAT map unavailable ({e}) — first-digit fallback only",
              file=sys.stderr)
    return out


def build_daily_dispatch(root: Path, as_of_iso: str):
    """Per-day NETWORK dispatch cases (Issue Cases + loose bottles folded via
    BPC) for the as-of month and the month before, read straight from the
    Secondary COMBINED DISPATCHES sheets by Inv/GTN Date — the same per-date
    basis the liquidation artifact's daily trend uses (the windowed
    inbound_history 'out' figures are snapshot-window lumped, so they can't
    feed a daily chart). Feeds the Warehouse Daily Sales section (24 Jul
    2026, Abhay). Non-fatal: a missing month is simply omitted."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    dt = datetime.strptime(as_of_iso, "%Y-%m-%d")
    prev = (dt.year - 1, 12) if dt.month == 1 else (dt.year, dt.month - 1)
    secondary_dir = root / "Secondary sales"
    cat_map = _shop_cat_map(root)
    out = {}
    for key, (yy, mm) in (("cur", (dt.year, dt.month)), ("pri", prev)):
        m = MONTH_FULL_UP[mm - 1]
        path, _kind = _find_secondary_month_file(secondary_dir, m)
        if path is None:
            print(f"WARN: dailyDispatch — no Secondary workbook for {m}",
                  file=sys.stderr)
            continue
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            print(f"WARN: dailyDispatch could not open {path.name}: {e}",
                  file=sys.stderr)
            continue
        sheet = _candidate_dispatch_sheet(wb.sheetnames, m)
        if not sheet:
            wb.close()
            continue
        days, kd, idd = {}, {}, {}
        unmatched = 0
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 14:
                continue
            d = _parse_gtn_date(row[11])
            if d is None or d.year != yy or d.month != mm:
                continue
            try:
                cases = float(row[12] or 0)
            except (TypeError, ValueError):
                cases = 0.0
            try:
                btl = float(row[13] or 0)
            except (TypeError, ValueError):
                btl = 0.0
            if btl:
                cases += btl / _bpc_from_pack(row[5])
            if cases <= 0:
                continue
            days[d.day] = days.get(d.day, 0.0) + cases
            # channel split (24 Jul 2026, Abhay: "mix of KSBC and Cfed just
            # like the liquidation artifact") — master CAT first, first-digit
            # rule as fallback (1xx=KSBC shops, 2xx=FED, 4xx=BAR)
            code = str(row[6] or "").strip()
            if code.endswith(".0"):
                code = code[:-2]
            cat = cat_map.get(code)
            if cat is None:
                unmatched += 1
                cat = 'k' if code.startswith('1') else 'i'
            tgt = kd if cat == 'k' else idd
            tgt[d.day] = tgt.get(d.day, 0.0) + cases
        wb.close()
        if days:
            out[key] = {"m": m, "y": yy, "mi": mm,
                        "days": {str(k): round(v, 2)
                                 for k, v in sorted(days.items())},
                        "ksbc": {str(k): round(v, 2)
                                 for k, v in sorted(kd.items())},
                        "inv": {str(k): round(v, 2)
                                for k, v in sorted(idd.items())}}
            print(f"  dailyDispatch[{m}]: {len(days)} days, "
                  f"{round(sum(days.values())):,} cs "
                  f"(KSBC {round(sum(kd.values())):,} / "
                  f"invoice {round(sum(idd.values())):,}"
                  + (f" · {unmatched} rows off-master" if unmatched else "")
                  + ")", file=sys.stderr)
    return out or None


def read_factory_loads(root: Path):
    """Factory truck loads from Load dispatch/LOAD DISPATCH TRACKER.xlsx
    (sheet DISPATCH LOG). Abhay updates it daily (22 Jul 2026); the artifact
    cross-checks derived inbound against it and flags big gaps. Non-fatal:
    missing/locked tracker just omits the payload key."""
    out = []
    _bad_dates = []
    path = root / "Load dispatch" / "LOAD DISPATCH TRACKER.xlsx"
    ALIAS = {"KADAVANTRA": "KADAVANTHRA", "KOTHMAGALAM": "KOTHAMANGALAM",
             "KOTHAMANGLAM": "KOTHAMANGALAM", "PATHANAMTHITTA": "PATHANAMTHITA",
             "NEDUMANGADU": "NEDUMANGAD", "TRIPUNITURA": "TRIPUNITHURA"}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "DISPATCH LOG" not in wb.sheetnames:
            return out
        ws = wb["DISPATCH LOG"]

        def d2s(v):
            if v is None:
                return None
            try:
                return v.strftime("%Y-%m-%d")
            except Exception:
                pass
            s = str(v).strip()
            if not s:
                return None
            if len(s) >= 10 and s[4:5] == "-":
                return s[:10]
            for _fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y"):
                try:
                    return datetime.strptime(s, _fmt).strftime("%Y-%m-%d")
                except ValueError:
                    continue
            _bad_dates.append(s)   # non-empty but unparseable -> flag, never silent
            return None

        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r or len(r) < 4 or r[2] is None or r[3] is None:
                continue
            w = ALIAS.get(str(r[2]).strip().upper(), str(r[2]).strip().upper())
            dd = d2s(r[1])
            ud = d2s(r[4]) if len(r) > 4 else None
            try:
                cs = float(r[3])
            except Exception:
                continue
            if not dd or cs <= 0:
                continue
            out.append({"d": dd, "w": w, "cs": round(cs, 1), "u": ud,
                        "st": str(r[5] or "").strip() if len(r) > 5 else ""})
    except Exception as e:  # noqa: BLE001 — never block the build on the tracker
        print(f"WARN: factory load tracker skipped ({e})")
    if _bad_dates:
        import collections as _c
        _samp = ", ".join(f"{v} (x{n})" for v, n in _c.Counter(_bad_dates).most_common(5))
        print(f"WARN: {len(_bad_dates)} unparseable date cell(s) in LOAD DISPATCH TRACKER "
              f"— a bad Dispatch date drops the load, a bad Unload date reads it as in-transit: {_samp}")
    return out


# ------------------------------------------------------ GRS date matching --
# Abhay, 10 Aug 2026. Bevco's GRS books a receipt on its own clock — a truck
# unloaded at the yard on Friday shows up in the Bevco export the NEXT working
# snapshot. Across the whole register that lag runs 0-4 days. Almost always
# both dates land in the same month and nothing is visible, but a truck
# unloaded near month-end books in the FOLLOWING month, and then the truck log
# and the derived figure disagree in BOTH months (KOTTAYAM 700 cs, unloaded
# 31 Jul, booked 4 Aug — the 699-case mismatch on the Inbound tile).
#
# The rule Abhay set: take the QUANTITY from the truck log, take the TIMING
# from the raw data. Each load is matched to the Bevco inbound event at the
# same warehouse whose quantity agrees, on/after its unload date, and carries
# that booking date as its effective date ("g"). A load the books have not
# picked up yet is held OUT of inbound and surfaces as an awaiting-GRS chip;
# past GRS_STALE_DAYS it flags red, because a truck that unloaded and never
# got booked is a real problem, not a timing artefact.
GRS_TOL_CASES = 2.0     # Bevco rounds: a 720 truck books as 719 often enough
GRS_TOL_PCT = 0.01
GRS_LOOKAHEAD = 10      # max observed lag is 4d; beyond 10 it is not this load
GRS_STALE_DAYS = 7      # unloaded this long ago and still unbooked -> flag


def _dnum(s):
    y, m, d = int(s[0:4]), int(s[5:7]), int(s[8:10])
    return date(y, m, d).toordinal()


def _grs_tol(total, parts=1):
    """Rounding tolerance for a match. Bevco rounds each booking independently,
    so a match made of `parts` pieces can drift by `parts` roundings."""
    return max(GRS_TOL_CASES * parts, GRS_TOL_PCT * total)


def match_grs_dates(loads, inbound, as_of):
    """Stamp each unloaded truck load with the Bevco booking date that
    represents it. Mutates and returns `loads`. Fields added per load:
        g     effective GRS date (None while the books have not caught up)
        gl    lag in days from unload to booking
        gp    how the match was made: 1:1, N:1 (combined) or 1:N (split)
        stale True when unloaded > GRS_STALE_DAYS ago and still unbooked

    THREE PASSES, in order of confidence — a later pass only ever sees loads
    and bookings the earlier ones left behind, so adding them can never steal a
    booking from a clean one-to-one match:

      1. 1:1  one truck, one booking of the same size. Greedy by unload date,
              so two same-size trucks to one warehouse consume two separate
              bookings in order.
      2. N:1  several trucks booked by Bevco as ONE larger receipt. Real case
              that forced this (31 Aug 2026): two ALAPPUZHA 720s unloaded
              23 Aug were booked as a single 1,440 on 24 Aug, so neither found
              a same-size partner, both were held out of inbound, and at 8 days
              both flagged red — while the stock had plainly landed (physical
              39 -> 1,472 that day). The tile read 9,170 cs / -15.8% when the
              books said 10,608 / -2.5%.
      3. 1:N  the mirror case — one truck the books break into several smaller
              receipts. Its effective date is the LAST of them, the point the
              booking is actually complete.
    """
    if not loads or not inbound:
        return loads
    events = [{"d": r["d"], "w": (r["w"] or "").strip().upper(),
               "cs": float(r["in"]), "used": False}
              for r in inbound if float(r["in"] or 0) > 0]
    events.sort(key=lambda e: e["d"])
    # Coverage start = the first day the books cover AT ALL, which is the first
    # row of any kind in inbound_history — NOT the first day that happens to
    # carry a positive receipt. Using the latter made every load unloaded
    # before the first booking look like an out-of-coverage load and silently
    # fall back to its own unload date, defeating the whole re-dating rule.
    first = min((r["d"] for r in inbound), default=None)
    m11 = mn1 = m1n = pending = stale = 0

    ordered = sorted([x for x in loads if x.get("u")],
                     key=lambda x: (x["u"], x["w"]))
    for l in ordered:
        l["g"] = None
        l["gl"] = None
        l["gp"] = None
        l["stale"] = False

    def _assign(load, when):
        load["g"] = when
        load["gl"] = _dnum(when) - _dnum(load["u"])

    # --- pass 1: one truck, one booking ----------------------------------
    for l in ordered:
        # loads that unloaded before the inbound history begins can never be
        # matched — fall back to the unload date rather than stranding them
        if first and l["u"] < first:
            l["g"] = l["u"]
            l["gl"] = 0
            l["gp"] = "1:1"
            continue
        tol = _grs_tol(l["cs"])
        for e in events:            # sorted, so the first hit is the earliest
            if e["used"] or e["w"] != l["w"] or e["d"] < l["u"]:
                continue
            if _dnum(e["d"]) - _dnum(l["u"]) > GRS_LOOKAHEAD:
                break
            if abs(e["cs"] - l["cs"]) > tol:
                continue
            e["used"] = True
            _assign(l, e["d"])
            l["gp"] = "1:1"
            m11 += 1
            break

    # --- pass 2: several trucks booked as ONE larger receipt -------------
    # Walk the unused bookings earliest-first. For each, accumulate the
    # still-unmatched loads at that warehouse that could feed it (unloaded on
    # or before it, inside the lookahead), oldest first, and take the group the
    # moment its running total reaches the booking. Requires 2+ loads — a
    # single one would have matched in pass 1.
    for e in events:
        if e["used"]:
            continue
        feeders = [l for l in ordered
                   if l["g"] is None and l["w"] == e["w"]
                   and l["u"] <= e["d"]
                   and _dnum(e["d"]) - _dnum(l["u"]) <= GRS_LOOKAHEAD]
        if len(feeders) < 2:
            continue
        run, group = 0.0, []
        for l in feeders:
            run += l["cs"]
            group.append(l)
            if len(group) >= 2 and abs(run - e["cs"]) <= _grs_tol(run, len(group)):
                e["used"] = True
                for g in group:
                    _assign(g, e["d"])
                    g["gp"] = f"{len(group)}:1"
                mn1 += len(group)
                break
            if run > e["cs"] + _grs_tol(run, len(group) + 1):
                break   # overshot; no subset starting here can land on it

    # --- pass 3: one truck the books split across several receipts -------
    # Effective date is the LAST piece: that is when the load is fully booked.
    for l in ordered:
        if l["g"] is not None:
            continue
        parts, run = [], 0.0
        for e in events:
            if e["used"] or e["w"] != l["w"] or e["d"] < l["u"]:
                continue
            if _dnum(e["d"]) - _dnum(l["u"]) > GRS_LOOKAHEAD:
                break
            run += e["cs"]
            parts.append(e)
            if len(parts) >= 2 and abs(run - l["cs"]) <= _grs_tol(run, len(parts)):
                for p in parts:
                    p["used"] = True
                _assign(l, parts[-1]["d"])
                l["gp"] = f"1:{len(parts)}"
                m1n += 1
                break
            if run > l["cs"] + _grs_tol(run, len(parts) + 1):
                break

    # --- whatever is left is genuinely not on the books yet --------------
    for l in ordered:
        if l["g"] is None:
            pending += 1
            if as_of and _dnum(as_of) - _dnum(l["u"]) > GRS_STALE_DAYS:
                l["stale"] = True
                stale += 1

    crossed = sum(1 for l in loads if l.get("g") and l["g"][:7] != l["u"][:7])
    extra = []
    if mn1:
        extra.append(f"{mn1} via combined booking")
    if m1n:
        extra.append(f"{m1n} via split booking")
    print(f"GRS match: {m11 + mn1 + m1n} load(s) booked"
          f"{' (' + ', '.join(extra) + ')' if extra else ''}"
          f", {pending} awaiting GRS"
          f"{f' ({stale} STALE >{GRS_STALE_DAYS}d)' if stale else ''}"
          f", {crossed} crossing a month boundary")
    for l in loads:
        if l.get("stale"):
            print(f"  WARN: {l['w']} {l['cs']:.0f} cs unloaded {l['u']} "
                  f"— still not booked by Bevco as of {as_of}")
    return loads
def build_payload(root: Path):
    hist = history_dir(root)
    stock_csv = hist / "stock_history.csv"
    inbound_csv = hist / "inbound_history.csv"
    if not stock_csv.exists():
        raise SystemExit(f"stock_history.csv not found at {stock_csv}")
    stock = read_stock(stock_csv)
    inbound = read_inbound(inbound_csv)
    if not stock:
        raise SystemExit(f"stock_history.csv is empty at {stock_csv}")
    as_of = max(r["d"] for r in stock)
    sales = build_sales_avg(root, as_of)
    sku_detail = read_sku_detail(hist / "brand_pack_history.csv")
    return {
        "stock": stock,
        "inbound": inbound,
        "asOf": as_of,
        "sales": sales,
        "skuDetail": sku_detail,
        "brandTrend": read_brand_trend(hist / "brand_pack_history.csv"),
        "packTrend": read_pack_trend(hist / "brand_pack_history.csv"),
        "mixByWh": read_mix_by_wh(hist / "brand_pack_history.csv"),
        "factoryLoads": match_grs_dates(read_factory_loads(root), inbound, as_of),
        "dailyDispatch": build_daily_dispatch(root, as_of),
    }


def write_artifact(payload, out_path: Path) -> Path:
    tpl = template_path()
    if not tpl.exists():
        raise SystemExit(f"Template missing: {tpl}")
    html = tpl.read_text()
    if "{{PAYLOAD}}" not in html:
        raise SystemExit("Template has no {{PAYLOAD}} placeholder")
    payload_js = json.dumps(payload, separators=(",", ":"))
    rendered = html.replace("{{PAYLOAD}}", payload_js)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(rendered)
    return out_path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=None,
                    help="Output HTML path. Defaults to "
                         "$HOME/mnt/outputs/warehouse_live.html or a sibling "
                         "of this script if outputs/ isn't mounted.")
    args = ap.parse_args()

    root = claude_root()
    payload = build_payload(root)
    out_path = args.out or default_out_path()
    written = write_artifact(payload, out_path)

    dates = sorted({r["d"] for r in payload["stock"]})
    warehouses = sorted({r["w"] for r in payload["stock"]
                         if r["d"] == payload["asOf"]})
    sales_wh = payload["sales"]["byWarehouse"]
    sales_total = sum(v["avg"] for v in sales_wh.values())
    print(f"As of:          {payload['asOf']}")
    print(f"Days history:   {len(dates)}")
    print(f"Warehouses:     {len(warehouses)}")
    print(f"Stock rows:     {len(payload['stock'])}")
    print(f"Inbound rows:   {len(payload['inbound'])}")
    print(f"Sales period:   {payload['sales']['periodLabel']} "
          f"({len(payload['sales']['monthsUsed'])} months)")
    print(f"Sales warehouses: {len(sales_wh)}  "
          f"·  3-mo avg total: {sales_total:,} cs/mo")
    sku = payload.get("skuDetail")
    if sku:
        n_sku = sum(len(v) for v in sku["byWarehouse"].values())
        print(f"SKU detail:     {sku['date']} · {len(sku['byWarehouse'])} warehouses · {n_sku} SKU rows")
    else:
        print("SKU detail:     none yet (brand_pack_history.csv starts with next stock upload)")
    print(f"Brand trend:    {len(payload.get('brandTrend') or [])} date-brand rows")
    print(f"Pack trend:     {len(payload.get('packTrend') or [])} date-pack rows")
    mix = payload.get("mixByWh")
    if mix:
        print(f"Mix by WH:      {len(mix['b'])} brand + {len(mix['p'])} pack rows "
              f"across {len(mix['wh'])} warehouses, {len(mix['date'])} days")
    else:
        print("Mix by WH:      none yet (brand_pack_history.csv absent)")
    print(f"Template:       {template_path()}")
    print(f"Bytes written:  {written.stat().st_size}")
    print(f"OUT: {written}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
