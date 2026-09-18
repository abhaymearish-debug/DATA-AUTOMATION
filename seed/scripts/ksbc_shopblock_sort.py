"""
============================================================================
DEPRECATED (1 Jun 2026) — DO NOT USE. NOT part of any pipeline.
----------------------------------------------------------------------------
This script targets the LEGACY two-table (BRAND-WISE + PACK-WISE) DETAIL
layout, which was retired 23 Apr 2026 in favour of the nested Brand→Pack
layout. On a current nested DETAIL sheet it finds no sub-tables and does
nothing; worse, if it ever DID match an old block it would regenerate the
Rating formula with only 4 tiers (no "— No activity"), re-introducing the
24 Apr KOZHIKODE blank-rating regression. Inner ordering is handled by the
nested rebuild (ksbc_detail_nested.py) + ksbc_sort_by_rating.py. The CLI is
gated off below — pass --force only if you know exactly why.
============================================================================

KSBC inner sub-table sort (BRAND-WISE + PACK-WISE inside each DETAIL shop block).

Pair with `ksbc_sort_by_rating.py` (Pass A — outer shops + bonds). This script
is Pass B: the inner sub-tables inside each shop block. Pass A sorts blocks
whole; Pass B sorts the rows INSIDE each block.

Each shop block on a DETAIL sheet has two sub-tables:

    - BRAND-WISE: column-header row -> brand rows -> "TOTAL" row
    - PACK-WISE:  column-header row -> pack  rows -> "TOTAL" row

Sub-table rows carry formulas in cols F (Sell-Through %) and G (Rating) that
reference their own row number. After sorting by sell-through % descending
we regenerate those formulas against the new row indices so computed values
track the moved rows. The TOTAL row (SUM ranges) is NOT moved — only the
data rows in between are shuffled.

Usage:

    python3 .claude/scripts/ksbc_shopblock_sort.py <workbook.xlsx>

Rating order: 🚀 High Performance (≥80%) → ✅ Balanced (≥60%) →
              ⚠️ Inventory Heavy (≥40%) → 🚫 Critical Overstock (<40%)

within each tier, strictly by sell-through % descending.
"""

from __future__ import annotations
import sys
from copy import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_COL = 7            # A..G
DETAIL_START = 4


def _capture(c):
    return {
        "value": c.value,
        "number_format": c.number_format,
        "font": copy(c.font),
        "fill": copy(c.fill),
        "alignment": copy(c.alignment),
        "border": copy(c.border),
        "protection": copy(c.protection),
    }


def _restore(cell, data):
    cell.value = data["value"]
    cell.number_format = data["number_format"]
    cell.font = data["font"]
    cell.fill = data["fill"]
    cell.alignment = data["alignment"]
    cell.border = data["border"]
    cell.protection = data["protection"]


def _sellthrough(opening, receipts, sales):
    denom = (opening or 0) + (receipts or 0)
    return (sales or 0) / denom if denom else 0.0


def _find_blocks(ws):
    headers = []
    for r in range(DETAIL_START, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and " — " in v and "Staff:" in v:
            headers.append(r)
    blocks = []
    for i, h in enumerate(headers):
        end = (headers[i + 1] - 1) if i + 1 < len(headers) else ws.max_row
        blocks.append((h, end))
    return blocks


def _find_subtable_bounds(ws, start, end):
    """Return (brand_data_start, brand_data_end, pack_data_start, pack_data_end)
    where data rows EXCLUDE the column-header row above and the TOTAL row below.
    """
    brand_hdr = start + 2        # after "CODE — …" (row N) and "BRAND-WISE" (row N+1)
    # scan for first TOTAL after brand_hdr
    brand_total = None
    for r in range(brand_hdr + 1, end + 1):
        if ws.cell(r, 1).value == "TOTAL":
            brand_total = r
            break
    if brand_total is None:
        return None
    # find "PACK-WISE"
    pack_label = None
    for r in range(brand_total + 1, end + 1):
        if ws.cell(r, 1).value == "PACK-WISE":
            pack_label = r
            break
    if pack_label is None:
        return None
    pack_hdr = pack_label + 1
    pack_total = None
    for r in range(pack_hdr + 1, end + 1):
        if ws.cell(r, 1).value == "TOTAL":
            pack_total = r
            break
    if pack_total is None:
        return None
    return (
        brand_hdr + 1, brand_total - 1,   # brand data rows
        pack_hdr + 1,  pack_total - 1,    # pack data rows
    )


def _sort_inner_table(ws, data_start, data_end):
    """Sort rows [data_start..data_end] by sell-through % desc and rebuild
    cols F/G formulas for each new row index. Returns True if any reorder
    happened, False otherwise (idempotence check)."""
    if data_end < data_start:
        return False

    # Capture each row (cells 1..7) + numeric sort key
    rows = []
    for r in range(data_start, data_end + 1):
        cells = [_capture(ws.cell(r, c)) for c in range(1, MAX_COL + 1)]
        opening  = ws.cell(r, 2).value or 0
        receipts = ws.cell(r, 3).value or 0
        sales    = ws.cell(r, 4).value or 0
        rows.append((cells, _sellthrough(opening, receipts, sales)))

    original_order = [id(r[0]) for r in rows]
    rows.sort(key=lambda t: t[1], reverse=True)
    new_order = [id(r[0]) for r in rows]
    changed = original_order != new_order

    # Rewrite
    for offset, (cells, _st) in enumerate(rows):
        new_r = data_start + offset
        for c_idx, data in enumerate(cells, start=1):
            cell = ws.cell(new_r, c_idx)
            _restore(cell, data)
        # Regenerate col F (Sell-Through %) + col G (Rating) formulas for this row
        ws.cell(new_r, 6).value = f"=IFERROR(D{new_r}/(B{new_r}+C{new_r}),0)"
        ws.cell(new_r, 7).value = (
            f'=IF(F{new_r}>=0.8,"🚀 High Performance",'
            f'IF(F{new_r}>=0.6,"✅ Balanced",'
            f'IF(F{new_r}>=0.4,"⚠️ Inventory Heavy","🚫 Critical Overstock")))'
        )

    return changed


def sort_detail_sheet(ws):
    blocks = _find_blocks(ws)
    table_count = 0
    changed_count = 0
    for (start, end) in blocks:
        bounds = _find_subtable_bounds(ws, start, end)
        if not bounds:
            continue
        b_start, b_end, p_start, p_end = bounds
        if _sort_inner_table(ws, b_start, b_end):
            changed_count += 1
        table_count += 1
        if _sort_inner_table(ws, p_start, p_end):
            changed_count += 1
        table_count += 1
    return len(blocks), table_count, changed_count


def sort_workbook(path):
    wb = load_workbook(path)
    detail_sheets = [n for n in wb.sheetnames if n.endswith(" DETAIL")]
    summary = {}
    for sn in detail_sheets:
        blocks, tables, changed = sort_detail_sheet(wb[sn])
        summary[sn] = (blocks, tables, changed)
    wb.save(path)
    return summary


if __name__ == "__main__":
    # Deprecation gate (1 Jun 2026): refuse to run without --force so this
    # legacy two-table sorter can't accidentally rewrite a nested DETAIL sheet
    # (or stamp the obsolete 4-tier rating formula). See module docstring.
    args = [a for a in sys.argv[1:] if a != "--force"]
    if "--force" not in sys.argv:
        sys.stderr.write(
            "ksbc_shopblock_sort.py is DEPRECATED and not part of the pipeline. "
            "The nested layout handles inner ordering. Refusing to run. "
            "Pass --force only if you specifically need the legacy two-table sort.\n")
        sys.exit(3)
    if len(args) != 1:
        print("usage: ksbc_shopblock_sort.py <workbook.xlsx> --force")
        sys.exit(2)
    s = sort_workbook(args[0])
    total_changed = 0
    for sheet, (blocks, tables, changed) in s.items():
        print(f"  {sheet}: {blocks} blocks, {tables} sub-tables, {changed} reordered")
        total_changed += changed
    print(f"done. total sub-tables reordered: {total_changed}")
