"""
KSBC month bootstrap (added 2 Jun 2026).

The locked daily pipeline (`ksbc_daily_update.py`) only ever UPDATES an existing
same-month analysis workbook — it has no "start a fresh month" step. So the
first raw of a new month (e.g. `june 1st.xlsx` when only `MAY SHOP SALES
ANALYSIS.xlsx` exists) cannot be ingested until a month workbook exists.

This script creates that base by templating from the most recent full-month
workbook: it copies the prior month's workbook and STRIPS every month-specific
data sheet (daily `<MONTH> <N>`, `*CUMULATIVE*`, `*COMBINED*`), keeping only the
structural sheets (DASHBOARD, BOND PERFORMANCE, 15 region + 15 DETAIL,
MASTER DATA / MASTER DUPLICATES). Opening balances are NOT carried manually —
KSBC's first-of-month raw already reports each shop's opening stock, so the
normal driver picks them up when it ingests day 1.

After running this, ingest the day's raw the normal way:
    python3 ksbc_daily_update.py "<MONTH> 1st - 1st ANALYSIS.xlsx" "<month> 1st.xlsx"
then run the usual formatting/guard tail.

Usage:
    python3 ksbc_bootstrap_month.py --month JUNE \
        [--from "<prior full-month workbook>.xlsx"] \
        [--out "<MONTH> 1st - 1st ANALYSIS.xlsx"] \
        [--folder "<KSBC shop sales folder>"]

If --from is omitted, the newest `* SHOP SALES ANALYSIS.xlsx` in --folder is used.
Refuses to overwrite an existing --out (a month workbook already exists → no
bootstrap needed).
"""

from __future__ import annotations
import argparse
import os
import re
import shutil
import sys
from openpyxl import load_workbook

MONTHS = {"JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
          "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"}

# A month-specific DATA sheet: a daily/range sheet `<MONTH> <N>` or `<MONTH> <M>-<N>`,
# or any CUMULATIVE / COMBINED sheet. Everything else is structural and is kept.
_DAILY_RE = re.compile(r'^[A-Z]+\s+\d+(?:\s*-\s*\d+)?$', re.IGNORECASE)


def _is_data_sheet(name: str) -> bool:
    up = name.upper()
    if 'CUMULATIVE' in up or 'COMBINED' in up:
        return True
    return bool(_DAILY_RE.match(name.strip()))


def _find_prior_workbook(folder: str) -> str | None:
    """Newest `* SHOP SALES ANALYSIS.xlsx` in folder (by mtime)."""
    cands = []
    for fn in os.listdir(folder):
        if fn.startswith('~'):
            continue
        if fn.upper().endswith('SHOP SALES ANALYSIS.XLSX'):
            cands.append(os.path.join(folder, fn))
    if not cands:
        return None
    return max(cands, key=os.path.getmtime)


def bootstrap_month(month_u: str, prior_path: str, out_path: str) -> dict:
    if os.path.exists(out_path):
        raise FileExistsError(
            f"{out_path} already exists — a {month_u} workbook is already present, "
            f"so no bootstrap is needed. Ingest the raw with ksbc_daily_update.py instead.")
    shutil.copyfile(prior_path, out_path)
    wb = load_workbook(out_path)
    deleted = [n for n in list(wb.sheetnames) if _is_data_sheet(n)]
    for n in deleted:
        del wb[n]
    kept = list(wb.sheetnames)
    if not kept:
        raise RuntimeError("Bootstrap aborted — stripping data sheets left no structural "
                           "sheets; the prior workbook layout was unexpected.")
    wb.save(out_path)
    return {'month': month_u, 'from': prior_path, 'out': out_path,
            'deleted_data_sheets': len(deleted), 'kept_structural_sheets': len(kept),
            'kept': kept}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--month', required=True, help='New month, e.g. JUNE')
    ap.add_argument('--from', dest='prior', default=None, help='Prior full-month workbook to template from')
    ap.add_argument('--out', default=None, help='Output base workbook path')
    ap.add_argument('--folder', default=None, help='KSBC shop sales folder (for auto-detect / default out)')
    a = ap.parse_args()

    month_u = a.month.strip().upper()
    if month_u not in MONTHS:
        print(f"ERROR: --month '{a.month}' is not a valid month name"); sys.exit(2)

    folder = a.folder or (os.path.dirname(a.prior) if a.prior else os.getcwd())
    prior = a.prior or _find_prior_workbook(folder)
    if not prior or not os.path.exists(prior):
        print(f"ERROR: no prior full-month workbook found in {folder} (pass --from)"); sys.exit(3)
    out = a.out or os.path.join(folder, f"{month_u} 1st - 1st ANALYSIS.xlsx")

    try:
        r = bootstrap_month(month_u, prior, out)
    except FileExistsError as e:
        print(f"SKIP: {e}"); sys.exit(0)

    print(f"Bootstrapped {r['month']} base from {os.path.basename(r['from'])}")
    print(f"  stripped {r['deleted_data_sheets']} data sheets, kept {r['kept_structural_sheets']} structural")
    print(f"  -> {r['out']}")
    print("Next: ingest day 1 with ksbc_daily_update.py, then run the formatting/guard tail.")


if __name__ == '__main__':
    main()
