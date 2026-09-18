#!/usr/bin/env python3
"""
Build a standalone PERMIT BREAKUP workbook from the current warehouse
requirement workbook(s).

Pulls one or more warehouse rows out of a
`Warehouse requirement/WAREHOUSE REQUIREMENT - <MONTH> ... .xlsx`
and writes a 2-sheet permit file (WAREHOUSE REQUIREMENT + BRAND x PACK MATRIX)
with no blend / ENA / physical-stock sections.

Styling is inherited by cloning an existing permit workbook as a template, so
the output is visually identical to every permit file already on disk.

Usage
-----
  python3 build_permit_breakup.py --warehouse KOZHIKODE
  python3 build_permit_breakup.py --warehouse NEDUMANGAD --warehouse ATTINGAL \
      --out "PERMIT FOR TOMORROW.xlsx"
  python3 build_permit_breakup.py --warehouse PALAKKAD --scale "PALAKKAD=0.5"
  python3 build_permit_breakup.py --warehouse KOTTARAKKARA \
      --set "KOTTARAKKARA|KS|500=30" --drop "KOTTARAKKARA|MBR|500"

Quantities are copied VERBATIM from the source row unless --scale / --set /
--drop / --only say otherwise.
"""

from __future__ import annotations

import argparse
import copy
import math
import os
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- constants

BRANDS = ["BCB", "BLENDERS", "MWB", "KS", "OPR", "ROFR", "MBR"]
PACKS = ["180ml", "375ml", "500ml", "1000ml"]

# Brand aliases seen in field messages / older workbooks.
BRAND_ALIAS = {
    "BCBC": "BCB", "B C B": "BCB", "BCB": "BCB",
    "BLEND": "BLENDERS", "BLENDER": "BLENDERS", "BLENDERS": "BLENDERS",
    "MWB": "MWB", "MW B": "MWB",
    "KS99": "KS", "K.S.": "KS", "K S": "KS", "KS": "KS",
    "OP": "OPR", "OPR": "OPR", "OLDPEARL": "OPR", "OLD PEARL": "OPR",
    "ROF": "ROFR", "ROFR": "ROFR", "ROYALOLDFORT": "ROFR",
    "MBR": "MBR",
}

# Pack aliases -> canonical pack header text.
PACK_ALIAS = {
    "180": "180ml", "180ML": "180ml", "N": "180ml", "NIP": "180ml",
    "375": "375ml", "375ML": "375ml", "P": "375ml", "PINT": "375ml",
    "500": "500ml", "500ML": "500ml", "H": "500ml", "HL": "500ml",
    "HLF": "500ml", "HLTR": "500ml",
    "1000": "1000ml", "1000ML": "1000ml", "L": "1000ml", "LTR": "1000ml",
    "1L": "1000ml",
}

# Warehouse name variants. Key = normalised alias, value = normalised canonical.
WAREHOUSE_ALIAS = {
    "CALICUT": "KOZHIKODE",
    "KOZHIKKODE": "KOZHIKODE",
    "CALLICUT": "KOZHIKODE",
    "TRIPUNITHURA": "THRIPPUNNITHURA",
    "TRIPUNITHARA": "THRIPPUNNITHURA",
    "THRIPUNITHURA": "THRIPPUNNITHURA",
    "KALPETTA": "KALPATTA",
    "BATTATHUR": "BETTATHUR",
    "BATHATHUR": "BETTATHUR",
    "KOTTARAKARA": "KOTTARAKKARA",
    "PATHANAMTHITA": "PATHANAMTHITTA",
    "PERUMBAVOOR": "PERUMBAVOOR",
}

# Rows in the source sheet that are never warehouses.
SKIP_ROW_RE = re.compile(
    r"^(CLUSTER|TOTAL|LESS|PRODUCTION|BLEND|ENA|GRAND)", re.I
)

TEMPLATE_CANDIDATES = [
    "PERMIT FOR TOMORROW.xlsx",       # 4 data rows: has both zebra variants
    "PERMIT BREAKUP - KOTTARAKKARA.xlsx",
    "PERMIT BREAKUP.xlsx",
]

TITLE = "K.S. DISTILLERY    •    PERMIT BREAKUP"
REQ_SHEET = "WAREHOUSE REQUIREMENT"
MATRIX_SHEET = "BRAND × PACK MATRIX"


# ---------------------------------------------------------------- helpers

def norm(s) -> str:
    """Normalise a name for comparison: upper, alphanumerics only."""
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def canon_warehouse(name: str) -> str:
    n = norm(name)
    return WAREHOUSE_ALIAS.get(n, n)


def display_label(source_label: str) -> str:
    """
    The name to print on the permit.

    Defaults to the workbook's own spelling, EXCEPT where that spelling is a
    known alias (CALICUT, KALPETTA, TRIPUNITHURA...) — then the canonical
    warehouse name is used, because that is what the permit desk expects.
    """
    raw = str(source_label).strip()
    if norm(raw) != canon_warehouse(raw):
        return canon_warehouse(raw)
    return raw


def canon_brand(b: str) -> str:
    key = norm(b)
    for k, v in BRAND_ALIAS.items():
        if norm(k) == key:
            return v
    raise SystemExit(f"ABORT: unknown brand {b!r}. Known: {', '.join(BRANDS)}")


def canon_pack(p: str) -> str:
    key = norm(p)
    for k, v in PACK_ALIAS.items():
        if norm(k) == key:
            return v
    raise SystemExit(f"ABORT: unknown pack {p!r}. Known: {', '.join(PACKS)}")


def default_base() -> Path:
    """Walk up from this script to find the Claude root."""
    here = Path(__file__).resolve()
    for parent in here.parents:
        if (parent / "Warehouse requirement").is_dir():
            return parent
    return Path.cwd()


# ---------------------------------------------------------------- source

class Source:
    """A parsed warehouse-requirement workbook."""

    def __init__(self, path: Path):
        self.path = path
        wb = openpyxl.load_workbook(path, data_only=False)
        if REQ_SHEET not in wb.sheetnames:
            raise ValueError(f"{path.name}: no '{REQ_SHEET}' sheet")
        ws = wb[REQ_SHEET]
        self.ws = ws

        # Header row = the row whose col A reads WAREHOUSE.
        self.hdr = None
        for r in range(1, 12):
            if norm(ws.cell(r, 1).value) == "WAREHOUSE":
                self.hdr = r
                break
        if self.hdr is None:
            raise ValueError(f"{path.name}: could not find WAREHOUSE header row")
        self.pack_row = self.hdr + 1
        self.data_start = self.hdr + 2

        # Column -> (brand, pack). Brand headers are merged across 4 columns.
        self.colmap = {}
        brand = None
        for c in range(2, 30):  # B..AC
            v = ws.cell(self.hdr, c).value
            if v:
                brand = canon_brand(v)
            pack = ws.cell(self.pack_row, c).value
            if brand and pack:
                self.colmap[c] = (brand, canon_pack(pack))
        missing = {(b, p) for b in BRANDS for p in PACKS} - set(self.colmap.values())
        if missing:
            raise ValueError(f"{path.name}: missing brand/pack columns {sorted(missing)}")

        # Warehouse rows.
        self.rows = {}   # canonical name -> (row, label)
        self.labels = []
        for r in range(self.data_start, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if not label or SKIP_ROW_RE.match(str(label).strip()):
                continue
            key = canon_warehouse(label)
            self.labels.append(str(label).strip())
            self.rows.setdefault(key, (r, str(label).strip()))

    def get(self, name: str):
        """Return (label, {(brand,pack): cases}) for a warehouse, or None."""
        key = canon_warehouse(name)
        if key not in self.rows:
            return None
        r, label = self.rows[key]
        data = {}
        for c, bp in self.colmap.items():
            v = self.ws.cell(r, c).value
            if isinstance(v, (int, float)) and v:
                data[bp] = float(v)
        return label, data


def find_sources(base: Path, explicit: str | None):
    folder = base / "Warehouse requirement"
    if explicit:
        for cand in (Path(explicit), folder / explicit, base / explicit,
                     folder / Path(explicit).name):
            if cand.exists():
                return [cand]
        raise SystemExit(f"ABORT: source workbook not found: {explicit}")
    cands = [
        p for p in folder.glob("WAREHOUSE REQUIREMENT*.xlsx")
        if not p.name.startswith("~$")
    ]
    if not cands:
        raise SystemExit(f"ABORT: no WAREHOUSE REQUIREMENT*.xlsx in {folder}")
    return sorted(cands, key=lambda p: p.stat().st_mtime, reverse=True)


# ---------------------------------------------------------------- overrides

def parse_kv(spec: str, expect_value: bool):
    """Parse 'WAREHOUSE|BRAND|PACK=VAL' or 'WAREHOUSE|BRAND|PACK'."""
    if expect_value:
        if "=" not in spec:
            raise SystemExit(f"ABORT: --set needs WAREHOUSE|BRAND|PACK=VALUE, got {spec!r}")
        left, val = spec.rsplit("=", 1)
        try:
            value = float(val)
        except ValueError:
            raise SystemExit(f"ABORT: --set value must be a number, got {val!r}")
    else:
        left, value = spec, 0.0
    parts = [p.strip() for p in left.split("|")]
    if len(parts) != 3:
        raise SystemExit(
            f"ABORT: expected WAREHOUSE|BRAND|PACK, got {left!r}"
        )
    wh, brand, pack = parts
    return canon_warehouse(wh), canon_brand(brand), canon_pack(pack), value


def apply_overrides(plan, args):
    """Mutate plan[wh]['data'] per --scale / --set / --drop / --only."""
    log = []

    for spec in args.scale or []:
        if "=" not in spec:
            raise SystemExit(f"ABORT: --scale needs WAREHOUSE=FACTOR, got {spec!r}")
        wh, fac = spec.rsplit("=", 1)
        key = canon_warehouse(wh)
        if key not in plan:
            raise SystemExit(f"ABORT: --scale names {wh!r}, not in this permit")
        try:
            f = float(fac)
        except ValueError:
            raise SystemExit(f"ABORT: --scale factor must be a number, got {fac!r}")
        if f <= 0:
            raise SystemExit("ABORT: --scale factor must be > 0")
        d = plan[key]["data"]
        for bp in list(d):
            before = d[bp]
            d[bp] = float(round(before * f))
            if d[bp] == 0:
                del d[bp]
            log.append(f"scale {plan[key]['label']} {bp[0]} {bp[1]}: "
                       f"{before:g} -> {d.get(bp, 0):g}")

    for spec in args.only or []:
        if "|" not in spec:
            raise SystemExit(f"ABORT: --only needs WAREHOUSE|BRAND PACK,BRAND PACK, got {spec!r}")
        wh, keep = spec.split("|", 1)
        key = canon_warehouse(wh)
        if key not in plan:
            raise SystemExit(f"ABORT: --only names {wh!r}, not in this permit")
        keeps = set()
        for item in keep.split(","):
            bits = item.strip().split()
            if len(bits) != 2:
                raise SystemExit(f"ABORT: --only item must be 'BRAND PACK', got {item!r}")
            keeps.add((canon_brand(bits[0]), canon_pack(bits[1])))
        d = plan[key]["data"]
        for bp in list(d):
            if bp not in keeps:
                log.append(f"only {plan[key]['label']}: dropped {bp[0]} {bp[1]} ({d[bp]:g})")
                del d[bp]

    for spec in args.set or []:
        key, brand, pack, value = parse_kv(spec, True)
        if key not in plan:
            raise SystemExit(f"ABORT: --set names a warehouse not in this permit: {spec}")
        d = plan[key]["data"]
        before = d.get((brand, pack), 0.0)
        if value == 0:
            d.pop((brand, pack), None)
        else:
            d[(brand, pack)] = value
        log.append(f"set {plan[key]['label']} {brand} {pack}: {before:g} -> {value:g}")

    for spec in args.drop or []:
        key, brand, pack, _ = parse_kv(spec, False)
        if key not in plan:
            raise SystemExit(f"ABORT: --drop names a warehouse not in this permit: {spec}")
        d = plan[key]["data"]
        before = d.pop((brand, pack), None)
        if before is None:
            print(f"  note: --drop {spec} had nothing to drop")
        else:
            log.append(f"drop {plan[key]['label']} {brand} {pack}: {before:g} -> 0")

    return log


# ---------------------------------------------------------------- builder

def build(plan, order, template: Path, out: Path, compiled: str, title: str):
    wb = openpyxl.load_workbook(template)
    ws = wb[REQ_SHEET]

    # Locate template geometry.
    hdr = None
    for r in range(1, 12):
        if norm(ws.cell(r, 1).value) == "WAREHOUSE":
            hdr = r
            break
    if hdr is None:
        raise SystemExit(f"ABORT: template {template.name} has no WAREHOUSE header")
    pack_row = hdr + 1
    first = hdr + 2

    # Template column map (by brand/pack, so column order can never desync).
    tcol = {}
    brand = None
    for c in range(2, 30):
        v = ws.cell(hdr, c).value
        if v:
            brand = canon_brand(v)
        pack = ws.cell(pack_row, c).value
        if brand and pack:
            tcol[(brand, canon_pack(pack))] = c
    total_col = 30  # AD

    # Capture styles before we clear anything.
    old_total_row = None
    for r in range(first, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").upper().startswith("TOTAL"):
            old_total_row = r
            break
    if old_total_row is None:
        raise SystemExit(f"ABORT: template {template.name} has no TOTAL row")

    n_data = old_total_row - first
    style_odd = [copy.copy(ws.cell(first, c)._style) for c in range(1, total_col + 1)]
    style_even = [
        copy.copy(ws.cell(first + 1, c)._style) for c in range(1, total_col + 1)
    ] if n_data >= 2 else style_odd
    style_total = [copy.copy(ws.cell(old_total_row, c)._style) for c in range(1, total_col + 1)]

    # The quantity cells follow a value-driven rule, NOT a positional one:
    # a cell with a figure is 11pt bold, an empty cell is 10pt regular.
    # Copying the template row positionally would carry the template
    # warehouse's bolding onto whatever columns the new one populates.
    font_filled = font_empty = None
    for r in range(first, old_total_row):
        for c in range(2, total_col):
            cell = ws.cell(r, c)
            if cell.value not in (None, "") and font_filled is None:
                font_filled = copy.copy(cell.font)
            elif cell.value in (None, "") and font_empty is None:
                font_empty = copy.copy(cell.font)
    if font_filled is None or font_empty is None:
        raise SystemExit(
            f"ABORT: template {template.name} has no filled/empty pair to "
            "read the quantity-cell font rule from"
        )
    h_data = ws.row_dimensions[first].height
    h_total = ws.row_dimensions[old_total_row].height

    # Clear everything from the first data row down.
    ws.delete_rows(first, ws.max_row - first + 1)

    # Write data rows.
    n = len(order)
    for i, key in enumerate(order):
        r = first + i
        styles = style_odd if i % 2 == 0 else style_even
        for c in range(1, total_col + 1):
            ws.cell(r, c)._style = copy.copy(styles[c - 1])
        ws.row_dimensions[r].height = h_data
        ws.cell(r, 1).value = plan[key]["label"]
        for bp, val in plan[key]["data"].items():
            ws.cell(r, tcol[bp]).value = int(val) if float(val).is_integer() else val
        for c in range(2, total_col):
            cell = ws.cell(r, c)
            cell.font = copy.copy(
                font_filled if cell.value not in (None, "") else font_empty
            )
        ws.cell(r, total_col).value = (
            f"=SUM(B{r}:{get_column_letter(total_col - 1)}{r})"
        )

    # TOTAL row.
    tr = first + n
    for c in range(1, total_col + 1):
        ws.cell(tr, c)._style = copy.copy(style_total[c - 1])
    ws.row_dimensions[tr].height = h_total
    ws.cell(tr, 1).value = "TOTAL  REQUIREMENT"
    for c in range(2, total_col):
        L = get_column_letter(c)
        ws.cell(tr, c).value = f"=SUM({L}{first}:{L}{first + n - 1})"
    ws.cell(tr, total_col).value = (
        f"=SUM(B{tr}:{get_column_letter(total_col - 1)}{tr})"
    )

    # Drop leftover row-height records below the TOTAL row, otherwise the
    # template's tall empty rows survive as blank bands under the table.
    for r in list(ws.row_dimensions):
        if r > tr:
            del ws.row_dimensions[r]

    ws["A1"] = title
    ws.freeze_panes = f"B{first}"

    # Matrix sheet: re-point every reference at the new TOTAL row.
    if MATRIX_SHEET in wb.sheetnames:
        m = wb[MATRIX_SHEET]
        m["A2"] = (
            "Aggregate requirement by brand & pack size   |   "
            f"Compiled {compiled}   |   All quantities in CASES"
        )
        for row in m.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = re.sub(
                        r"('WAREHOUSE REQUIREMENT'!\$?[A-Z]{1,2}\$?)\d+",
                        lambda mo: f"{mo.group(1)}{tr}",
                        cell.value,
                    )

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return first, tr, tcol


# ---------------------------------------------------------------- verify

def verify(out: Path, plan, order, first, tr, tcol):
    """Reload and prove every figure matches the plan. Raises on any drift."""
    wb = openpyxl.load_workbook(out)
    ws = wb[REQ_SHEET]
    problems = []

    for i, key in enumerate(order):
        r = first + i
        if str(ws.cell(r, 1).value).strip() != plan[key]["label"]:
            problems.append(f"row {r} label {ws.cell(r,1).value!r} != {plan[key]['label']!r}")
        written = {}
        for bp, c in tcol.items():
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and v:
                written[bp] = float(v)
        if written != {k: float(v) for k, v in plan[key]["data"].items()}:
            extra = set(written) - set(plan[key]["data"])
            miss = set(plan[key]["data"]) - set(written)
            problems.append(
                f"{plan[key]['label']}: cell mismatch (extra={sorted(extra)}, missing={sorted(miss)})"
            )

    if str(ws.cell(tr, 1).value).strip() != "TOTAL  REQUIREMENT":
        problems.append("TOTAL row label wrong")
    for c in range(2, 30):
        L = get_column_letter(c)
        want = f"=SUM({L}{first}:{L}{first + len(order) - 1})"
        got = ws.cell(tr, c).value
        if got != want:
            problems.append(f"TOTAL {L}{tr}: {got!r} != {want!r}")

    # Nothing may survive below the TOTAL row (no blend / ENA / stock).
    for r in range(tr + 1, ws.max_row + 1):
        for c in range(1, 31):
            if ws.cell(r, c).value not in (None, ""):
                problems.append(f"stray content below TOTAL at {get_column_letter(c)}{r}")

    if problems:
        raise SystemExit("ABORT: verification failed —\n  " + "\n  ".join(problems))


# ---------------------------------------------------------------- main

def main(argv=None):
    ap = argparse.ArgumentParser(description="Build a PERMIT BREAKUP workbook.")
    ap.add_argument("--warehouse", action="append", required=True,
                    help="Warehouse to include (repeatable, order preserved).")
    ap.add_argument("--source", help="Force a specific requirement workbook.")
    ap.add_argument("--base", help="Claude root folder.")
    ap.add_argument("--out", help="Output path or filename.")
    ap.add_argument("--template", help="Permit workbook to clone styling from.")
    ap.add_argument("--title", default=TITLE)
    ap.add_argument("--label", action="append",
                    help="WAREHOUSE=NAME  (force the name printed on the row)")
    ap.add_argument("--compiled", help="Compiled date text, e.g. '6 August 2026'.")
    ap.add_argument("--scale", action="append", help="WAREHOUSE=FACTOR")
    ap.add_argument("--set", action="append", help="WAREHOUSE|BRAND|PACK=VALUE")
    ap.add_argument("--drop", action="append", help="WAREHOUSE|BRAND|PACK")
    ap.add_argument("--only", action="append",
                    help="WAREHOUSE|BRAND PACK,BRAND PACK  (keep only these)")
    args = ap.parse_args(argv)

    base = Path(args.base) if args.base else default_base()
    folder = base / "Warehouse requirement"
    sources = find_sources(base, args.source)

    label_override = {}
    for spec in args.label or []:
        if "=" not in spec:
            raise SystemExit(f"ABORT: --label needs WAREHOUSE=NAME, got {spec!r}")
        wh, shown = spec.split("=", 1)
        label_override[canon_warehouse(wh)] = shown.strip()

    # Resolve each warehouse against the newest workbook that actually has it.
    plan, order, used = {}, [], {}
    for wh in args.warehouse:
        key = canon_warehouse(wh)
        if key in plan:
            raise SystemExit(f"ABORT: {wh} requested twice")
        hit = None
        also = []
        for src_path in sources:
            try:
                src = Source(src_path)
            except Exception as e:                     # unreadable / wrong shape
                print(f"  skip {src_path.name}: {e}")
                continue
            got = src.get(wh)
            if got:
                if hit is None:
                    hit = (src_path, src, got)
                else:
                    also.append(src_path.name)
        if hit is None:
            known = sorted({lbl for p in sources
                            for lbl in (Source(p).labels if p.exists() else [])})
            raise SystemExit(
                f"ABORT: warehouse {wh!r} not found in any requirement workbook.\n"
                f"  Available: {', '.join(known)}"
            )
        src_path, src, (label, data) = hit
        if not data:
            raise SystemExit(f"ABORT: {label} has an all-zero row in {src_path.name}")
        shown = label_override.get(key, display_label(label))
        plan[key] = {"label": shown, "src_label": label,
                     "data": data, "source": src_path.name}
        order.append(key)
        used[shown] = (src_path.name, also, label)

    log = apply_overrides(plan, args)

    # Output path.
    if args.out:
        out = Path(args.out)
        if not out.is_absolute():
            out = folder / args.out
    else:
        if len(order) == 1:
            out = folder / f"PERMIT BREAKUP - {plan[order[0]]['label']}.xlsx"
        else:
            import datetime as _dt
            out = folder / (
                "PERMIT BREAKUP - "
                + _dt.date.today().strftime("%d %b").lstrip("0") + ".xlsx"
            )

    # Template.
    if args.template:
        tpl = Path(args.template)
        if not tpl.is_absolute():
            tpl = folder / args.template
    else:
        tpl = next((folder / t for t in TEMPLATE_CANDIDATES if (folder / t).exists()), None)
    if tpl is None or not tpl.exists():
        raise SystemExit(
            "ABORT: no permit template found. Expected one of: "
            + ", ".join(TEMPLATE_CANDIDATES)
        )

    if args.compiled:
        compiled = args.compiled
    else:
        import datetime as _dt
        compiled = _dt.date.today().strftime("%d %B %Y").lstrip("0")

    # Never clobber a locked file.
    if (out.parent / f"~${out.name}").exists():
        raise SystemExit(f"ABORT: {out.name} is open in Excel — close it first.")

    first, tr, tcol = build(plan, order, tpl, out, compiled, args.title)
    verify(out, plan, order, first, tr, tcol)

    # ---- report
    print(f"\nSOURCE   : " + ", ".join(sorted({v['source'] for v in plan.values()})))
    for label, (src_name, also, src_label) in used.items():
        if src_label != label:
            print(f"  note   : source workbook calls it {src_label!r} — "
                  f"printed on the permit as {label!r}")
    dupes = sorted(lbl for lbl, (_, also, _) in used.items() if also)
    if dupes:
        others = sorted({n for _, (_, also, _) in used.items() for n in also})
        print(f"  note   : {', '.join(dupes)} also appear in "
              f"{', '.join(others)} — used the newest workbook. "
              f"Different (SELECTED) files are different despatch rounds, "
              f"so confirm this is the right one.")
    print(f"TEMPLATE : {tpl.name}")
    print(f"OUTPUT   : {out}")
    if log:
        print("EDITS    :")
        for line in log:
            print(f"  {line}")
    print("\nPERMIT CONTENT (cases)")
    grand = 0.0
    for key in order:
        d = plan[key]["data"]
        tot = sum(d.values())
        grand += tot
        bits = ", ".join(
            f"{b} {p} {v:g}" for (b, p), v in sorted(d.items())
        )
        print(f"  {plan[key]['label']:<18} {tot:>7,.0f}   {bits}")
    print(f"  {'TOTAL':<18} {grand:>7,.0f}")
    if grand > 720:
        print(f"  ! {grand:,.0f} cs exceeds one 720-case truck "
              f"({math.ceil(grand / 720)} loads)")
    print("\nverification: PASS (labels, cells, TOTAL formulas, nothing below TOTAL)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
