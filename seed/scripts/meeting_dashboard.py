#!/usr/bin/env python3
import openpyxl, json, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict
SRC=sys.argv[1]; OUT=sys.argv[2]
import os
DATA=os.environ.get('BOND_DATA') or (os.path.join(os.environ['MEETING_OUT'],'bond_data.json') if os.environ.get('MEETING_OUT') else '')
if not DATA or not os.path.exists(DATA): sys.exit('BOND_DATA not set / bond_data.json missing')
data=json.load(open(DATA))
META=data.get('__meta__',{'cur':'MAY 2026','prev':'APRIL 2026','curm':'May','prevm':'Apr'})
BONDS=[b for b in ['KOLLAM','KOZHIKODE','ATTINGAL','PALAKKAD','KOTTARAKARA','KANNUR','ALAPPUZHA','NEDUMANGAD','ALUVA','PERINTHALMANNA','THODUPUZHA','PATHANAMTHITTA','KOTTAYAM','THRISSUR','TRIPUNITHURA'] if b in data]
# DARK palette
BG='FF0B1730'; BAND='FF0A1430'; SUBF='FF13244A'; CARD='FF1B2C54'; KHEAD='FF22386B'
ROWA='FF15264A'; ROWB='FF0E1C3A'; TXT='FFEAF0FB'; DIM='FFA9B8D6'
GOLD='FFFFD700'; WHITE='FFFFFFFF'; NAVY_DEEP='FF0A1430'
RHI='FF4FC3F7'; RBAL='FF81C784'; RINV='FFFFB74D'; RCRI='FFE57373'
DGREEN='FF66BB6A'; DRED='FFEF5350'
def f(h): return PatternFill('solid',fgColor=h)
thin=Side(style='thin',color='FF2A3F6B'); gold_b=Side(style='thin',color=GOLD); gold_m=Side(style='medium',color=GOLD)
C=Alignment('center','center',wrap_text=True); L=Alignment('left','center',wrap_text=True,indent=1)
B0=2; BN=19
def fillr(ws,r,c1,c2,fill=None,border=False,botgold=False):
    for c in range(c1,c2+1):
        cell=ws.cell(row=r,column=c); bd={}
        if fill: cell.fill=fill
        if border: bd=dict(left=Side(style='thin',color='FF24385F'),right=Side(style='thin',color='FF24385F'),top=Side(style='thin',color='FF24385F'),bottom=Side(style='thin',color='FF24385F'))
        if botgold: bd['bottom']=gold_b
        if bd: cell.border=Border(**bd)
def band(ws,r,text,fill=BAND,font=GOLD,sz=13,h=28):
    ws.merge_cells(start_row=r,start_column=B0,end_row=r,end_column=BN)
    fillr(ws,r,B0,BN,fill=f(fill),botgold=True)
    cc=ws.cell(row=r,column=B0,value=text); cc.font=Font(bold=True,size=sz,color=font); cc.alignment=C
    ws.row_dimensions[r].height=h
def kpi(ws,r,idx,title,value,sub,fmt='#,##0',vcolor=TXT,card=CARD):
    c1=B0+idx*3; c2=c1+2
    ws.merge_cells(start_row=r,start_column=c1,end_row=r,end_column=c2); fillr(ws,r,c1,c2,fill=f(KHEAD))
    t=ws.cell(row=r,column=c1,value=title); t.font=Font(bold=True,size=9,color=GOLD); t.alignment=C; ws.row_dimensions[r].height=17
    ws.merge_cells(start_row=r+1,start_column=c1,end_row=r+1,end_column=c2); fillr(ws,r+1,c1,c2,fill=f(card),border=True)
    v=ws.cell(row=r+1,column=c1,value=value); v.font=Font(bold=True,size=15,color=vcolor); v.alignment=C; v.number_format=fmt; ws.row_dimensions[r+1].height=24
    ws.merge_cells(start_row=r+2,start_column=c1,end_row=r+2,end_column=c2); fillr(ws,r+2,c1,c2,fill=f(card),border=True)
    s=ws.cell(row=r+2,column=c1,value=sub); s.font=Font(size=8,italic=True,color=DIM); s.alignment=C; ws.row_dimensions[r+2].height=14
def hdr(ws,r,cols,h=20):
    c=B0; ws.row_dimensions[r].height=h
    for lab,span in cols:
        c2=c+span-1
        if span>1: ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c2)
        cell=ws.cell(row=r,column=c,value=lab); cell.font=Font(bold=True,size=9,color=GOLD); cell.alignment=C
        fillr(ws,r,c,c2,fill=f(KHEAD),botgold=True); c=c2+1
def row(ws,r,cells,tint=ROWB,hero=False,h=17):
    c=B0; ws.row_dimensions[r].height=h
    for val,span,fmt,al,opt in cells:
        c2=c+span-1
        if span>1: ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c2)
        cell=ws.cell(row=r,column=c,value=val)
        fg=NAVY_DEEP if hero else opt.get('color',TXT)
        cell.font=Font(bold=hero or opt.get('bold',False),size=(11 if hero else 10),color=fg)
        cell.alignment=L if al=='l' else C
        if fmt: cell.number_format=fmt
        if hero:
            for cc in range(c,c2+1):
                hc=ws.cell(row=r,column=cc); hc.fill=f(GOLD)
                hc.border=Border(top=gold_m,bottom=gold_m,left=(gold_m if cc==B0 else None),right=(gold_m if cc==BN else None))
        else:
            fillr(ws,r,c,c2,fill=f(tint),border=True)
        c=c2+1
def tier(st):
    if st>=0.8: return ('🚀 High Performance',RHI)
    if st>=0.6: return ('✅ Balanced',RBAL)
    if st>=0.4: return ('⚠️ Inventory Heavy',RINV)
    return ('🚫 Critical Overstock',RCRI)
def dpc(n,o): return f"{'+' if n>=o else ''}{(n-o)/o*100:.0f}% vs "+META['prevm'] if o else "—"
def dcol(v): return DGREEN if v>=0 else DRED

def S(k,s): return sum(data[b][k].get(s,0) for b in BONDS)
shop_m=sum(data[b]['tert_may']['S'] for b in BONDS); shop_a=sum(data[b]['tert_apr'].get('S',0) for b in BONDS)
openm=sum(data[b]['tert_may']['O'] for b in BONDS); recvm=sum(data[b]['tert_may']['R'] for b in BONDS)
st_m=shop_m/(openm+recvm) if (openm+recvm) else 0
fed_m=S('sec_may','fed');fed_a=S('sec_apr','fed');bar_m=S('sec_may','bar');bar_a=S('sec_apr','bar')
seck_m=S('sec_may','ksbc');seck_a=S('sec_apr','ksbc');sect_m=S('sec_may','total');sect_a=S('sec_apr','total')
liq_m=shop_m+fed_m+bar_m; liq_a=shop_a+fed_a+bar_a
visits=sum(data[b]['tert_may']['visits'] for b in BONDS)   # KSBC shop visits by SEs (assigned-shop basis)
covered=sum(data[b]['tert_may']['visited'] for b in BONDS); shops=sum(data[b]['tert_may']['nshops'] for b in BONDS)
_fn=data.get('__field__',{})   # D1/D2: network SE totals -- ASM-excluded, de-duplicated (no per-bond double count)
active_sm=_fn.get('active', sum(1 for b in BONDS if data[b].get('vis')))
field_h=float(_fn.get('field_h',0)); wh=int(_fn.get('wh',0))

wb=openpyxl.load_workbook(SRC)
idx=wb.sheetnames.index('DASHBOARD'); del wb['DASHBOARD']
ws=wb.create_sheet('DASHBOARD', idx); ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=2
for col in range(B0,BN+1): ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=10.5
# dark canvas
for rr in range(1,95):
    for c in range(1,25): ws.cell(row=rr,column=c).fill=f(BG)

r=2
band(ws,r,'K.S. DISTILLERY  —  BOND MEETING DASHBOARD',fill=BAND,font=GOLD,sz=18,h=40); r+=1
ws.merge_cells(start_row=r,start_column=B0,end_row=r,end_column=BN); fillr(ws,r,B0,BN,fill=f(SUBF))
sc=ws.cell(row=r,column=B0,value=META['cur']+'   ·   All channels — Shop · Invoice · Secondary · Total Liquidation · Field Coverage   ·   figures in cases unless noted')
sc.font=Font(size=10,italic=True,color=DIM); sc.alignment=C; ws.row_dimensions[r].height=20; r+=2

kpi(ws,r,0,'TOTAL LIQUIDATION',liq_m,dpc(liq_m,liq_a),vcolor=GOLD)
kpi(ws,r,1,'SHOP SALES (tertiary)',shop_m,dpc(shop_m,shop_a))
kpi(ws,r,2,'SECONDARY SALES',sect_m,dpc(sect_m,sect_a))
kpi(ws,r,3,'INVOICE (CFD+BAR)',fed_m+bar_m,dpc(fed_m+bar_m,fed_a+bar_a))
kpi(ws,r,4,'SELL-THROUGH %',st_m,'shop sales basis',fmt='0%')
kpi(ws,r,5,'KSBC SHOP VISITS',visits,f'{active_sm} sales executives active')
r+=3
kpi(ws,r,0,'SHOPS COVERED',covered,f'of {shops} active',fmt='0')
kpi(ws,r,1,'COVERAGE %',covered/shops if shops else 0,'visited / active',fmt='0%')
kpi(ws,r,2,'FIELD TIME (h)',field_h,'total SE field time · '+META['curm'],fmt='#,##0')
kpi(ws,r,3,'WAREHOUSE STOPS',wh,'sales executive WH visits',fmt='0')
kpi(ws,r,4,'KSBC DISPATCH',seck_m,'warehouse → shop',fmt='#,##0')
kpi(ws,r,5,'CONSUMER-FED',fed_m,dpc(fed_m,fed_a),fmt='#,##0')
r+=4

band(ws,r,'CHANNEL SUMMARY  —  NETWORK, '+META['curm']+' vs '+META['prevm'],fill='FF13244A',font=GOLD,sz=12,h=24); r+=1
hdr(ws,r,[('Channel',4),(META['prevm'],3),(META['curm'],3),('Δ',2),('% Δ',2),('Share of Liq.',4)]); r+=1
rr=[r]
for nm,a,m,sh,bold in [('Shop liquidation (tertiary)',shop_a,shop_m,shop_m/liq_m,True),
                  ('KSBC dispatch (secondary)',seck_a,seck_m,None,False),
                  ('Consumer-fed (invoice)',fed_a,fed_m,fed_m/liq_m,False),
                  ('BAR (invoice)',bar_a,bar_m,bar_m/liq_m,False),
                  ('Secondary TOTAL',sect_a,sect_m,None,True)]:
    dv=m-a; pd=(m-a)/a if a else 0
    row(ws,rr[0],[(nm,4,None,'l',{'bold':bold}),(a,3,'#,##0','c',{}),(m,3,'#,##0','c',{'bold':bold}),
        (dv,2,'+#,##0;-#,##0','c',{'color':dcol(dv)}),(pd,2,'+0%;-0%','c',{'color':dcol(dv)}),
        (sh if sh is not None else '',4,'0%','c',{})],tint=(ROWA if (rr[0]-r)%2 else ROWB)); rr[0]+=1
r=rr[0]
row(ws,r,[('TOTAL LIQUIDATION ✦',4,None,'l',{}),(liq_a,3,'#,##0','c',{}),(liq_m,3,'#,##0','c',{}),
    (liq_m-liq_a,2,'+#,##0;-#,##0','c',{}),(((liq_m-liq_a)/liq_a if liq_a else 0),2,'+0%;-0%','c',{}),(1.0,4,'0%','c',{})],hero=True,h=22); r+=2

band(ws,r,'BOND SCORECARD  —  all channels by bond (sorted by Total Liquidation)',fill='FF13244A',font=GOLD,sz=12,h=24); r+=1
hdr(ws,r,[('Bond',3),('Shop Sales',2),('Sell-Thr%',2),('Cons-fed',2),('Total Liq.',3),('Visits',2),('Cov%',2),('Rating',2)]); r+=1
order=sorted(BONDS,key=lambda b:-(data[b]['tert_may']['S']+data[b]['sec_may'].get('fed',0)+data[b]['sec_may'].get('bar',0)))
for i,b in enumerate(order):
    d=data[b]; tm=d['tert_may']; sm=d['sec_may']; lq=tm['S']+sm.get('fed',0)+sm.get('bar',0)
    cov=tm['visited']/tm['nshops'] if tm['nshops'] else 0; lab,rc=tier(tm['st'])
    row(ws,r,[(b,3,None,'l',{'bold':True}),(tm['S'],2,'#,##0','c',{}),(tm['st'],2,'0%','c',{}),(sm.get('fed',0),2,'#,##0','c',{}),
        (lq,3,'#,##0','c',{'bold':True,'color':GOLD}),(int(tm['visits']),2,'#,##0','c',{}),(cov,2,'0%','c',{}),
        (lab,2,None,'c',{'color':rc,'bold':True})],tint=(ROWA if i%2 else ROWB)); r+=1
row(ws,r,[('NETWORK TOTAL',3,None,'l',{}),(shop_m,2,'#,##0','c',{}),(st_m,2,'0%','c',{}),(fed_m,2,'#,##0','c',{}),
    (liq_m,3,'#,##0','c',{}),(int(visits),2,'#,##0','c',{}),(covered/shops if shops else 0,2,'0%','c',{}),('',2,None,'c',{})],hero=True,h=22); r+=2

band(ws,r,'FIELD COVERAGE BY BOND  —  sales executive activity, '+META['curm'],fill='FF13244A',font=GOLD,sz=12,h=24); r+=1
hdr(ws,r,[('Bond',3),('Sales Executive',4),('Visits',2),('Coverage%',2),('Active Days',2),('Avg/Day',2),('Field h',3)]); r+=1
for i,b in enumerate(BONDS):
    d=data[b]; tm=d['tert_may']; v=d.get('vis'); cov=tm['visited']/tm['nshops'] if tm['nshops'] else 0
    if v and tm['visits']>0:
        ad=v['active_days'] if v['active_days'] else '—'; avd=(round(float(v['avg_visits_day']),1) if v['avg_visits_day'] else '—')
        row(ws,r,[(b,3,None,'l',{'bold':True}),(v['staff'],4,None,'l',{}),(int(tm['visits']),2,'#,##0','c',{}),
            (cov,2,'0%','c',{}),(ad,2,('0' if ad!='—' else None),'c',{}),(avd,2,('0.0' if avd!='—' else None),'c',{}),
            (float(v['field_h']),3,'#,##0','c',{})],tint=(ROWA if i%2 else ROWB))
    else:
        who=(str(v['staff'])+' — no shop visits') if v else 'VACANT — no visit log'   # D2: never show exec hours beside 0 visits
        row(ws,r,[(b,3,None,'l',{'bold':True}),(who,4,None,'l',{'color':RCRI}),(int(tm['visits']),2,'#,##0','c',{}),(cov,2,'0%','c',{}),
            ('—',2,None,'c',{}),('—',2,None,'c',{}),('—',3,None,'c',{})],tint=(ROWA if i%2 else ROWB))
    r+=1
r+=1

bl=defaultdict(float)
for b in BONDS:
    for br in data[b]['brands']: bl[br['brand']]+=br['sales']
top=sorted(bl.items(),key=lambda x:-x[1]); tl=sum(v for _,v in top) or 1
band(ws,r,'TOP BRANDS BY TOTAL LIQUIDATION  —  network, '+META['curm'],fill='FF13244A',font=GOLD,sz=12,h=24); r+=1
hdr(ws,r,[('#',1),('Brand',8),('Liquidation (cs)',4),('% of Network',5)]); r+=1
for i,(br,v) in enumerate(top[:10],1):
    row(ws,r,[(i,1,'0','c',{}),(br,8,None,'l',{}),(v,4,'#,##0.0','c',{'bold':True,'color':GOLD}),(v/tl,5,'0%','c',{})],tint=(ROWA if i%2 else ROWB)); r+=1
r+=1
ws.merge_cells(start_row=r,start_column=B0,end_row=r,end_column=BN)
fc=ws.cell(row=r,column=B0,value='Generated from '+META['curm']+' shop-sales, secondary/invoice dispatch, and field-visit logs.  ·  Drill into bond tabs for full deep-dives.')
fc.font=Font(size=8,italic=True,color=DIM); fc.alignment=C
wb.save(OUT)
print('dark dashboard built through row', r)
# FINAL pass, must be the LAST write: openpyxl's loader discards each legacy comment's VML shape
# size and re-writes the hard-coded 144x79px default on every save, so the receipt notes the KSBC
# build sized to fit (260px wide) come out of this pipeline clipped -- an 11 cs / multi-date note
# renders as its first line only. Pure ZIP/VML surgery; cell values and note TEXT are untouched.
try:
    import resize_comment_boxes
    _nb,_hlo,_hhi = resize_comment_boxes.resize_boxes(OUT)
    if _nb: print(f'comment boxes sized to fit: {_nb} (height {_hlo}..{_hhi}px, width 260px)')
except Exception as _e:
    print('WARN comment-box resize skipped:', _e)
