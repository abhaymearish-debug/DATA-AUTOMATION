#!/usr/bin/env python3
"""Build KSBC shop map (Leaflet HTML) from latest KSBC analysis workbook + shop coordinates.

Coordinates source: .claude/data/shop_coordinates.csv (cached). To refresh the cache,
pass --clientlist <path to ClientList xlsx export> (FieldSense client export with
Latitude/Longitude columns). Match key = first 3-6 digit run in the shop/client name.
Usage: python3 build_shop_map.py [--clientlist PATH] [--out PATH]
"""
import openpyxl, re, json, glob, csv, os, sys, argparse
from collections import Counter

BASE = os.environ.get('KSD_BASE', '/sessions/funny-hopeful-cray/mnt/Claude')
COORD_CSV = os.path.join(BASE, '.claude/data/shop_coordinates.csv')
REGIONS = ['KOLLAM','KOZHIKODE','ATTINGAL','PALAKKAD','KOTTARAKARA','KANNUR','ALAPPUZHA',
           'NEDUMANGAD','ALUVA','THODUPUZHA','KOTTAYAM','THRISSUR','TRIPUNITHURA',
           'PERINTHALMANNA','PATHANAMTHITTA']

def key_of(name):
    m = re.search(r'\d{3,6}', str(name or ''))
    return m.group(0).lstrip('0') if m else None

def refresh_coords(clientlist_path):
    wb = openpyxl.load_workbook(clientlist_path, read_only=True)
    ws = wb['Clients']
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if str(r[0] or '') == 'ID')
    out, dupes = {}, []
    for r in rows[hdr_i+1:]:
        name = str(r[1] or '').strip()
        k = key_of(name)
        if not name or not k or r[8] in (None,'') or r[9] in (None,''): continue
        try: lat, lng = float(r[8]), float(r[9])
        except (TypeError, ValueError): continue
        if k in out: dupes.append(f'{k}: kept "{out[k][3]}", skipped "{name}"'); continue
        out[k] = (lat, lng, str(r[6] or '').strip(), name)
    os.makedirs(os.path.dirname(COORD_CSV), exist_ok=True)
    with open(COORD_CSV, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['key','lat','lng','address','client_name'])
        for k, v in sorted(out.items()): w.writerow([k, *v])
    print(f'coords cache refreshed: {len(out)} entries -> {COORD_CSV}')
    for d in dupes: print('  dupe key, first kept |', d)
    return {k: {'lat': v[0], 'lng': v[1], 'addr': v[2]} for k, v in out.items()}

def load_coords():
    out = {}
    with open(COORD_CSV, encoding='utf-8') as f:
        for r in csv.DictReader(f):
            out[r['key']] = {'lat': float(r['lat']), 'lng': float(r['lng']), 'addr': r['address']}
    return out

def load_shops(coords):
    cand = [x for x in glob.glob(os.path.join(BASE, 'KSBC shop sales', '*ANALYSIS.xlsx'))
            if not os.path.basename(x).startswith('~$')]
    f = sorted(cand, key=os.path.getmtime)[-1]
    label = re.sub(r'\s*(SHOP SALES )?ANALYSIS\.xlsx$', '', os.path.basename(f))
    wb = openpyxl.load_workbook(f, read_only=True, data_only=False)
    shops, missing = [], []
    for region in REGIONS:
        for r in wb[region].iter_rows(min_row=5, values_only=True):
            code, name, staff = r[0], r[1], r[2]
            if code in (None,'') or name in (None,'') or str(name).strip().upper().startswith('TOTAL'): continue
            name = str(name).strip()
            o, rec, s, c = [float(x or 0) for x in (r[3], r[4], r[5], r[6])]
            st = s/(o+rec) if (o+rec) > 0 else 0
            if max(o, rec, s, c) == 0: tier = 'No activity'
            elif st >= .8: tier = 'High Performance'
            elif st >= .6: tier = 'Balanced'
            elif st >= .4: tier = 'Inventory Heavy'
            else: tier = 'Critical Overstock'
            cc = coords.get(key_of(name))
            if not cc:
                missing.append({'name': name, 'bond': region, 's': round(s,2)}); continue
            shops.append({'code': str(code), 'name': name, 'bond': region, 'staff': str(staff or ''),
                          'o': round(o,2), 'r': round(rec,2), 's': round(s,2), 'c': round(c,2),
                          'st': round(st*100,1), 'tier': tier, **cc})
    return shops, missing, label, os.path.basename(f)

def build_html(shops, missing, label, src, out_path):
    payload = {'period': label, 'source': src, 'built': __import__('datetime').date.today().isoformat(),
               'shops': shops, 'missing': missing}
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'shop_map_template.html'), encoding='utf-8') as f:
        html = f.read()
    html = html.replace('__PAYLOAD__', json.dumps(payload, ensure_ascii=False))
    with open(out_path, 'w', encoding='utf-8') as f: f.write(html)
    print(f'map written: {out_path} ({os.path.getsize(out_path)//1024} KB)')

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--clientlist'); ap.add_argument('--out')
    a = ap.parse_args()
    coords = refresh_coords(a.clientlist) if a.clientlist else load_coords()
    shops, missing, label, src = load_shops(coords)
    out = a.out or os.path.join(BASE, 'KSBC shop sales', f'KSBC SHOP MAP - {label}.html')
    tot = sum(s['s'] for s in shops); den = sum(s['o']+s['r'] for s in shops)
    print(f'period {label} | mapped {len(shops)} shops | unmapped {len(missing)} '
          f'({sum(m["s"] for m in missing):.2f} cs) | mapped sales {tot:.2f} cs | network ST {tot/den*100:.1f}%')
    print(' tiers:', dict(Counter(s['tier'] for s in shops)))
    for m in missing: print('  NO COORDS:', m['bond'], '-', m['name'], f"({m['s']} cs)")
    bad = [s for s in shops if not (8.0 <= s['lat'] <= 13.0 and 74.5 <= s['lng'] <= 77.8)]
    if bad: print(' !! coords outside Kerala:', [(s['name'], s['lat'], s['lng']) for s in bad]); sys.exit(1)
    build_html(shops, missing, label, src, out)
